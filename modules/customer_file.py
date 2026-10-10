"""Customer copy of delivery files (Talimat 43).

One clean sheet of firms with only confident contact values, an information
sheet and, for a full list, a CSV copy. Internal sheets, sources and
confidence columns never reach a customer. A sample keeps the first letter
and domain of each e-mail and the last two digits of each phone.
"""

from __future__ import annotations

import csv
from datetime import date
from pathlib import Path
import re
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from modules import field_merge, run_report


LANGUAGES = ("tr", "pl", "en")
HEADERS = {
    "tr": ("Firma", "Fuar", "Salon", "Stant", "Web sitesi", "E-posta", "Telefon"),
    "pl": ("Firma", "Targi", "Hala", "Stoisko", "Strona www", "E-mail", "Telefon"),
    "en": ("Company", "Fair", "Hall", "Stand", "Website", "E-mail", "Phone"),
}
SHEET_NAMES = {
    "tr": ("Firmalar", "Bilgi"),
    "pl": ("Firmy", "Informacje"),
    "en": ("Companies", "Information"),
}
INFO_LINES = {
    "tr": (
        "Kaynak: fuarın yayımladığı katılımcı listesi ve firmaların kendi web siteleri.",
        "Yalnız firma bilgisi vardır; kişi adı ve kişisel e-posta adresi yoktur.",
        "Boş hücre: o bilgi güvenilir biçimde bulunamadı.",
        "Listede yer almak istemeyen firmalar bildirildiğinde sonraki listelere alınmaz.",
    ),
    "pl": (
        "Źródło: lista wystawców opublikowana przez organizatora targów i strony internetowe firm.",
        "Tylko dane firmowe; bez imion i nazwisk oraz imiennych adresów e-mail.",
        "Pusta komórka: tej informacji nie udało się wiarygodnie ustalić.",
        "Firmy, które zgłoszą sprzeciw, nie trafiają do kolejnych list.",
    ),
    "en": (
        "Source: the exhibitor list published by the fair organiser and the companies' own websites.",
        "Company data only; no personal names or personal e-mail addresses.",
        "Empty cell: that detail could not be found reliably.",
        "Companies that object are left out of later lists.",
    ),
}
DATE_LABEL = {"tr": "Tarih", "pl": "Data", "en": "Date"}
COUNT_LABEL = {"tr": "Firma sayısı", "pl": "Liczba firm", "en": "Companies"}
SAMPLE_NOTE = {
    "tr": "Örnek: e-posta ve telefon kısmen gizlidir. Tam listede tamamı görünür.",
    "pl": "Próbka: e-mail i telefon są częściowo ukryte. Pełna lista pokazuje je w całości.",
    "en": "Sample: e-mail and phone are partly hidden. The full list shows them in full.",
}
TABLE_KEYS = ("Firma", "Fuar", "Salon", "Stant", "Web sitesi", "E-posta", "Telefon")
CONFIDENCE_KEYS = {"Web sitesi": "Web güven", "E-posta": "E-posta güven", "Telefon": "Telefon güven"}
COLUMN_WIDTHS = (48, 26, 10, 12, 34, 32, 18)
HEADER_FILL = "DDEBF7"


def customer_row(table_row: dict, *, fair: str = "") -> dict[str, str]:
    """One firm as a customer sees it: contact values only where the system is confident."""
    row = {}
    for key in TABLE_KEYS:
        value = str(table_row.get(key, "") or "").strip()
        confidence_key = CONFIDENCE_KEYS.get(key)
        if confidence_key and str(table_row.get(confidence_key, "") or "").upper() not in field_merge.CONFIDENT:
            value = ""
        row[key] = value
    # A fair name given for the run wins over the one the list extractor wrote.
    if fair:
        row["Fuar"] = fair
    return row


def customer_rows(
    rows: list[dict], *, keep_firm=None, opt_out_entries: tuple[str, ...] = (), fair: str = "",
) -> tuple[list[dict[str, str]], int, int]:
    """Customer rows of one run; also the counts left out by the firm rule and the opt-out list."""
    kept, left_out, opted_out = run_report.deliver(
        run_report.prepare_rows(rows), opt_out_entries=opt_out_entries, keep_firm=keep_firm,
    )
    return [customer_row(table_row, fair=fair) for _source, table_row in kept], left_out or 0, opted_out


def mask_email(value: str) -> str:
    if "@" not in value:
        return ""
    local, domain = value.split("@", 1)
    return f"{local[:1]}***@{domain}"


def mask_phone(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    return f"*** *** *{digits[-2:]}" if len(digits) >= 7 else ""


def sample_rows(rows: list[dict[str, str]], size: int) -> list[dict[str, str]]:
    """The first firms with a website and an e-mail or phone, contacts partly hidden."""
    chosen = [row for row in rows if row["Web sitesi"] and (row["E-posta"] or row["Telefon"])][:size]
    return [{**row, "E-posta": mask_email(row["E-posta"]), "Telefon": mask_phone(row["Telefon"])} for row in chosen]


def write_customer_file(
    path: Path, rows: list[dict[str, str]], *, lang: str, sample: bool = False, today: date | None = None,
) -> list[Path]:
    """Write the customer workbook (and a CSV for a full list) atomically; return the written paths."""
    if lang not in LANGUAGES:
        raise ValueError(f"unknown language: {lang}")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    firms_name, info_name = SHEET_NAMES[lang]
    sheet = book.active
    sheet.title = firms_name
    sheet.append(list(HEADERS[lang]))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append([run_report._safe(row[key]) for key in TABLE_KEYS])
    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    info = book.create_sheet(info_name)
    lines = [f"{DATE_LABEL[lang]}: {(today or date.today()).strftime('%d.%m.%Y')}", f"{COUNT_LABEL[lang]}: {len(rows)}"]
    lines += list(INFO_LINES[lang]) + ([SAMPLE_NOTE[lang]] if sample else [])
    for index, line in enumerate(lines, start=1):
        info.cell(row=index, column=1, value=line)
    info.column_dimensions["A"].width = 100
    written = [path]
    temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
    if not sample:
        csv_path = path.with_suffix(".csv")
        temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.csv")
        try:
            with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, delimiter=";")
                writer.writerow(HEADERS[lang])
                for row in rows:
                    writer.writerow([run_report._safe(row[key]) for key in TABLE_KEYS])
            temporary.replace(csv_path)
        finally:
            temporary.unlink(missing_ok=True)
        written.append(csv_path)
    return written
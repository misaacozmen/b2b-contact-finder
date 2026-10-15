"""Prospect lists: firms that serve fair exhibitors, found by web search (Talimat 46).

Each service category (hostess, catering, promotional products and print,
photo and video, logistics, furniture and equipment hire) is searched in the
main fair cities with the free search engines.  Every company website becomes
one row; rows that are not a service firm (directories, marketplaces, fair
and news sites, "best 10" lists, foreign or public sites, firms already in an
earlier list) are marked "hayır" with the reason.  The home page of every
remaining site must name the service and fairs or events.  The "evet" rows
become the input of an ordinary free run, which reads the e-mail and phone
from each firm's own website.
"""

from __future__ import annotations

import re
import time
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

import config
from modules import extractor, scorer

CATEGORIES = {
    "TR": {
        "Hostes": ("fuar hostes ajansı {city}", "fuar hostesi hizmeti {city}"),
        "Catering": ("fuar catering firması {city}", "stand catering hizmeti {city}"),
        "Promosyon ve baskı": ("promosyon ürünleri firması {city}", "fuar baskı reklam firması {city}"),
        "Fotoğraf ve video": ("fuar fotoğrafçısı {city}", "kurumsal video prodüksiyon şirketi {city}"),
        "Lojistik": ("fuar lojistiği firması {city}", "fuar taşımacılığı firması {city}"),
        "Mobilya ve ekipman kiralama": ("fuar mobilya kiralama {city}", "fuar ekipman kiralama {city}"),
    },
    "GB": {
        "Hostess": ("exhibition staffing agency {city}", "event hostess agency {city}"),
        "Catering": ("exhibition stand catering {city}", "event catering company {city}"),
        "Promotional products and print": ("promotional products company {city}", "exhibition printing company {city}"),
        "Photo and video": ("exhibition photographer {city}", "event video production company {city}"),
        "Logistics": ("exhibition logistics company {city}", "trade show freight forwarder {city}"),
        "Furniture and equipment hire": ("exhibition furniture hire {city}", "event AV equipment hire {city}"),
    },
}
CITIES = {"TR": ("İstanbul", "Ankara", "İzmir", "Antalya"), "GB": ("London", "Birmingham", "Manchester")}
COUNTRY_NAMES = {"TR": "Türkiye", "GB": "United Kingdom"}
BACKENDS = ("yandex", "yahoo", "duckduckgo", "brave", "mojeek")
BACKENDS_PER_QUERY = 2
RESULTS_PER_QUERY = 20
QUERY_PAUSE_SEC = 3.0
# Marketplaces, directories and venues that show up for service searches but are no service firm.
NON_FIRM_DOMAINS = frozenset({
    "armut.com", "bionluk.com", "bulurum.com", "ciceksepeti.com", "cnrexpo.com", "eventbrite.com",
    "expoassist.net", "expointurkey.org", "firmasec.com", "hepsiburada.com", "ifm.com.tr", "n11.com",
    "sahibinden.com", "sikayetvar.com", "tradefairdates.com", "trendyol.com", "10times.com", "amazon.com.tr",
    "hitched.co.uk", "addtoevent.co.uk", "poptop.uk.com", "tagvenue.com", "hirespace.com", "fiverr.com",
    "upwork.com", "peopleperhour.com", "passo.com.tr", "biletix.com", "biletimgo.com", "bilettakvimi.com",
    "eventpal.co", "restaurantguru.com", "flickr.com", "jooble.org", "webstatsdomain.org", "firmafihristi.com.tr",
    "esnafsiteleri.com", "wordpress.com", "blogspot.com", "wixsite.com", "anfas.com.tr", "fuarizmir.com.tr",
    "izfas.com.tr", "mywed.com", "kurumsal.web.tr", "sanalsirket.com", "refbinder.com", "web.app",
    "firebaseapp.com",
})
# "X Fuarı" names a fair; a service firm writes "fuar catering", "fuar hostesi".
EVENT_RE = re.compile(
    r"(?i)\bfuar[ıi]\w*|\bfuarlar[ıi]\w*|\bfuar merkezi\b|\bexpo cent(?:re|er)\b|\bexhibition cent(?:re|er)\b"
    r"|\btrade (?:fair|show)\b|\b(?:expo|show|summit)\b[^|]*\b20\d\d\b"
)
LISTING_RE = re.compile(
    r"(?i)(?:\ben iyi\b|\ben ucuz\b|\bfirmaları\b|\bfirmalari\b|\blistesi\b|\bfiyatları\b|\bfiyatlari\b|\brehberi\b"
    r"|\byorumları\b|\btop \d+\b|\bbest \d+\b|\b\d+ best\b|\bnear me\b|\bdirectory\b|\bcompare\b"
    r"|\bnedir\b|\bblog\b|\bdetay sayfas|\bşirket profili\b|\bfirma profili\b|\brehberler\b|\barchives\b"
    r"|\bnelere dikkat\b)"
)
WEDDING_RE = re.compile(r"(?i)\b(?:düğün|gelin|nikah|wedding|bride)\w*")
ASSOCIATION_RE = re.compile(r"(?i)\b(?:derneği|odası|birliği|association|federation)\b")
NEWS_WORDS = (
    "haber", "gazete", "news", "dunyasi", "magazin", "blog", "sondakika", "gundem", "times", "postasi", "yanki",
    "havadis", "manset",
)
# Words the home page of a firm in each category uses (folded: ı → i, ş → s ...).
SERVICE_WORDS = {
    "Hostes": ("hostes", "host", "manken", "karsilama", "promotor"),
    "Catering": ("catering", "ikram", "kokteyl", "yemek", "coffee break", "canape", "food", "drinks", "menu"),
    "Promosyon ve baskı": ("promosyon", "baski", "matbaa", "katalog", "brosur", "kartvizit"),
    "Fotoğraf ve video": ("fotograf", "video", "cekim", "produksiyon", "drone"),
    "Lojistik": ("lojistik", "tasima", "nakliye", "gumruk", "forwarding"),
    "Mobilya ve ekipman kiralama": ("kiralama", "kiralik", "mobilya", "ekipman"),
    "Hostess": ("hostess", "staff", "promo staff", "brand ambassador"),
    "Promotional products and print": ("promotional", "branded", "printing", "print", "merchandise"),
    "Photo and video": ("photograph", "video", "film", "production"),
    "Logistics": ("logistics", "freight", "forwarding", "shipping", "transport"),
    "Furniture and equipment hire": ("hire", "rental", "furniture", "equipment"),
}
FAIR_WORDS = (
    "fuar", "stand", "etkinlik", "organizasyon", "kongre", "lansman", "exhibition", "event", "trade show",
    "conference",
)
SEPARATORS = re.compile(r"\s+[|–—•·:-]\s+|\s*\|\s*|;\s+")
# A website found under this many categories is an aggregator, not a firm.
MAX_CATEGORIES = 3


def free_search(query: str, region: str = "") -> list[dict]:
    """Results of the first two free engines that answer."""
    from ddgs import DDGS

    results: list[dict] = []
    answered = 0
    for backend in BACKENDS:
        try:
            with DDGS() as client:
                found = list(client.text(query, max_results=RESULTS_PER_QUERY, backend=backend, **({"region": region} if region else {})))
        except Exception:
            continue
        if found:
            results.extend(found)
            answered += 1
        if answered >= BACKENDS_PER_QUERY:
            break
    return results


def queries(country: str, categories: tuple[str, ...] | None = None) -> list[tuple[str, str]]:
    """(category, query) pairs of the country, in category order."""
    pairs = []
    for category, templates in CATEGORIES[country].items():
        if categories and category not in categories:
            continue
        for template in templates:
            pairs.extend((category, template.format(city=city)) for city in CITIES[country])
    return pairs


def _compact(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", scorer.normalize_text(value))


def firm_name(title: str, website: str) -> str:
    """The part of the page title that names the firm, else the web address's name."""
    core = scorer.compact_domain_core(website)
    for part in (part.strip() for part in SEPARATORS.split(str(title or "")) if part.strip()):
        compact = _compact(part)
        if compact and core and (compact in core or core in compact or compact[:5] == core[:5]):
            return part
    return core[:1].upper() + core[1:] if core else ""


def domains_in(paths: list) -> set[str]:
    """Company domains named in earlier lists (any cell holding a web address or an e-mail)."""
    found: set[str] = set()
    for path in paths:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for cell in row:
                        text = str(cell or "").strip()
                        if "@" in text:
                            text = text.rsplit("@", 1)[-1]
                        if "." in text and " " not in text:
                            domain = scorer.registrable_domain(text)
                            if domain:
                                found.add(domain)
        finally:
            book.close()
    return found - {*config.NON_CORPORATE_EMAIL_DOMAINS, *config.EXTRA_FREE_MAIL_DOMAINS}


def skip_reason(row: dict, known: set[str]) -> str:
    domain = row["domain"]
    title = row["title"]
    if domain in known:
        return "zaten listede"
    if scorer.is_excluded_domain(domain) or any(domain == other or domain.endswith(f".{other}") for other in NON_FIRM_DOMAINS):
        return "rehber, pazar yeri ya da büyük site"
    if scorer.is_public_body_domain(domain):
        return "kamu sitesi"
    if scorer.is_foreign_country_domain(domain):
        return "yabancı uzantı"
    if any(word in domain for word in NEWS_WORDS) or "haberleri" in title.casefold():
        return "haber sitesi"
    if EVENT_RE.search(title):
        return "fuar ya da etkinlik sayfası"
    if LISTING_RE.search(title):
        return "liste ya da rehber sayfası"
    if WEDDING_RE.search(title):
        return "düğün ya da özel gün sayfası"
    if ASSOCIATION_RE.search(title):
        return "dernek ya da oda"
    if len(row["categories"]) >= MAX_CATEGORIES:
        return "çok kategoride çıktı"
    return ""


def site_reason(row: dict, fetch_page) -> str:
    """Why the home page does not look like a firm of the category serving fairs; "" when it does."""
    html = fetch_page(row["website"]) or ""
    text = scorer.normalize_text(extractor._visible_text(html)) if html else ""
    if not text.strip():
        return "site açılmadı"
    words = tuple(word for category in row["categories"] for word in SERVICE_WORDS.get(category, ()))
    if not any(word in text for word in words):
        return "sitede bu hizmet geçmiyor"
    if not any(word in text for word in FAIR_WORDS):
        return "sitede fuar ya da etkinlik geçmiyor"
    return ""


def find(country: str, *, search=None, fetch_page=None, sleep=time.sleep, categories: tuple[str, ...] | None = None,
         known: set[str] | None = None, report=print) -> tuple[list[dict], dict]:
    """Search every query once; return one row per website and the counts."""
    search = search or (lambda query: free_search(query, config.DDGS_REGION))
    if fetch_page is None:
        from modules import crawler

        fetch_page = crawler.fetch_page
    rows: dict[str, dict] = {}
    stats = {"queries": 0, "answered": 0, "results": 0}
    pairs = queries(country, categories)
    for index, (category, query) in enumerate(pairs):
        if index:
            sleep(QUERY_PAUSE_SEC)
        results = search(query)
        stats["queries"] += 1
        stats["answered"] += bool(results)
        stats["results"] += len(results)
        report(f"{index + 1}/{len(pairs)} {query}: {len(results)}")
        for result in results:
            url = str(result.get("href") or result.get("url") or "")
            parts = urlsplit(url)
            domain = scorer.registrable_domain(url)
            if parts.scheme not in {"http", "https"} or not domain:
                continue
            row = rows.setdefault(domain, {
                "domain": domain, "website": f"https://{parts.hostname}/", "title": str(result.get("title") or ""),
                "categories": [], "queries": [],
            })
            if category not in row["categories"]:
                row["categories"].append(category)
            if query not in row["queries"]:
                row["queries"].append(query)
    known = known or set()
    result = []
    cores: set[str] = set()
    for row in rows.values():
        reason = skip_reason(row, known)
        core = scorer.compact_domain_core(row["website"])
        if not reason and core in cores:
            reason = "aynı firma, başka uzantı"
        if not reason:
            cores.add(core)
        result.append({**row, "name": firm_name(row["title"], row["website"]), "ok": not reason, "reason": reason})
    checked = [row for row in result if row["ok"]]
    for index, row in enumerate(checked, start=1):
        row["reason"] = site_reason(row, fetch_page)
        row["ok"] = not row["reason"]
        if index % 25 == 0 or index == len(checked):
            report(f"site kontrolü {index}/{len(checked)}")
    stats["sites_checked"] = len(checked)
    result.sort(key=lambda row: (not row["ok"], row["categories"][0], row["domain"]))
    return result, stats


REVIEW_COLUMNS = ("Uygun", "Neden", "Kategori", "Firma", "Web sitesi", "Başlık", "Sorgu")


def write_review(path, rows: list[dict]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Adaylar"
    sheet.append(list(REVIEW_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in rows:
        sheet.append([
            "evet" if row["ok"] else "hayır", row["reason"], ", ".join(row["categories"]), row["name"],
            row["website"], row["title"], " / ".join(row["queries"][:3]),
        ])
    for letter, width in zip("ABCDEFG", (8, 28, 26, 32, 36, 60, 50)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)


def write_input(path, rows: list[dict], country: str) -> int:
    """The run input: the "evet" rows with company, website and country."""
    book = Workbook()
    sheet = book.active
    sheet.append(["company", "website", "country"])
    count = 0
    for row in rows:
        if row["ok"]:
            sheet.append([row["name"], row["website"], COUNTRY_NAMES[country]])
            count += 1
    book.save(path)
    return count

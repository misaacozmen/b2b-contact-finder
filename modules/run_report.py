"""Best-effort, atomic run summaries alongside the immutable run artifacts."""

from __future__ import annotations

import csv
from datetime import datetime, timezone
import hashlib
import json
import logging
from pathlib import Path
import uuid

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from modules import excel, field_merge, redaction, runtime, scorer


logger = logging.getLogger(__name__)

CANDIDATE_WEBSITE_HEADER = "Aday web sitesi (kontrol edin)"
HEADERS = [
    "Firma", "Web sitesi", "Web kaynağı", "Web güven", CANDIDATE_WEBSITE_HEADER, "E-posta",
    "E-posta kaynağı", "E-posta güven", "Telefon", "Telefon kaynağı",
    "Telefon güven", "Diğer telefonlar", "Referans web sitesi",
    "Referans durumu", "Referans sinyalleri", "Referans telefonu",
    "Girdi web durumu", "Eksik alanlar", "Yayına hazır", "Durum",
    "Ücretli durum", "Not",
]
CONTACT_SHEET = "İletişim"
SUMMARY_SHEET = "Özet"
DETAIL_SHEET = "Detaylar"
CONTACT_HEADERS = ["Firma", "Web sitesi", "E-posta", "Telefon"]
CONTACT_HEADER_FILL = "DDEBF7"
LOW_CONFIDENCE_FILL = "FFF2CC"
CONFIDENCE_COLUMNS = {"Web sitesi": "Web güven", "E-posta": "E-posta güven", "Telefon": "Telefon güven"}
DETAIL_GROUPS = (
    ("DDEBF7", ("Firma", "Yayına hazır", "Eksik alanlar", CANDIDATE_WEBSITE_HEADER, "Diğer telefonlar")),
    ("E2EFDA", ("Web kaynağı", "Web güven", "E-posta kaynağı", "E-posta güven", "Telefon kaynağı", "Telefon güven")),
    ("FCE4D6", ("Referans web sitesi", "Referans telefonu", "Referans durumu", "Referans sinyalleri", "Girdi web durumu")),
    ("EDEDED", ("Durum", "Ücretli durum", "Not")),
)
DETAIL_HEADERS = [header for _fill, headers in DETAIL_GROUPS for header in headers]
TECHNICAL_HEADERS = DETAIL_GROUPS[-1][1]
COLUMN_WIDTHS = {
    "Firma": 45, "Web sitesi": 38, "E-posta": 36, "Telefon": 22, CANDIDATE_WEBSITE_HEADER: 38,
    "Referans web sitesi": 34, "Ücretli durum": 40, "Not": 60,
}
DEFAULT_COLUMN_WIDTH = 18

CONFIDENCE_LABELS = {"HIGH": "Yüksek", "MEDIUM": "Orta", "LOW": "Düşük", "NONE": ""}
FIELD_LABELS = {"website": "web sitesi", "email": "e-posta", "phone": "telefon"}
SIGNAL_LABELS = {"S1_phone": "telefon", "S2_email": "e-posta", "S3_name": "ad", "S4_country": "ülke"}
WEBSITE_SOURCE_LABELS = {
    "OWN_SEARCH": "Arama",
    "OWN_SEARCH+REFERENCE": "Arama + fuar listesi",
    "OWN_SEARCH_CALIBRATED": "Arama (kalibre kural)",
    "REFERENCE_VERIFIED": "Fuar listesi (doğrulandı)",
    "REFERENCE_ACCEPTED": "Fuar listesi (kabul edildi)",
    "REFERENCE_UNREACHABLE": "Fuar listesi (site açılmadı)",
    "REFERENCE_THIN": "Fuar listesi (zayıf kanıt)",
    "PAID_BRIGHTDATA": "Bright Data (ücretli)",
}
CONTACT_SOURCE_LABELS = {
    "SITE": "Firma sitesi",
    "SITE_FREEMAIL": "Firma sitesi (gmail vb.)",
    "REFERENCE_LISTING": "Fuar listesi",
    "PAID_HUNTER": "Hunter (ücretli)",
    "PAID_GOOGLE_PLACES": "Google Places (ücretli)",
}
REFERENCE_TIER_LABELS = {
    "REFERENCE_VERIFIED": "Doğrulandı",
    "REFERENCE_MATCHES_OWN_SEARCH": "Aramayla aynı",
    "REFERENCE_ACCEPTED": "Kabul edildi",
    "REFERENCE_THIN": "Zayıf kanıt",
    "REFERENCE_UNREACHABLE": "Site açılmadı",
    "REFERENCE_UNUSABLE": "Kullanılamaz",
    "REFERENCE_CONFLICT": "Çelişkili",
    "NO_REFERENCE": "Referans yok",
}
LISTED_WEBSITE_LABELS = {"OK": "Geçerli", "FIXED": "Düzeltildi", "INVALID": "Geçersiz", "EMPTY": "Boş"}
STATUS_LABELS = {
    "OK_HIGH_CONFIDENCE": "Doğrulandı (yüksek)",
    "OK_MEDIUM_CONFIDENCE": "Doğrulandı (orta)",
    "REVIEW_NEEDED": "Kontrol gerekli",
    "WEBSITE_NOT_FOUND": "Site bulunamadı",
    "WEBSITE_FETCH_FAILED": "Site açılamadı",
    "WEBSITE_AMBIGUOUS": "Birden çok aday",
    "PROCESSING_FAILED": "İşlenemedi",
}
DETAIL_LABELS = {
    "Web kaynağı": WEBSITE_SOURCE_LABELS,
    "Web güven": CONFIDENCE_LABELS,
    "E-posta güven": CONFIDENCE_LABELS,
    "Telefon güven": CONFIDENCE_LABELS,
    "E-posta kaynağı": CONTACT_SOURCE_LABELS,
    "Telefon kaynağı": CONTACT_SOURCE_LABELS,
    "Referans durumu": REFERENCE_TIER_LABELS,
    "Girdi web durumu": LISTED_WEBSITE_LABELS,
    "Durum": STATUS_LABELS,
}
ALLOWED_STATUSES = {
    "TAMAMLANDI", "KISMI_UCRETLI_ONAY_BEKLIYOR", "KISMI_ZAMANLAYICI_DURDU",
    "KISMI_KOSU_HATASI", "KISMI_KULLANICI_DURDURDU",
}


def _safe(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return excel._safe_cell_value(value)
    return excel._safe_cell_value(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str))


def _candidate_website(row: dict) -> str:
    """Best unverified search candidate, shown only when no website was accepted."""
    if str(row.get("website", "") or "").strip():
        return ""
    for index in (1, 2, 3):
        url = str(row.get(f"candidate_{index}_url", "") or "").strip()
        domain = scorer.normalize_domain(url) if url else ""
        if domain and not scorer.is_excluded_domain(domain) and not scorer.is_foreign_country_domain(domain):
            return url
    return ""


def _display_time(value: object) -> str:
    try:
        return datetime.fromisoformat(str(value).strip()).astimezone().strftime("%d.%m.%Y %H:%M")
    except ValueError:
        return str(value)


def _detail_value(header: str, value: object) -> object:
    """Human-readable Turkish label for the Detaylar sheet; unknown codes stay unchanged."""
    if value is None or value == "":
        return ""
    if header == "Yayına hazır":
        text = str(value).strip().lower()
        if value is True or text == "true":
            return "Evet"
        if value is False or text == "false":
            return "Hayır"
        return value
    if header == "Eksik alanlar":
        return ", ".join(FIELD_LABELS.get(part.strip(), part.strip()) for part in str(value).split(";") if part.strip())
    if header == "Referans sinyalleri":
        return ", ".join(SIGNAL_LABELS.get(part.strip(), part.strip()) for part in str(value).split(",") if part.strip())
    labels = DETAIL_LABELS.get(header)
    if labels is not None:
        return labels.get(str(value).strip(), value)
    return value


def _link_target(value: object) -> str:
    text = str(value or "").strip()
    if text.lower().startswith(("http://", "https://")) and excel._safe_cell_value(text) == text:
        return text
    return ""


def _set_link(cell, value: object) -> None:
    target = _link_target(value)
    if target:
        cell.hyperlink = target
        cell.style = "Hyperlink"


def _write_header(sheet, headers: list[str], fills: dict[str, str]) -> None:
    sheet.append(list(headers))
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fills[header])
        sheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS.get(header, DEFAULT_COLUMN_WIDTH)
    sheet.freeze_panes = "B2"


def _is_confident(row: dict, confidence_header: str) -> bool:
    return str(row.get(confidence_header, "") or "").strip().upper() in field_merge.CONFIDENT


def _write_contact_sheet(sheet, table_rows: list[dict]) -> None:
    _write_header(sheet, CONTACT_HEADERS, {header: CONTACT_HEADER_FILL for header in CONTACT_HEADERS})
    for row_index, row in enumerate(table_rows, start=2):
        for column_index, header in enumerate(CONTACT_HEADERS, start=1):
            value = row.get(header, "")
            cell = sheet.cell(row=row_index, column=column_index, value=_safe(value))
            if header == "Web sitesi":
                _set_link(cell, value)
            confidence_header = CONFIDENCE_COLUMNS.get(header)
            if confidence_header and str(value or "").strip() and not _is_confident(row, confidence_header):
                cell.fill = PatternFill("solid", fgColor=LOW_CONFIDENCE_FILL)
    sheet.auto_filter.ref = sheet.dimensions


def _write_detail_sheet(sheet, table_rows: list[dict]) -> None:
    fills = {header: fill for fill, headers in DETAIL_GROUPS for header in headers}
    _write_header(sheet, DETAIL_HEADERS, fills)
    for row_index, row in enumerate(table_rows, start=2):
        for column_index, header in enumerate(DETAIL_HEADERS, start=1):
            value = row.get(header, "")
            cell = sheet.cell(row=row_index, column=column_index, value=_safe(_detail_value(header, value)))
            if header == CANDIDATE_WEBSITE_HEADER:
                _set_link(cell, value)
    sheet.auto_filter.ref = sheet.dimensions
    for header in TECHNICAL_HEADERS:
        dimension = sheet.column_dimensions[get_column_letter(DETAIL_HEADERS.index(header) + 1)]
        dimension.outlineLevel = 1
        dimension.hidden = True
    sheet.column_dimensions[get_column_letter(len(DETAIL_HEADERS) + 1)].collapsed = True


def _write_summary_sheet(
    sheet, table_rows: list[dict], stage_rows: list[dict], *, run_status: str, generated_at: str,
) -> None:
    total = len(table_rows)

    def found(value_header: str, confidence_header: str) -> int:
        return sum(
            1 for row in table_rows
            if str(row.get(value_header, "") or "").strip() and _is_confident(row, confidence_header)
        )

    counts = (
        ("Web sitesi", found("Web sitesi", "Web güven")),
        ("E-posta", found("E-posta", "E-posta güven")),
        ("Telefon", found("Telefon", "Telefon güven")),
        ("Yayına hazır (web sitesi + e-posta veya telefon)", sum(1 for row in table_rows if row.get("Yayına hazır") is True)),
        ("Aday web sitesi (kontrol edilecek)", sum(1 for row in table_rows if str(row.get(CANDIDATE_WEBSITE_HEADER, "") or "").strip())),
    )
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=CONTACT_HEADER_FILL)
    sheet["A1"] = "Sonuç özeti"
    sheet["A1"].font = Font(bold=True, size=14)
    for row_number, (label, value) in enumerate(
        (("Firma sayısı", total), ("Koşu durumu", run_status), ("Rapor tarihi", _display_time(generated_at))),
        start=3,
    ):
        sheet.cell(row=row_number, column=1, value=label).font = bold
        sheet.cell(row=row_number, column=2, value=_safe(value))
    for column, title in enumerate(("Alan", "Bulunan firma", "Oran"), start=1):
        cell = sheet.cell(row=7, column=column, value=title)
        cell.font = bold
        cell.fill = header_fill
    for offset, (label, count) in enumerate(counts):
        sheet.cell(row=8 + offset, column=1, value=label)
        sheet.cell(row=8 + offset, column=2, value=count)
        ratio = sheet.cell(row=8 + offset, column=3, value=(count / total) if total else 0)
        ratio.number_format = "0.0%"
    sheet.cell(row=14, column=1, value="Aşama tablosu").font = bold
    for column, title in enumerate(("Aşama", "Web sitesi", "E-posta", "Telefon", "Yayına hazır"), start=1):
        cell = sheet.cell(row=15, column=column, value=title)
        cell.font = bold
        cell.fill = header_fill
    for offset, item in enumerate(stage_rows):
        for column, key in enumerate(("stage", "website", "email", "phone", "ready"), start=1):
            sheet.cell(row=16 + offset, column=column, value=_safe(item.get(key, "")))
    note_row = 16 + len(stage_rows) + 1
    sheet.cell(row=note_row, column=1, value="Açıklama").font = bold
    notes = (
        "İletişim: her firmanın web sitesi, e-postası ve telefonu. Satır sırası girdi listesiyle aynıdır.",
        "Sarı hücre: düşük güvenli bilgi. Kullanmadan önce kontrol edin; yukarıdaki sayılara dahil değildir.",
        "Detaylar: aynı firmalar aynı sırada; aday web sitesi, kaynak, güven ve fuar listesi bilgileri.",
        "Detaylar sayfasındaki gri başlıklı teknik sütunlar gizlidir; sütun başlıklarının üstündeki + işaretiyle açılır.",
    )
    for offset, text in enumerate(notes, start=1):
        sheet.cell(row=note_row + offset, column=1, value=text)
    for letter, width in (("A", 50), ("B", 24), ("C", 12), ("D", 12), ("E", 16)):
        sheet.column_dimensions[letter].width = width


def write_results_workbook(
    path: Path, table_rows: list[dict], stage_rows: list[dict], *, run_status: str, generated_at: str,
) -> None:
    """Write the three-sheet sonuclar.xlsx atomically."""
    book = Workbook()
    contact = book.active
    contact.title = CONTACT_SHEET
    _write_contact_sheet(contact, table_rows)
    _write_summary_sheet(
        book.create_sheet(SUMMARY_SHEET), table_rows, stage_rows,
        run_status=run_status, generated_at=generated_at,
    )
    _write_detail_sheet(book.create_sheet(DETAIL_SHEET), table_rows)
    book.active = 0
    path = Path(path)
    temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _table_row(row: dict) -> dict[str, object]:
    row = dict(row)
    field_merge.annotate(row, None)
    last_error = str(row.get("last_error", "") or "")
    paid_status = str(row.get("paid_state", "") or "")
    if last_error:
        paid_status = f"{paid_status} — {last_error}" if paid_status else last_error
    gaps = ";".join(field for field in ("website", "email", "phone") if field in field_merge.field_gaps(row))
    ready = field_merge.ready_for_publication(row)
    return {
        "Firma": row.get("company", ""),
        "Web sitesi": row.get("website", ""),
        "Web kaynağı": row.get("website_source", ""),
        "Web güven": row.get("website_confidence", ""),
        CANDIDATE_WEBSITE_HEADER: _candidate_website(row),
        "E-posta": row.get("email", ""),
        "E-posta kaynağı": row.get("email_source_tier", row.get("email_source", "")),
        "E-posta güven": row.get("email_confidence", ""),
        "Telefon": row.get("phone", ""),
        "Telefon kaynağı": row.get("phone_source_tier", row.get("phone_source", "")),
        "Telefon güven": row.get("phone_confidence", ""),
        "Diğer telefonlar": row.get("alternative_phones", row.get("alternative_phone", "")),
        "Referans web sitesi": row.get("reference_website", ""),
        "Referans durumu": row.get("reference_tier", ""),
        "Referans sinyalleri": row.get("reference_signals", ""),
        "Referans telefonu": row.get("reference_phone", row.get("listed_phone", "")),
        "Girdi web durumu": row.get("listed_website_status", ""),
        "Eksik alanlar": gaps,
        "Yayına hazır": bool(ready),
        "Durum": row.get("status", ""),
        "Ücretli durum": paid_status,
        "Not": str(row.get("reason", "") or "")[:300],
    }


def _atomic_text(path: Path, value: str, *, encoding: str = "utf-8") -> None:
    temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp{path.suffix}")
    try:
        temporary.write_text(value, encoding=encoding, newline="")
        temporary.replace(path)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def _sheet_value(stage: dict, field: str) -> object:
    return stage.get(field, "") if isinstance(stage, dict) else ""


def _confidence(stage: dict, field: str) -> str:
    value = _sheet_value(stage, f"{field}_confidence")
    return str(value or "").upper()


def _ready(stage: dict) -> bool:
    if not isinstance(stage, dict):
        return False
    return (
        _confidence(stage, "website") in field_merge.CONFIDENT
        and any(_confidence(stage, field) in field_merge.CONFIDENT for field in ("email", "phone"))
    )


def _stage_summary(rows: list[dict]) -> list[dict[str, object]]:
    stages = (
        ("A — Bağımsız arama (referanssız)", "stage_a"),
        ("A + Referans", "stage_ab"),
        ("Final (A + Referans + Ücretli)", "final"),
    )
    fields = ("website", "email", "phone", "ready")
    result: list[dict[str, object]] = []
    total = len(rows)
    for stage_label, stage_key in stages:
        values = []
        for field in fields:
            count = 0
            for row in rows:
                stage = row if stage_key == "final" else row.get(stage_key, {})
                if field == "ready":
                    count += int(_ready(stage))
                    continue
                count += int(bool(_sheet_value(stage, field)) and _confidence(stage, field) in {"HIGH", "MEDIUM"})
            values.append(f"{count} ({count / total * 100:.1f}%)" if total else "0 (0.0%)")
        result.append({"stage": stage_label, **dict(zip(fields, values))})
    return result


def _md(value: object) -> str:
    text = str(value if value is not None else "")
    return text.replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def _build_report(
    rows: list[dict], *, run_id: str, run_status: str, status_detail: str,
    elapsed_seconds: float | None, telemetry: dict | None, generated_at: str,
    file_names: list[str],
) -> str:
    started_at = "bilinmiyor"
    try:
        from modules import checkpoint
        state = checkpoint.load_run_state_by_id(run_id)
        if state:
            context = state.get("context") or {}
            started_at = str(context.get("started_at", started_at))
    except Exception:
        logger.exception("run start time could not be read for report run_id=%s", run_id)
    elapsed = f"{float(elapsed_seconds):.1f} sn" if elapsed_seconds is not None else "bilinmiyor"
    table = _stage_summary(rows)
    lines = [
        "# Koşu raporu", "", f"- Koşu kimliği: `{_md(run_id)}`",
        f"- Durum: `{_md(run_status)}`", f"- Durum ayrıntısı: {_md(status_detail)}",
        f"- Başlangıç: {_md(started_at)}", f"- Bitiş: {_md(generated_at)}",
        f"- Süre: {_md(elapsed)}", "", "## Oran tablosu", "",
        "| Aşama | Web sitesi | E-posta | Telefon | Yayına hazır |",
        "|---|---:|---:|---:|---:|",
    ]
    lines.extend(
        f"| {_md(row['stage'])} | {_md(row['website'])} | {_md(row['email'])} | {_md(row['phone'])} | {_md(row['ready'])} |"
        for row in table
    )
    def hit_rate(*, brand: bool, verified_only: bool = False) -> tuple[int, int, str]:
        matches = 0
        comparable = 0
        for row in rows:
            stage_a = row.get("stage_a") if isinstance(row.get("stage_a"), dict) else {}
            if verified_only and not str(stage_a.get("status", "")).startswith("OK_"):
                continue
            site = scorer.normalize_domain(str(stage_a.get("website", "") or ""))
            reference = scorer.normalize_domain(str(row.get("reference_website", "") or ""))
            if not site or not reference:
                continue
            comparable += 1
            strict_match = scorer.same_registrable_domain(site, reference)
            brand_match = strict_match or (
                scorer.compact_domain_core(site) == scorer.compact_domain_core(reference)
            )
            matches += int(brand_match if brand else strict_match)
        percentage = f"{matches / comparable * 100:.1f}%" if comparable else "0.0%"
        return matches, comparable, percentage

    strict = hit_rate(brand=False)
    brand = hit_rate(brand=True)
    strict_ok = hit_rate(brand=False, verified_only=True)
    brand_ok = hit_rate(brand=True, verified_only=True)
    lines.extend([
        "", "## Bağımsız arama isabet vekili", "",
        f"A isabeti — aynı alan adı (katı): {strict[0]} / {strict[1]} ({strict[2]})",
        f"A isabeti — aynı marka çekirdeği (TLD farkı dahil, ör. firma.com ↔ firma.com.tr): {brand[0]} / {brand[1]} ({brand[2]})",
        f"  - Yalnız doğrulanmış (OK) A satırları — katı: {strict_ok[0]} / {strict_ok[1]} ({strict_ok[2]})",
        f"  - Yalnız doğrulanmış (OK) A satırları — marka: {brand_ok[0]} / {brand_ok[1]} ({brand_ok[2]})",
    ])

    confidence_fields = ("website", "email", "phone")
    lines.extend(["", "## Güven dağılımı", "", "| Alan | HIGH | MEDIUM | LOW | NONE |", "|---|---:|---:|---:|---:|"])
    for field in confidence_fields:
        counts = {level: 0 for level in ("HIGH", "MEDIUM", "LOW", "NONE")}
        for row in rows:
            level = str(row.get(f"{field}_confidence", "NONE") or "NONE").upper()
            counts[level if level in counts else "NONE"] += 1
        lines.append(f"| {_md(field)} | {counts['HIGH']} | {counts['MEDIUM']} | {counts['LOW']} | {counts['NONE']} |")

    lines.extend(["", "## Kaynak dağılımı", ""])
    for field, key in (("Web sitesi", "website_source"), ("E-posta", "email_source_tier"), ("Telefon", "phone_source_tier")):
        counts: dict[str, int] = {}
        for row in rows:
            source = str(row.get(key, "") or "NONE")
            counts[source] = counts.get(source, 0) + 1
        lines.append(f"- {field}: " + ", ".join(f"{_md(name)}={count}" for name, count in sorted(counts.items())))

    lines.extend(["", "## Referans durumu ve girdi temizliği", ""])
    for key, label in (("reference_tier", "Referans"), ("listed_website_status", "Girdi web")):
        counts: dict[str, int] = {}
        for row in rows:
            value = str(row.get(key, "") or "NONE")
            counts[value] = counts.get(value, 0) + 1
        lines.append(f"- {label}: " + (", ".join(f"{_md(name)}={count}" for name, count in sorted(counts.items())) or "veri yok"))

    paid_indexes = sum(bool(row.get("paid_required")) and int(row.get("paid_attempts", 0) or 0) > 0 for row in rows)
    skipped = sum(bool(row.get("paid_required")) and int(row.get("paid_attempts", 0) or 0) == 0 for row in rows)
    lines.extend(["", "## Ücretli aşama", "", f"- İşlem gören firma: {paid_indexes}", f"- Atlanan ücretli firma: {skipped}"])
    provider_budgets = (telemetry or {}).get("provider_budgets", {}) if isinstance(telemetry, dict) else {}
    if isinstance(provider_budgets, dict) and provider_budgets:
        lines.extend(["", "| Sağlayıcı | Çağrı | Başarılı | Başarısız | Belirsiz | Bütçe |", "|---|---:|---:|---:|---:|---:|"])
        for provider, values in sorted(provider_budgets.items()):
            if not isinstance(values, dict):
                continue
            total_calls = sum(int(values.get(key, 0) or 0) for key in ("done", "failed", "unknown", "reserved"))
            lines.append(f"| {_md(provider)} | {total_calls} | {values.get('done', 0)} | {values.get('failed', 0)} | {values.get('unknown', 0)} | {values.get('effective_limit', 0)} |")
    else:
        lines.append("- Sağlayıcı bütçe ve çağrı verisi: ölçülmedi")

    lines.extend(["", "## Ücretsiz arama canary ve motor sağlığı", ""])
    canary = (telemetry or {}).get("free_search_canary") if isinstance(telemetry, dict) else None
    if isinstance(canary, dict):
        alive = ", ".join(str(value) for value in canary.get("alive", [])) or "yok"
        dead = ", ".join(str(value) for value in canary.get("dead", [])) or "yok"
        lines.append(f"- Canary durumu: {canary.get('status', 'bilinmiyor')}; yanıt veren: {alive}; yanıt vermeyen: {dead}")
    else:
        lines.append("- Canary durumu: ölçülmedi")
    lines.append("")
    health = (telemetry or {}).get("free_search_backend_health") if isinstance(telemetry, dict) else None
    if not isinstance(health, dict):
        try:
            health = runtime.free_backend_health_snapshot()
        except Exception:
            health = {}
    if health:
        lines.extend(["", "| Motor | OK | Boş | Hata |", "|---|---:|---:|---:|"])
        for backend, values in sorted(health.items()):
            if isinstance(values, dict):
                lines.append(
                    f"| {_md(backend)} | {int(values.get('ok', 0) or 0)} | "
                    f"{int(values.get('empty', 0) or 0)} | {int(values.get('error', 0) or 0)} |"
                )
    else:
        lines.append("Motor sağlık ölçümleri: ölçülmedi")

    lines.extend(["", "## Hatalar", ""])
    categories = (
        ("İzole edilen firmalar", lambda row: str(row.get("last_error", "")).startswith("invariant:")),
        ("PROCESSING_FAILED firmalar", lambda row: str(row.get("status", "")).upper() == "PROCESSING_FAILED"),
        ("Belirsiz/bütçesi yetmeyen ücretli firmalar", lambda row: str(row.get("last_error", "")) in {"provider_call_unknown_outcome", "paid_budget_exhausted"}),
    )
    for label, predicate in categories:
        lines.append(f"### {label}")
        matches_rows = [row for row in rows if predicate(row)]
        if matches_rows:
            lines.extend(f"- {_md(row.get('company', ''))}: {_md(str(row.get('last_error') or row.get('reason') or '')[:200])}" for row in matches_rows)
        else:
            lines.append("- Yok")
    lines.extend(["", "## Dosyalar", "", *[f"- `{_md(name)}`" for name in file_names], ""])
    return redaction.redact_text("\n".join(lines))


def write_run_report(
    rows: list[dict],
    *,
    output_root: Path,
    run_status: str,
    status_detail: str,
    elapsed_seconds: float | None,
    telemetry: dict | None = None,
) -> dict[str, str]:
    """Write all available summary files; collect every filesystem error."""
    errors: dict[str, str] = {}
    files: dict[str, str] = {}
    safe_status = str(run_status) if str(run_status) in ALLOWED_STATUSES else "KISMI_KOSU_HATASI"
    target_root = Path(output_root)
    normalized_rows = [dict(row) for row in rows if isinstance(row, dict)]
    for row in normalized_rows:
        field_merge.annotate(row, None)
    table_rows = [_table_row(row) for row in normalized_rows]
    try:
        target_root.mkdir(parents=True, exist_ok=True)
    except Exception as exc:
        logger.exception("run report output directory could not be created: %s", target_root)
        errors["output_root"] = f"{type(exc).__name__}:{str(exc)[:300]}"

    generated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    workbook_path = target_root / "sonuclar.xlsx"
    try:
        write_results_workbook(
            workbook_path, table_rows, _stage_summary(normalized_rows),
            run_status=safe_status, generated_at=generated_at,
        )
        files[workbook_path.name] = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    except Exception as exc:
        logger.exception("run summary workbook could not be written")
        errors["sonuclar.xlsx"] = f"{type(exc).__name__}:{str(exc)[:300]}"
        csv_path = target_root / "sonuclar.csv"
        try:
            temporary = target_root / f".sonuclar.{uuid.uuid4().hex}.tmp.csv"
            try:
                with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
                    writer = csv.DictWriter(handle, fieldnames=HEADERS, delimiter=";", extrasaction="ignore")
                    writer.writeheader()
                    for row in table_rows:
                        writer.writerow({key: _safe(row.get(key, "")) for key in HEADERS})
                temporary.replace(csv_path)
            finally:
                temporary.unlink(missing_ok=True)
            files[csv_path.name] = hashlib.sha256(csv_path.read_bytes()).hexdigest()
        except Exception as csv_exc:
            logger.exception("run summary CSV fallback could not be written")
            errors["sonuclar.csv"] = f"{type(csv_exc).__name__}:{str(csv_exc)[:300]}"

    run_id = str((normalized_rows[0].get("run_id", "") if normalized_rows else "") or target_root.parent.name)
    output_names = list(files)
    output_names.append("rapor.md")
    report_path = target_root / "rapor.md"
    try:
        markdown = _build_report(
            normalized_rows, run_id=run_id, run_status=safe_status,
            status_detail=str(status_detail)[:300], elapsed_seconds=elapsed_seconds,
            telemetry=telemetry, generated_at=generated_at, file_names=output_names + ["run_status.json"],
        )
        _atomic_text(report_path, markdown)
        files[report_path.name] = hashlib.sha256(report_path.read_bytes()).hexdigest()
    except Exception as exc:
        logger.exception("run markdown report could not be written")
        errors["rapor.md"] = f"{type(exc).__name__}:{str(exc)[:300]}"

    status_payload = {
        "run_status": safe_status,
        "status_detail": str(status_detail)[:300],
        "generated_at": generated_at,
        "row_count": len(normalized_rows),
        "files": files,
    }
    status_path = target_root / "run_status.json"
    try:
        _atomic_text(status_path, json.dumps(status_payload, ensure_ascii=False, sort_keys=True, indent=2) + "\n")
        files[status_path.name] = hashlib.sha256(status_path.read_bytes()).hexdigest()
    except Exception as exc:
        logger.exception("run status file could not be written")
        errors["run_status.json"] = f"{type(exc).__name__}:{str(exc)[:300]}"
    return errors

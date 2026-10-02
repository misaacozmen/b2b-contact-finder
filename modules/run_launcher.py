"""Launch and watch fair runs for the desktop panel (Talimat 21).

The panel runs the same main.py commands as FUAR_KOSU_KILAVUZU.md. This module
holds the testable parts: input checks, commands, paid limits, progress and
archiving. It never changes how a run searches or accepts results.
"""

from __future__ import annotations

import json
import os
from datetime import date, datetime
from pathlib import Path
import re
import shutil
import subprocess
import time

import config
from modules import excel, run_budget, run_report


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RUNTIME_PYTHON = PROJECT_ROOT / ".runtime" / "python3147-sqlite3534" / "python.exe"
ARCHIVE_ROOT = Path.home() / "Documents" / "FUAR_SONUCLARI"
PANEL_LOG_DIR = PROJECT_ROOT / "state" / "panel"
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)

MODE_FREE = "free"
MODE_PLACES_HUNTER = "places_hunter"
MODE_BRIGHTDATA = "brightdata"
MODE_LABELS = {
    MODE_FREE: "Ücretsiz (önerilen)",
    MODE_PLACES_HUNTER: "Ücretli: Places + Hunter (eksik telefon ve e-posta)",
    MODE_BRIGHTDATA: "Ücretli: Bright Data (sitesi olmayan listeler)",
}
BRIGHTDATA_BUDGET_PER_FIRM = 4
# Worst-case list price; failed queries are not billed, so real cost is lower.
BRIGHTDATA_MAX_USD_PER_1000 = 3.0
PROVIDER_LABELS = {"brightdata": "Bright Data", "google_places": "Google Places", "hunter": "Hunter"}
PHASE_LABELS = {
    "ÜCRETSİZ ARAMA": "Ücretsiz arama",
    "ÜCRETLİ TAMAMLAMA": "Ücretli tamamlama",
    "ÜCRETLİ TAMAMLAMA ATLANDI": "Ücretli tamamlama atlandı",
    "RAPOR": "Rapor",
}
# Inherited settings that would silently change the run; the panel sets them itself.
CONTROLLED_ENV = ("SEARCH_PROVIDER", "ENABLE_LLM_ARBITER", "ENABLE_LINKEDIN_COMPANY_LOOKUP", "B2B_TEST_OFFLINE")

_PHASE = re.compile(r"=== AŞAMA (\d)/3: ([^=—]+)")
_PROGRESS = re.compile(r"\b(?:free|paid)_completed=(\d+) (?:free|paid)_total=(\d+)\b.*?\bcompany=(.*)$")
_WINDOWS_BAD_NAME = re.compile(r'[<>:"/\\|?*]')


def inspect_input(path: Path) -> dict:
    """Check a fair list before starting; never raises."""
    path = Path(path)
    result = {"ok": False, "error": "", "company_count": 0, "has_website": False, "has_phone": False, "has_email": False}
    if not path.is_file():
        result["error"] = "Dosya bulunamadı."
        return result
    if path.suffix.lower() != ".xlsx":
        result["error"] = "Yalnız .xlsx dosyası seçin."
        return result
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            header = next(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True), ())
        finally:
            workbook.close()
        if "company" not in [str(value or "").strip().lower() for value in header]:
            result["error"] = 'İlk satırda "company" başlığı yok.'
            return result
        records = excel.read_company_records(path)
    except Exception as exc:
        result["error"] = f"Dosya okunamadı ({type(exc).__name__})."
        return result
    if not records:
        result["error"] = "Dosyada firma yok."
        return result
    result.update({
        "ok": True,
        "company_count": len(records),
        "has_website": any(record.get("website") for record in records),
        "has_phone": any(record.get("listed_phone") for record in records),
        "has_email": any(record.get("listed_email") for record in records),
    })
    return result


def brightdata_budget(company_count: int) -> int:
    return max(0, int(company_count)) * BRIGHTDATA_BUDGET_PER_FIRM


def build_command(mode: str, input_path: Path, company_count: int) -> tuple[list[str], dict[str, str]]:
    """Return the guide's command for a mode and the environment it needs."""
    command = [str(RUNTIME_PYTHON), str(PROJECT_ROOT / "main.py"), "--input", str(Path(input_path)), "--non-interactive"]
    if mode == MODE_FREE:
        return command + ["--no-allow-paid", "--finalize-without-paid"], {}
    paid_env = {"ENABLE_LLM_ARBITER": "0", "ENABLE_LINKEDIN_COMPANY_LOOKUP": "0"}
    closed = ["--brandfetch-budget", "0", "--linkedin-company-budget", "0", "--llm-budget", "0"]
    if mode == MODE_PLACES_HUNTER:
        return command + ["--allow-paid", "--brightdata-budget", "0", *closed], paid_env
    if mode == MODE_BRIGHTDATA:
        budget = str(brightdata_budget(company_count))
        return command + ["--allow-paid", "--brightdata-budget", budget, *closed], {**paid_env, "SEARCH_PROVIDER": "brightdata"}
    raise ValueError(f"unknown mode: {mode}")


def child_env(overrides: dict[str, str], base: dict[str, str] | None = None) -> dict[str, str]:
    env = dict(os.environ if base is None else base)
    for name in CONTROLLED_ENV:
        env.pop(name, None)
    env.update(overrides)
    env["PYTHONIOENCODING"] = "utf-8"
    return env


def paid_limits(mode: str, company_count: int) -> dict[str, int]:
    """Upper call limits the run itself enforces for each paid provider."""
    if mode == MODE_FREE:
        return {}
    caps = {
        "brightdata": brightdata_budget(company_count) if mode == MODE_BRIGHTDATA else 0,
        "google_places": None,
        "hunter": None,
        "brandfetch": 0,
    }
    budgets = run_budget.calculate_paid_api_budgets(company_count, caps)
    return {provider: int(calls) for provider, calls in budgets.items() if provider in PROVIDER_LABELS and calls > 0}


def brightdata_cost_ceiling(calls: int) -> float:
    return round(max(0, int(calls)) * BRIGHTDATA_MAX_USD_PER_1000 / 1000, 2)


def paid_confirmation_text(mode: str, company_count: int) -> str:
    limits = paid_limits(mode, company_count)
    lines = ["Bu koşu ücretlidir. Sistem şu sınırları aşmaz:"]
    for provider, calls in limits.items():
        line = f"- {PROVIDER_LABELS[provider]}: en çok {calls} çağrı"
        if provider == "brightdata":
            line += f" (üst sınır ~{brightdata_cost_ceiling(calls):.2f} $; yalnız sonuç veren sorgular ücretlidir, genellikle çok daha az tutar)"
        lines.append(line)
    lines.append("Places ve Hunter yalnız anahtarları kayıtlıysa ve eksik alan varsa çağrılır.")
    lines.append("Devam edilsin mi?")
    return "\n".join(lines)


def phase_label(phase: str) -> str:
    return PHASE_LABELS.get(str(phase or "").strip(), str(phase or "").strip())


def new_progress() -> dict:
    return {"stage": 0, "phase": "", "phase_started": 0.0, "done": 0, "total": 0, "company": ""}


def _log_time(line: str) -> float:
    try:
        return time.mktime(datetime.strptime(line[:19], "%Y-%m-%d %H:%M:%S").timetuple())
    except ValueError:
        return 0.0


def parse_progress(text: str, state: dict | None = None) -> dict:
    """Fold complete log lines into the progress state."""
    state = dict(state or new_progress())
    for line in text.splitlines():
        phase = _PHASE.search(line)
        if phase:
            state.update({
                "stage": int(phase.group(1)), "phase": phase.group(2).strip(),
                "phase_started": _log_time(line), "done": 0, "total": 0, "company": "",
            })
            continue
        progress = _PROGRESS.search(line)
        if progress:
            state.update({
                "done": int(progress.group(1)), "total": int(progress.group(2)),
                "company": progress.group(3).strip(),
            })
    return state


def estimate_remaining(elapsed: float, done: int, total: int) -> float | None:
    if done <= 0 or total <= 0 or done >= total or elapsed <= 0:
        return None
    return elapsed / done * (total - done)


def find_new_run_dir(runs_root: Path, before: set[str], started_at: float) -> Path | None:
    runs_root = Path(runs_root)
    if not runs_root.is_dir():
        return None
    directories = [path for path in runs_root.iterdir() if path.is_dir()]
    created = [path for path in directories if path.name not in before]
    if created:
        return max(created, key=lambda path: path.stat().st_mtime)
    touched = [
        path for path in directories
        if (path / "output" / "logs.txt").is_file()
        and (path / "output" / "logs.txt").stat().st_mtime >= started_at
    ]
    return max(touched, key=lambda path: (path / "output" / "logs.txt").stat().st_mtime) if touched else None


def read_run_status(run_dir: Path) -> dict:
    path = Path(run_dir) / "output" / "run_status.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def summarize_results(run_dir: Path) -> dict:
    """Count firms and filled contact cells on the results workbook's contact sheet."""
    path = Path(run_dir) / "output" / "sonuclar.xlsx"
    summary = {"firms": 0, "website": 0, "email": 0, "phone": 0}
    if not path.is_file():
        return summary
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if run_report.CONTACT_SHEET not in workbook.sheetnames:
            return summary
        rows = workbook[run_report.CONTACT_SHEET].iter_rows(values_only=True)
        header = [str(value or "") for value in next(rows, ())]
        columns = {key: header.index(name) for key, name in zip(("firms", "website", "email", "phone"), run_report.CONTACT_HEADERS) if name in header}
        for row in rows:
            for key, index in columns.items():
                if index < len(row) and str(row[index] or "").strip():
                    summary[key] += 1
    finally:
        workbook.close()
    return summary


def archive_run(run_dir: Path, input_path: Path, fair_name: str, *, archive_root: Path = ARCHIVE_ROOT, today: date | None = None) -> Path:
    """Copy results, report and input into FUAR_SONUCLARI; never overwrites."""
    name = " ".join(str(fair_name or "").split())
    prefix = re.sub(r"[^\w]", "", name)
    if not prefix or _WINDOWS_BAD_NAME.search(name):
        raise ValueError("Geçerli bir fuar adı girin (harf, rakam ve boşluk).")
    today = today or date.today()
    target = Path(archive_root) / f"{today:%Y-%m} {name}"
    input_path = Path(input_path)
    files = {
        Path(run_dir) / "output" / "sonuclar.xlsx": target / f"{prefix}_{today:%Y}_sonuclar.xlsx",
        Path(run_dir) / "output" / "rapor.md": target / f"{prefix}_{today:%Y}_rapor.md",
        input_path: target / f"{prefix}_{today:%Y}_girdi{input_path.suffix.lower()}",
    }
    missing = [source.name for source in files if not source.is_file()]
    if missing:
        raise FileNotFoundError("Eksik dosya: " + ", ".join(missing))
    existing = [destination.name for destination in files.values() if destination.exists()]
    if existing:
        raise FileExistsError("Arşivde zaten var: " + ", ".join(existing))
    target.mkdir(parents=True, exist_ok=True)
    for source, destination in files.items():
        shutil.copy2(source, destination)
    return target


def pipeline_running() -> bool | None:
    """True when another main.py run is active; None when it cannot be checked."""
    if os.name != "nt":
        return None
    script = "Get-CimInstance Win32_Process -Filter \"Name='python.exe'\" | ForEach-Object { $_.CommandLine }"
    try:
        completed = subprocess.run(
            ["powershell", "-NoProfile", "-Command", script],
            capture_output=True, text=True, errors="replace", timeout=30, creationflags=NO_WINDOW,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    return any("main.py" in line and "--input" in line for line in completed.stdout.splitlines())


def keep_awake(enabled: bool) -> bool:
    """Ask Windows not to sleep while a run is active."""
    try:
        import ctypes

        flags = 0x80000000 | (0x00000001 if enabled else 0)
        return bool(ctypes.windll.kernel32.SetThreadExecutionState(flags))
    except (AttributeError, OSError):
        return False


def tail_text(path: Path, lines: int = 8) -> str:
    try:
        data = Path(path).read_bytes()[-20_000:]
    except OSError:
        return ""
    return "\n".join(data.decode("utf-8", errors="replace").splitlines()[-lines:])


class RunSession:
    """One panel-launched main.py run. poll() is cheap; call it every few seconds."""

    def __init__(self, mode: str, input_path: Path, company_count: int, *,
                 runs_root: Path | None = None, log_dir: Path = PANEL_LOG_DIR, popen=subprocess.Popen):
        self.mode = mode
        self.input_path = Path(input_path)
        self.company_count = int(company_count)
        self.runs_root = Path(runs_root or config.RUNS_DIR)
        self.log_dir = Path(log_dir)
        self._popen = popen
        self.process = None
        self.run_dir: Path | None = None
        self.console_path: Path | None = None
        self.started_at = 0.0
        self.progress = new_progress()
        self._before: set[str] = set()
        self._offset = 0
        self._rest = b""

    def start(self) -> None:
        if self.process is not None:
            raise RuntimeError("run already started")
        command, overrides = build_command(self.mode, self.input_path, self.company_count)
        self._before = {path.name for path in self.runs_root.iterdir() if path.is_dir()} if self.runs_root.is_dir() else set()
        self.started_at = time.time()
        self.log_dir.mkdir(parents=True, exist_ok=True)
        self.console_path = self.log_dir / f"kosu_{datetime.now():%Y%m%d_%H%M%S}.log"
        with open(self.console_path, "wb") as console:
            self.process = self._popen(
                command, cwd=str(PROJECT_ROOT), env=child_env(overrides),
                stdout=console, stderr=subprocess.STDOUT, stdin=subprocess.DEVNULL,
                creationflags=NO_WINDOW,
            )

    def _read_log(self) -> None:
        path = self.run_dir / "output" / "logs.txt"
        if not path.is_file():
            return
        with open(path, "rb") as stream:
            stream.seek(self._offset)
            data = stream.read()
        self._offset += len(data)
        data = self._rest + data
        cut = data.rfind(b"\n") + 1
        self._rest = data[cut:]
        self.progress = parse_progress(data[:cut].decode("utf-8", errors="replace"), self.progress)

    def poll(self) -> dict:
        exit_code = self.process.poll() if self.process is not None else None
        if self.run_dir is None and self.started_at:
            self.run_dir = find_new_run_dir(self.runs_root, self._before, self.started_at)
        if self.run_dir is not None:
            self._read_log()
        now = time.time()
        started = self.progress["phase_started"]
        remaining = estimate_remaining(now - started, self.progress["done"], self.progress["total"]) if started else None
        return {
            **self.progress,
            "running": self.process is not None and exit_code is None,
            "exit_code": exit_code,
            "run_dir": str(self.run_dir) if self.run_dir else "",
            "elapsed": now - self.started_at if self.started_at else 0.0,
            "remaining": remaining,
        }

    def stop(self) -> None:
        if self.process is not None and self.process.poll() is None:
            subprocess.run(
                ["taskkill", "/PID", str(self.process.pid), "/T", "/F"],
                capture_output=True, creationflags=NO_WINDOW,
            )

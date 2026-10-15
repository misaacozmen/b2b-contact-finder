"""Cold e-mail sending from a tracking workbook (Talimat 45).

One e-mail per firm, only to a generic company mailbox, from the user's own
Zoho address.  The tool sends one e-mail at a time with a pause between them,
only on weekdays inside the recipients' office hours, and stops at a daily
limit counted over every tracking workbook.  A firm that was written to, that
answered, that carries a note or that is on the opt-out list is never written
to.  Each sent e-mail goes to the log in state/ first and then to the
tracking workbook, so neither a crash nor an open workbook leads to a second
e-mail.
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from email.message import EmailMessage
from email.utils import formataddr, formatdate, make_msgid, parseaddr
import json
from pathlib import Path
import random
import re
import smtplib
import time
import uuid
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

import config
from modules import customer_file, email_kind, legal_form, opt_out, scorer


COLUMNS = (
    "Firma", "Hitap adı", "Web sitesi", "E-posta", "Telefon", "Gönderim tarihi", "Gönderim sonucu",
    "Yanıt", "Yanıt tarihi", "Not",
)
# Columns the tool fills; an older tracking workbook gets them at its end.
TOOL_COLUMNS = ("Hitap adı", "Gönderim sonucu")
COLUMN_WIDTHS = {"Firma": 44, "Hitap adı": 24, "Web sitesi": 32, "E-posta": 32, "Not": 40}
SENT = "gönderildi"
TEST = "deneme"
TIMEZONES = {"TR": "Europe/Istanbul", "PL": "Europe/Warsaw", "GB": "Europe/London"}
# Legal forms and trade words that end the short name used in the greeting.
SALUTATION_STOP = frozenset({
    "san", "sanayi", "sanayii", "tic", "ticaret", "ve", "ltd", "limited", "sti", "sirketi", "sirket", "as",
    "a", "s", "anonim", "ith", "ihr", "ithalat", "ihracat", "paz", "pazarlama", "ins", "insaat", "taah",
    "imalat", "ic", "dis", "plc", "llp", "inc", "gmbh", "co", "company", "sp", "z", "o", "oo", "spolka", "sa",
    "srl", "bv",
})
EMAIL_RE = re.compile(r"[^@\s]+@[^@\s]+\.[a-z]{2,}", re.IGNORECASE)
PLACEHOLDER_RE = re.compile(r"\[[^\]\n]+\]")
FIELD_RE = re.compile(r"\{([^{}\n]*)\}")


class OutreachError(Exception):
    """A problem that stops the whole sending run."""


def _text(value: object) -> str:
    return " ".join(str(value or "").split())


def _capital(word: str, country: str) -> str:
    # A dotted capital İ shows Turkish spelling, where I is the dotless ı;
    # catalogues often write Turkish words in plain ASCII capitals (HIZMETLERI).
    if country == "TR" and "İ" in word:
        lower = word.replace("I", "ı").replace("İ", "i").lower()
        return ("İ" if lower[:1] == "i" else lower[:1].upper()) + lower[1:]
    return word[:1].upper() + word[1:].lower()


def salutation(name: str, country: str = "TR") -> str:
    """A short firm name for the greeting: the words before the legal form, at most three."""
    words = []
    for token in str(name or "").split():
        folded = legal_form._fold(token)
        if not folded:
            continue
        if all(part in SALUTATION_STOP for part in folded.split()):
            break
        letters = re.sub(r"[^\w]", "", token)
        if letters.isupper() and len(letters) > 3 and not any(ch.isdigit() for ch in letters):
            token = _capital(token, country)
        words.append(token.strip(",;"))
        if len(words) == 3:
            break
    return " ".join(words) or _text(name)


# --- tracking workbook -------------------------------------------------------

def read_customer_list(path: Path) -> list[dict[str, str]]:
    """Rows of a customer file (any header language) keyed by the Turkish column names."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book.worksheets[0].iter_rows(values_only=True))
    finally:
        book.close()
    if not rows:
        return []
    names = {header: key for headers in customer_file.HEADERS.values() for header, key in zip(headers, customer_file.TABLE_KEYS)}
    keys = [names.get(_text(cell), "") for cell in rows[0]]
    if "Firma" not in keys or "E-posta" not in keys:
        raise OutreachError("Listede Firma ve E-posta sütunları bulunamadı.")
    return [{key: _text(value) for key, value in zip(keys, row) if key} for row in rows[1:]]


def new_tracker(path: Path, rows: list[dict[str, str]], country: str) -> int:
    """A new tracking workbook with the firms that have an e-mail address; returns their number."""
    path = Path(path)
    if path.exists():
        raise OutreachError(f"Takip dosyası zaten var: {path}")
    book = Workbook()
    sheet = book.active
    sheet.title = "Takip"
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    count = 0
    for row in rows:
        if not row.get("E-posta"):
            continue
        values = {**row, "Hitap adı": salutation(row.get("Firma", ""), country)}
        sheet.append([values.get(column, "") for column in COLUMNS])
        count += 1
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = COLUMN_WIDTHS.get(column, 16)
    sheet.freeze_panes = "A2"
    save_tracker(book, path)
    return count


class Tracker:
    """A tracking workbook read by its header names; tool columns are added when missing."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.book = load_workbook(self.path)
        self.sheet = self.book.worksheets[0]
        self.columns = {_text(cell.value): cell.column for cell in self.sheet[1] if _text(cell.value)}
        missing = [column for column in ("Firma", "E-posta", "Gönderim tarihi") if column not in self.columns]
        if missing:
            raise OutreachError(f"Takip dosyasında sütun eksik: {', '.join(missing)}")
        for column in TOOL_COLUMNS:
            if column not in self.columns:
                index = self.sheet.max_column + 1
                self.sheet.cell(row=1, column=index, value=column).font = Font(bold=True)
                self.columns[column] = index

    def rows(self) -> list[dict]:
        result = []
        for number in range(2, self.sheet.max_row + 1):
            values = {column: self.sheet.cell(row=number, column=index).value for column, index in self.columns.items()}
            if _text(values.get("Firma")) or _text(values.get("E-posta")):
                result.append({"row": number, **values})
        return result

    def set(self, row: int, column: str, value) -> None:
        self.sheet.cell(row=row, column=self.columns[column], value=value)

    def save(self) -> None:
        save_tracker(self.book, self.path)


def save_tracker(book, path: Path) -> None:
    path = Path(path)
    temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def writable(path: Path) -> bool:
    """False while another program (Excel) holds the workbook open."""
    try:
        with Path(path).open("r+b"):
            return True
    except PermissionError:
        return False


def fill_salutations(tracker: Tracker, country: str) -> int:
    filled = 0
    for row in tracker.rows():
        if not _text(row.get("Hitap adı")) and _text(row.get("Firma")):
            tracker.set(row["row"], "Hitap adı", salutation(row["Firma"], country))
            filled += 1
    return filled


def sync_from_log(tracker: Tracker, log: list[dict]) -> int:
    """Write e-mails the log shows as sent but the workbook does not (it was open at the time)."""
    sent = {entry["email"]: entry for entry in log if entry.get("result") == SENT}
    synced = 0
    for row in tracker.rows():
        entry = sent.get(_text(row.get("E-posta")).casefold())
        if entry and not _text(row.get("Gönderim tarihi")):
            tracker.set(row["row"], "Gönderim tarihi", date.fromisoformat(entry["date"]))
            tracker.set(row["row"], "Gönderim sonucu", SENT)
            synced += 1
    return synced


# --- rules ---------------------------------------------------------------------

def _free_mail_domains() -> set[str]:
    return {*config.NON_CORPORATE_EMAIL_DOMAINS, *config.EXTRA_FREE_MAIL_DOMAINS}


def written_addresses(log: list[dict]) -> tuple[set[str], set[str]]:
    """Addresses and company mail domains already written to (test e-mails do not count)."""
    emails = {entry["email"] for entry in log if entry.get("result") != TEST}
    domains = {scorer.registrable_domain(email.rsplit("@", 1)[-1]) for email in emails} - _free_mail_domains()
    return emails, domains


def skip_reason(row: dict, *, opt_out_entries: tuple[str, ...], emails: set[str], domains: set[str]) -> str:
    """Why this firm gets no e-mail now; "" when it may get one."""
    email = _text(row.get("E-posta")).casefold()
    if not email:
        return "e-posta yok"
    if not EMAIL_RE.fullmatch(email):
        return "geçersiz e-posta"
    if email_kind.is_excluded(email):
        return "kişi ya da kişisel veri adresi"
    if _text(row.get("Gönderim tarihi")) or _text(row.get("Gönderim sonucu")):
        return "daha önce gönderildi"
    if _text(row.get("Yanıt")):
        return "yanıt var"
    if _text(row.get("Not")):
        return "Not sütunu dolu"
    if opt_out.matches(opt_out_entries, company=row.get("Firma"), website=row.get("Web sitesi"), email=email):
        return "ret listesinde"
    if email in emails:
        return "bu adrese daha önce yazıldı"
    if scorer.registrable_domain(email.rsplit("@", 1)[-1]) in domains:
        return "bu firmaya daha önce yazıldı"
    return ""


def office_hours(now: datetime, country: str) -> bool:
    local = now.astimezone(ZoneInfo(TIMEZONES.get(country, TIMEZONES["TR"])))
    start, end = config.OUTREACH_OFFICE_HOURS
    return local.weekday() < 5 and start <= local.hour < end


# --- template and message --------------------------------------------------------

def load_template(path: Path) -> dict:
    """Read a template: "Gönderen:" and "Konu:" lines (or From/Subject), a blank line, the body."""
    lines = Path(path).read_text(encoding="utf-8-sig").replace("\r\n", "\n").split("\n")
    headers: dict[str, str] = {}
    index = 0
    while index < len(lines) and lines[index].strip():
        key, _, value = lines[index].partition(":")
        headers[key.strip().casefold()] = value.strip()
        index += 1
    sender = headers.get("gönderen") or headers.get("from") or ""
    name, address = parseaddr(sender)
    template = {
        "sender_name": name, "sender_address": address.casefold(),
        "subject": headers.get("konu") or headers.get("subject") or "",
        "body": "\n".join(lines[index + 1:]).strip("\n") + "\n",
    }
    problems = []
    if not EMAIL_RE.fullmatch(template["sender_address"]):
        problems.append("Gönderen satırında geçerli bir e-posta adresi yok.")
    if not template["subject"]:
        problems.append("Konu satırı yok.")
    if "{hitap}" not in template["body"]:
        problems.append("Metinde {hitap} yok.")
    if PLACEHOLDER_RE.search(f"{sender}\n{template['subject']}\n{template['body']}"):
        problems.append("Şablonda doldurulmamış köşeli parantez var (örnek: [Ad Soyad]).")
    if any(field != "hitap" for field in FIELD_RE.findall(template["body"])):
        problems.append("Metinde {hitap} dışında süslü parantez var.")
    if problems:
        raise OutreachError(" ".join(problems))
    return template


def render(template: dict, row: dict) -> tuple[str, str]:
    greeting = _text(row.get("Hitap adı")) or salutation(_text(row.get("Firma")))
    return template["subject"], template["body"].replace("{hitap}", greeting)


def build_message(template: dict, to: str, subject: str, body: str) -> EmailMessage:
    message = EmailMessage()
    sender = template["sender_address"]
    message["From"] = formataddr((template["sender_name"], sender))
    message["To"] = to
    message["Subject"] = subject
    message["Date"] = formatdate(localtime=True)
    message["Message-ID"] = make_msgid(domain=sender.rsplit("@", 1)[-1])
    message["List-Unsubscribe"] = f"<mailto:{sender}?subject=unsubscribe>"
    message.set_content(body)
    return message


# --- log ---------------------------------------------------------------------------

def load_log() -> list[dict]:
    try:
        lines = config.OUTREACH_LOG_FILE.read_text(encoding="utf-8").splitlines()
    except FileNotFoundError:
        return []
    entries = []
    for line in lines:
        try:
            entry = json.loads(line)
        except ValueError:
            continue
        if isinstance(entry, dict) and entry.get("email"):
            entries.append(entry)
    return entries


def append_log(entry: dict) -> None:
    config.OUTREACH_LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    with config.OUTREACH_LOG_FILE.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(entry, ensure_ascii=False) + "\n")


def sent_on(log: list[dict], day: date) -> int:
    """E-mails the server accepted on that day, test e-mails included."""
    return sum(1 for entry in log if entry.get("date") == day.isoformat() and entry.get("result") in {SENT, TEST})


def daily_room(log: list[dict], limit: int, today: date) -> int:
    return max(0, min(limit, config.OUTREACH_MAX_DAILY_LIMIT) - sent_on(log, today))


# --- sending ----------------------------------------------------------------------

def smtp_password() -> str:
    from modules import api_configuration

    return config.OUTREACH_SMTP_PASSWORD or api_configuration.load_saved_api_keys().get("zoho_smtp", "")


def send_one(message: EmailMessage, *, address: str, password: str, smtp_factory=None) -> None:
    """One connection per e-mail; Zoho may close a connection left idle during the pause."""
    factory = smtp_factory or (lambda: smtplib.SMTP_SSL(config.OUTREACH_SMTP_HOST, config.OUTREACH_SMTP_PORT, timeout=30))
    with factory() as server:
        server.login(address, password)
        server.send_message(message)


def _save(tracker: Tracker) -> None:
    try:
        tracker.save()
    except PermissionError:
        raise OutreachError(
            "Takip dosyası başka bir programda açıldı; gönderim durdu. Excel'i kapatıp aynı komutu "
            "yeniden çalıştırın; gönderilen e-postalar kayıttan dosyaya işlenir."
        ) from None


def next_firms(tracker: Tracker, *, opt_out_entries: tuple[str, ...], log: list[dict]) -> tuple[list[dict], dict[str, int]]:
    """Firms that may get an e-mail, in workbook order, and how many are skipped for each reason."""
    emails, domains = written_addresses(log)
    ready, reasons = [], {}
    for row in tracker.rows():
        reason = skip_reason(row, opt_out_entries=opt_out_entries, emails=emails, domains=domains)
        if reason:
            reasons[reason] = reasons.get(reason, 0) + 1
            continue
        ready.append(row)
        email = _text(row["E-posta"]).casefold()
        emails.add(email)
        domain = scorer.registrable_domain(email.rsplit("@", 1)[-1])
        if domain not in _free_mail_domains():
            domains.add(domain)
    return ready, reasons


def preview(template: dict, rows: list[dict]) -> str:
    """The e-mails exactly as they would be sent."""
    parts = []
    for row in rows:
        subject, body = render(template, row)
        parts.append(
            f"Kime: {_text(row['E-posta'])}\nKonu: {subject}\n\n{body}\n" + "=" * 60 + "\n"
        )
    return "\n".join(parts)


def send_test(template: dict, row: dict, *, to: str, password: str, smtp_factory=None, now=None) -> None:
    """The first firm's e-mail sent to the user's own address; the workbook does not change."""
    subject, body = render(template, row)
    message = build_message(template, to, f"[DENEME] {subject}", body)
    stamp = (now or (lambda: datetime.now(timezone.utc)))()
    try:
        send_one(message, address=template["sender_address"], password=password, smtp_factory=smtp_factory)
    except smtplib.SMTPAuthenticationError:
        raise OutreachError("Zoho girişi reddedildi. Uygulama şifresini kontrol edin (tools\\api_anahtarlari.py).") from None
    except (smtplib.SMTPException, OSError) as error:
        raise OutreachError(f"Deneme e-postası gönderilemedi ({type(error).__name__}).") from None
    append_log({
        "time": stamp.isoformat(timespec="seconds"), "date": stamp.astimezone().date().isoformat(),
        "tracker": "", "company": _text(row.get("Firma")), "email": to.casefold(),
        "message_id": message["Message-ID"], "result": TEST,
    })


def send_batch(
    tracker: Tracker, template: dict, *, country: str, limit: int, password: str, opt_out_entries: tuple[str, ...],
    smtp_factory=None, now=None, sleep=time.sleep, report=print,
) -> dict:
    """Send to the next firms of the workbook until the daily room, office hours or a server refusal ends it."""
    now = now or (lambda: datetime.now(timezone.utc))
    log = load_log()
    sync_from_log(tracker, log)
    ready, _reasons = next_firms(tracker, opt_out_entries=opt_out_entries, log=log)
    result = {"sent": 0, "refused": 0, "stopped": "" if ready else "gönderilecek firma yok"}
    for row in ready:
        if daily_room(log, limit, now().astimezone().date()) <= 0:
            result["stopped"] = "günlük sınır doldu"
            break
        if result["sent"] or result["refused"]:
            sleep(random.uniform(*config.OUTREACH_DELAY_SEC))
        if not office_hours(now(), country):
            result["stopped"] = "mesai saati dışında"
            break
        email = _text(row["E-posta"]).casefold()
        subject, body = render(template, row)
        message = build_message(template, email, subject, body)
        stamp = now()
        entry = {
            "time": stamp.isoformat(timespec="seconds"), "date": stamp.astimezone().date().isoformat(),
            "tracker": tracker.path.name, "company": _text(row.get("Firma")), "email": email,
            "message_id": message["Message-ID"],
        }
        try:
            send_one(message, address=template["sender_address"], password=password, smtp_factory=smtp_factory)
        except smtplib.SMTPRecipientsRefused as error:
            code = next(iter(error.recipients.values()), (0, b""))[0]
            outcome = f"adres reddedildi ({code})"
            append_log({**entry, "result": outcome})
            log.append({**entry, "result": outcome})
            tracker.set(row["row"], "Gönderim sonucu", outcome)
            _save(tracker)
            result["refused"] += 1
            report(f"Adres reddedildi: {entry['company']} ({code})")
            continue
        except smtplib.SMTPAuthenticationError:
            raise OutreachError("Zoho girişi reddedildi. Uygulama şifresini kontrol edin (tools\\api_anahtarlari.py).") from None
        except smtplib.SMTPResponseException as error:
            raise OutreachError(
                f"Zoho gönderimi durdurdu (kod {error.smtp_code}). Bugün başka gönderim yapmayın; "
                "Zoho'ya girip hesabın durumuna bakın."
            ) from None
        except (smtplib.SMTPException, OSError) as error:
            raise OutreachError(f"Zoho'ya bağlanılamadı ({type(error).__name__}). Biraz sonra yeniden deneyin.") from None
        append_log({**entry, "result": SENT})
        log.append({**entry, "result": SENT})
        tracker.set(row["row"], "Gönderim tarihi", date.fromisoformat(entry["date"]))
        tracker.set(row["row"], "Gönderim sonucu", SENT)
        _save(tracker)
        result["sent"] += 1
        report(f"{result['sent']}. gönderildi: {entry['company']}")
    return result

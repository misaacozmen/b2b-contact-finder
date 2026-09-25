from __future__ import annotations

import hashlib
import inspect
import json
import os
import shutil
import uuid
from copy import deepcopy
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

import config
from modules import (
    contact_publication,
    contact_identity,
    discovery_coverage,
    entity_memory,
    entity_registry,
    evidence,
    excel,
    publication_policy,
    quality_audit,
    redaction,
    report,
    runtime,
    scorer,
)


class ArtifactResult(str):
    """String-compatible writer result carrying immutable artifact metadata."""

    def __new__(cls, report_text: str, artifacts: dict, entity_memory_rows: list[dict]):
        result = str.__new__(cls, report_text)
        result.artifacts = artifacts
        result.entity_memory_rows = entity_memory_rows
        return result


def is_quarantined_row(row: dict) -> bool:
    return bool(
        row.get("quarantine_state")
        or row.get("quarantine_status")
        or "legacy_recovery_provisional" in str(row.get("publication_blockers", ""))
        or str(row.get("source_record_id_quality", "")).casefold() == "legacy_recovery"
    )


def apply_quarantine(row: dict, *, state: str, status: str = "", blockers: str = "") -> dict:
    """Reapply durable item quarantine to every newly saved result."""
    row["quarantine_state"] = str(state)
    row["quarantine_status"] = str(status)
    row["publication_eligible"] = False
    marker = "HANDOFF_PENDING" if str(state) == "HANDOFF_PENDING" else "legacy_recovery_provisional"
    row["publication_blockers"] = "; ".join(sorted(set(filter(None, [str(blockers), str(row.get("publication_blockers", "")), marker]))))
    if marker != "HANDOFF_PENDING":
        suppress_all_contacts(row, marker)
    return row


def is_publishable_row(row: dict) -> bool:
    return publication_policy.is_publishable_row(row)


def apply_global_identity_collision_gate(rows: list[dict]) -> list[dict]:
    """Fail closed when distinct entities share a contact or website."""
    def has_relationship(row: dict) -> bool:
        if row.get("human_verified") is True and row.get("entity_relationship"):
            return True
        for relation in row.get("human_verified_relationships", []) or []:
            if isinstance(relation, dict) and relation.get("human_verified") is True:
                return True
        return False
    fields = ("website", "email", "phone")
    indexed: dict[tuple[str, str], list[dict]] = {}
    for row in rows:
        entity = str(row.get("entity_id") or row.get("source_record_id") or row.get("company") or "").strip()
        for field in fields:
            value = contact_identity.contact_key(field, row.get(field))
            if value:
                indexed.setdefault((field, value), []).append((entity, row))
    for (field, value), matches in indexed.items():
        entities = {entity for entity, _row in matches}
        if len(entities) < 2 or all(has_relationship(row) for _entity, row in matches):
            continue
        reason = "cross_entity_contact_collision"
        for _entity, row in matches:
            row["status"] = "REVIEW_NEEDED"
            row["publication_eligible"] = False
            row["collision_reason"] = reason
            row.setdefault("publication_blockers", "")
            row["publication_blockers"] = "; ".join(filter(None, [row.get("publication_blockers", ""), reason]))
            row[field] = ""
    # Pending rows and same-name rows from different sources are always review-only.
    by_name: dict[str, list[dict]] = {}
    for row in rows:
        key = str(row.get("company") or "").strip().casefold()
        by_name.setdefault(key, []).append(row)
    for matches in by_name.values():
        source_ids = {str(row.get("source_record_id") or "") for row in matches}
        if len(matches) > 1 and len(source_ids) == len(matches):
            for row in matches:
                row["status"] = "REVIEW_NEEDED"
                row["publication_eligible"] = False
                row["collision_reason"] = "cross_source_name_collision"
                row["publication_blockers"] = "; ".join(filter(None, [row.get("publication_blockers", ""), "cross_source_name_collision"]))
                suppress_all_contacts(row, "cross_source_name_collision")
    for row in rows:
        if str(row.get("status", "")).upper() in {"PENDING", "PENDING_PAID", "PENDING_FREE_RETRY", "PROCESSING_FAILED"}:
            row["publication_eligible"] = False
            row["publication_blockers"] = "; ".join(filter(None, [row.get("publication_blockers", ""), "pending_run_state"]))
            suppress_all_contacts(row, "pending_run_state")
    return rows


def _atomic_excel(path: Path, writer: Callable[[Path, object], None], rows: list[dict], *, frozen_timestamp: str = "2000-01-01T00:00:00+00:00") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}.{os.getpid()}.staging.xlsx")
    try:
        parameters = inspect.signature(writer).parameters
        if "frozen_timestamp" in parameters or any(p.kind == inspect.Parameter.VAR_KEYWORD for p in parameters.values()):
            writer(temporary, rows, frozen_timestamp=frozen_timestamp)
        else:
            writer(temporary, rows)
        if not temporary.exists():
            raise RuntimeError(f"artifact writer produced no file: {path.name}")
        workbook = load_workbook(temporary, read_only=True, data_only=True)
        header = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
        if not header or any(value is None for value in header):
            workbook.close()
            raise RuntimeError(f"artifact has invalid header: {path.name}")
        if workbook.active.max_row - 1 != len(rows):
            workbook.close()
            raise RuntimeError(f"artifact row count mismatch: {path.name}")
        workbook.close()
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def attach_candidates(row: dict, candidates: list[dict]) -> dict:
    row["selected_website"] = row.get("selected_website") or row.get("website", "")
    selected_domain = scorer.normalize_domain(row.get("selected_website", ""))
    status = str(row.get("status", ""))
    candidate_evaluations = []
    for candidate in candidates:
        history = candidate.setdefault("_stage_history", [])
        if not any(item.get("stage") == "discovered" for item in history):
            history.insert(0, {
                "stage": "discovered",
                "source": candidate.get("query", ""),
                "score": candidate.get("score", 0),
            })
        candidate_domain = scorer.normalize_domain(candidate.get("url", ""))
        if selected_domain and candidate_domain == selected_domain:
            final_stage = "published" if publication_policy.is_publishable_row(row) else "selected_for_review"
            if not any(item.get("stage") == final_stage for item in history):
                history.append({"stage": final_stage, "status": status})
        elif not any(item.get("stage") in {"rejected", "not_evaluated"} for item in history):
            evaluated = any(
                item.get("stage") in {"identity_evaluated", "full_evaluated"}
                for item in history
            )
            history.append({
                "stage": "rejected" if evaluated else "not_evaluated",
                "reason": "lower_identity_rank_or_failed_gate" if evaluated else "candidate_limit_or_lower_rank",
            })
        candidate_evaluations.append({
            "domain": candidate.get("domain") or candidate_domain,
            "url": candidate.get("url", ""),
            "source": candidate.get("query", ""),
            "stages": history,
        })
    row["__candidates"] = candidates
    row["__candidate_evaluations"] = candidate_evaluations
    row["__search_trace"] = getattr(candidates, "trace", [])
    row["__source_health"] = getattr(candidates, "source_health", {})
    for idx, candidate in enumerate(candidates[:3], start=1):
        row[f"candidate_{idx}_url"] = candidate.get("url", "")
        row[f"candidate_{idx}_score"] = candidate.get("score", "")
        row[f"candidate_{idx}_reason"] = candidate.get("reason", "")
        row[f"candidate_{idx}_query"] = candidate.get("query", "")
        row[f"candidate_{idx}_role"] = candidate.get("role", "")
    return row


def contact_output_fields(evaluation: dict) -> dict:
    alternative_phones = evaluation.get("alternative_phones", [])
    return {
        "email": evaluation.get("email", ""),
        "email_source": evaluation.get("email_source", ""),
        "email_source_url": evaluation.get("email_source_url", ""),
        "alternative_emails": "; ".join(evaluation.get("alternative_emails", [])),
        "alternative_email_sources": "; ".join(
            f"{item.get('value', '')} | {item.get('source_url', '')}"
            for item in evaluation.get("alternative_email_records", [])
        ),
        "email_verification": evaluation.get("email_verification", "not_checked"),
        "email_verification_reason": evaluation.get("email_verification_reason", ""),
        "email_publication_status": evaluation.get(
            "email_publication_status", "suppressed",
        ),
        "email_publication_reason": evaluation.get(
            "email_publication_reason", "",
        ),
        "phone": evaluation.get("phone", ""),
        "phone_source": evaluation.get("phone_source", ""),
        "phone_source_url": evaluation.get("phone_source_url", ""),
        "phone_label": evaluation.get("phone_label", ""),
        "alternative_phones": "; ".join(
            f"{item.get('value', '')} [{item.get('label', 'general')}]"
            for item in alternative_phones
        ),
        "alternative_phone_sources": "; ".join(
            f"{item.get('value', '')} | {item.get('source_url', '')}"
            for item in alternative_phones
        ),
        "phone_publication_status": evaluation.get(
            "phone_publication_status", "suppressed",
        ),
        "phone_publication_reason": evaluation.get(
            "phone_publication_reason", "",
        ),
        "contact_policy_version": evaluation.get(
            "contact_publication", {},
        ).get("policy_version", contact_publication.POLICY_VERSION),
    }


def evaluation_evidence(
    evaluation: dict,
    *,
    contact_output_fields_fn: Callable[[dict], dict] = contact_output_fields,
) -> dict:
    crawl_result = evaluation.get("crawl_result", {})
    return {
        "candidate": evaluation.get("candidate", {}),
        "final_score": evaluation.get("final_score", 0),
        "reasons": evaluation.get("reasons", []),
        "structured_identity": evaluation.get("structured_identity", {}),
        "semantic_identity": evaluation.get("semantic_identity", {}),
        "identity_assessment": evaluation.get("identity_assessment", {}),
        "publication_policy": evaluation.get("publication_policy", {}),
        "rerank_evidence": evaluation.get("rerank_evidence", {}),
        "linkedin_company_evidence": evaluation.get(
            "linkedin_company_evidence", {}
        ),
        "llm_arbiter_evidence": evaluation.get("llm_arbiter_evidence", {}),
        "llm_arbiter_decisions": evaluation.get("_llm_arbiter_decisions", []),
        "contact_publication": evaluation.get("contact_publication", {}),
        "identity_resolution": evaluation.get("_identity_resolution", ""),
        "automation": evaluation.get("_automation", {}),
        "crawl": {
            "url": crawl_result.get("url", ""),
            "cache_status": crawl_result.get("cache_status", ""),
            "error": crawl_result.get("error", ""),
            "pages": [page.get("url", "") for page in crawl_result.get("pages", [])],
            "recovery_trace": crawl_result.get("recovery_trace", []),
        },
        "contacts": contact_output_fields_fn(evaluation),
    }


def policy_output_fields(evaluation: dict) -> dict:
    decision = evaluation.get("publication_policy", {})
    blockers = decision.get("blockers", decision.get("hard_blockers", []))
    return {
        "publication_policy_version": decision.get(
            "policy_version", publication_policy.POLICY_VERSION,
        ),
        "publication_policy_action": decision.get("action", ""),
        "publication_eligible": bool(decision.get("publishable", decision.get("eligible", False))),
        "publication_advisory_eligible": bool(decision.get("advisory_eligible", decision.get("eligible", False))),
        "website_identity_verified": bool(decision.get("website_identity_verified", False)),
        "allowed_contact_fields": "; ".join(decision.get("allowed_contact_fields", [])),
        "publication_safety_score": int(decision.get("safety_score", 0) or 0),
        "publication_risk_index": int(decision.get("risk_index", 100) or 0),
        "publication_risk_tier": decision.get("risk_tier", "blocked"),
        "publication_blockers": "; ".join(blockers),
    }


def suppress_all_contacts(row: dict, reason: str) -> None:
    row["website"] = ""
    row["website_source"] = ""
    suppress_contact_field(row, "email", reason)
    suppress_contact_field(row, "phone", reason)


def suppress_contact_field(row: dict, field: str, reason: str) -> None:
    """Project one disallowed contact field out of every primary artifact."""
    field = str(field).casefold()
    if field == "email":
        for name in publication_policy.CONTACT_PROJECTION_FIELDS["email"]:
            row[name] = ""
        row["email_verification"] = "not_checked"
        row["email_verification_reason"] = reason
        row["email_publication_status"] = "suppressed"
        row["email_publication_reason"] = reason
    elif field == "phone":
        for name in publication_policy.CONTACT_PROJECTION_FIELDS["phone"]:
            row[name] = ""
        row["phone_publication_status"] = "suppressed"
        row["phone_publication_reason"] = reason


def project_publication_row(row: dict) -> dict:
    """Create a fail-closed export projection without rewriting source evidence."""
    projected = deepcopy(row)
    decision = publication_policy.decide_row(row)
    allowed = [str(value) for value in decision.get("allowed_contact_fields", [])]
    for field in ("email", "phone"):
        if field not in allowed:
            suppress_contact_field(projected, field, f"{field}_not_published")
    projected["publication_eligible"] = bool(decision.get("publishable"))
    projected["publication_advisory_eligible"] = bool(decision.get("advisory_eligible"))
    projected["website_identity_verified"] = bool(decision.get("website_identity_verified"))
    projected["allowed_contact_fields"] = "; ".join(allowed)
    content = row.get("content_decision") or row.get("__content_decision") or {}
    projected["publication_field_gaps"] = "; ".join(dict.fromkeys(
        value for value in (content.get("missing_evidence", []) if isinstance(content, dict) else [])
        if value in {"email_evidence_missing", "phone_evidence_missing", "email_not_allowed", "phone_not_allowed"}
    ))
    projected["contact_status"] = (
        "complete" if projected.get("email") and projected.get("phone")
        else "partial" if projected.get("email") or projected.get("phone")
        else "missing"
    )
    if not isinstance(projected.get("content_evidence_records"), list) and isinstance(content, dict):
        projected["content_evidence_records"] = deepcopy(content.get("evidence_records", []))
    projected["publication_projection_receipt"] = publication_policy.create_projection_receipt(
        row, projected, decision,
    )
    projected["publication_projection_status"] = "PROJECTED"
    return projected


def validate_publication_projection(row: dict) -> bool:
    return publication_policy.validate_projection(row)


def clear_unpublished_contacts(row: dict) -> None:
    """Compatibility alias for the fail-closed contact suppression."""
    suppress_all_contacts(row, "website_not_published")


def partition_output_rows(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Return publishable and review rows without changing the input list."""
    published = [row for row in rows if is_publishable_row(row)]
    review = [row for row in rows if not is_publishable_row(row)]
    return published, review


def apply_publication_policy(
    company: str,
    evaluation: dict,
    status: str,
    confidence: str,
    reasons: list[str],
) -> tuple[str, str]:
    decision = publication_policy.evaluate(
        company,
        evaluation,
        status,
        minimum_safety_score=config.PUBLICATION_POLICY_MIN_SAFETY_SCORE,
    )
    mode = str(config.PUBLICATION_POLICY_MODE or "enforce_downgrade_only").strip().casefold()
    hard_security_blockers = {"tls_certificate_unverified"}
    force_enforcement = bool(
        hard_security_blockers.intersection(decision.get("hard_blockers", []))
    )
    if mode == "shadow" and not force_enforcement:
        decision["mode"] = "shadow"
        if decision.get("action") == "downgrade_to_review":
            decision["action"] = "would_downgrade_to_review"
    else:
        status, confidence = publication_policy.enforce(
            decision, status, confidence, reasons,
        )
    evaluation["publication_policy"] = decision
    return status, confidence


def confidence_status(score: int, has_contact: bool, reasons: list[str], identity_verified: bool = True) -> tuple[str, str]:
    if not identity_verified:
        reasons.append("website_identity_not_independently_verified")
        return "REVIEW_NEEDED", "review"
    if score >= config.HIGH_CONFIDENCE_SCORE and has_contact:
        return "OK_HIGH_CONFIDENCE", "high"
    if score >= config.MEDIUM_CONFIDENCE_SCORE and has_contact:
        return "OK_MEDIUM_CONFIDENCE", "medium"
    if score >= config.REVIEW_SCORE:
        reasons.append("needs_manual_review")
        return "REVIEW_NEEDED", "review"
    # A successfully crawled, independently verified official website must not
    # disappear merely because a sector/context penalty pulled its numeric score
    # below the review threshold.  Keep it quarantined for a human decision.
    reasons.append("trusted_website_below_score_preserved_for_review")
    return "REVIEW_NEEDED", "review"


def write_outputs(rows: list[dict], elapsed_seconds: float, *, telemetry_snapshot: dict | None = None,
                  operational_metrics: dict | None = None) -> ArtifactResult:
    frozen_snapshot = dict(telemetry_snapshot) if telemetry_snapshot is not None else runtime.snapshot()
    frozen_timestamp = str(frozen_snapshot.get("generated_at", "2000-01-01T00:00:00+00:00"))
    apply_global_identity_collision_gate(rows)
    raw_rows = deepcopy(rows)
    output_root = Path(config.OUTPUT_DIR)
    staging_root = output_root / ".staging" / uuid.uuid4().hex
    staging_root.mkdir(parents=True, exist_ok=False)
    def staged(path: Path) -> Path:
        return staging_root / Path(path).name

    all_results_path = staged(output_root / "all_results.xlsx")
    for row in raw_rows:
        decision = publication_policy.decide_row(row)
        row["publication_eligible"] = decision["publishable"]
        row["publication_advisory_eligible"] = decision["advisory_eligible"]
        row["website_identity_verified"] = decision["website_identity_verified"]
        row["allowed_contact_fields"] = "; ".join(decision["allowed_contact_fields"])
        row["publication_field_gaps"] = "; ".join(dict.fromkeys(
            value for value in (row.get("content_decision") or {}).get("missing_evidence", [])
            if value in {"email_evidence_missing", "phone_evidence_missing", "email_not_allowed", "phone_not_allowed"}
        ))
        if not decision["publishable"]:
            row["publication_blockers"] = "; ".join(dict.fromkeys(
                value for value in [str(row.get("publication_blockers", "")), *decision["blockers"]] if value
            ))
        evaluation = row.get("__evaluation")
        if isinstance(evaluation, dict):
            policy = evaluation.setdefault("publication_policy", {})
            policy.update({
                "publishable": decision["publishable"],
                "blockers": decision["blockers"],
                "website_identity_verified": decision["website_identity_verified"],
                "allowed_contact_fields": decision["allowed_contact_fields"],
                "policy_version": decision["policy_version"],
                "advisory_eligible": decision["advisory_eligible"],
            })
        row["website_status"] = (
            "verified" if row.get("website") and is_publishable_row(row)
            else "review" if row.get("website") or row.get("status") == "WEBSITE_AMBIGUOUS"
            else "not_found"
        )
        row["contact_status"] = (
            "complete" if row.get("email") and row.get("phone")
            else "partial" if row.get("email") or row.get("phone")
            else "missing"
        )
    evidence.write_jsonl(staged(config.EVIDENCE_FILE), raw_rows)
    entity_registry.write_observations(staged(config.ENTITY_RELATIONSHIPS_FILE), raw_rows, observed_at=frozen_timestamp)
    quality_audit.write(staged(config.QUALITY_AUDIT_FILE), raw_rows, runtime_snapshot=frozen_snapshot)
    projected_rows = [project_publication_row(row) for row in raw_rows]
    published_rows, review_rows = partition_output_rows(projected_rows)
    for row in published_rows:
        discovery_coverage.mark_published(
            row.get("company", ""),
            row.get("source_record_id", ""),
            row.get("original_index"),
        )
    failed_output_rows = report.failed_rows(projected_rows)
    report_text = redaction.redact_text(report.build_report(
        projected_rows, elapsed_seconds, runtime_snapshot=frozen_snapshot,
        operational_metrics=operational_metrics,
    ))
    memory_rows = [deepcopy(row) for row in published_rows]
    for row in raw_rows:
        row.pop("__index", None)
        row.pop("__candidates", None)
        row.pop("__evaluation", None)
        row.pop("__candidate_evaluations", None)
        row.pop("__search_trace", None)
        row.pop("__source_health", None)
        row.pop("__paid_escalation_complete", None)
    for row in projected_rows:
        row.pop("__index", None)
        row.pop("__candidates", None)
        row.pop("__evaluation", None)
        row.pop("__candidate_evaluations", None)
        row.pop("__search_trace", None)
        row.pop("__source_health", None)
        row.pop("__paid_escalation_complete", None)
    # contacts.xlsx is the publication surface. Review/abstain rows remain in
    # the dedicated audit artifacts and must never look like published firms.
    sanitized_published_rows = redaction.sanitize(published_rows)
    sanitized_review_rows = redaction.sanitize(review_rows)
    sanitized_all_rows = redaction.sanitize(projected_rows)
    _atomic_excel(all_results_path, excel.write_contacts, sanitized_all_rows, frozen_timestamp=frozen_timestamp)
    _atomic_excel(staged(config.CONTACTS_FILE), excel.write_contacts, sanitized_published_rows, frozen_timestamp=frozen_timestamp)
    _atomic_excel(staged(config.VERIFIED_CONTACTS_FILE), excel.write_contacts, sanitized_published_rows, frozen_timestamp=frozen_timestamp)
    _atomic_excel(staged(config.REVIEW_QUEUE_FILE), excel.write_contacts, sanitized_review_rows, frozen_timestamp=frozen_timestamp)
    _atomic_excel(staged(config.FAILED_FILE), excel.write_failed, redaction.sanitize(failed_output_rows), frozen_timestamp=frozen_timestamp)
    _atomic_excel(staged(config.CANDIDATES_FILE), excel.write_website_candidates, sanitized_all_rows, frozen_timestamp=frozen_timestamp)
    report_path = staged(config.REPORT_FILE)
    report_tmp = report_path.with_name(f".{report_path.name}.{os.getpid()}.tmp")
    try:
        report_tmp.write_text(report_text, encoding="utf-8")
        if report_tmp.exists():
            report_tmp.replace(report_path)
    finally:
        report_tmp.unlink(missing_ok=True)
    discovery_coverage.write(
        staged(config.DISCOVERY_COVERAGE_FILE),
        config.DISCOVERY_ACQUISITION_QUERIES_PER_COMPANY,
    )
    runtime.write(staged(config.TELEMETRY_FILE), frozen_snapshot)
    mandatory = [
        staged(path) for path in (
            output_root / "all_results.xlsx", config.CONTACTS_FILE,
            config.VERIFIED_CONTACTS_FILE, config.REVIEW_QUEUE_FILE,
            config.FAILED_FILE, config.CANDIDATES_FILE, config.EVIDENCE_FILE,
            config.ENTITY_RELATIONSHIPS_FILE, config.QUALITY_AUDIT_FILE,
            config.DISCOVERY_COVERAGE_FILE, config.REPORT_FILE, config.TELEMETRY_FILE,
        )
    ]
    if any(not path.exists() for path in mandatory):
        raise RuntimeError("finalization missing mandatory artifact")
    artifact_set_hash = hashlib.sha256("".join(
        f"{path.name}:{hashlib.sha256(path.read_bytes()).hexdigest()}\n" for path in sorted(mandatory)
    ).encode("utf-8")).hexdigest()
    artifact_dir = output_root / "artifacts" / artifact_set_hash
    artifact_dir.parent.mkdir(parents=True, exist_ok=True)
    artifact_staging = artifact_dir.parent / f".{artifact_set_hash}.{uuid.uuid4().hex}.staging"
    artifact_staging.mkdir(parents=False, exist_ok=False)
    artifact_info = {}
    for path in mandatory:
        target = artifact_staging / path.name
        shutil.move(str(path), str(target))
        info = {"sha256": hashlib.sha256(target.read_bytes()).hexdigest(), "bytes": target.stat().st_size}
        if target.suffix == ".xlsx":
            workbook = load_workbook(target, read_only=True, data_only=True)
            info["rows"] = workbook.active.max_row - 1
            info["columns"] = workbook.active.max_column
            workbook.close()
        artifact_info[target.name] = info
    if artifact_dir.exists():
        existing_info = {
            path.name: hashlib.sha256(path.read_bytes()).hexdigest()
            for path in artifact_dir.iterdir() if path.is_file()
        }
        if existing_info != {name: info["sha256"] for name, info in artifact_info.items()}:
            raise RuntimeError("existing artifact set hash has different contents")
        shutil.rmtree(artifact_staging)
    else:
        artifact_staging.replace(artifact_dir)
    for source_path, target in zip(mandatory, (
        output_root / "all_results.xlsx", config.CONTACTS_FILE,
        config.VERIFIED_CONTACTS_FILE, config.REVIEW_QUEUE_FILE,
        config.FAILED_FILE, config.CANDIDATES_FILE, config.EVIDENCE_FILE,
        config.ENTITY_RELATIONSHIPS_FILE, config.QUALITY_AUDIT_FILE,
        config.DISCOVERY_COVERAGE_FILE, config.REPORT_FILE, config.TELEMETRY_FILE,
    )):
        shutil.copy2(artifact_dir / source_path.name, target)
    artifacts = {"artifact_set_sha256": artifact_set_hash, "files": artifact_info, "artifact_dir": str(artifact_dir)}
    shutil.rmtree(staging_root, ignore_errors=True)
    return ArtifactResult(report_text, artifacts, memory_rows)


def write_partial_recovery(
    rows: list[dict], *, output_root: Path, metadata: dict,
) -> dict[str, object]:
    """Write a clearly non-final recovery set for a stalled paid run."""
    target = Path(output_root) / "partial_recovery"
    target.mkdir(parents=True, exist_ok=True)
    safe_rows = redaction.sanitize(deepcopy(rows))
    payload = {
        "partial": True,
        "accepted_as_complete": False,
        "metadata": redaction.sanitize(metadata),
        "rows": safe_rows,
    }
    files = {
        "partial_results.json": json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True),
        "partial_report.txt": redaction.redact_text(
            "SCHEDULER_STALLED\n" + json.dumps(redaction.sanitize(metadata), ensure_ascii=False, sort_keys=True)
        ),
    }
    receipts = {}
    for name, content in files.items():
        path = target / name
        temporary = path.with_name(f".{path.name}.{os.getpid()}.{uuid.uuid4().hex}.tmp")
        temporary.write_text(content, encoding="utf-8")
        temporary.replace(path)
        receipts[name] = {"sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size}
    manifest = {"partial": True, "accepted_as_complete": False, "files": receipts, "metadata": redaction.sanitize(metadata)}
    manifest_path = target / "partial_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    receipts[manifest_path.name] = {"sha256": hashlib.sha256(manifest_path.read_bytes()).hexdigest(), "bytes": manifest_path.stat().st_size}
    return {"directory": str(target), "files": receipts, "metadata": redaction.sanitize(metadata)}


def publish_manifest(*, path: Path, run_id: str, input_hash: str, config_sha256: str,
                     counts: dict, artifacts: dict, complete: bool = True,
                     telemetry: dict | None = None,
                     operational_metrics: dict | None = None,
                     finalized: bool | None = None,
                     status: str | None = None) -> None:
    if not complete or not artifacts:
        raise RuntimeError("cannot publish incomplete artifact set")
    payload = {}
    if path.exists():
        try:
            payload.update(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, json.JSONDecodeError):
            pass
    payload.update({
        "complete": True, "artifact_schema_version": 3, "run_id": run_id,
        "input_sha256": input_hash, "config_sha256": config_sha256,
        "counts": counts, **artifacts,
    })
    if telemetry is not None:
        payload["telemetry"] = telemetry
        telemetry_json = json.dumps(telemetry, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        payload["telemetry_sha256"] = hashlib.sha256(telemetry_json.encode("utf-8")).hexdigest()
    if operational_metrics is not None:
        payload["operational_metrics"] = operational_metrics
        operational_json = json.dumps(operational_metrics, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        payload["operational_metrics_sha256"] = hashlib.sha256(operational_json.encode("utf-8")).hexdigest()
    if finalized is not None:
        payload["finalized"] = bool(finalized)
    if status is not None:
        payload["status"] = str(status)
    payload["phase"] = "COMPLETE"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.staging.json")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


_attach_candidates = attach_candidates
_contact_output_fields = contact_output_fields
_evaluation_evidence = evaluation_evidence
_policy_output_fields = policy_output_fields
_suppress_all_contacts = suppress_all_contacts
_clear_unpublished_contacts = clear_unpublished_contacts
_apply_publication_policy = apply_publication_policy
_confidence_status = confidence_status
_write_outputs = write_outputs

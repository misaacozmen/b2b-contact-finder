"""Repair Golden6 workbooks with deterministic source_record_id values."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import TableColumn


SOURCE = "automechanika_istanbul_2026"
UUID_RE = re.compile(
    r"/company/([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})/"
)
ID_HEADER = "source_record_id"
EXPECTED_ROWS = 20


class Golden6RepairError(RuntimeError):
    """A fail-closed Golden6 repair validation error."""


def _text(value: object) -> str:
    return str(value or "").strip()


def _headers(ws) -> list[str]:
    return [_text(cell.value).casefold() for cell in ws[1]]


def _column(ws, *names: str) -> int:
    wanted = {name.casefold() for name in names}
    for index, value in enumerate(_headers(ws), start=1):
        if value in wanted:
            return index
    raise Golden6RepairError(f"missing_header:{'/'.join(names)}")


def _companies(ws) -> list[str]:
    column = _column(ws, "company", "Company")
    values = [_text(ws.cell(row, column).value) for row in range(2, ws.max_row + 1)]
    if len(values) != EXPECTED_ROWS or any(not value for value in values):
        raise Golden6RepairError("company_population_must_be_20")
    return values


def _snapshot(wb) -> dict:
    snapshot: dict = {"sheetnames": list(wb.sheetnames), "sheets": {}}
    for ws in wb.worksheets:
        cells = {}
        for row in range(1, ws.max_row + 1):
            for column in range(1, ws.max_column + 1):
                cell = ws.cell(row, column)
                if cell.value is not None or cell.has_style:
                    cells[(row, column)] = (
                        cell.value,
                        cell.data_type,
                        copy.copy(cell._style),
                    )
        snapshot["sheets"][ws.title] = {
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "cells": cells,
            "merged": tuple(str(value) for value in ws.merged_cells.ranges),
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "show_grid_lines": ws.sheet_view.showGridLines,
            "row_dimensions": {
                key: (value.height, value.hidden, value.outlineLevel)
                for key, value in ws.row_dimensions.items()
            },
            "column_dimensions": {
                key: (value.width, value.hidden, value.outlineLevel)
                for key, value in ws.column_dimensions.items()
            },
        }
    return snapshot


def _assert_untouched(wb, before: dict, additions: dict[str, int]) -> None:
    if list(wb.sheetnames) != before["sheetnames"]:
        raise Golden6RepairError("sheet_layout_changed")
    for ws in wb.worksheets:
        original = before["sheets"][ws.title]
        expected_max_column = original["max_column"] + additions.get(ws.title, 0)
        if ws.max_row != original["max_row"] or ws.max_column != expected_max_column:
            raise Golden6RepairError(f"sheet_dimensions_changed:{ws.title}")
        for key, expected in original["cells"].items():
            row, column = key
            cell = ws.cell(row, column)
            actual = (cell.value, cell.data_type, cell._style)
            if actual != expected:
                raise Golden6RepairError(f"existing_cell_changed:{ws.title}:{cell.coordinate}")
        current_merged = tuple(str(value) for value in ws.merged_cells.ranges)
        if current_merged != original["merged"]:
            raise Golden6RepairError(f"merged_ranges_changed:{ws.title}")
        if (str(ws.freeze_panes) if ws.freeze_panes else None) != original["freeze_panes"]:
            raise Golden6RepairError(f"freeze_panes_changed:{ws.title}")
        if ws.sheet_view.showGridLines != original["show_grid_lines"]:
            raise Golden6RepairError(f"gridline_setting_changed:{ws.title}")
        for key, expected in original["row_dimensions"].items():
            value = ws.row_dimensions[key]
            if (value.height, value.hidden, value.outlineLevel) != expected:
                raise Golden6RepairError(f"row_layout_changed:{ws.title}:{key}")
        for key, expected in original["column_dimensions"].items():
            value = ws.column_dimensions[key]
            if (value.width, value.hidden, value.outlineLevel) != expected:
                raise Golden6RepairError(f"column_layout_changed:{ws.title}:{key}")


def _no_replacement(wb, label: str) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if "\ufffd" in _text(cell.value):
                    raise Golden6RepairError(f"replacement_character:{label}:{ws.title}:{cell.coordinate}")


def _ordered_ids(ws) -> list[str]:
    column = _column(ws, ID_HEADER)
    return [_text(ws.cell(row, column).value) for row in range(2, ws.max_row + 1)]


def _copy_column_style(ws, source_column: int, target_column: int) -> None:
    for row in range(1, ws.max_row + 1):
        source = ws.cell(row, source_column)
        target = ws.cell(row, target_column)
        target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
    source_letter = get_column_letter(source_column)
    target_letter = get_column_letter(target_column)
    source_dimension = ws.column_dimensions[source_letter]
    target_dimension = ws.column_dimensions[target_letter]
    target_dimension.width = source_dimension.width
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel


def _extend_tables(ws, old_max_column: int, new_column: int, header: str) -> None:
    for table in ws.tables.values():
        if not table.ref:
            continue
        if table.ref != f"A1:{get_column_letter(old_max_column)}{ws.max_row}":
            continue
        table.ref = f"A1:{get_column_letter(new_column)}{ws.max_row}"
        if len(table.tableColumns) < new_column:
            table.tableColumns.append(TableColumn(id=new_column, name=header))


def _fill_ids(ws, ids: list[str], label: str) -> int:
    if ws.max_row != EXPECTED_ROWS + 1:
        raise Golden6RepairError(f"row_count_invalid:{label}")
    existing = next(
        (index for index, value in enumerate(_headers(ws), start=1) if value == ID_HEADER),
        None,
    )
    if existing is not None:
        actual = _ordered_ids(ws)
        if actual != ids:
            raise Golden6RepairError(f"existing_ids_mismatch:{label}")
        return 0
    old_max_column = ws.max_column
    new_column = old_max_column + 1
    _copy_column_style(ws, old_max_column, new_column)
    ws.cell(1, new_column).value = ID_HEADER
    for row, source_id in enumerate(ids, start=2):
        ws.cell(row, new_column).value = source_id
    _extend_tables(ws, old_max_column, new_column, ID_HEADER)
    if _ordered_ids(ws) != ids:
        raise Golden6RepairError(f"id_column_write_failed:{label}")
    return 1


def _load_paths(root: Path) -> dict[str, Path]:
    directory = root / "outputs" / "golden_6_20260718"
    return {
        "source_assisted": directory / "golden_6_pipeline_input_20.xlsx",
        "blind": directory / "golden_6_discovery_blind_input_20.xlsx",
        "manual": directory / "golden_6_manual_validation_20_ready.xlsx",
    }


def repair(root: Path) -> dict:
    paths = _load_paths(Path(root).resolve())
    if any(not path.is_file() for path in paths.values()):
        missing = next(str(path) for path in paths.values() if not path.is_file())
        raise Golden6RepairError(f"workbook_missing:{missing}")

    workbooks = {name: load_workbook(path, data_only=False) for name, path in paths.items()}
    try:
        before = {name: _snapshot(wb) for name, wb in workbooks.items()}
        source_ws = workbooks["source_assisted"]["Pipeline Input"]
        source_companies = _companies(source_ws)
        source_column = _column(source_ws, "source")
        if any(_text(source_ws.cell(row, source_column).value) != SOURCE for row in range(2, 22)):
            raise Golden6RepairError("source_value_mismatch")
        profile_column = _column(source_ws, "profile_url")
        ids: list[str] = []
        for row in range(2, 22):
            url = _text(source_ws.cell(row, profile_column).value)
            match = UUID_RE.search(url)
            if not match:
                raise Golden6RepairError(f"profile_uuid_missing:row_{row}")
            ids.append(f"{SOURCE}:{match.group(1).lower()}")
        if len(set(ids)) != EXPECTED_ROWS:
            raise Golden6RepairError("profile_uuid_duplicate")

        blind_ws = workbooks["blind"]["Pipeline Input"]
        manual_pipeline_ws = workbooks["manual"]["Pipeline Input"]
        manual_report_ws = workbooks["manual"]["Manual Report"]
        target_lists = {
            "blind": _companies(blind_ws),
            "manual_pipeline": _companies(manual_pipeline_ws),
            "manual_report": _companies(manual_report_ws),
        }
        for label, companies in target_lists.items():
            if companies != source_companies:
                raise Golden6RepairError(f"company_order_mismatch:{label}")
        blind_profile_column = _column(blind_ws, "profile_url")
        if any(_text(blind_ws.cell(row, blind_profile_column).value) for row in range(2, 22)):
            raise Golden6RepairError("blind_profile_url_present")
        for name, wb in workbooks.items():
            _no_replacement(wb, f"before:{name}")

        additions = {
            "source_assisted": {"Pipeline Input": _fill_ids(source_ws, ids, "source_assisted_pipeline")},
            "blind": {"Pipeline Input": _fill_ids(blind_ws, ids, "blind_pipeline")},
            "manual": {
                "Pipeline Input": _fill_ids(manual_pipeline_ws, ids, "manual_pipeline"),
                "Manual Report": _fill_ids(manual_report_ws, ids, "manual_report"),
            },
        }
        for name, wb in workbooks.items():
            _assert_untouched(wb, before[name], additions.get(name, {}))
            _no_replacement(wb, f"after:{name}")

        staged: list[tuple[Path, Path]] = []
        try:
            for name, wb in workbooks.items():
                temporary = paths[name].with_name(f".{paths[name].name}.{os.getpid()}.tmp")
                if temporary.exists():
                    temporary.unlink()
                wb.save(temporary)
                staged.append((temporary, paths[name]))
            for temporary, destination in staged:
                os.replace(temporary, destination)
        finally:
            for temporary, _destination in staged:
                temporary.unlink(missing_ok=True)
    finally:
        for wb in workbooks.values():
            wb.close()

    for name, path in paths.items():
        check = load_workbook(path, read_only=True, data_only=False)
        try:
            _no_replacement(check, f"saved:{name}")
            for sheet in ("Pipeline Input", "Manual Report") if name == "manual" else ("Pipeline Input",):
                ws = check[sheet]
                if _ordered_ids(ws) != ids:
                    raise Golden6RepairError(f"saved_ids_mismatch:{name}:{sheet}")
        finally:
            check.close()
    return {
        "source_record_count": len(ids),
        "source_record_id_sha256": hashlib.sha256("\n".join(ids).encode("utf-8")).hexdigest(),
        "ordered_source_record_ids": ids,
        "files": {name: str(path) for name, path in paths.items()},
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Repair Golden6 source_record_id columns.")
    parser.add_argument("--repo-root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args(argv)
    try:
        print(json.dumps(repair(args.repo_root), ensure_ascii=False, indent=2))
    except Exception as exc:
        print(f"GOLDEN6_ID_REPAIR_BLOCKED:{exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

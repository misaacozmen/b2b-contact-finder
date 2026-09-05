"""Export a completed free-only production run as an auditable A8 actual."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sqlite3
from pathlib import Path
from typing import Any

import sys
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from modules import excel, publication_policy


PAID_PROVIDER_BUDGETS = {
    "brightdata": 0,
    "google_places": 0,
    "brandfetch": 0,
    "hunter": 0,
    "linkedin": 0,
    "llm": 0,
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical_hash(value: Any) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise RuntimeError(f"JSON object required: {path}")
    return payload


def _read_final_decisions(
    evidence_path: Path,
    expected_ids: list[str],
    *,
    run_id: str,
) -> dict[str, dict[str, Any]]:
    decisions: dict[str, dict[str, Any]] = {}
    with evidence_path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            record = json.loads(line)
            if not isinstance(record, dict):
                raise RuntimeError(f"final evidence line {line_number} is not an object")
            source_id = str(record.get("source_record_id") or "").strip()
            decision = record.get("publication_decision")
            if not source_id or not isinstance(decision, dict) or source_id in decisions:
                raise RuntimeError(f"final evidence decision coverage is invalid at line {line_number}")
            publication_policy.verify_publication_decision(
                decision,
                source_record_id=source_id,
                run_id=run_id,
            )
            decisions[source_id] = decision
    if list(decisions) != expected_ids:
        raise RuntimeError("final evidence decision order/set does not match expected workbook")
    return decisions


def _attach_decisions_to_workbook(
    workbook_path: Path,
    expected_ids: list[str],
    decisions: dict[str, dict[str, Any]],
) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook.active
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        if "source_record_id" not in headers:
            raise RuntimeError("actual workbook has no source_record_id column")
        source_index = headers.index("source_record_id") + 1
        decision_index = headers.index("publication_decision") + 1 if "publication_decision" in headers else sheet.max_column + 1
        if "publication_decision" not in headers:
            sheet.cell(row=1, column=decision_index, value="publication_decision")
        rows_by_id: dict[str, int] = {}
        for row_number in range(2, sheet.max_row + 1):
            source_id = str(sheet.cell(row=row_number, column=source_index).value or "").strip()
            if not source_id or source_id in rows_by_id:
                raise RuntimeError("actual workbook source ID coverage is invalid")
            rows_by_id[source_id] = row_number
        if list(rows_by_id) != expected_ids:
            raise RuntimeError("actual workbook source ID order/set changed while attaching decisions")
        for source_id in expected_ids:
            sheet.cell(
                row=rows_by_id[source_id],
                column=decision_index,
                value=json.dumps(decisions[source_id], ensure_ascii=False, sort_keys=True, separators=(",", ":")),
            )
        workbook.save(workbook_path)
    finally:
        workbook.close()


def _runtime_snapshot(connection: sqlite3.Connection, run_id: str) -> dict[str, Any]:
    row = connection.execute("SELECT runtime_json FROM runs WHERE run_id=?", (run_id,)).fetchone()
    if not row or not row[0]:
        raise RuntimeError("completed run has no durable runtime snapshot")
    snapshot = json.loads(str(row[0]))
    if not isinstance(snapshot, dict):
        raise RuntimeError("durable runtime snapshot is not an object")
    return snapshot


def export_actual(
    run_dir: Path,
    expected: Path,
    destination: Path,
    *,
    runtime_sha256: str,
) -> dict[str, Any]:
    run_dir = run_dir.resolve()
    expected = expected.resolve()
    destination = destination.resolve()
    manifest_path = run_dir / "manifest.json"
    state_path = run_dir / "state" / "progress.sqlite3"
    manifest = _read_json(manifest_path)
    if manifest.get("complete") is not True or manifest.get("phase") != "COMPLETE":
        raise RuntimeError("actual export requires a complete production run")
    if manifest.get("handoff") is True or manifest.get("checkpoint_materialized"):
        raise RuntimeError("handoff/checkpoint run cannot be exported as final actual")
    if not state_path.is_file():
        raise RuntimeError("completed run checkpoint is missing")
    if destination.exists():
        raise RuntimeError(f"destination already exists: {destination}")

    artifact_hash = str(manifest.get("artifact_set_sha256") or "")
    artifact_dir = run_dir / "output" / "artifacts" / artifact_hash
    source_workbook = artifact_dir / "all_results.xlsx"
    evidence_path = artifact_dir / "evidence.jsonl"
    if not artifact_hash or not source_workbook.is_file():
        raise RuntimeError("completed run all_results.xlsx artifact is missing")
    if not evidence_path.is_file():
        raise RuntimeError("completed run final evidence.jsonl artifact is missing")

    expected_ids = [str(row.get("source_record_id") or "").strip() for row in excel.read_company_records(expected)]
    actual_ids = [str(row.get("source_record_id") or "").strip() for row in excel.read_company_records(source_workbook)]
    if not expected_ids or any(not value for value in expected_ids) or len(expected_ids) != len(set(expected_ids)):
        raise RuntimeError("expected workbook source IDs are not exact")
    if actual_ids != expected_ids:
        raise RuntimeError("completed actual source ID order does not match expected workbook")
    if list(manifest.get("ordered_source_record_ids") or []) != expected_ids:
        raise RuntimeError("run manifest source ID order does not match expected workbook")
    if str(manifest.get("runtime_source_tree_sha256") or "") != runtime_sha256:
        raise RuntimeError("run runtime source-tree hash does not match current runtime")

    connection = sqlite3.connect(state_path)
    try:
        run_id = str(manifest.get("run_id") or "")
        if not run_id:
            raise RuntimeError("completed run has no run_id")
        runtime = _runtime_snapshot(connection, run_id)
        result_rows = connection.execute(
            "SELECT run_id,item_index,source_record_id,payload_sha256,payload "
            "FROM run_items JOIN results USING(run_id,item_index) "
            "WHERE run_id=? ORDER BY item_index",
            (run_id,),
        ).fetchall()
        item_count = int(manifest.get("item_count", -1))
        if len(result_rows) != item_count or len(result_rows) != len(expected_ids):
            raise RuntimeError("completed checkpoint result coverage is not exact")
        checkpoint_records: list[dict[str, Any]] = []
        for run_value, item_index, source_id, payload_sha256, payload_text in result_rows:
            if str(run_value) != run_id or str(source_id) != expected_ids[int(item_index)]:
                raise RuntimeError(f"checkpoint result identity mismatch at item_index={item_index}")
            payload_text = str(payload_text)
            if hashlib.sha256(payload_text.encode("utf-8")).hexdigest() != str(payload_sha256):
                raise RuntimeError(f"checkpoint payload hash mismatch at item_index={item_index}")
            payload = json.loads(payload_text)
            if not isinstance(payload, dict) or str(payload.get("source_record_id") or "") != str(source_id):
                raise RuntimeError(f"checkpoint payload source ID mismatch at item_index={item_index}")
            checkpoint_records.append({
                "run_id": run_id,
                "item_index": int(item_index),
                "source_record_id": str(source_id),
                "payload_sha256": str(payload_sha256),
                "payload": payload,
            })
        provider_calls = int(connection.execute("SELECT count(*) FROM provider_calls WHERE run_id=?", (run_id,)).fetchone()[0])
    finally:
        connection.close()

    counters = runtime.get("counters") if isinstance(runtime.get("counters"), dict) else {}
    run_config_sha256 = str(manifest.get("config_sha256") or "")
    decisions = _read_final_decisions(
        evidence_path,
        expected_ids,
        run_id=str(manifest["run_id"]),
    )
    decision_config_hashes = {str(decision.get("config_sha256") or "") for decision in decisions.values()}
    if len(decision_config_hashes) != 1 or not next(iter(decision_config_hashes), ""):
        raise RuntimeError("final evidence decisions do not share one config SHA-256")
    config_sha256 = next(iter(decision_config_hashes))
    actual_manifest = {
        "schema_version": 3,
        "status": "complete_free_only",
        "complete": True,
        "finalized": True,
        "phase": "COMPLETE",
        "run_id": str(manifest["run_id"]),
        "source_record_ids": expected_ids,
        "source_record_ids_sha256": _canonical_hash(expected_ids),
        "expected_sha256": _sha256(expected),
        "config_sha256": config_sha256,
        "run_config_sha256": run_config_sha256,
        "runtime_source_tree_sha256": runtime_sha256,
        "artifact_set_sha256": artifact_hash,
        "provider_calls": provider_calls,
        "physical_http_requests": int(counters.get("http.crawler.requests", 0) or 0) + int(counters.get("http.search.physical_http_requests", 0) or 0),
        "logical_request_count": int(counters.get("http.search.requests", 0) or 0),
        "browser_page_count": int(counters.get("recovery.browser_attempts", 0) or 0),
        "ocr_page_count": int(counters.get("recovery.pdf_ocr_attempts", 0) or 0),
        "elapsed_seconds": float(runtime.get("elapsed_seconds", 0) or 0),
        "cost": 0.0,
        "paid_enabled": False,
        "paid_provider_budgets": PAID_PROVIDER_BUDGETS,
        "capability_profile": {
            "browser_dependency_available": bool((manifest.get("run_config") or {}).get("effective_settings", {}).get("capability_browser_dependency_available", False)),
            "browser_enabled": bool((manifest.get("run_config") or {}).get("effective_settings", {}).get("capability_browser_enabled", False)),
            "ocr_dependency_available": bool((manifest.get("run_config") or {}).get("effective_settings", {}).get("capability_ocr_dependency_available", False)),
            "ocr_enabled": bool((manifest.get("run_config") or {}).get("effective_settings", {}).get("capability_ocr_enabled", False)),
            "search_provider": str((manifest.get("run_config") or {}).get("search_provider") or "ddgs"),
        },
    }

    destination.mkdir(parents=True)
    shutil.copy2(source_workbook, destination / "all_results.xlsx")
    _attach_decisions_to_workbook(destination / "all_results.xlsx", expected_ids, decisions)
    (destination / "checkpoint_results.jsonl").write_text(
        "".join(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n" for record in checkpoint_records),
        encoding="utf-8",
    )
    actual_manifest["files"] = {
        name: {"sha256": _sha256(destination / name), "bytes": (destination / name).stat().st_size}
        for name in ("all_results.xlsx", "checkpoint_results.jsonl")
    }
    (destination / "actual_manifest.json").write_text(
        json.dumps(actual_manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return actual_manifest


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--expected", type=Path, required=True)
    parser.add_argument("--destination", type=Path, required=True)
    parser.add_argument("--runtime-sha256", required=True)
    args = parser.parse_args()
    payload = export_actual(args.run_dir, args.expected, args.destination, runtime_sha256=args.runtime_sha256)
    print(json.dumps({"destination": str(args.destination), "source_record_count": len(payload["source_record_ids"])}))


if __name__ == "__main__":
    main()

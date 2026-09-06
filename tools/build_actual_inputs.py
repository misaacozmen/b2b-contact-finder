"""Build benchmark actual inputs from immutable source selections only."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


OUTPUT_HEADERS = [
    "Company", "website", "source_listed_website", "listed_website", "source_listed_website_status",
    "source", "source_record_id", "selection_group", "selection_reason", "selection_score", "selection_status",
    "profile_url", "listing_url", "listed_legal_name", "brand", "sector",
    "description", "listed_email", "listed_phone", "listed_address",
]
FORBIDDEN_TOKENS = (
    "expected", "publication", "review", "label_status", "score", "reason",
    "status", "quality", "provider", "scheduler", "run", "__",
)


def _xlsx_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = workbook.active.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        return [dict(zip(headers, row)) for row in iterator if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()


def _selection_records(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload.get("records"), list):
        return [dict(row) for row in payload["records"]]
    records: list[dict[str, Any]] = []
    for source_key in ("hometex", "ambiente"):
        section = payload.get(source_key) or {}
        records.extend(dict(row) for row in section.get("selected", []))
    if not records:
        raise RuntimeError("selection contains no source records")
    return records


def _source_rows(path: Path | None) -> dict[str, dict[str, Any]]:
    if path is None:
        return {}
    rows = _xlsx_rows(path)
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        source_id = str(row.get("source_record_id") or "").strip()
        if source_id:
            if source_id in result:
                raise RuntimeError(f"duplicate source_record_id in original source input: {source_id}")
            result[source_id] = row
    return result


def _value(record: dict[str, Any], original: dict[str, Any], *names: str) -> str:
    for name in names:
        value = record.get(name)
        if value not in (None, ""):
            return str(value).strip()
        value = original.get(name)
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _profile_url(record: dict[str, Any], source: str) -> str:
    if source == "zuchex_2026":
        return ""
    return _value(record, {}, "official_profile_url", "profile_url")


def _listing_url(record: dict[str, Any], source: str) -> str:
    value = _value(record, {}, "listing_url")
    if value:
        return value
    if source == "hometex_2026":
        return "https://hometex.com.tr/en/2026-exhibitor-list"
    if source == "ambiente_2026":
        return "https://ambiente.messefrankfurt.com/frankfurt/en/exhibitor-search.html"
    if source == "zuchex_2026":
        return "https://zuchex.com/"
    return ""


def _validate_input_headers(headers: list[str]) -> None:
    for header in headers:
        folded = header.casefold().replace("-", "_")
        if not folded.startswith(("selection_", "source_listed_")) and any(token in folded for token in FORBIDDEN_TOKENS):
            raise RuntimeError(f"forbidden non-source input column: {header}")
    if len({header.casefold() for header in headers}) != len(headers):
        raise RuntimeError("actual input header case collision")


def build(selection_path: Path, output_path: Path, original_input: Path | None = None) -> None:
    records = _selection_records(selection_path)
    original = _source_rows(original_input)
    seen: set[str] = set()
    rows: list[list[str]] = []
    for record in records:
        source_id = str(record.get("source_record_id") or "").strip()
        if not source_id or source_id in seen:
            raise RuntimeError(f"missing or duplicate selected source_record_id: {source_id}")
        seen.add(source_id)
        original_row = original.get(source_id, {})
        source = _value(record, original_row, "source")
        if not source:
            raise RuntimeError(f"selected source has no source name: {source_id}")
        profile = _profile_url(record, source)
        listing = _listing_url(record, source)
        # Benchmark input is discovery-only: the actual website column must not
        # contain the exhibitor profile/listing URL or a fair organiser URL.
        rows.append([
            _value(record, original_row, "display_name", "Company", "company"),
            "",
            _value(record, original_row, "source_listed_website"),
            _value(record, original_row, "listed_website", "official_website"),
            _value(record, original_row, "source_listed_website_status"),
            source,
            source_id,
            _value(record, original_row, "selection_group"),
            _value(record, original_row, "selection_reason"),
            _value(record, original_row, "selection_score"),
            _value(record, original_row, "selection_status"),
            profile,
            listing,
            _value(record, original_row, "legal_name", "listed_legal_name"),
            _value(record, original_row, "brand", "brands"),
            _value(record, original_row, "sector"),
            _value(record, original_row, "description"),
            _value(record, original_row, "listed_email", "email"),
            _value(record, original_row, "listed_phone", "phone"),
            _value(record, original_row, "listed_address", "address"),
        ])
    _validate_input_headers(OUTPUT_HEADERS)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Companies"
    sheet.append(OUTPUT_HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(output_path)
    workbook.close()
    print(json.dumps({"rows": len(rows), "output": str(output_path), "website_nonempty": 0}, sort_keys=True))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--selection", type=Path, required=True)
    parser.add_argument("--original-input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.selection, args.output, args.original_input)


if __name__ == "__main__":
    main()

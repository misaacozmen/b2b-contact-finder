"""Build a portable, immutable Golden6 replay package from one live run."""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from modules import excel, replay_snapshot, run_context
from tools.free_only_contract import validate_database as validate_free_only_database
from tools.free_only_contract import validate_manifest as validate_free_only_manifest


EXPECTED_COUNT = 20


class ReplayPackageError(RuntimeError):
    """A fail-closed replay package validation error."""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _ids_hash(ids: list[str]) -> str:
    return hashlib.sha256("\n".join(ids).encode("utf-8")).hexdigest()


def _json(path: Path) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ReplayPackageError(f"invalid_json:{path.name}") from exc
    if not isinstance(value, dict):
        raise ReplayPackageError(f"json_not_object:{path.name}")
    return value


def _git(repo: Path, *args: str) -> str:
    result = subprocess.run(
        ["git", *args], cwd=repo, check=True, capture_output=True, text=True,
    )
    return result.stdout.strip()


def _workbook_ids(path: Path, sheet_name: str | None = None) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip().casefold() for value in rows[0]]
        try:
            index = headers.index("source_record_id")
        except ValueError as exc:
            raise ReplayPackageError(f"source_record_id_missing:{path.name}") from exc
        return [str(row[index] or "").strip() for row in rows[1:] if any(value is not None for value in row)]
    finally:
        workbook.close()


def _source_ids(input_path: Path) -> list[str]:
    records = excel.read_company_records(input_path)
    ids = [str(record.get("source_record_id") or "").strip() for record in records]
    if len(ids) != EXPECTED_COUNT or any(not value for value in ids) or len(set(ids)) != EXPECTED_COUNT:
        raise ReplayPackageError("input_source_record_id_population_invalid")
    return ids


def _find_run(repo: Path, input_hash: str, run_id: str | None) -> tuple[Path, dict]:
    candidates: list[tuple[Path, dict]] = []
    if run_id:
        paths = [repo / "runs" / run_id / "manifest.json"]
    else:
        paths = sorted((repo / "runs").glob("*/manifest.json")) if (repo / "runs").is_dir() else []
    for manifest_path in paths:
        if not manifest_path.is_file():
            continue
        manifest = _json(manifest_path)
        if manifest.get("complete") is True and manifest.get("input_sha256") == input_hash:
            candidates.append((manifest_path.parent, manifest))
    if len(candidates) != 1:
        raise ReplayPackageError(f"live_run_ambiguous_or_missing:{len(candidates)}")
    return candidates[0]


def _source_integrity(db: Path, shard_root: Path) -> dict:
    paths = [db, *sorted(shard_root.glob("*.json.gz"))] if shard_root.is_dir() else [db]
    if not db.is_file():
        raise ReplayPackageError("live_run_db_missing")
    return {
        "database": {"path": str(db), "sha256": _sha256(db), "bytes": db.stat().st_size},
        "shards": {
            str(path.relative_to(shard_root).as_posix()): {
                "sha256": _sha256(path), "bytes": path.stat().st_size,
            }
            for path in paths if path != db
        },
    }


def _check_run_db(db: Path, run_id: str, ids: list[str]) -> None:
    try:
        with sqlite3.connect(db) as connection:
            phase = connection.execute("SELECT phase FROM runs WHERE run_id=?", (run_id,)).fetchone()
            items = [row[0] for row in connection.execute("SELECT source_record_id FROM run_items WHERE run_id=? ORDER BY item_index", (run_id,))]
            result_rows = [row[0] for row in connection.execute("SELECT payload FROM results WHERE run_id=? ORDER BY item_index", (run_id,))]
    except sqlite3.Error as exc:
        raise ReplayPackageError(f"live_run_db_invalid:{exc.__class__.__name__}") from exc
    if not phase or phase[0] != "COMPLETE":
        raise ReplayPackageError("live_run_not_complete")
    if items != ids or len(result_rows) != EXPECTED_COUNT:
        raise ReplayPackageError("live_run_population_mismatch")
    payload_ids = []
    for payload in result_rows:
        try:
            payload_ids.append(str(json.loads(payload).get("source_record_id") or "").strip())
        except (TypeError, json.JSONDecodeError) as exc:
            raise ReplayPackageError("live_result_payload_invalid") from exc
    if payload_ids != ids:
        raise ReplayPackageError("live_result_source_record_id_order_mismatch")


def _artifact_dir(run_root: Path, manifest: dict, ids: list[str]) -> Path:
    artifact_hash = str(manifest.get("artifact_set_sha256") or "")
    directory = run_root / "output" / "artifacts" / artifact_hash
    required = ("all_results.xlsx", "contacts.xlsx", "website_candidates.xlsx")
    if not artifact_hash or not directory.is_dir() or any(not (directory / name).is_file() for name in required):
        raise ReplayPackageError("live_artifact_set_incomplete")
    if _workbook_ids(directory / "all_results.xlsx") != ids:
        raise ReplayPackageError("live_all_results_population_mismatch")
    return directory


def _copy_body_shards(source: Path, target: Path) -> list[dict]:
    target.mkdir(parents=True, exist_ok=False)
    rows = []
    for path in sorted(source.glob("*.json.gz")):
        destination = target / path.name
        shutil.copy2(path, destination)
        rows.append({
            "path": f"replay_shards/{path.name}",
            "sha256": _sha256(destination),
            "bytes": destination.stat().st_size,
        })
    return rows


def _snapshot_markers(path: Path) -> set[str]:
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        payload = json.load(handle)

    def walk(value):
        found = set()
        if isinstance(value, dict):
            if "__replay_shard__" in value:
                found.add(str(value["__replay_shard__"]))
            for child in value.values():
                found.update(walk(child))
        elif isinstance(value, list):
            for child in value:
                found.update(walk(child))
        return found

    markers = walk(payload.get("entries", []))
    if payload.get("entry_count") != len(payload.get("entries", [])):
        raise ReplayPackageError("snapshot_entry_count_invalid")
    if hashlib.sha256(json.dumps(payload.get("entries", []), ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest() != payload.get("entries_sha256"):
        raise ReplayPackageError("snapshot_integrity_invalid")
    return markers


def _snapshot_entry_count(path: Path) -> int:
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        return int(json.load(handle).get("entry_count", 0))


def build(repo_root: Path, input_path: Path, package_dir: Path, run_id: str | None = None) -> dict:
    repo = Path(repo_root).resolve()
    input_path = Path(input_path).resolve()
    package_dir = Path(package_dir).resolve()
    if package_dir.exists():
        raise ReplayPackageError(f"package_exists:{package_dir}")
    ids = _source_ids(input_path)
    input_hash = _sha256(input_path)
    live_root, live_manifest = _find_run(repo, input_hash, run_id)
    live_id = str(live_manifest.get("run_id") or live_root.name)
    if live_root.name != live_id:
        raise ReplayPackageError("live_run_identity_mismatch")
    run_config = live_manifest.get("run_config") if isinstance(live_manifest.get("run_config"), dict) else {}
    config_hash = str(live_manifest.get("config_sha256") or "")
    if not config_hash or run_config.get("search_cache_mode") != "refresh" or run_config.get("crawl_cache_mode") != "refresh":
        raise ReplayPackageError("live_run_modes_not_refresh")
    try:
        validate_free_only_manifest(live_manifest)
    except ValueError as exc:
        raise ReplayPackageError(str(exc)) from exc
    try:
        resolved_config = run_context.RunConfig.from_dict(run_config)
    except Exception as exc:
        raise ReplayPackageError(f"live_run_config_invalid:{type(exc).__name__}") from exc
    if resolved_config.sha256 != config_hash:
        raise ReplayPackageError("live_run_config_hash_mismatch")
    db = live_root / "state" / "progress.sqlite3"
    source_shards = live_root / "state" / "replay_shards"
    before_integrity = _source_integrity(db, source_shards)
    _check_run_db(db, live_id, ids)
    try:
        paid_activity = validate_free_only_database(db, live_id, EXPECTED_COUNT)
    except ValueError as exc:
        raise ReplayPackageError(str(exc)) from exc
    artifacts = _artifact_dir(live_root, live_manifest, ids)
    head_sha = _git(repo, "rev-parse", "HEAD")

    staging = package_dir.parent / f".{package_dir.name}.{uuid.uuid4().hex}.staging"
    staging.mkdir(parents=True, exist_ok=False)
    try:
        replay_snapshot.reset()
        replay_snapshot.configure_run_store(db, live_id, read_only=True)
        export_root = staging / "_replay_export"
        replay_snapshot.export_shards(export_root, run_id=live_id, config_hash=config_hash)
        replay_snapshot.load_shards(export_root / "manifest.json", expected_run_id=live_id, expected_config_hash=config_hash)
        snapshot_path = staging / "replay_snapshot.json.gz"
        replay_snapshot.write(snapshot_path)
        body_shards = _copy_body_shards(export_root / "replay_shards", staging / "replay_shards")
        markers = _snapshot_markers(snapshot_path)
        available = {Path(item["path"]).stem.removesuffix(".json") for item in body_shards}
        if not markers.issubset(available):
            raise ReplayPackageError("replay_body_shard_closure_incomplete")
        shutil.rmtree(export_root)
        replay_snapshot.reset()
        replay_snapshot.load(snapshot_path, max_uncompressed_bytes=64 * 1024 * 1024)
        after_integrity = _source_integrity(db, source_shards)
        if before_integrity != after_integrity:
            raise ReplayPackageError("source_run_mutated_during_export")
        files = {}
        for path in sorted(staging.rglob("*")):
            if path.is_file() and path.name != "package_manifest.json":
                files[path.relative_to(staging).as_posix()] = {"sha256": _sha256(path), "bytes": path.stat().st_size}
        package_manifest = {
            "schema_version": 1,
            "paid_enabled": False,
            "finalize_without_paid": True,
            "head_sha": head_sha,
            "input": {
                "path": str(input_path), "sha256": input_hash,
                "ordered_id_sha256": _ids_hash(ids), "record_count": len(ids),
            },
            "ordered_source_record_ids": ids,
            "run_config": run_config,
            "live_run": {
                "run_id": live_id, "run_dir": str(live_root),
                "manifest_sha256": _sha256(live_root / "manifest.json"),
                "config_sha256": config_hash,
                "artifact_dir": str(artifacts),
                "artifact_set_sha256": str(live_manifest.get("artifact_set_sha256") or ""),
                "search_cache_mode": run_config.get("search_cache_mode"),
                "crawl_cache_mode": run_config.get("crawl_cache_mode"),
                "finalize_without_paid": True,
                "paid_enabled": False,
                "paid_activity": paid_activity,
            },
            "free_only": {
                "finalize_without_paid": True,
                "paid_enabled": False,
                "paid_budget_zero": True,
                "provider_calls_zero": True,
            },
            "replay": {
                "snapshot": "replay_snapshot.json.gz",
                "snapshot_sha256": _sha256(snapshot_path),
                "body_shards": body_shards,
            "entry_count": _snapshot_entry_count(snapshot_path),
            },
            "source_integrity_before": before_integrity,
            "source_integrity_after": after_integrity,
            "files": files,
            "package_manifest_hash_scope": "all package files except package_manifest.json",
        }
        temporary = staging / ".package_manifest.json.tmp"
        temporary.write_text(json.dumps(package_manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(staging / "package_manifest.json")
        staging.replace(package_dir)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    return package_manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the immutable Golden6 replay package.")
    parser.add_argument("--repo-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--package", type=Path, required=True)
    parser.add_argument("--run-id", default=None)
    args = parser.parse_args(argv)
    try:
        print(json.dumps(build(args.repo_root, args.input, args.package, args.run_id), ensure_ascii=False, indent=2))
    except Exception as exc:
        print(f"GOLDEN6_REPLAY_PACKAGE_BLOCKED:{exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

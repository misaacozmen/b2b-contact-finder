"""Prepare deterministic A8 review selections without creating source labels."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(iterator)]
        return [dict(zip(headers, row)) for row in iterator if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()


def _truth(value: object) -> bool:
    return value is True or str(value or "").strip().casefold() in {"true", "1", "yes", "evet"}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _ordered_source_ids(path: Path) -> list[str]:
    ids = []
    for row in _rows(path):
        source_id = str(row.get("source_record_id", "") or "").strip()
        if not source_id or source_id in ids:
            raise ValueError(f"invalid source ID sequence in {path}: {source_id}")
        ids.append(source_id)
    return ids


def _priority(row: dict[str, Any]) -> str:
    reason = str(row.get("reason", "") or "").casefold()
    score = int(row.get("score") or 0)
    if score >= 65 and not _truth(row.get("publication_eligible")):
        return "high_score_gate_reject"
    if any(token in reason for token in ("identity", "collision", "homonym", "wrong", "owner", "website")):
        return "identity_ownership"
    return "contact_access"


def _ordered(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(
        rows,
        key=lambda row: hashlib.sha256(
            f"{str(row['source_record_id']).strip()}architect-v1".encode("utf-8")
        ).hexdigest(),
    )


def prepare(source_root: Path) -> dict[str, Any]:
    input_rows = _rows(source_root / "input" / "firms.xlsx")
    review_rows = _rows(source_root / "output" / "exhibition_893_reconciled_v3" / "review_queue.xlsx")
    if len(review_rows) != 477:
        raise ValueError(f"expected 477 review rows, got {len(review_rows)}")
    parent_root = source_root / "output" / "exhibition_893_reconciled_v3"
    original_input = source_root / "input" / "firms.xlsx"
    parent_all_results = parent_root / "all_results.xlsx"
    parent_evidence = parent_root / "evidence.jsonl"
    parent_delivery = parent_root / "delivery_manifest.json"
    original_ids = _ordered_source_ids(parent_all_results)
    if len(original_ids) != 893:
        raise ValueError(f"expected 893 parent IDs, got {len(original_ids)}")
    review_ids = [str(row.get("source_record_id", "") or "").strip() for row in review_rows]
    if len(set(review_ids)) != 477 or not all(review_ids):
        raise ValueError("review queue source IDs are not exactly 477 unique nonempty IDs")
    by_index = {index: row for index, row in enumerate(input_rows)}
    all_ids: list[str] = []
    selected: list[dict[str, Any]] = []
    counts: dict[str, dict[str, int]] = {}
    for review in review_rows:
        source_id = str(review.get("source_record_id", "") or "").strip()
        if not source_id or source_id in all_ids:
            raise ValueError(f"review source_record_id missing or duplicate: {source_id}")
        all_ids.append(source_id)
        index = int(review.get("legacy_index") or 0)
        source_row = by_index.get(index, {})
        source = str(source_row.get("source") or ("texhibition_2026" if index < 477 else "zuchex_2026"))
        group = _priority(review)
        counts.setdefault(source, {})[group] = counts.setdefault(source, {}).get(group, 0) + 1
        review_copy = {
            "source_record_id": source_id,
            "Company": str(review.get("company") or source_row.get("company") or "").strip(),
            "source": source,
            "official_profile_url": str(source_row.get("profile_url") or "").strip(),
            "listed_legal_name": str(source_row.get("listed_legal_name") or "").strip(),
            "listed_address": str(source_row.get("listed_address") or "").strip(),
            "listed_phone": str(source_row.get("listed_phone") or "").strip(),
            "selection_group": group,
            "selection_score": int(review.get("score") or 0),
            "selection_status": str(review.get("status") or ""),
            "selection_reason": str(review.get("reason") or ""),
        }
        review_copy["listing_url"] = str(source_row.get("listing_url") or "").strip()
        selected.append(review_copy)

    diagnostic: list[dict[str, Any]] = []
    for source in ("texhibition_2026", "zuchex_2026"):
        source_rows = [row for row in selected if row["source"] == source]
        for group in ("high_score_gate_reject", "identity_ownership", "contact_access"):
            candidates = _ordered([row for row in source_rows if row["selection_group"] == group])
            if len(candidates) < 16:
                raise ValueError(f"not enough {group} rows for {source}: {len(candidates)}")
            diagnostic.extend(candidates[:16])
    if len({row["source_record_id"] for row in diagnostic}) != 96:
        raise ValueError("diagnostic selection is not unique")

    return {
        "schema_version": 2,
        "selection_algorithm": "sha256(source_record_id + architect-v1)",
        "source_root": str(source_root.resolve()),
        "input_path": str((source_root / "input" / "firms.xlsx").resolve()),
        "input_sha256": hashlib.sha256((source_root / "input" / "firms.xlsx").read_bytes()).hexdigest(),
        "review_path": str((source_root / "output" / "exhibition_893_reconciled_v3" / "review_queue.xlsx").resolve()),
        "review_count": len(review_rows),
        "review_source_record_ids": all_ids,
        "selection_counts": counts,
        "diagnostic_rows": diagnostic,
        "parent_artifacts": {
            "original_input": {
                "path": str(original_input.resolve()),
                "sha256": _sha256(original_input),
                "size": original_input.stat().st_size,
            },
            "review_queue": {
                "path": str((source_root / "output" / "exhibition_893_reconciled_v3" / "review_queue.xlsx").resolve()),
                "sha256": _sha256(source_root / "output" / "exhibition_893_reconciled_v3" / "review_queue.xlsx"),
                "size": (source_root / "output" / "exhibition_893_reconciled_v3" / "review_queue.xlsx").stat().st_size,
                "ordered_source_id_sha256": hashlib.sha256("\n".join(review_ids).encode("utf-8")).hexdigest(),
                "ordered_source_id_count": len(review_ids),
            },
            "parent_all_results": {
                "path": str(parent_all_results.resolve()),
                "sha256": _sha256(parent_all_results),
                "size": parent_all_results.stat().st_size,
                "ordered_source_id_sha256": hashlib.sha256("\n".join(original_ids).encode("utf-8")).hexdigest(),
                "ordered_source_id_count": len(original_ids),
            },
            "parent_evidence": {
                "path": str(parent_evidence.resolve()),
                "sha256": _sha256(parent_evidence),
                "size": parent_evidence.stat().st_size,
            },
            "parent_delivery_manifest": {
                "path": str(parent_delivery.resolve()),
                "sha256": _sha256(parent_delivery),
                "size": parent_delivery.stat().st_size,
            },
        },
        "original_source_id_order": original_ids,
        "independent_research": {
            "hometex_listing_url": "https://hometex.com.tr/en/2026-exhibitor-list",
            "ambiente_search_url": "https://ambiente.messefrankfurt.com/frankfurt/en/exhibitor-search.html",
            "hometex_status": "official_listing_observed",
            "ambiente_status": "official_public_api_observed",
            "hometex_unique_detail_slug_count": 580,
            "ambiente_hits_total": 96,
            "ambiente_unique_rewrite_id_count": 96,
            "independent_pool_status": "acquisition_required_before_selection",
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    payload = prepare(args.source_root.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"review_count": payload["review_count"], "diagnostic_count": len(payload["diagnostic_rows"])}, sort_keys=True))


if __name__ == "__main__":
    main()

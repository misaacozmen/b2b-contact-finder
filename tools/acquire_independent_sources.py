"""Acquire the two independent A8 source pools from their official surfaces only."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from modules.scorer import normalize_text, registrable_domain


HOMETEX_LISTING_URL = "https://hometex.com.tr/en/2026-exhibitor-list"
AMBIENTE_SEARCH_PAGE = "https://ambiente.messefrankfurt.com/frankfurt/en/exhibitor-search.html"
AMBIENTE_FILTERS_URL = "https://api.messefrankfurt.com/service/esb_api/exhibitor-service/api/2.1/public/exhibitor/searchfilters/en-GB/AMBIENTE"
AMBIENTE_SEARCH_URL = "https://api.messefrankfurt.com/service/esb_api/exhibitor-service/api/2.1/public/exhibitor/search"
ALLOWED_HOSTS = {
    "hometex.com.tr",
    "www.hometex.com.tr",
    "ambiente.messefrankfurt.com",
    "exhibitorsearch.messefrankfurt.com",
    "api.messefrankfurt.com",
}
CATEGORY_NAMES = {
    "mf_ppx_import_9789309": "Table & Table Select",
    "mf_ppx_import_9789424": "Cook & Cut",
    "mf_ppx_import_9789553": "Modern Kitchen & Baking",
    "mf_ppx_import_9789677": "Clean Home & Storage Solutions",
    "mf_ppx_import_9789939": "Global Sourcing Dining",
    "mf_ppx_import_9790106": "Interior Design",
    "mf_ppx_import_9790189": "Interiors & Decoration",
    "mf_ppx_import_9790250": "Global Sourcing Living",
}
EXCLUDED_HOMETEX_PATHS = {"/en/exhibitor", "/en/exhibitor-profile", "/en/2026-exhibitor-list", "/en/"}
NON_COMPANY_HOSTS = {
    "facebook.com", "instagram.com", "linkedin.com", "twitter.com", "x.com",
    "youtube.com", "pinterest.com",
}
NON_COMPANY_HOSTS |= {
    "apps.apple.com", "play.google.com", "informa.com", "xing.com", "tobb.org.tr",
    "old.texhibitionist.com", "cdn.texhibitionist.com",
}
ASSET_SUFFIXES = (".pdf", ".jpg", ".jpeg", ".png", ".gif", ".webp", ".doc", ".docx")


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _sha256_file(path: Path) -> str:
    return _sha256_bytes(path.read_bytes())


def _strict_text(response: requests.Response, *, reject_replacement: bool = True) -> str:
    body = response.content
    text = body.decode("utf-8", "strict")
    if reject_replacement and "\ufffd" in text:
        raise RuntimeError(f"replacement character in official response: {response.url}")
    return text


def _clean_string(value: Any) -> str:
    return unicodedata.normalize("NFC", html.unescape(str(value or "")).strip())


def _valid_company_website(value: Any) -> bool:
    raw = _clean_string(value)
    parsed = urlparse(raw)
    host = (parsed.hostname or "").casefold().removeprefix("www.")
    if parsed.scheme not in {"http", "https"} or not host or host in {"-", "www"}:
        return False
    if host in NON_COMPANY_HOSTS or host.endswith(".hometex.com.tr") or host.endswith(".messefrankfurt.com"):
        return False
    if raw.casefold() in {"http://-", "https://-", "-"} or any(parsed.path.casefold().endswith(suffix) for suffix in ASSET_SUFFIXES):
        return False
    return True


class _OfficialRequestLimiter:
    def __init__(self, minimum_interval: float = 1.0) -> None:
        self.minimum_interval = minimum_interval
        self._last_started: dict[str, float] = {}

    def wait(self, host: str) -> None:
        now = time.monotonic()
        previous = self._last_started.get(host)
        if previous is not None:
            delay = self.minimum_interval - (now - previous)
            if delay > 0:
                time.sleep(delay)
        self._last_started[host] = time.monotonic()


def _get(session: requests.Session, url: str, **kwargs: Any) -> requests.Response:
    host = (urlparse(url).hostname or "").casefold()
    if host not in ALLOWED_HOSTS:
        raise RuntimeError(f"official-source host is not allowlisted: {host}")
    kwargs.setdefault("timeout", 20)
    telemetry = getattr(session, "_a8_telemetry", None)
    limiter = getattr(session, "_a8_limiter", None)
    retryable = {429, 500, 502, 503, 504}
    deadline = time.monotonic() + 30.0
    for attempt in range(1, 4):
        if limiter is not None:
            limiter.wait(host)
        started = time.monotonic()
        try:
            response = session.get(url, **kwargs)
        except Exception:
            elapsed = time.monotonic() - started
            if isinstance(telemetry, list):
                telemetry.append({"attempt": attempt, "status": None, "host": host, "elapsed": elapsed, "response_sha256": None})
            if attempt >= 3:
                raise
            delay = min(max(0.0, deadline - time.monotonic()), min(2 ** (attempt - 1), 10.0))
            if delay:
                time.sleep(delay)
            continue
        elapsed = time.monotonic() - started
        status = int(response.status_code)
        response_hash = _sha256_bytes(response.content)
        if isinstance(telemetry, list):
            telemetry.append({"attempt": attempt, "status": status, "host": host, "elapsed": elapsed, "response_sha256": response_hash})
        if status not in retryable:
            response.raise_for_status()
            return response
        if attempt >= 3:
            response.raise_for_status()
        try:
            retry_after = min(float(response.headers.get("Retry-After", "0") or 0), 10.0)
        except (TypeError, ValueError):
            retry_after = 0.0
        delay = min(max(0.0, deadline - time.monotonic()), max(retry_after, min(2 ** (attempt - 1), 10.0)))
        if delay <= 0:
            response.raise_for_status()
        time.sleep(delay)
    raise RuntimeError("official request retry loop exhausted")


def _candidate_scalar(value: Any) -> str:
    if isinstance(value, str):
        return _clean_string(value)
    if isinstance(value, list):
        return "; ".join(_candidate_scalar(item) for item in value if _candidate_scalar(item))
    if isinstance(value, dict):
        for key in ("name", "label", "value", "text"):
            if value.get(key):
                return _candidate_scalar(value[key])
    return ""


def _labelled_website(soup: BeautifulSoup) -> str:
    """Read only a link contained by an explicit Website/Web Sitesi field."""
    labels = re.compile(r"^(website|web site|web sitesi|internet sitesi|official site)$", re.I)
    for label in soup.find_all(string=labels):
        parent = label.parent
        if parent is None:
            continue
        containers = [parent, parent.parent, parent.parent.parent if parent.parent else None]
        for container in containers:
            if container is None:
                continue
            for link in container.find_all("a", href=True):
                candidate = _clean_string(link.get("href"))
                if _valid_company_website(candidate):
                    return candidate
    return ""


def _find_category_objects(value: Any, target: str) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    if isinstance(value, dict):
        if any(str(item) == target for item in value.values()):
            found.append(value)
        for child in value.values():
            found.extend(_find_category_objects(child, target))
    elif isinstance(value, list):
        for child in value:
            found.extend(_find_category_objects(child, target))
    return found


def _extract_api_key(session: requests.Session, page_text: str, page_url: str) -> dict[str, str]:
    soup = BeautifulSoup(page_text, "html.parser")
    bundle_urls = []
    for script in soup.find_all("script", src=True):
        url = urljoin(page_url, script["src"])
        host = (urlparse(url).hostname or "").casefold()
        if host in ALLOWED_HOSTS and any(token in url.casefold() for token in ("main", "app")):
            bundle_urls.append(url)
    for bundle_url in dict.fromkeys(bundle_urls):
        bundle_response = _get(session, bundle_url, timeout=20)
        bundle_text = _strict_text(bundle_response, reject_replacement=False)
        for key_match in re.finditer(r'externalDataUrlApiKey="([^"]+)"', bundle_text):
            preceding = bundle_text[max(0, key_match.start() - 2000):key_match.start()]
            urls = re.findall(r'externalDataUrl="(https://api[.]messefrankfurt[.]com[^"]+)"', preceding)
            if urls and key_match.group(1):
                return {
                    "bundle_url": bundle_url,
                    "api_base_url": urls[-1],
                    "api_key": key_match.group(1),
                }
    raise RuntimeError("production Ambiente API key was not found in the official bundle")


def _hometex_detail(session: requests.Session, href: str, listing_label: str) -> dict[str, Any]:
    detail_url = urljoin(HOMETEX_LISTING_URL, href)
    response = _get(session, detail_url)
    text = _strict_text(response)
    soup = BeautifulSoup(text, "html.parser")
    title = _clean_string(soup.title.get_text(" ", strip=True) if soup.title else "")
    if not title or title.casefold() in {"hometex", "hometex | international home textile fair"}:
        raise RuntimeError(f"HOMETEX detail has no verified company title: {detail_url}")
    website = _labelled_website(soup)
    slug = href.split("?", 1)[0].rstrip("/").rsplit("/", 1)[-1]
    return {
        "source": "hometex_2026",
        "source_record_id": f"hometex_2026:{slug}",
        "display_name": _clean_string(listing_label),
        "legal_name": title.strip(),
        "brand": _clean_string(listing_label) if _clean_string(listing_label) and not _clean_string(listing_label).isdigit() else title.strip(),
        "official_profile_url": detail_url,
        "listed_website": website,
        "website": "",
        "response_sha256": _sha256_bytes(response.content),
        "observed_at": datetime.now(timezone.utc).isoformat(),
    }


def acquire_hometex(session: requests.Session) -> dict[str, Any]:
    response = _get(session, HOMETEX_LISTING_URL)
    text = _strict_text(response)
    soup = BeautifulSoup(text, "html.parser")
    links: dict[str, tuple[str, str]] = {}
    for link in soup.select('div.list ul li a[href^="/en/"]'):
        href = str(link.get("href", "")).strip().split("?", 1)[0].rstrip("/")
        if href in EXCLUDED_HOMETEX_PATHS:
            continue
        slug = href.rsplit("/", 1)[-1]
        if not slug:
            continue
        links.setdefault(slug, (href, link.get_text(" ", strip=True)))
    if len(links) != 580:
        raise RuntimeError(f"HOMETEX unique detail slug count changed: {len(links)} (expected 580)")

    records = []
    for slug, (href, label) in sorted(links.items()):
        records.append({
            "source": "hometex_2026",
            "source_record_id": f"hometex_2026:{slug}",
            "display_name": label.strip(),
            "legal_name": None,
            "brand": label.strip() if label.strip() and not label.strip().isdigit() else "",
            "official_profile_url": urljoin(HOMETEX_LISTING_URL, href),
            "website": "",
            "listing_label_verified": False,
            "listing_response_sha256": _sha256_bytes(response.content),
            "observed_at": datetime.now(timezone.utc).isoformat(),
        })
    return {
        "listing_url": HOMETEX_LISTING_URL,
        "listing_response_sha256": _sha256_bytes(response.content),
        "unique_slug_count": len(links),
        "records": records,
    }


def verify_selected_hometex(session: requests.Session, records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    verified = []
    for index, record in enumerate(records):
        detail = _hometex_detail(session, urlparse(record["official_profile_url"]).path, record["display_name"])
        record.update(detail)
        record["listing_label_verified"] = True
        verified.append(record)
        if index + 1 < len(records):
            time.sleep(0.15)
    return verified


def acquire_ambiente(session: requests.Session) -> dict[str, Any]:
    page_response = _get(session, AMBIENTE_SEARCH_PAGE)
    page_text = _strict_text(page_response)
    key_info = _extract_api_key(session, page_text, AMBIENTE_SEARCH_PAGE)
    headers = {"Apikey": key_info["api_key"]}
    filter_response = _get(session, AMBIENTE_FILTERS_URL, headers=headers)
    filter_text = _strict_text(filter_response)
    filters = json.loads(filter_text)
    category_matches: dict[str, str] = {}
    for category_id, expected_name in CATEGORY_NAMES.items():
        objects = _find_category_objects(filters, category_id)
        names = {_candidate_scalar(item.get("name") or item.get("label") or item.get("title")) for item in objects}
        names.discard("")
        if normalize_text(expected_name) not in {normalize_text(name) for name in names}:
            raise RuntimeError(f"Ambiente category mismatch for {category_id}: {sorted(names)}")
        category_matches[category_id] = expected_name

    params: list[tuple[str, str | int]] = [
        ("country", "TUR"),
        ("language", "en-GB"),
        ("q", ""),
        ("orderBy", "name"),
        ("pageNumber", 1),
        ("pageSize", 100),
        ("orSearchFallback", "true"),
        ("showJumpLabels", "true"),
        ("findEventVariable", "AMBIENTE"),
    ] + [("categoryId", category_id) for category_id in CATEGORY_NAMES]
    search_response = _get(session, AMBIENTE_SEARCH_URL, headers=headers, params=params)
    search_text = _strict_text(search_response)
    payload = json.loads(search_text)
    result = payload.get("result") if isinstance(payload, dict) else None
    metadata = result.get("metaData") if isinstance(result, dict) else None
    hits = result.get("hits") if isinstance(result, dict) else None
    total = metadata.get("hitsTotal") if isinstance(metadata, dict) else None
    if total != 96 or not isinstance(hits, list) or len(hits) != 96:
        raise RuntimeError(f"Ambiente result changed: hitsTotal={total}, hits={len(hits) if isinstance(hits, list) else None}")

    records = []
    seen: set[str] = set()
    for hit in hits:
        exhibitor = hit.get("exhibitor") if isinstance(hit, dict) else None
        if not isinstance(exhibitor, dict):
            raise RuntimeError("Ambiente hit has no exhibitor object")
        rewrite_id = str(exhibitor.get("rewriteId") or "").strip()
        if not rewrite_id or rewrite_id in seen:
            raise RuntimeError(f"Ambiente rewriteId missing or duplicate: {rewrite_id}")
        seen.add(rewrite_id)
        detail_url = f"https://ambiente.messefrankfurt.com/frankfurt/en/exhibitor-search.detail.html/{rewrite_id}.html"
        address = exhibitor.get("address") if isinstance(exhibitor.get("address"), dict) else {}
        records.append({
            "source": "ambiente_2026",
            "source_record_id": f"ambiente_2026:{rewrite_id}",
            "display_name": _clean_string(exhibitor.get("name")),
            "legal_name": None,
            "brand": _candidate_scalar(exhibitor.get("brands")),
            "official_profile_url": detail_url,
            "website": "",
            "listed_website": _clean_string(exhibitor.get("homepage")),
            "listed_email": _clean_string(address.get("email")),
            "listed_phone": _clean_string(address.get("tel")),
            "listed_address": _clean_string(address),
            "observed_first_party_fields": {
                "id": exhibitor.get("id"),
                "rewriteId": exhibitor.get("rewriteId"),
                "name": _clean_string(exhibitor.get("name")),
                "href": _clean_string(exhibitor.get("href")),
                "address": address,
                "homepage": _clean_string(exhibitor.get("homepage")),
                "description": _clean_string(exhibitor.get("description")),
                "brands": _candidate_scalar(exhibitor.get("brands")),
                "categories": exhibitor.get("categories"),
            },
            "response_sha256": _sha256_bytes(search_response.content),
            "observed_at": datetime.now(timezone.utc).isoformat(),
        })
    return {
        "search_page_url": AMBIENTE_SEARCH_PAGE,
        "filters_url": AMBIENTE_FILTERS_URL,
        "search_url": AMBIENTE_SEARCH_URL,
        "search_page_response_sha256": _sha256_bytes(page_response.content),
        "filters_response_sha256": _sha256_bytes(filter_response.content),
        "search_response_sha256": _sha256_bytes(search_response.content),
        "hits_total": total,
        "unique_rewrite_id_count": len(seen),
        "category_matches": category_matches,
        "records": records,
    }


def _xlsx_rows(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        return [dict(zip(headers, row)) for row in iterator if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()


def _known_records(original_root: Path, diagnostic_manifest: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for path in [
        original_root / "input" / "firms.xlsx",
        original_root / "output" / "exhibition_893_reconciled_v3" / "all_results.xlsx",
    ]:
        for row in _xlsx_rows(path):
            source_id = str(row.get("source_record_id") or "").strip()
            name = str(row.get("listed_legal_name") or row.get("company") or "").strip()
            if name or source_id:
                records.append({
                    "source_record_id": source_id,
                    "legal_name": _clean_string(name),
                    "brand": _clean_string(row.get("brands")),
                    "website": _clean_string(row.get("website") or row.get("listed_website")),
                })
    for path in sorted((original_root / "outputs").glob("golden*/*.xlsx")):
        for row in _xlsx_rows(path):
            source_id = str(row.get("source_record_id") or "").strip()
            name = _clean_string(row.get("listed_legal_name") or row.get("company"))
            if name or source_id:
                records.append({
                    "source_record_id": source_id,
                    "legal_name": name,
                    "brand": _clean_string(row.get("brands")),
                    "website": _clean_string(row.get("website") or row.get("listed_website")),
                })
    if diagnostic_manifest.is_file():
        payload = json.loads(diagnostic_manifest.read_text(encoding="utf-8"))
        for row in payload.get("diagnostic_rows", []):
            records.append({
                "source_record_id": str(row.get("source_record_id") or "").strip(),
                "legal_name": _clean_string(row.get("Company")),
                "brand": _clean_string(row.get("Company")),
                "website": _clean_string(row.get("official_profile_url")),
            })
    return records


def _overlap_inputs(original_root: Path, diagnostic_manifest: Path) -> list[dict[str, Any]]:
    paths = [
        original_root / "input" / "firms.xlsx",
        original_root / "output" / "exhibition_893_reconciled_v3" / "all_results.xlsx",
        *sorted((original_root / "outputs").glob("golden*/*.xlsx")),
        diagnostic_manifest,
    ]
    result: list[dict[str, Any]] = []
    for path in paths:
        entry = {"path": str(path.resolve()), "sha256": None, "record_count": 0, "parse_status": "missing"}
        if path.is_file():
            entry["sha256"] = _sha256_file(path)
            if path.suffix.casefold() == ".json":
                payload = json.loads(path.read_text(encoding="utf-8"))
                rows = payload.get("diagnostic_rows", [])
            else:
                rows = _xlsx_rows(path)
            entry["record_count"] = len(rows)
            entry["parse_status"] = "ok"
        result.append(entry)
    if any(entry["parse_status"] != "ok" for entry in result):
        raise RuntimeError(f"overlap input manifest contains unreadable paths: {result}")
    return result


def _match_keys(record: dict[str, Any]) -> dict[str, str]:
    keys = {"source_record_id": _clean_string(record.get("source_record_id"))}
    for field in ("legal_name", "brand"):
        value = normalize_text(str(record.get(field) or "").strip())
        if value:
            keys[field] = value
    domain = registrable_domain(str(record.get("website") or "").strip())
    if domain:
        keys["registrable_domain"] = domain
    return keys


def _select_pool(records: list[dict[str, Any]], known: list[dict[str, Any]], source: str, limit: int = 60) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    known_by_key: dict[str, set[str]] = {}
    for old in known:
        old_id = _clean_string(old.get("source_record_id"))
        for key_type, value in _match_keys(old).items():
            if value:
                known_by_key.setdefault(f"{key_type}:{value}", set()).add(old_id)
    excluded: list[dict[str, Any]] = []
    remaining: list[dict[str, Any]] = []
    for record in records:
        matches = []
        for key_type, value in _match_keys(record).items():
            previous_ids = sorted(known_by_key.get(f"{key_type}:{value}", set()))
            matches.extend((key_type, previous_id) for previous_id in previous_ids)
        if matches:
            for key_type, previous_id in matches:
                excluded.append({
                    "source": source,
                    "source_record_id": record["source_record_id"],
                    "matched_key": key_type,
                    "matched_previous_id": previous_id,
                    "reason": "overlap_with_previous_893_golden_or_diagnostic",
                })
        else:
            remaining.append(record)
    remaining.sort(key=lambda item: hashlib.sha256(
        f"{item['source_record_id']}architect-independent-v1".encode("utf-8")
    ).hexdigest())
    if len(remaining) < limit:
        raise RuntimeError(f"{source} has only {len(remaining)} independent candidates; required {limit}")
    selected = [dict(record) for record in remaining[:limit]]
    listing_url = HOMETEX_LISTING_URL if source == "hometex_2026" else AMBIENTE_SEARCH_PAGE
    for record in selected:
        record["listing_url"] = listing_url
    return selected, excluded


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--original-root", type=Path, required=True)
    parser.add_argument("--diagnostic-manifest", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    session = requests.Session()
    session.headers.update({"User-Agent": "A8-source-review/1.0"})
    request_telemetry: list[dict[str, Any]] = []
    session._a8_telemetry = request_telemetry
    session._a8_limiter = _OfficialRequestLimiter(1.0)
    hometex = acquire_hometex(session)
    ambiente = acquire_ambiente(session)
    known = _known_records(args.original_root.resolve(), args.diagnostic_manifest.resolve())
    overlap_inputs = _overlap_inputs(args.original_root.resolve(), args.diagnostic_manifest.resolve())
    hometex_verified = verify_selected_hometex(session, hometex["records"])
    hometex_selected, hometex_excluded = _select_pool(hometex_verified, known, "hometex_2026")
    ambiente_selected, ambiente_excluded = _select_pool(ambiente["records"], known, "ambiente_2026")
    payload = {
        "schema_version": 1,
        "selection_algorithm": "SHA256(source_record_id + architect-independent-v1)",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "hometex": {**{key: value for key, value in hometex.items() if key != "records"}, "selected": hometex_selected},
        "ambiente": {**{key: value for key, value in ambiente.items() if key != "records"}, "selected": ambiente_selected},
        "exclusion_manifest": hometex_excluded + ambiente_excluded,
        "overlap_inputs": {
            "inputs": overlap_inputs,
            "known_record_count": len(known),
        },
        "request_telemetry": request_telemetry,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "hometex_slugs": hometex["unique_slug_count"],
        "hometex_selected": len(hometex_selected),
        "ambiente_hits": ambiente["hits_total"],
        "ambiente_selected": len(ambiente_selected),
        "exclusions": len(hometex_excluded) + len(ambiente_excluded),
    }, sort_keys=True))


if __name__ == "__main__":
    main()

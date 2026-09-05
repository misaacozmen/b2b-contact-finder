"""Copy the immutable 893-row source input and add its derived IDs in a new file."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from modules import run_context


def build(source: Path, output: Path, source_id_manifest: Path | None = None) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not rows:
        raise RuntimeError("source input is empty")
    headers = [str(value or "").strip() for value in rows[0]]
    folded = [header.casefold() for header in headers]
    if "company" not in folded or "source_record_id" in folded:
        raise RuntimeError("source input must have company and no existing source_record_id")
    source_index = {header.casefold(): index for index, header in enumerate(headers)}
    manifest_ids: list[str] | None = None
    if source_id_manifest is not None:
        manifest = json.loads(Path(source_id_manifest).read_text(encoding="utf-8"))
        declared_input_hash = str(manifest.get("input_sha256") or "").strip().casefold()
        observed_input_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        if not declared_input_hash or declared_input_hash != observed_input_hash:
            raise RuntimeError("source input hash does not match source-ID manifest")
        manifest_ids = list(manifest.get("ordered_source_record_ids") or manifest.get("source_record_ids") or [])
        if len(manifest_ids) != 893 or len(set(manifest_ids)) != 893 or any(not str(value).strip() for value in manifest_ids):
            raise RuntimeError("source-ID manifest must contain exactly 893 unique nonempty IDs")
    out = Workbook()
    sheet = out.active
    sheet.title = "Companies"
    sheet.append(headers + ["source_record_id"])
    ids: list[str] = []
    for row in rows[1:]:
        company = str(row[source_index["company"]] or "").strip()
        if not company:
            continue
        values = {header: (row[index] if index < len(row) else "") for index, header in enumerate(headers)}
        record = {
            "company": company,
            "source": values.get("source", ""),
            "profile_url": values.get("profile_url", ""),
            "listing_url": values.get("listing_url", ""),
            "hall": values.get("hall", ""),
            "stand": values.get("stand", ""),
        }
        source_id = manifest_ids[len(ids)] if manifest_ids is not None else run_context.source_record_identity(record)[0]
        if source_id in ids:
            raise RuntimeError(f"duplicate derived source_record_id: {source_id}")
        ids.append(source_id)
        sheet.append(list(row) + [source_id])
    if len(ids) != 893:
        raise RuntimeError(f"expected 893 source records, got {len(ids)}")
    output.parent.mkdir(parents=True, exist_ok=True)
    out.save(output)
    out.close()
    print(f"rows={len(ids)} output={output}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--source-id-manifest", type=Path)
    args = parser.parse_args()
    build(args.source, args.output, args.source_id_manifest)


if __name__ == "__main__":
    main()

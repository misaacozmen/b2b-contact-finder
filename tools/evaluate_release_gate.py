"""Evidence-only Golden6 release evaluator.

The three release flags are derived here.  They are never accepted from CLI,
environment, configuration, or a supplied JSON override.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import urllib.parse
import urllib.request
from pathlib import Path, PurePosixPath, PureWindowsPath

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from modules import checkpoint, run_context

from tools.free_only_contract import validate_database as validate_free_only_database
from tools.free_only_contract import validate_manifest as validate_free_only_manifest
from tools.free_only_contract import expected_offline_run_config
from tools.free_only_contract import ACTIVITY_COUNTER_KEYS
from tools.free_only_contract import classify_replay_receipt
from tools.free_only_contract import compare_run_configs
from tools.free_only_contract import inspect_budget_config
from tools.free_only_contract import inspect_database
from tools.free_only_contract import validate_manifest_config_sha


REQUIRED_CI_JOBS = ("test", "browser-smoke", "ocr-smoke")
REQUEST_COUNTER_PREFIXES = ("api.", "http.search.", "http.crawler.")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json(path: Path) -> dict:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"json_not_object:{path}")
    return value


def _git(repo: Path, *args: str) -> tuple[int, str, str]:
    result = subprocess.run(
        ["git", *args], cwd=repo, capture_output=True, text=True,
    )
    return result.returncode, result.stdout.strip(), result.stderr.strip()


def _ids(path: Path, sheet_name: str | None = None) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip().casefold() for value in rows[0]]
        if "source_record_id" not in headers:
            return []
        index = headers.index("source_record_id")
        return [str(row[index] or "").strip() for row in rows[1:] if any(value is not None for value in row)]
    finally:
        workbook.close()


def _provider_http_calls(artifact_dir: Path) -> dict:
    """Read every present telemetry source without coercing malformed data to zero."""
    def invalid(reason: str) -> dict:
        return {"valid": False, "count": None, "failure": reason}

    try:
        telemetry = _json(artifact_dir / "telemetry.json")
        try:
            checkpoint.validate_free_only_canonical_telemetry(telemetry)
        except Exception as exc:
            return invalid(str(exc))
        provider_budgets = telemetry["provider_budgets"]
        physical_total = sum(values["physical_http_attempts"] for values in provider_budgets.values())

        legacy_total = 0
        if "counters" in telemetry:
            counters = telemetry["counters"]
            if type(counters) is not dict:
                return invalid("telemetry_counters_invalid")
            for key, value in counters.items():
                if type(key) is not str:
                    return invalid("telemetry_counter_key_invalid")
                if any(key.startswith(prefix) for prefix in REQUEST_COUNTER_PREFIXES) and key.endswith(".requests"):
                    if type(value) is not int or value < 0:
                        return invalid(f"telemetry_counter_value_invalid:{key}")
                    if value != 0:
                        return invalid(f"telemetry_counter_nonzero:{key}")
                    legacy_total += value
        return {"valid": True, "count": physical_total + legacy_total, "failure": ""}
    except (OSError, TypeError, ValueError, json.JSONDecodeError) as exc:
        return invalid(f"telemetry_invalid:{type(exc).__name__}")


def _replay_coverage(artifact_dir: Path) -> dict:
    try:
        coverage = _json(artifact_dir / "discovery_coverage.json")
        count = coverage.get("replay_miss_count")
        if type(count) is not int or count < 0:
            return {"valid": False, "count": None, "failure": "replay_miss_count_invalid"}
        return {"valid": True, "count": count, "failure": ""}
    except (OSError, TypeError, ValueError, json.JSONDecodeError) as exc:
        return {"valid": False, "count": None, "failure": f"replay_coverage_invalid:{type(exc).__name__}"}


def _telemetry_zero(observation: object) -> bool:
    return (
        type(observation) is dict
        and observation.get("valid") is True
        and type(observation.get("count")) is int
        and observation.get("count") == 0
    )


def _telemetry_failure_reason(prefix: str, exc: BaseException) -> str:
    reason = str(exc)
    if reason.startswith("TELEMETRY_CANONICAL_"):
        return reason
    return f"{prefix}:{reason}"


def _coverage_zero(observation: object) -> bool:
    return (
        type(observation) is dict
        and observation.get("valid") is True
        and type(observation.get("count")) is int
        and observation.get("count") == 0
    )


def _package_files_valid(package_dir: Path, manifest: dict) -> tuple[bool, list[str]]:
    """Validate the complete, path-safe package file registry."""
    package_dir = Path(package_dir).resolve()
    failures: list[str] = []
    if type(manifest) is not dict:
        return False, ["package_manifest_not_object"]
    if (package_dir / "package_manifest.json").is_symlink():
        failures.append("package_manifest_symlink")

    def safe_relative(value: object) -> str | None:
        if type(value) is not str or not value or "\x00" in value or "\\" in value:
            return None
        posix = PurePosixPath(value)
        windows = PureWindowsPath(value)
        if posix.is_absolute() or windows.is_absolute() or windows.drive:
            return None
        parts = value.split("/")
        if value != posix.as_posix() or any(part in {"", ".", ".."} for part in parts):
            return None
        return value

    files = manifest.get("files")
    registry: dict[str, dict] = {}
    if type(files) is not dict or not files:
        failures.append("package_files_registry_invalid")
    else:
        for relative, info in files.items():
            canonical = safe_relative(relative)
            if canonical is None:
                failures.append(f"package_file_path_invalid:{relative!r}")
                continue
            if canonical == "package_manifest.json" or canonical in registry:
                failures.append(f"package_file_path_collision:{canonical}")
                continue
            if (
                type(info) is not dict
                or set(info) != {"sha256", "bytes"}
                or type(info.get("sha256")) is not str
                or not re.fullmatch(r"[0-9a-f]{64}", info.get("sha256", ""))
                or type(info.get("bytes")) is not int
                or info.get("bytes") < 0
            ):
                failures.append(f"package_file_record_invalid:{canonical}")
                continue
            registry[canonical] = info

    actual: dict[str, Path] = {}
    for path in package_dir.rglob("*") if package_dir.is_dir() else ():
        if path == package_dir / "package_manifest.json":
            continue
        if path.is_symlink():
            failures.append(f"package_file_symlink:{path}")
            continue
        if not path.is_file():
            continue
        try:
            relative = path.relative_to(package_dir).as_posix()
        except ValueError:
            failures.append(f"package_file_outside_root:{path}")
            continue
        canonical = safe_relative(relative)
        if canonical is None:
            failures.append(f"package_file_path_invalid:{relative!r}")
            continue
        resolved = path.resolve()
        try:
            resolved.relative_to(package_dir)
        except ValueError:
            failures.append(f"package_file_outside_root:{relative}")
            continue
        actual[canonical] = path

    for relative in sorted(set(registry) - set(actual)):
        failures.append(f"package_file_missing:{relative}")
    for relative in sorted(set(actual) - set(registry)):
        failures.append(f"package_file_unregistered:{relative}")
    for relative in sorted(set(registry) & set(actual)):
        info = registry[relative]
        path = actual[relative]
        if _sha256(path) != info["sha256"] or path.stat().st_size != info["bytes"]:
            failures.append(f"package_file_hash_mismatch:{relative}")

    replay = manifest.get("replay")
    if type(replay) is not dict:
        failures.append("replay_metadata_invalid")
        replay = {}
    snapshot_name = safe_relative(replay.get("snapshot"))
    if snapshot_name is None or snapshot_name not in registry:
        failures.append("replay_snapshot_unregistered")
    else:
        snapshot = actual.get(snapshot_name)
        if snapshot is None:
            failures.append("replay_snapshot_missing")
        elif type(replay.get("snapshot_sha256")) is not str or _sha256(snapshot) != replay.get("snapshot_sha256"):
            failures.append("replay_snapshot_hash_mismatch")

    body_shards = replay.get("body_shards", [])
    if type(body_shards) is not list:
        failures.append("replay_body_shards_invalid")
    else:
        seen_shards: set[str] = set()
        for shard in body_shards:
            if type(shard) is not dict:
                failures.append("replay_body_shard_record_invalid")
                continue
            shard_name = safe_relative(shard.get("path"))
            if (
                shard_name is None
                or shard_name in seen_shards
                or shard_name not in registry
                or type(shard.get("sha256")) is not str
                or type(shard.get("bytes")) is not int
                or shard.get("bytes") < 0
                or registry.get(shard_name, {}).get("sha256") != shard.get("sha256")
                or registry.get(shard_name, {}).get("bytes") != shard.get("bytes")
            ):
                failures.append(f"replay_body_shard_invalid:{shard.get('path')!r}")
            elif shard_name is not None:
                seen_shards.add(shard_name)

    if manifest.get("source_integrity_before") != manifest.get("source_integrity_after"):
        failures.append("source_run_hash_changed")
    package_input = manifest.get("input")
    if type(package_input) is not dict or type(package_input.get("record_count")) is not int or package_input.get("record_count") != 20:
        failures.append("package_record_count_not_20")
    return not failures, failures


def _validator_report(repo: Path, expected: Path, artifact_dir: Path, output: Path) -> tuple[dict, list[str]]:
    if not artifact_dir.is_dir():
        return {}, [f"artifact_dir_missing:{artifact_dir}"]
    command = [
        sys.executable, "validate_golden_xlsx.py",
        "--expected", str(expected),
        "--actual", str(artifact_dir / "contacts.xlsx"),
        "--candidates", str(artifact_dir / "website_candidates.xlsx"),
        "--all-results", str(artifact_dir / "all_results.xlsx"),
        "--json-output", str(output),
    ]
    result = subprocess.run(command, cwd=repo, capture_output=True, text=True)
    if not output.is_file():
        return {}, [f"validator_report_missing:{output}", f"validator_exit:{result.returncode}"]
    report = _json(output)
    failures = [] if result.returncode == 0 and report.get("status") == "PASS" else ["validator_status_not_pass"]
    failures.extend(str(value) for value in report.get("issues", []) if value)
    return report, failures


def _metrics(report: dict) -> dict:
    return {
        "fields": report.get("fields", {}),
        "stages": report.get("stages", {}),
        "records": report.get("records", {}),
    }


def _positive_denominators(report: dict) -> bool:
    fields = report.get("fields") if isinstance(report.get("fields"), dict) else {}
    required_fields = {"website", "email", "phone"}
    if set(fields) != required_fields:
        return False
    if any(
        values.get("tp", 0) + values.get("fp", 0) <= 0
        or values.get("tp", 0) + values.get("fn", 0) <= 0
        or values.get("precision") is None
        or values.get("recall") is None
        for values in fields.values() if isinstance(values, dict)
    ):
        return False
    stages = report.get("stages") if isinstance(report.get("stages"), dict) else {}
    return all(
        stages.get(key, 0) > 0
        for key in ("published_count", "expected_websites", "selected_count")
    )


REQUIRED_V2_RECEIPT_FIELDS = {
    "paid_budget_nonzero", "budget_check_completed", "budget_offenders",
    "threshold_mismatch", "threshold_mismatch_details", "config_mismatch",
    "config_mismatch_details", "paid_activity_nonzero", "paid_activity_check_completed",
    "config_invalid_details", "failure_reason", "exception", "evidence_class",
    "replay_evidence_eligible", "pipeline_exit_code", "exit_code", "paid_activity",
    "paid_provider_calls", "replay_network_events", "network_events", "free_only_config_valid",
    "config_invalid", "network_activity", "complete_run_missing", "replay_pipeline_failure",
    "database_validation_failed", "failure_reasons",
}


def _strict_zero_activity(value: object) -> bool:
    return (
        type(value) is dict
        and set(value) == set(ACTIVITY_COUNTER_KEYS)
        and all(type(value[key]) is int and value[key] == 0 for key in ACTIVITY_COUNTER_KEYS)
    )


def _validate_receipt_chain(
    receipt: dict,
    package_manifest: dict | None,
    package_manifest_path: Path | None,
    offline_manifest: dict | None,
    offline_manifest_path: Path | None,
) -> list[str]:
    if package_manifest is None:
        return []
    failures: list[str] = []
    if package_manifest_path is None or not package_manifest_path.is_file():
        failures.append("RECEIPT_PACKAGE_MANIFEST_MISSING")
    else:
        if receipt.get("package_manifest_sha256") != _sha256(package_manifest_path):
            failures.append("RECEIPT_PACKAGE_MANIFEST_HASH_MISMATCH")
    package_input = package_manifest.get("input") if isinstance(package_manifest.get("input"), dict) else {}
    package_replay = package_manifest.get("replay") if isinstance(package_manifest.get("replay"), dict) else {}
    input_path_value = package_input.get("path")
    input_path = None
    if isinstance(input_path_value, str) and input_path_value:
        input_path = Path(input_path_value)
        if not input_path.is_absolute():
            input_path = package_manifest_path.parent / input_path
        input_path = input_path.resolve()
    input_hash = _sha256(input_path) if input_path is not None and input_path.is_file() else None
    if (
        input_hash is None
        or input_hash != package_input.get("sha256")
        or type(receipt.get("input_sha256")) is not str
        or receipt.get("input_sha256") != input_hash
    ):
        failures.append("RECEIPT_INPUT_HASH_MISMATCH")
    if (
        type(receipt.get("replay_snapshot_sha256")) is not str
        or type(package_replay.get("snapshot_sha256")) is not str
        or receipt.get("replay_snapshot_sha256") != package_replay.get("snapshot_sha256")
    ):
        failures.append("RECEIPT_REPLAY_SNAPSHOT_HASH_MISMATCH")
    if offline_manifest is None or offline_manifest_path is None:
        failures.append("RECEIPT_OFFLINE_MANIFEST_MISSING")
        return failures
    run_id = offline_manifest.get("run_id")
    artifact_set = offline_manifest.get("artifact_set_sha256")
    expected_run_dir = offline_manifest_path.parent.resolve()
    expected_artifact_dir = (expected_run_dir / "output" / "artifacts" / str(artifact_set or "")).resolve()
    if receipt.get("run_id") != run_id:
        failures.append("RECEIPT_RUN_ID_MISMATCH")
    if Path(str(receipt.get("run_dir", ""))).resolve() != expected_run_dir:
        failures.append("RECEIPT_RUN_DIR_MISMATCH")
    if receipt.get("artifact_set_sha256") != artifact_set:
        failures.append("RECEIPT_ARTIFACT_SET_HASH_MISMATCH")
    if Path(str(receipt.get("artifact_dir", ""))).resolve() != expected_artifact_dir:
        failures.append("RECEIPT_ARTIFACT_DIR_MISMATCH")
    return failures


def receipt_eligibility(
    receipt: dict,
    *,
    package_manifest: dict | None = None,
    package_manifest_path: Path | None = None,
    offline_manifest: dict | None = None,
    offline_manifest_path: Path | None = None,
) -> tuple[bool, list[str]]:
    """Return whether a replay receipt can supply behavioral replay evidence."""
    failures: list[str] = []
    if type(receipt) is not dict:
        return False, ["RECEIPT_NOT_OBJECT"]
    if package_manifest is None or package_manifest_path is None or offline_manifest is None or offline_manifest_path is None:
        return False, ["RECEIPT_CONTEXT_REQUIRED"]
    if type(receipt.get("schema_version")) is int and receipt.get("schema_version") == 2:
        missing = sorted(REQUIRED_V2_RECEIPT_FIELDS - set(receipt))
        if missing:
            failures.append("RECEIPT_V2_FIELDS_MISSING:" + ",".join(missing))
        conditions = {
            "status_pass": receipt.get("status") == "PASS",
            "exit_zero": type(receipt.get("exit_code")) is int and receipt.get("exit_code") == 0,
            "pipeline_exit_zero": type(receipt.get("pipeline_exit_code")) is int and receipt.get("pipeline_exit_code") == 0,
            "evidence_class": receipt.get("evidence_class") == "REPLAY_PASS",
            "eligible_flag": receipt.get("replay_evidence_eligible") is True,
            "config_valid_flag": receipt.get("free_only_config_valid") is True,
            "failure_reason_empty": receipt.get("failure_reason") == "",
            "exception_empty": receipt.get("exception") == "",
            "paid_budget_zero": receipt.get("paid_budget_nonzero") is False,
            "budget_check_complete": receipt.get("budget_check_completed") is True,
            "threshold_zero": receipt.get("threshold_mismatch") is False,
            "config_zero": receipt.get("config_mismatch") is False,
            "config_invalid_details_empty": type(receipt.get("config_invalid_details")) is list and receipt.get("config_invalid_details") == [],
            "config_invalid_flag": receipt.get("config_invalid") is False,
            "network_activity_flag": receipt.get("network_activity") is False,
            "complete_run_missing_flag": receipt.get("complete_run_missing") is False,
            "pipeline_failure_flag": receipt.get("replay_pipeline_failure") is False,
            "database_validation_flag": receipt.get("database_validation_failed") is False,
            "failure_reasons_empty": type(receipt.get("failure_reasons")) is list and receipt.get("failure_reasons") == [],
            "paid_activity_zero": receipt.get("paid_activity_nonzero") is False,
            "paid_activity_check_complete": receipt.get("paid_activity_check_completed") is True,
            "network_zero": type(receipt.get("replay_network_events")) is int and receipt.get("replay_network_events") == 0,
            "provider_calls_zero": type(receipt.get("paid_provider_calls")) is int and receipt.get("paid_provider_calls") == 0,
            "budget_offenders_empty": type(receipt.get("budget_offenders")) is list and receipt.get("budget_offenders") == [],
            "threshold_details_empty": type(receipt.get("threshold_mismatch_details")) is list and receipt.get("threshold_mismatch_details") == [],
            "config_details_empty": type(receipt.get("config_mismatch_details")) is list and receipt.get("config_mismatch_details") == [],
            "network_events_empty": type(receipt.get("network_events")) is list and receipt.get("network_events") == [],
            "activity_counters_zero": _strict_zero_activity(receipt.get("paid_activity")),
        }
        for name, passed in conditions.items():
            if not passed:
                failures.append("RECEIPT_" + name.upper())
        failures.extend(_validate_receipt_chain(receipt, package_manifest, package_manifest_path, offline_manifest, offline_manifest_path))
        return not failures, failures
    if type(receipt.get("schema_version")) is int and receipt.get("schema_version") == 1:
        conditions = {
            "status_pass": receipt.get("status") == "PASS",
            "exit_zero": type(receipt.get("exit_code")) is int and receipt.get("exit_code") == 0,
            "exception_empty": type(receipt.get("exception")) is str and receipt.get("exception") == "",
            "config_valid_flag": receipt.get("free_only_config_valid") is True,
            "paid_budget_zero": receipt.get("paid_budget_nonzero") is False,
            "provider_calls_zero": type(receipt.get("paid_provider_calls")) is int and receipt.get("paid_provider_calls") == 0,
            "network_zero": type(receipt.get("replay_network_events")) is int and receipt.get("replay_network_events") == 0,
            "network_events_empty": type(receipt.get("network_events")) is list and receipt.get("network_events") == [],
            "activity_counters_zero": _strict_zero_activity(receipt.get("paid_activity")),
        }
        for name, passed in conditions.items():
            if not passed:
                failures.append("LEGACY_RECEIPT_" + name.upper())
        failures.extend(_validate_receipt_chain(receipt, package_manifest, package_manifest_path, offline_manifest, offline_manifest_path))
        return not failures, failures
    return False, ["RECEIPT_SCHEMA_UNSUPPORTED"]


def _threshold_rejection_invariant(receipt: dict) -> bool:
    return (
        type(receipt) is dict
        and receipt.get("schema_version") == 2
        and receipt.get("status") == "FAIL"
        and type(receipt.get("exit_code")) is int and receipt.get("exit_code") != 0
        and type(receipt.get("pipeline_exit_code")) is int and receipt.get("pipeline_exit_code") == 0
        and receipt.get("evidence_class") == "REJECTED_THRESHOLD_CANDIDATE"
        and receipt.get("replay_evidence_eligible") is False
        and receipt.get("failure_reason") == "REJECTED_THRESHOLD_CANDIDATE"
        and receipt.get("failure_reasons") == ["REJECTED_THRESHOLD_CANDIDATE"]
        and receipt.get("exception") == "REJECTED_THRESHOLD_CANDIDATE"
        and receipt.get("free_only_config_valid") is False
        and receipt.get("paid_budget_nonzero") is False
        and receipt.get("config_invalid") is False
        and receipt.get("network_activity") is False
        and receipt.get("complete_run_missing") is False
        and receipt.get("replay_pipeline_failure") is False
        and receipt.get("paid_activity_nonzero") is False
        and receipt.get("config_mismatch") is False
        and receipt.get("database_validation_failed") is False
        and receipt.get("threshold_mismatch") is True
        and receipt.get("budget_check_completed") is True
        and receipt.get("paid_activity_check_completed") is True
        and type(receipt.get("paid_provider_calls")) is int and receipt.get("paid_provider_calls") == 0
        and _strict_zero_activity(receipt.get("paid_activity"))
        and type(receipt.get("budget_offenders")) is list and receipt.get("budget_offenders") == []
        and type(receipt.get("config_invalid_details")) is list and receipt.get("config_invalid_details") == []
        and type(receipt.get("threshold_mismatch_details")) is list and receipt.get("threshold_mismatch_details") != []
        and type(receipt.get("config_mismatch_details")) is list and receipt.get("config_mismatch_details") == []
        and type(receipt.get("network_events")) is list and receipt.get("network_events") == []
        and type(receipt.get("replay_network_events")) is int and receipt.get("replay_network_events") == 0
    )


def _receipt_artifact_name(
    receipt: dict,
    eligible: bool,
    *,
    threshold_only: bool = False,
    threshold_candidate: bool | None = None,
) -> str:
    if eligible:
        return "replay_receipt"
    if threshold_only and (
        _threshold_rejection_invariant(receipt)
        if threshold_candidate is None else threshold_candidate
    ):
        return "rejected_threshold_candidate_receipt"
    return "failed_replay_receipt"


def _recomputed_threshold_candidate(
    receipt: dict,
    package_manifest: dict,
    package_manifest_path: Path,
    offline_manifest: dict | None,
    offline_manifest_path: Path | None,
    *,
    package_valid: bool,
    package_free_only_valid: bool,
    live_bundle_valid: bool,
    live_database_valid: bool,
    offline_bundle_valid: bool,
    offline_database_valid: bool,
    live_telemetry_valid: bool,
    offline_telemetry_valid: bool,
    offline_replay_coverage_valid: bool,
) -> bool:
    """Require the underlying package/run/DB to independently prove threshold-only."""
    if not (
        package_valid
        and package_free_only_valid
        and live_bundle_valid
        and live_database_valid
        and offline_bundle_valid
        and offline_database_valid
        and live_telemetry_valid
        and offline_telemetry_valid
        and offline_replay_coverage_valid
        and offline_manifest is not None
        and offline_manifest_path is not None
        and _threshold_rejection_invariant(receipt)
    ):
        return False
    package = package_manifest
    package_budget = inspect_budget_config(package)
    run_budget = inspect_budget_config(offline_manifest)
    details: list[dict] = []
    for label, candidate in (("package", package), ("run", offline_manifest)):
        try:
            validate_free_only_manifest(candidate, require_complete=False)
        except ValueError as exc:
            if str(exc) != "free_only_paid_budget_nonzero":
                details.append({"field": f"{label}_semantic", "reason": str(exc)})
        try:
            validate_manifest_config_sha(candidate)
        except Exception as exc:
            details.append({"field": f"{label}_config_hash", "reason": type(exc).__name__})
    try:
        expected_config = expected_offline_run_config(package["run_config"])
        threshold_info = compare_run_configs(expected_config, offline_manifest.get("run_config", {}))
    except Exception as exc:
        details.append({"field": "offline_run_config", "reason": type(exc).__name__})
        expected_config = {}
        threshold_info = {}
    db_path = offline_manifest_path.parent / "state" / "progress.sqlite3"
    try:
        database = inspect_database(
            db_path,
            str(offline_manifest.get("run_id", offline_manifest_path.parent.name)),
            int(package.get("input", {}).get("record_count", 0)),
        )
    except (OSError, TypeError, ValueError):
        return False
    actual = classify_replay_receipt(
        budget_checks=[package_budget, run_budget],
        threshold_info=threshold_info,
        database=database,
        run_present=True,
        complete_run_missing=not (
            offline_manifest.get("complete") is True
            and offline_manifest.get("phase") == "COMPLETE"
            and offline_manifest.get("finalized") is True
            and offline_manifest.get("status") == "complete_free_only"
        ),
        pipeline_exit_code=receipt.get("pipeline_exit_code", receipt.get("exit_code", 1)),
        network_event_count=receipt.get("replay_network_events"),
        config_invalid=bool(details),
        config_invalid_details=details,
        base={"network_events": receipt.get("network_events", [])},
    )
    return (
        not details
        and package_budget.get("budget_check_completed") is True
        and run_budget.get("budget_check_completed") is True
        and actual.get("status") == "FAIL"
        and actual.get("failure_reason") == "REJECTED_THRESHOLD_CANDIDATE"
        and actual.get("failure_reasons") == ["REJECTED_THRESHOLD_CANDIDATE"]
        and actual.get("evidence_class") == "REJECTED_THRESHOLD_CANDIDATE"
        and actual.get("replay_evidence_eligible") is False
    )


def evaluate_evidence(evidence: dict) -> dict:
    """Compute the flags strictly from evidence fields."""
    failures: list[str] = []
    behavioral = evidence.get("behavioral") if isinstance(evidence.get("behavioral"), dict) else {}
    checks = evidence.get("checks") if isinstance(evidence.get("checks"), dict) else {}
    git = evidence.get("git") if isinstance(evidence.get("git"), dict) else {}
    ci = checks.get("ci") if isinstance(checks.get("ci"), dict) else {}

    conditions = {
        "package_valid": behavioral.get("package_valid") is True,
        "replay_miss_zero": type(behavioral.get("replay_miss_count")) is int and behavioral.get("replay_miss_count") == 0,
        "replay_network_zero": type(behavioral.get("replay_network_events")) is int and behavioral.get("replay_network_events") == 0,
        "provider_http_zero": type(behavioral.get("provider_http_calls")) is int and behavioral.get("provider_http_calls") == 0,
        "provider_call_zero": type(behavioral.get("provider_call_count")) is int and behavioral.get("provider_call_count") == 0,
        "free_only_config_valid": behavioral.get("free_only_config_valid") is True,
        "all_results_unique_ids_20": type(behavioral.get("all_results_unique_ids")) is int and behavioral.get("all_results_unique_ids") == 20,
        "expected_all_results_order_match": behavioral.get("expected_all_results_order_match") is True,
        "actual_is_subset": behavioral.get("actual_is_subset") is True,
        "validator_pass": behavioral.get("validator_status") == "PASS",
        "issues_empty": behavioral.get("issues") == [],
        "positive_denominators": behavioral.get("positive_denominators") is True,
        "live_replay_metrics_equal": behavioral.get("live_replay_metrics_equal") is True,
        "artifact_hash_equal": behavioral.get("live_artifact_hash")
        and behavioral.get("live_artifact_hash") == behavioral.get("offline_artifact_hash"),
        "receipt_eligible": behavioral.get("receipt_eligible") is True,
    }
    for name, passed in conditions.items():
        if not passed:
            failures.append(f"BEHAVIORAL_{name.upper()}")
    behavioral_recall_validated = all(conditions.values())

    command_results = checks.get("command_results") if isinstance(checks.get("command_results"), dict) else {}
    required_commands_pass = bool(command_results) and all(
        type(value) is dict and type(value.get("returncode")) is int and value.get("returncode") == 0
        for value in command_results.values()
    ) and all(name in command_results for name in ("compileall", "pytest", "benchmark", "pip_check", "help", "diff_check"))
    release_conditions = {
        "behavioral_recall_validated": behavioral_recall_validated,
        "js_default_unset_true": checks.get("js_default_unset_true") is True,
        "browser_smoke": checks.get("browser_smoke") is True,
        "required_commands_pass": required_commands_pass,
        "offline_test_env_used": checks.get("offline_test_env_used") is True,
        "offline_test_env_removed": checks.get("offline_test_env_removed") is True,
        "tests_intact": checks.get("tests_intact") is True,
        "tracked_clean": checks.get("tracked_clean") is True,
        "feature_remote_exact": checks.get("feature_remote_exact") is True,
        "ci_exact_head": ci.get("head_sha") == git.get("head_sha") and ci.get("jobs_success") is True,
        "ci_token_masked": checks.get("ci_token_masked") is True,
    }
    for name, passed in release_conditions.items():
        if not passed:
            failures.append(f"RELEASE_{name.upper()}")
    release_ready = all(release_conditions.values())

    merge_conditions = {
        "release_ready": release_ready,
        "active_branch": git.get("branch") == "codex/release-hardening",
        "origin_feature_exact": git.get("origin_feature_sha") == git.get("head_sha"),
        "gate_base_is_origin_main": git.get("gate_base_sha") == git.get("origin_main_sha"),
        "origin_main_ancestor": git.get("origin_main_ancestor") is True,
        "no_divergence_or_force": git.get("no_divergence_or_force") is True,
    }
    for name, passed in merge_conditions.items():
        if not passed:
            failures.append(f"MERGE_{name.upper()}")
    merge_allowed = all(merge_conditions.values())
    return {
        "behavioral_recall_validated": behavioral_recall_validated,
        "release_ready": release_ready,
        "merge_allowed": merge_allowed,
        "ci_exact_head_success": release_conditions["ci_exact_head"],
        "failure_reasons": sorted(set(failures)),
    }


def _repo_state(repo: Path) -> dict:
    branch_code, branch, _ = _git(repo, "branch", "--show-current")
    head_code, head, _ = _git(repo, "rev-parse", "HEAD")
    base_code, origin_main, _ = _git(repo, "rev-parse", "refs/remotes/origin/main")
    feature_code, origin_feature, _ = _git(repo, "rev-parse", "refs/remotes/origin/codex/release-hardening")
    ancestor = _git(repo, "merge-base", "--is-ancestor", "refs/remotes/origin/main", "HEAD")[0] == 0
    count_code, count, _ = _git(repo, "rev-list", "--left-right", "--count", "refs/remotes/origin/main...HEAD")
    left, right = (count.split() if count_code == 0 else ("-1", "-1"))
    status_code, status, _ = _git(repo, "status", "--porcelain", "--untracked-files=no")
    return {
        "branch": branch if branch_code == 0 else "",
        "head_sha": head if head_code == 0 else "",
        "origin_main_sha": origin_main if base_code == 0 else "",
        "origin_feature_sha": origin_feature if feature_code == 0 else "",
        "gate_base_sha": origin_main if base_code == 0 else "",
        "origin_main_ancestor": ancestor,
        "no_divergence_or_force": count_code == 0 and left == "0" and int(right) >= 0,
        "divergence": {"behind": int(left), "ahead": int(right)},
        "tracked_clean": status_code == 0 and not status,
    }


def _test_integrity(repo: Path) -> bool:
    code, diff, _ = _git(repo, "diff", "HEAD^!", "--", "tests")
    if code != 0:
        return False
    added_bad = any(
        line.startswith("+") and not line.startswith("+++") and re.search(r"skip|xfail", line, re.IGNORECASE)
        for line in diff.splitlines()
    )
    deleted_test = any(
        line.startswith("-") and not line.startswith("---") and re.search(r"def test_", line)
        for line in diff.splitlines()
    )
    return not added_bad and not deleted_test


def _run_command(repo: Path, name: str, command: list[str], env: dict[str, str]) -> dict:
    result = subprocess.run(command, cwd=repo, env=env, capture_output=True, text=True)
    return {
        "command": command,
        "returncode": result.returncode,
        "stdout_sha256": hashlib.sha256(result.stdout.encode("utf-8")).hexdigest(),
        "stderr_sha256": hashlib.sha256(result.stderr.encode("utf-8")).hexdigest(),
    }


def _github_repo(remote: str) -> tuple[str, str] | None:
    value = remote.strip()
    if value.startswith("git@"):
        value = value.split(":", 1)[1]
    elif "://" in value:
        value = urllib.parse.urlparse(value).path.lstrip("/")
    value = value.removesuffix(".git").strip("/")
    parts = value.split("/")
    return (parts[-2], parts[-1]) if len(parts) >= 2 else None


def _ci_auth_token(repo: Path) -> tuple[str, str]:
    """Resolve a GitHub token without ever returning it as evidence."""
    for name in ("GITHUB_TOKEN", "GH_TOKEN"):
        value = os.environ.get(name, "").strip()
        if value:
            return value, name
    command_env = os.environ.copy()
    command_env["GIT_TERMINAL_PROMPT"] = "0"
    try:
        gh = subprocess.run(
            ["gh", "auth", "token"], cwd=repo, env=command_env,
            capture_output=True, text=True, timeout=15,
        )
    except (OSError, subprocess.SubprocessError):
        gh = None
    if gh is not None and gh.returncode == 0 and gh.stdout.strip():
        return gh.stdout.strip(), "gh_auth"
    try:
        credential = subprocess.run(
            ["git", "credential", "fill"], cwd=repo, env=command_env,
            input="protocol=https\nhost=github.com\n\n",
            capture_output=True, text=True, timeout=15,
        )
    except (OSError, subprocess.SubprocessError):
        credential = None
    if credential is not None and credential.returncode == 0:
        fields = {
            line.split("=", 1)[0]: line.split("=", 1)[1]
            for line in credential.stdout.splitlines()
            if "=" in line
        }
        password = str(fields.get("password", "")).strip()
        if password:
            return password, "git_credential"
    return "", ""


def query_github_ci(repo: Path, head_sha: str, branch: str) -> dict:
    _code, remote, _ = _git(repo, "config", "--get", "remote.origin.url")
    identity = _github_repo(remote)
    if not identity:
        return {"jobs_success": False, "failure": "CI_REMOTE_NOT_GITHUB"}
    token, auth_method = _ci_auth_token(repo)
    if not token:
        return {"jobs_success": False, "failure": "CI_API_AUTH_UNAVAILABLE"}
    owner, repo_name = identity
    headers = {
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
        "Authorization": f"Bearer {token}",
        "User-Agent": "golden6-release-gate",
    }

    def get(url: str) -> dict:
        request = urllib.request.Request(url, headers=headers, method="GET")
        with urllib.request.urlopen(request, timeout=30) as response:
            value = json.loads(response.read().decode("utf-8"))
        if not isinstance(value, dict):
            raise ValueError("github_response_not_object")
        return value

    try:
        runs_url = f"https://api.github.com/repos/{urllib.parse.quote(owner)}/{urllib.parse.quote(repo_name)}/actions/runs?head={urllib.parse.quote(branch)}&per_page=50"
        runs = get(runs_url).get("workflow_runs", [])
        selected = next((run for run in runs if run.get("head_sha") == head_sha), None)
        if not selected:
            return {"jobs_success": False, "failure": "CI_EXACT_HEAD_NOT_FOUND", "head_sha": ""}
        jobs_url = f"https://api.github.com/repos/{urllib.parse.quote(owner)}/{urllib.parse.quote(repo_name)}/actions/runs/{selected['id']}/jobs?per_page=100"
        jobs = get(jobs_url).get("jobs", [])
        job_map = {str(job.get("name")): job for job in jobs}
        success = all(
            job_map.get(name, {}).get("conclusion") == "success"
            for name in REQUIRED_CI_JOBS
        )
        return {
            "run_id": selected.get("id"),
            "run_name": selected.get("name"),
            "head_sha": selected.get("head_sha"),
            "status": selected.get("status"),
            "conclusion": selected.get("conclusion"),
            "jobs": {name: {"id": job_map.get(name, {}).get("id"), "conclusion": job_map.get(name, {}).get("conclusion")} for name in REQUIRED_CI_JOBS},
            "jobs_success": success and selected.get("status") == "completed" and selected.get("conclusion") == "success",
            "repository": f"{owner}/{repo_name}",
            "token_used_for_auth_only": True,
            "auth_method": auth_method,
        }
    except Exception as exc:
        return {"jobs_success": False, "failure": f"CI_API_ERROR:{type(exc).__name__}"}


def evaluate(repo_root: Path, package_dir: Path, replay_receipt_path: Path, report_path: Path) -> dict:
    repo = Path(repo_root).resolve()
    package_dir = Path(package_dir).resolve()
    report_path = Path(report_path).resolve()
    evidence: dict = {"behavioral": {}, "checks": {}, "git": {}, "artifacts": {}, "metrics": {}}
    failures: list[str] = []
    live_artifact_hash = ""
    offline_artifact_hash = ""
    offline_artifact_hash_value = ""
    provider_call_count = -1
    live_artifact: Path | None = None
    offline_artifact: Path | None = None
    live_telemetry: dict = {"valid": False, "count": None, "failure": "live_artifact_missing"}
    offline_telemetry: dict = {"valid": False, "count": None, "failure": "offline_artifact_missing"}
    offline_coverage: dict = {"valid": False, "count": None, "failure": "offline_artifact_missing"}
    live_telemetry_chain_valid = False
    offline_telemetry_chain_valid = False
    try:
        package_manifest_path = package_dir / "package_manifest.json"
        package_manifest = _json(package_manifest_path)
        package_valid, package_failures = _package_files_valid(package_dir, package_manifest)
        failures.extend(package_failures)
        package_free_only_valid = True
        try:
            validate_free_only_manifest(package_manifest, require_complete=False)
            validate_manifest_config_sha(package_manifest)
        except Exception as exc:
            package_free_only_valid = False
            failures.append(f"FREE_ONLY_PACKAGE:{exc}")
        package_input = package_manifest.get("input") if isinstance(package_manifest.get("input"), dict) else {}
        input_path_value = package_input.get("path")
        input_path = None
        if isinstance(input_path_value, str) and input_path_value:
            input_path = Path(input_path_value)
            if not input_path.is_absolute():
                input_path = package_manifest_path.parent / input_path
            input_path = input_path.resolve()
        package_input_hash = package_input.get("sha256") if isinstance(package_input.get("sha256"), str) else ""
        if input_path is None or not input_path.is_file():
            package_valid = False
            failures.append("package_input_missing")
        elif _sha256(input_path) != package_input_hash:
            package_valid = False
            failures.append("package_input_hash_mismatch")
        expected_path = repo / "outputs" / "golden_6_20260718" / "golden_6_manual_validation_20_ready.xlsx"
        expected_ids = package_manifest.get("ordered_source_record_ids") if type(package_manifest.get("ordered_source_record_ids")) is list else []
        receipt_path = Path(replay_receipt_path).resolve()
        receipt = _json(receipt_path)
        offline_manifest_path = None
        offline_manifest = None
        receipt_run_dir = receipt.get("run_dir")
        if isinstance(receipt_run_dir, str) and receipt_run_dir:
            candidate_manifest_path = Path(receipt_run_dir).resolve() / "manifest.json"
            if candidate_manifest_path.is_file():
                try:
                    offline_manifest_path = candidate_manifest_path
                    offline_manifest = _json(candidate_manifest_path)
                except (OSError, ValueError, json.JSONDecodeError):
                    offline_manifest_path = None
                    offline_manifest = None
        receipt_shape_and_identity_valid, receipt_failures = receipt_eligibility(
            receipt,
            package_manifest=package_manifest,
            package_manifest_path=package_manifest_path,
            offline_manifest=offline_manifest,
            offline_manifest_path=offline_manifest_path,
        )
        failures.extend(receipt_failures)
        successful_receipt_invariants_valid = receipt_shape_and_identity_valid
        package_live_run = package_manifest.get("live_run")
        live_root = (
            Path(package_live_run.get("run_dir")).resolve()
            if type(package_live_run) is dict and type(package_live_run.get("run_dir")) is str and package_live_run.get("run_dir")
            else None
        )
        offline_root = offline_manifest_path.parent.resolve() if offline_manifest_path is not None else None
        live_bundle_valid = False
        live_database_valid = False
        offline_bundle_valid = False
        offline_database_valid = False
        offline_config_parity_valid = False
        offline_config_info = {
            "threshold_mismatch": False,
            "threshold_mismatch_details": [],
            "config_mismatch": True,
            "config_mismatch_details": [],
        }
        live_config = {}
        try:
            if live_root is None:
                raise ValueError("live_run_dir_invalid")
            live_manifest = _json(live_root / "manifest.json")
            validate_free_only_manifest(live_manifest)
            validate_manifest_config_sha(live_manifest)
            live_artifact_hash_value = live_manifest.get("artifact_set_sha256")
            if type(live_artifact_hash_value) is not str or not re.fullmatch(r"[0-9a-f]{64}", live_artifact_hash_value):
                raise ValueError("live_artifact_set_invalid")
            live_artifact = (live_root / "output" / "artifacts" / live_artifact_hash_value).resolve()
            if (
                type(package_live_run) is not dict
                or type(package_manifest.get("run_config")) is not dict
                or type(package_manifest.get("config_sha256")) is not str
                or type(package_live_run.get("config_sha256")) is not str
                or type(live_manifest.get("config_sha256")) is not str
                or package_live_run.get("config_sha256") != package_manifest.get("config_sha256")
                or live_manifest.get("config_sha256") != package_live_run.get("config_sha256")
                or package_live_run.get("run_id") != live_manifest.get("run_id")
                or package_live_run.get("manifest_sha256") != _sha256(live_root / "manifest.json")
                or package_live_run.get("artifact_set_sha256") != live_manifest.get("artifact_set_sha256")
                or type(package_live_run.get("artifact_dir")) is not str
                or Path(package_live_run.get("artifact_dir")).resolve() != live_artifact
            ):
                raise ValueError("live_package_config_link_mismatch")
            live_package_config_info = compare_run_configs(
                package_manifest.get("run_config"), live_manifest.get("run_config"),
            )
            if live_package_config_info["threshold_mismatch"] or live_package_config_info["config_mismatch"]:
                raise ValueError("live_package_config_parity_mismatch")
            run_context.validate_run_bundle(
                live_root,
                expected_input_hash=package_input_hash,
                expected_config_hash=str(package_manifest.get("live_run", {}).get("config_sha256") or ""),
                expected_source_ids=expected_ids,
                require_artifacts=True,
                profile="COMPLETE",
            )
            live_bundle_valid = True
            live_config = live_manifest.get("run_config", {})
            live_artifact_hash = live_artifact_hash_value
            if str(live_manifest.get("config_sha256", "")) != str(package_manifest.get("live_run", {}).get("config_sha256", "")):
                raise ValueError("live_config_hash_mismatch")
            validate_free_only_database(live_root / "state" / "progress.sqlite3", str(live_manifest.get("run_id", "")), 20)
            live_database_valid = True
        except Exception as exc:
            failures.append(f"FREE_ONLY_RUN:{exc}")
        if live_root is not None and isinstance(live_manifest, dict):
            try:
                checkpoint.validate_external_finalization_telemetry(
                    live_root / "state" / "progress.sqlite3", live_root,
                    str(live_manifest.get("run_id", "")), manifest=live_manifest,
                    require_free_only=True,
                )
                live_telemetry_chain_valid = True
            except Exception as exc:
                failures.append(_telemetry_failure_reason("LIVE_TELEMETRY_REPLICA", exc))
        if offline_root is not None:
            try:
                offline_manifest = _json(offline_root / "manifest.json")
                validate_free_only_manifest(offline_manifest)
                validate_manifest_config_sha(offline_manifest)
                offline_artifact_hash_value = offline_manifest.get("artifact_set_sha256")
                if type(offline_artifact_hash_value) is not str or not re.fullmatch(r"[0-9a-f]{64}", offline_artifact_hash_value):
                    raise ValueError("offline_artifact_set_invalid")
                offline_artifact = (offline_root / "output" / "artifacts" / offline_artifact_hash_value).resolve()
                run_context.validate_run_bundle(
                    offline_root,
                    expected_input_hash=package_input_hash,
                    expected_config_hash=str(offline_manifest.get("config_sha256") or ""),
                    expected_source_ids=expected_ids,
                    require_artifacts=True,
                    profile="COMPLETE",
                )
                offline_bundle_valid = True
            except Exception as exc:
                failures.append(f"FREE_ONLY_OFFLINE_RUN:{exc}")
            try:
                validate_free_only_database(
                    offline_root / "state" / "progress.sqlite3",
                    str(offline_manifest.get("run_id", "")), 20,
                )
                offline_database_valid = True
            except Exception as exc:
                failures.append(f"FREE_ONLY_OFFLINE_DATABASE:{exc}")
            if isinstance(offline_manifest, dict):
                try:
                    checkpoint.validate_external_finalization_telemetry(
                        offline_root / "state" / "progress.sqlite3", offline_root,
                        str(offline_manifest.get("run_id", "")), manifest=offline_manifest,
                        require_free_only=True,
                    )
                    offline_telemetry_chain_valid = True
                except Exception as exc:
                    failures.append(_telemetry_failure_reason("OFFLINE_TELEMETRY_REPLICA", exc))
            if live_bundle_valid and isinstance(offline_manifest, dict):
                try:
                    offline_config_info = compare_run_configs(
                        expected_offline_run_config(live_config),
                        offline_manifest.get("run_config"),
                    )
                    offline_config_parity_valid = offline_config_info["config_mismatch"] is False
                    if not offline_config_parity_valid:
                        failures.append("FREE_ONLY_OFFLINE_CONFIG_MISMATCH")
                except Exception as exc:
                    failures.append(f"FREE_ONLY_OFFLINE_CONFIG:{exc}")
        if live_telemetry_chain_valid:
            live_telemetry = {"valid": True, "count": 0, "failure": ""}
        if offline_telemetry_chain_valid:
            offline_telemetry = {"valid": True, "count": 0, "failure": ""}
        if offline_artifact is not None and offline_artifact.is_dir():
            offline_coverage = _replay_coverage(offline_artifact)
        live_telemetry_valid = _telemetry_zero(live_telemetry)
        offline_telemetry_valid = _telemetry_zero(offline_telemetry)
        offline_replay_coverage_valid = _coverage_zero(offline_coverage)
        if not live_telemetry_valid:
            failures.append("LIVE_PROVIDER_TELEMETRY_INVALID")
        if not offline_telemetry_valid:
            failures.append("OFFLINE_PROVIDER_TELEMETRY_INVALID")
        if not offline_replay_coverage_valid:
            failures.append("OFFLINE_REPLAY_COVERAGE_INVALID")
        successful_receipt_invariants_valid = (
            receipt_shape_and_identity_valid
            and live_database_valid
            and offline_config_parity_valid
            and offline_config_info["threshold_mismatch"] is False
            and live_telemetry_valid
            and offline_telemetry_valid
            and offline_replay_coverage_valid
        )
        receipt_eligible = all((
            receipt_shape_and_identity_valid,
            package_valid,
            package_free_only_valid,
            live_bundle_valid,
            offline_bundle_valid,
            offline_database_valid,
            successful_receipt_invariants_valid,
        ))
        if receipt_eligible:
            offline_artifact_hash = offline_artifact_hash_value
            provider_call_count = receipt.get("paid_provider_calls") if type(receipt.get("paid_provider_calls")) is int else -1
        free_only_config_valid = all((
            package_free_only_valid,
            live_bundle_valid,
            offline_bundle_valid,
            offline_database_valid,
            offline_config_parity_valid,
            offline_config_info["threshold_mismatch"] is False,
        ))
        if live_artifact is not None:
            live_report, live_failures = _validator_report(repo, expected_path, live_artifact, report_path.with_name(report_path.stem + ".live.validator.json"))
        else:
            live_report, live_failures = {}, ["live_artifact_missing"]
        if receipt_eligible and offline_artifact is not None:
            offline_report, offline_failures = _validator_report(repo, expected_path, offline_artifact, report_path.with_name(report_path.stem + ".offline.validator.json"))
        else:
            offline_report, offline_failures = {}, []
        failures.extend(live_failures + offline_failures)
        live_ids = _ids(live_artifact / "all_results.xlsx") if live_artifact is not None and live_artifact.is_dir() else []
        behavioral_offline_artifact = offline_artifact if receipt_eligible else None
        offline_ids = _ids(behavioral_offline_artifact / "all_results.xlsx") if behavioral_offline_artifact is not None and behavioral_offline_artifact.is_dir() else []
        live_contacts = _ids(live_artifact / "contacts.xlsx") if live_artifact is not None and live_artifact.is_dir() else []
        offline_contacts = _ids(behavioral_offline_artifact / "contacts.xlsx") if behavioral_offline_artifact is not None and behavioral_offline_artifact.is_dir() else []
        live_metrics = _metrics(live_report)
        offline_metrics = _metrics(offline_report) if receipt_eligible else {}
        live_telemetry_calls = live_telemetry.get("count") if live_telemetry_valid else -1
        offline_telemetry_calls = offline_telemetry.get("count") if receipt_eligible and offline_telemetry_valid else -1
        behavioral = {
            "package_valid": package_valid,
            "replay_miss_count": offline_coverage.get("count") if receipt_eligible and offline_replay_coverage_valid else -1,
            "replay_network_events": receipt.get("replay_network_events", -1) if receipt_eligible else -1,
            "provider_http_calls": offline_telemetry_calls,
            "provider_call_count": provider_call_count if receipt_eligible else -1,
            "live_provider_http_calls": live_telemetry_calls,
            "free_only_config_valid": free_only_config_valid,
            "all_results_unique_ids": len(live_ids) if len(live_ids) == len(set(live_ids)) else -1,
            "expected_all_results_order_match": live_ids == expected_ids and offline_ids == expected_ids,
            "actual_is_subset": set(live_contacts).issubset(set(live_ids)) and set(offline_contacts).issubset(set(offline_ids)),
            "validator_status": "PASS" if live_report.get("status") == "PASS" and offline_report.get("status") == "PASS" else "FAIL",
            "issues": list(live_report.get("issues", [])) + list(offline_report.get("issues", [])),
            "positive_denominators": _positive_denominators(live_report) and _positive_denominators(offline_report),
            "live_replay_metrics_equal": live_metrics == offline_metrics,
            "live_artifact_hash": live_artifact_hash,
            "offline_artifact_hash": offline_artifact_hash if receipt_eligible else "",
            "receipt_eligible": receipt_eligible,
            "receipt_rejection_reasons": receipt_failures,
        }
        git = _repo_state(repo)
        command_env = os.environ.copy()
        command_env["B2B_TEST_OFFLINE"] = "1"
        command_results = {
            "compileall": _run_command(repo, "compileall", [sys.executable, "-m", "compileall", "-q", "config.py", "main.py", "modules", "tools", "tests"], command_env),
            "pytest": _run_command(repo, "pytest", [sys.executable, "-m", "pytest", "-q", "-p", "no:cacheprovider"], command_env),
            "benchmark": _run_command(repo, "benchmark", [sys.executable, "validate_benchmark_suite.py"], command_env),
            "pip_check": _run_command(repo, "pip_check", [sys.executable, "-m", "pip", "check"], command_env),
            "help": _run_command(repo, "help", [sys.executable, "main.py", "--help"], command_env),
            "diff_check": _run_command(repo, "diff_check", ["git", "diff", "--check"], command_env),
        }
        clean_env = os.environ.copy()
        clean_env.pop("B2B_TEST_OFFLINE", None)
        clean_env.pop("ENABLE_JS_FALLBACK", None)
        js_result = subprocess.run([sys.executable, "-c", "import config; assert config.ENABLE_JS_FALLBACK is True"], cwd=repo, env=clean_env, capture_output=True, text=True)
        browser_result = subprocess.run([sys.executable, "-c", "from playwright.sync_api import sync_playwright; p=sync_playwright().start(); b=p.chromium.launch(); page=b.new_page(); page.set_content('<h1>Test</h1>'); assert 'Test' in page.content(); b.close(); p.stop()"], cwd=repo, env=clean_env, capture_output=True, text=True)
        ci = query_github_ci(repo, git.get("head_sha", ""), git.get("branch", ""))
        evidence = {
            "behavioral": behavioral,
            "checks": {
                "js_default_unset_true": js_result.returncode == 0,
                "browser_smoke": browser_result.returncode == 0,
                "command_results": command_results,
                "offline_test_env_used": True,
                "offline_test_env_removed": "B2B_TEST_OFFLINE" not in clean_env,
                "tests_intact": _test_integrity(repo),
                "tracked_clean": git.get("tracked_clean") is True,
                "feature_remote_exact": git.get("origin_feature_sha") == git.get("head_sha"),
                "ci_token_masked": ci.get("token_used_for_auth_only") is True,
                "ci": ci,
            },
            "git": git,
            "artifacts": {
                "package_manifest": {"path": str(package_manifest_path), "sha256": _sha256(package_manifest_path)},
                **({"live_artifact_dir": str(live_artifact)} if live_artifact is not None else {}),
                **({"offline_artifact_dir": str(behavioral_offline_artifact)} if behavioral_offline_artifact is not None else {}),
            },
            "metrics": {"live": live_metrics, "offline": offline_metrics},
        }
        chain_failures = _validate_receipt_chain(
            receipt, package_manifest, package_manifest_path, offline_manifest, offline_manifest_path,
        )
        threshold_candidate = _recomputed_threshold_candidate(
            receipt,
            package_manifest,
            package_manifest_path,
            offline_manifest,
            offline_manifest_path,
            package_valid=package_valid,
            package_free_only_valid=package_free_only_valid,
            live_bundle_valid=live_bundle_valid,
            live_database_valid=live_database_valid,
            offline_bundle_valid=offline_bundle_valid,
            offline_database_valid=offline_database_valid,
            live_telemetry_valid=live_telemetry_valid,
            offline_telemetry_valid=offline_telemetry_valid,
            offline_replay_coverage_valid=offline_replay_coverage_valid,
        ) if not chain_failures else False
        receipt_artifact_name = _receipt_artifact_name(
            receipt,
            receipt_eligible,
            threshold_only=not chain_failures,
            threshold_candidate=threshold_candidate,
        )
        evidence["artifacts"][receipt_artifact_name] = {
            "path": str(receipt_path), "sha256": _sha256(receipt_path),
        }
    except Exception as exc:
        failures.append(f"EVALUATOR_ERROR:{type(exc).__name__}:{exc}")
        evidence["behavioral"] = {"package_valid": False}
    computed = evaluate_evidence(evidence)
    computed["failure_reasons"] = sorted(set(computed["failure_reasons"] + failures))
    report = {
        "schema_version": 1,
        "behavioral_recall_validated": computed["behavioral_recall_validated"],
        "release_ready": computed["release_ready"],
        "merge_allowed": computed["merge_allowed"],
        "ci_exact_head_success": computed["ci_exact_head_success"],
        "head_sha": evidence.get("git", {}).get("head_sha", ""),
        "base_sha": evidence.get("git", {}).get("gate_base_sha", ""),
        "origin_main_sha": evidence.get("git", {}).get("origin_main_sha", ""),
        "source_record_count": _json(package_dir / "package_manifest.json").get("input", {}).get("record_count", 0) if (package_dir / "package_manifest.json").is_file() else 0,
        "source_record_id_sha256": _json(package_dir / "package_manifest.json").get("input", {}).get("ordered_id_sha256", "") if (package_dir / "package_manifest.json").is_file() else "",
        "replay_miss_count": evidence.get("behavioral", {}).get("replay_miss_count"),
        "replay_network_events": evidence.get("behavioral", {}).get("replay_network_events"),
        "provider_call_count": evidence.get("behavioral", {}).get("provider_call_count"),
        "live_artifact_hash": evidence.get("behavioral", {}).get("live_artifact_hash"),
        "offline_artifact_hash": evidence.get("behavioral", {}).get("offline_artifact_hash"),
        "validator_status": evidence.get("behavioral", {}).get("validator_status", "FAIL"),
        "publication_precision": evidence.get("metrics", {}).get("live", {}).get("stages", {}).get("publication_precision"),
        "publication_recall": evidence.get("metrics", {}).get("live", {}).get("stages", {}).get("publication_recall"),
        "ci": evidence.get("checks", {}).get("ci", {}),
        "artifacts": evidence.get("artifacts", {}),
        "metrics": evidence.get("metrics", {}),
        "failure_reasons": computed["failure_reasons"],
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = report_path.with_name(f".{report_path.name}.{os.getpid()}.tmp")
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(report_path)
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Evaluate the evidence-backed Golden6 release gate.")
    parser.add_argument("--repo-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--package", type=Path, required=True)
    parser.add_argument("--replay-receipt", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        report = evaluate(args.repo_root, args.package, args.replay_receipt, args.report)
    except Exception as exc:
        print(f"RELEASE_GATE_BLOCKED:{exc}", file=sys.stderr)
        return 2
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["release_ready"] and report["merge_allowed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())

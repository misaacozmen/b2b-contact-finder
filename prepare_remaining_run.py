"""Build the immutable, offline 159-row continuation input.

The source recovery database is opened read-only and immutable.  This module
never reads recovery result payloads; source values come only from firms.xlsx.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shlex
import sqlite3
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook

from modules import excel, run_context


_SCRIPT_ROOT = Path(__file__).resolve().parent
_PRODUCTION_ROOT = _SCRIPT_ROOT
if not (_PRODUCTION_ROOT / "input" / "firms.xlsx").exists():
    _PRODUCTION_ROOT = _SCRIPT_ROOT.parent / "Python b2b"
_RECOVERY_ROOT = Path(os.getenv(
    "B2B_FINAL10_R2_ROOT",
    str(_SCRIPT_ROOT.parent / "Python b2b_recovery_final10_r2"),
))
_RECOVERY_RUN = _RECOVERY_ROOT / "runs" / "5ab042f129d9d31108cccc8cecb8549c32f0965c4de418283256284b44f2a36b"
PARENT_MANIFEST = _RECOVERY_RUN / "manifest.json"
RECOVERY_DB = _RECOVERY_RUN / "output" / "artifacts" / "22a2993366d4bbbdebdf5d9ee805b59bc8949506f0f8b2d4edefdc964b4df008" / "recovery_state.sqlite3"
SOURCE_MAPPING = RECOVERY_DB.with_name("source_id_mapping.json")
ORIGINAL_INPUT = _PRODUCTION_ROOT / "input" / "firms.xlsx"
PAID_PROVIDERS = ("brightdata", "google_places", "brandfetch", "hunter", "linkedin", "llm")
FORBIDDEN_RESULT_FIELDS = frozenset({
    "email", "phone", "website", "website_source", "email_source", "phone_source",
    "selected_website", "candidate_1_url", "candidate_2_url", "candidate_3_url",
})


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _normalize_xlsx(path: Path) -> None:
    temp = path.with_name(f".{path.name}.normalized")
    with ZipFile(path, "r") as source, ZipFile(temp, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            data = source.read(name)
            if name == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>", data,
                )
            target.writestr(info, data)
    temp.replace(path)


def _write_remaining_workbook(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    workbook = Workbook()
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    workbook.properties.lastModifiedBy = "B2B Contact Finder"
    sheet = workbook.active
    sheet.title = "Sheet"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 12), 60)
    workbook.save(path)
    workbook.close()
    _normalize_xlsx(path)


def _read_source(path: Path) -> tuple[list[str], list[list[Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not values:
        raise ValueError("original workbook is empty")
    headers = [str(value or "").strip() for value in values[0]]
    if "company" not in {value.casefold() for value in headers}:
        raise ValueError("original workbook has no company column")
    return headers, values[1:]


def _load_selection(
    *,
    original_input: Path | None = None,
    recovery_db: Path | None = None,
    parent_manifest: Path | None = None,
    source_mapping: Path | None = None,
) -> tuple[dict, dict[str, Any], str, str]:
    original_input = Path(original_input or ORIGINAL_INPUT).resolve()
    recovery_db = Path(recovery_db or RECOVERY_DB).resolve()
    parent_manifest = Path(parent_manifest or PARENT_MANIFEST).resolve()
    source_mapping = Path(source_mapping or SOURCE_MAPPING).resolve()
    manifest = json.loads(parent_manifest.read_text(encoding="utf-8"))
    if manifest.get("run_id") != parent_manifest.parent.name:
        raise ValueError("parent manifest is not the expected final10-r2 manifest")
    if _sha256(original_input) != str(manifest.get("input_sha256", "")):
        raise ValueError("original input hash does not match parent manifest")
    mapping = json.loads(source_mapping.read_text(encoding="utf-8"))
    if _sha256(recovery_db) != str(manifest.get("checkpoint_sha256", "")):
        raise ValueError("recovery database hash does not match parent manifest checkpoint")
    mapping_without_digest = {key: value for key, value in mapping.items() if key != "mapping_sha256"}
    if hashlib.sha256(_canonical(mapping_without_digest).encode("utf-8")).hexdigest() != str(mapping.get("mapping_sha256", "")):
        raise ValueError("source mapping digest mismatch")
    mappings = {int(row["legacy_index"]): row for row in mapping.get("mappings", [])}
    if len(mappings) != len(mapping.get("mappings", [])):
        raise ValueError("duplicate legacy index in source mapping")
    headers, source_rows = _read_source(original_input)
    input_records = excel.read_company_records(original_input)
    if len(source_rows) != len(mappings):
        raise ValueError("source workbook and mapping row counts differ")
    db_uri = f"file:{recovery_db.as_posix()}?mode=ro&immutable=1"
    connection = sqlite3.connect(db_uri, uri=True)
    try:
        if connection.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise ValueError("recovery database integrity_check failed")
        selected = connection.execute(
            "SELECT item_index,source_record_id,free_state,paid_required,paid_state FROM run_items "
            "WHERE free_state='PENDING' OR (paid_required=1 AND paid_state='PENDING') ORDER BY item_index"
        ).fetchall()
    finally:
        connection.close()
    if len(selected) != 159:
        raise ValueError(f"selection count is {len(selected)}, expected 159")
    selected_ids = [str(row[1]) for row in selected]
    if len(selected_ids) != len(set(selected_ids)):
        raise ValueError("selected source IDs are not unique")
    free = [row for row in selected if row[2] == "PENDING"]
    paid = [row for row in selected if int(row[3]) == 1 and row[4] == "PENDING"]
    if len(free) != 2 or len(paid) != 157 or {str(row[1]) for row in free} & {str(row[1]) for row in paid}:
        raise ValueError("selection does not contain the exact disjoint 2 free + 157 paid sets")
    output_rows: list[list[Any]] = []
    mapping_rows: list[dict[str, Any]] = []
    id_index: list[dict[str, Any]] = []
    for item_index, db_source_id, free_state, paid_required, paid_state in selected:
        item_index = int(item_index)
        mapped = mappings.get(item_index)
        if not mapped or not (0 <= item_index < len(source_rows)):
            raise ValueError(f"missing source mapping for index {item_index}")
        source_row = source_rows[item_index]
        source_record = dict(zip(headers, source_row + [None] * (len(headers) - len(source_row))))
        source_record = {str(key): value for key, value in source_record.items()}
        input_row_hash = hashlib.sha256(run_context.canonical_json(input_records[item_index]).encode("utf-8")).hexdigest()
        if input_row_hash != str(mapped.get("input_row_sha256", "")):
            raise ValueError(f"source row hash mismatch at index {item_index}")
        canonical_id = str(mapped.get("canonical_source_record_id", ""))
        if str(db_source_id) != canonical_id:
            raise ValueError(f"source ID mapping mismatch at index {item_index}")
        output_rows.append(list(source_row) + [canonical_id])
        mapping_rows.append({"item_index": item_index, "source_record_id": canonical_id, "free_state": str(free_state), "paid_required": int(paid_required), "paid_state": str(paid_state)})
        id_index.append({"item_index": item_index, "source_record_id": canonical_id})
    output_headers = list(headers) + ["source_record_id"]
    leaked = FORBIDDEN_RESULT_FIELDS & {header.casefold() for header in output_headers}
    if leaked - {"website", "email", "phone"}:
        raise ValueError(f"result fields leaked into source package: {sorted(leaked)}")
    runtime_tree_hash = run_context.source_tree_sha256()
    return manifest, {"headers": output_headers, "rows": output_rows, "mapping_rows": mapping_rows, "id_index": id_index}, runtime_tree_hash, _sha256(source_mapping)


def _command_tokens(*, input_path: Path, run_root: Path) -> list[str]:
    return [
        str(Path(sys.executable).resolve()), str(Path(__file__).resolve().with_name("main.py")),
        "--input", str(input_path.resolve()), "--run-dir", str(run_root.resolve()),
        "--no-allow-paid", "--search-cache", "use", "--crawl-cache", "use",
        "--brightdata-budget", "0", "--google-places-budget", "0",
        "--brandfetch-budget", "0", "--hunter-budget", "0",
        "--linkedin-company-budget", "0", "--llm-budget", "0",
        "--non-interactive",
    ]


def _command_text(tokens: list[str]) -> str:
    return subprocess.list2cmdline(tokens)


def _parse_exact_command(command: str) -> tuple[list[str], dict[str, Any]]:
    tokens = shlex.split(command, posix=False)
    if tokens and tokens[0] == "&":
        tokens = tokens[1:]
    if not tokens:
        raise ValueError("exact command is empty")
    if len(tokens) < 2:
        raise ValueError("exact command is incomplete")
    tokens = [token[1:-1] if len(token) >= 2 and token[0] == token[-1] == '"' else token for token in tokens]
    if Path(tokens[1]).name.casefold() != "main.py":
        raise ValueError("exact command does not target main.py")
    import main

    args = main.parse_args(tokens[2:])
    resolved = main.resolve_cli_run_config(tokens[2:])
    return tokens, {"args": args, "effective_config": resolved.as_dict()}


def validate_remaining_plan(workbook_path: Path, plan_path: Path) -> dict[str, Any]:
    """Fail closed on count, order, hash, command, or run-ID tampering."""
    plan = json.loads(Path(plan_path).read_text(encoding="utf-8"))
    recorded_plan_hash = str(plan.get("plan_payload_sha256", ""))
    unsigned_plan = dict(plan)
    unsigned_plan.pop("plan_payload_sha256", None)
    if hashlib.sha256(_canonical(unsigned_plan).encode("utf-8")).hexdigest() != recorded_plan_hash:
        raise ValueError("remaining plan payload hash mismatch")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        values = [list(row) for row in workbook.active.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if len(values) != 160 or len(values[0]) != int(plan["workbook"]["column_count"]):
        raise ValueError("remaining workbook count or column tamper")
    rows = values[1:]
    if len(rows) != int(plan["selection"]["count"]) != 159:
        raise ValueError("remaining selection count tamper")
    if [hashlib.sha256(_canonical(row).encode("utf-8")).hexdigest() for row in rows] != plan["selection"]["row_sha256"]:
        raise ValueError("remaining workbook row reorder or content tamper")
    ordered = plan["selection"]["ordered_index_id"]
    ids = [str(item["source_record_id"]) for item in ordered]
    if len(ids) != len(set(ids)) or hashlib.sha256(_canonical(ordered).encode("utf-8")).hexdigest() != plan["selection"]["ordered_index_id_sha256"]:
        raise ValueError("remaining selection order or duplicate ID tamper")
    workbook_hash = _sha256(Path(workbook_path))
    if workbook_hash != plan["workbook"]["sha256"]:
        raise ValueError("remaining workbook byte tamper")
    parsed_tokens, parsed = _parse_exact_command(str(plan["exact_command"]))
    powershell_tokens, powershell_parsed = _parse_exact_command(str(plan["powershell_command"]))
    if powershell_tokens != parsed_tokens or _canonical(powershell_parsed["effective_config"]).encode("utf-8") != _canonical(parsed["effective_config"]).encode("utf-8"):
        raise ValueError("PowerShell command does not reconcile with exact command")
    effective = parsed["effective_config"]
    if _canonical(effective).encode("utf-8") != _canonical(plan["effective_config"]).encode("utf-8"):
        raise ValueError("remaining effective config does not match exact command")
    runtime_hash = str(plan["runtime_source_tree_sha256"])
    run_id = run_context.canonical_run_id(
        input_sha256=workbook_hash,
        ordered_source_record_ids=ids,
        effective_config=effective,
        runtime_source_tree_hash=runtime_hash,
    )
    if run_id != str(plan["expected_run_id"]) or run_id != str(plan["exact_command_run_id"]):
        raise ValueError("remaining exact command run ID mismatch")
    if "<" in str(plan["exact_command"]) or ">" in str(plan["exact_command"]):
        raise ValueError("remaining exact command contains a placeholder")
    if Path(parsed_tokens[1]).name.casefold() != "main.py":
        raise ValueError("remaining exact command is not main.py")
    return plan


def prepare_remaining_run(
    destination: Path,
    *,
    original_input: Path | None = None,
    recovery_db: Path | None = None,
    parent_manifest: Path | None = None,
    source_mapping: Path | None = None,
) -> dict[str, Any]:
    destination = Path(destination).resolve()
    original_input = Path(original_input or ORIGINAL_INPUT).resolve()
    recovery_db = Path(recovery_db or RECOVERY_DB).resolve()
    parent_manifest = Path(parent_manifest or PARENT_MANIFEST).resolve()
    source_mapping = Path(source_mapping or SOURCE_MAPPING).resolve()
    output_dir = destination / "input"
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest, bundle, runtime_tree_hash, mapping_hash = _load_selection(
        original_input=original_input,
        recovery_db=recovery_db,
        parent_manifest=parent_manifest,
        source_mapping=source_mapping,
    )
    workbook_temp = output_dir / f".remaining_159_fresh.{os.getpid()}.xlsx"
    plan_temp = output_dir / f".remaining_159_plan.{os.getpid()}.json"
    workbook_path = output_dir / "remaining_159_fresh.xlsx"
    plan_path = output_dir / "remaining_159_plan.json"
    created_workbook = False
    try:
        _write_remaining_workbook(workbook_temp, bundle["headers"], bundle["rows"])
        workbook_hash = _sha256(workbook_temp)
        workbook_bytes = workbook_temp.stat().st_size
        if workbook_path.exists() and _sha256(workbook_path) != workbook_hash:
            raise FileExistsError("remaining_159_fresh.xlsx exists with a different hash")
        if not workbook_path.exists():
            workbook_temp.replace(workbook_path)
            created_workbook = True
        else:
            workbook_temp.unlink(missing_ok=True)
        provisional_tokens = _command_tokens(input_path=workbook_path, run_root=destination / "runs" / "pending")
        provisional_config = _parse_exact_command(_command_text(provisional_tokens))[1]["effective_config"]
        expected_run_id = run_context.canonical_run_id(
            input_sha256=workbook_hash,
            ordered_source_record_ids=[row["source_record_id"] for row in bundle["id_index"]],
            effective_config=provisional_config, runtime_source_tree_hash=runtime_tree_hash,
        )
        exact_tokens = _command_tokens(input_path=workbook_path, run_root=destination / "runs" / expected_run_id)
        exact_command = _command_text(exact_tokens)
        parsed_tokens, parsed = _parse_exact_command(exact_command)
        effective_config = parsed["effective_config"]
        reparsed_run_id = run_context.canonical_run_id(
            input_sha256=workbook_hash,
            ordered_source_record_ids=[row["source_record_id"] for row in bundle["id_index"]],
            effective_config=effective_config, runtime_source_tree_hash=runtime_tree_hash,
        )
        if (
            _canonical(effective_config).encode("utf-8") != _canonical(provisional_config).encode("utf-8")
            or reparsed_run_id != expected_run_id
            or Path(parsed_tokens[1]).resolve() != Path(__file__).resolve().with_name("main.py")
        ):
            raise RuntimeError("exact CLI command/config/run ID reconciliation failed")
        row_hashes = [hashlib.sha256(_canonical(row).encode("utf-8")).hexdigest() for row in bundle["rows"]]
        selected_digest = hashlib.sha256(_canonical(bundle["id_index"]).encode("utf-8")).hexdigest()
        plan = {
            "schema_version": 1,
            "predicate": "free_state='PENDING' OR (paid_required=1 AND paid_state='PENDING')",
            "parent": {"manifest": str(parent_manifest), "manifest_sha256": _sha256(parent_manifest), "run_id": manifest["run_id"], "artifact_set_sha256": manifest["artifact_set_sha256"], "recovery_db_sha256": _sha256(recovery_db)},
            "sources": {"original_input": str(original_input), "original_input_sha256": _sha256(original_input), "source_mapping": str(source_mapping), "source_mapping_sha256": mapping_hash, "recovery_db": str(recovery_db), "recovery_db_sha256": _sha256(recovery_db)},
            "selection": {"count": len(bundle["id_index"]), "free_count": sum(1 for row in bundle["mapping_rows"] if row["free_state"] == "PENDING"), "paid_count": sum(1 for row in bundle["mapping_rows"] if row["paid_required"] == 1 and row["paid_state"] == "PENDING"), "ordered_index_id": bundle["id_index"], "ordered_index_id_sha256": selected_digest, "row_sha256": row_hashes},
            "workbook": {"relative_path": "input/remaining_159_fresh.xlsx", "bytes": workbook_bytes, "sha256": workbook_hash, "row_count": len(bundle["rows"]), "column_count": len(bundle["headers"])},
            "effective_config": effective_config,
            "paid_budgets": {provider: 0 for provider in PAID_PROVIDERS},
            "runtime_source_tree_sha256": runtime_tree_hash,
            "expected_run_id": expected_run_id,
            "expected_run_root": f"runs/{expected_run_id}",
            "exact_command": exact_command,
            "exact_command_argv": exact_tokens,
            "powershell_command": f"& {exact_command}",
            "exact_command_effective_config_sha256": hashlib.sha256(_canonical(effective_config).encode("utf-8")).hexdigest(),
            "exact_command_run_id": reparsed_run_id,
            "result_field_leakage": sorted(FORBIDDEN_RESULT_FIELDS & {header.casefold() for header in bundle["headers"]} - {"website", "email", "phone"}),
        }
        plan["plan_payload_sha256"] = hashlib.sha256(_canonical(plan).encode("utf-8")).hexdigest()
        plan_temp.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
        validate_remaining_plan(workbook_path, plan_temp)
        if plan_path.exists() and _sha256(plan_path) != _sha256(plan_temp):
            raise FileExistsError("remaining_159_plan.json exists with a different hash")
        verify_only = workbook_path.exists() and plan_path.exists()
        if not verify_only:
            plan_temp.replace(plan_path)
        return {"verify_only": verify_only, "workbook": str(workbook_path), "plan": str(plan_path), "workbook_sha256": workbook_hash, "plan_sha256": _sha256(plan_path if plan_path.exists() else plan_temp)}
    finally:
        workbook_temp.unlink(missing_ok=True)
        plan_temp.unlink(missing_ok=True)
        if created_workbook and not plan_path.exists():
            workbook_path.unlink(missing_ok=True)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--destination", type=Path, required=True)
    parser.add_argument("--original-input", type=Path, default=None)
    parser.add_argument("--recovery-db", type=Path, default=None)
    parser.add_argument("--parent-manifest", type=Path, default=None)
    parser.add_argument("--source-mapping", type=Path, default=None)
    args = parser.parse_args(argv)
    print(json.dumps(prepare_remaining_run(
        args.destination,
        original_input=args.original_input,
        recovery_db=args.recovery_db,
        parent_manifest=args.parent_manifest,
        source_mapping=args.source_mapping,
    ), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

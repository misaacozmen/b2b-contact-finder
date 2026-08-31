"""Network-free, fail-closed reconciliation of the 893-row recovery output."""

from __future__ import annotations

import argparse
import contextlib
import hashlib
import json
import os
import socket
import sqlite3
from collections import Counter
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

import config
from modules import contact_identity, excel, output_artifacts, publication_policy, run_context
from prepare_remaining_run import validate_remaining_plan


EXPECTED_INPUT_SHA256 = "29f966aa42be30375b1e52532140f0fa36cef094ae0301e7248c74dc275e5515"
EXPECTED_LEGACY_DB_SHA256 = "e33de2fcfa7cff9c82552ff4b73377bc1a49b550a61e7527e9ac323ad3b53654"
EXPECTED_SELECTED_COUNT = 159
EXPECTED_FREE_COUNT = 2
EXPECTED_PAID_COUNT = 157
PAID_PROVIDERS = ("brightdata", "google_places", "brandfetch", "hunter", "linkedin", "llm")
NONTERMINAL_PAID = {"PENDING", "RUNNING", "UNKNOWN", "BLOCKED_BUDGET", "RESERVED"}
MARKER = "HANDOFF_PENDING"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _ro(path: Path) -> sqlite3.Connection:
    connection = sqlite3.connect(f"file:{Path(path).resolve().as_posix()}?mode=ro&immutable=1", uri=True)
    if connection.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
        connection.close()
        raise ValueError(f"SQLite integrity check failed: {path}")
    return connection


def _xlsx_info(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return {"sha256": _sha256(path), "bytes": path.stat().st_size, "rows": sheet.max_row - 1, "columns": sheet.max_column}
    finally:
        workbook.close()


def _validate_recovery_parent(recovery_run: Path, original_input: Path) -> dict[str, Any]:
    manifest_path = recovery_run / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("run_id") != recovery_run.name or manifest.get("input_sha256") != EXPECTED_INPUT_SHA256:
        raise ValueError("recovery manifest identity mismatch")
    if _sha256(original_input) != EXPECTED_INPUT_SHA256:
        raise ValueError("original input hash mismatch")
    artifact_hash = str(manifest.get("artifact_set_sha256", ""))
    artifact_dir = recovery_run / "output" / "artifacts" / artifact_hash
    files = manifest.get("files")
    if not artifact_hash or not isinstance(files, dict) or not artifact_dir.is_dir():
        raise ValueError("recovery artifact manifest is incomplete")
    actual_set = hashlib.sha256("".join(f"{name}:{_sha256(artifact_dir / name)}\n" for name in sorted(files)).encode()).hexdigest()
    if actual_set != artifact_hash:
        raise ValueError("recovery artifact set hash mismatch")
    for name, info in files.items():
        path = artifact_dir / name
        if not path.is_file() or _sha256(path) != str(info.get("sha256")) or path.stat().st_size != int(info.get("bytes", -1)):
            raise ValueError(f"recovery artifact hash mismatch: {name}")
    return {"manifest": manifest, "manifest_sha256": _sha256(manifest_path), "artifact_dir": artifact_dir, "files": files}


def _validate_legacy(legacy_db: Path) -> tuple[str, list[tuple[int, str]], dict[str, Any]]:
    if _sha256(legacy_db) != EXPECTED_LEGACY_DB_SHA256:
        raise ValueError("legacy DB hash mismatch")
    with _ro(legacy_db) as connection:
        runs = connection.execute("SELECT run_id,input_hash,run_signature FROM runs").fetchall()
        if len(runs) != 1 or str(runs[0][1]) != EXPECTED_INPUT_SHA256 or not str(runs[0][2]).strip():
            raise ValueError("legacy DB run identity mismatch")
        rows = connection.execute("SELECT item_index,payload FROM results ORDER BY item_index").fetchall()
    if [int(row[0]) for row in rows] != list(range(893)):
        raise ValueError("legacy DB result order/count mismatch")
    return str(runs[0][0]), [(int(index), str(payload)) for index, payload in rows], {"run_id": str(runs[0][0]), "run_signature": str(runs[0][2])}


def _validate_recovery_db(artifact_dir: Path, manifest: dict[str, Any]) -> tuple[list[dict], dict[int, dict]]:
    db = artifact_dir / "recovery_state.sqlite3"
    with _ro(db) as connection:
        items = connection.execute("SELECT item_index,source_record_id,free_state,paid_required,paid_state FROM run_items ORDER BY item_index").fetchall()
        payloads = {int(index): json.loads(payload) for index, payload in connection.execute("SELECT item_index,payload FROM results ORDER BY item_index")}
        if int(connection.execute("SELECT COUNT(*) FROM replay_entries").fetchone()[0]) != 0:
            raise ValueError("recovery DB unexpectedly contains replay entries")
    if len(items) != 893 or len(payloads) != 893:
        raise ValueError("recovery DB item/result count mismatch")
    selected = [row for row in items if str(row[2]) == "PENDING" or (int(row[3]) == 1 and str(row[4]) in NONTERMINAL_PAID)]
    if len(selected) != EXPECTED_SELECTED_COUNT:
        raise ValueError("recovery selection is not exactly 159")
    if sum(str(row[2]) == "PENDING" for row in selected) != EXPECTED_FREE_COUNT or sum(int(row[3]) == 1 and str(row[4]) in NONTERMINAL_PAID for row in selected) != EXPECTED_PAID_COUNT:
        raise ValueError("recovery selection is not disjoint 2 free + 157 paid")
    if str(manifest.get("counts", {}).get("paid_pending")) != str(EXPECTED_PAID_COUNT):
        raise ValueError("recovery manifest paid selection mismatch")
    return [{"item_index": int(row[0]), "source_record_id": str(row[1]), "free_state": str(row[2]), "paid_required": bool(row[3]), "paid_state": str(row[4])} for row in items], payloads


def _validate_mapping(artifact_dir: Path, records: list[dict], legacy_rows: list[tuple[int, str]]) -> tuple[dict[int, dict], dict[int, dict]]:
    mapping_path = artifact_dir / "source_id_mapping.json"
    mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    unsigned = {key: value for key, value in mapping.items() if key != "mapping_sha256"}
    if hashlib.sha256(_canonical(unsigned).encode()).hexdigest() != str(mapping.get("mapping_sha256", "")):
        raise ValueError("source ID mapping digest mismatch")
    rows = {int(row["legacy_index"]): row for row in mapping.get("mappings", [])}
    if len(rows) != 893:
        raise ValueError("source ID mapping count mismatch")
    by_index = {}
    for index, record in enumerate(records):
        source_id, quality = run_context.source_record_identity(record, default_source="input")
        mapped = rows.get(index)
        if not mapped or str(mapped.get("canonical_source_record_id")) != source_id or hashlib.sha256(run_context.canonical_json(record).encode()).hexdigest() != str(mapped.get("input_row_sha256")):
            raise ValueError(f"source ID mapping mismatch at index {index}")
        by_index[index] = {"source_record_id": source_id, "quality": quality, "company": str(record.get("company", ""))}
    if len({item["source_record_id"] for item in by_index.values()}) != 893:
        raise ValueError("canonical source IDs are not unique")
    return rows, by_index


def _validate_forensic(artifact_dir: Path, legacy_rows: list[tuple[int, str]]) -> dict[int, dict]:
    path = artifact_dir / "forensic_legacy_payloads.jsonl"
    forensic: dict[int, dict] = {}
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            item = json.loads(line)
            index = int(item["legacy_index"])
            if index in forensic or hashlib.sha256(legacy_rows[index][1].encode()).hexdigest() != str(item.get("payload_sha256")) or json.loads(legacy_rows[index][1]) != item.get("payload"):
                raise ValueError(f"forensic payload mismatch at index {index}")
            forensic[index] = item
    if set(forensic) != set(range(893)):
        raise ValueError("forensic payload coverage mismatch")
    return forensic


def _evaluation(row: dict[str, Any]) -> dict[str, Any]:
    evaluation = dict(row.get("__evaluation") if isinstance(row.get("__evaluation"), dict) else {})
    if row.get("identity_resolution"):
        evaluation.setdefault("_identity_resolution", row.get("identity_resolution"))
    evaluation.setdefault("candidate", row.get("candidate") if isinstance(row.get("candidate"), dict) else {})
    evaluation.setdefault("identity_assessment", row.get("identity_assessment") if isinstance(row.get("identity_assessment"), dict) else {})
    evaluation.setdefault("structured_domain_relation", bool(row.get("structured_domain_relation")))
    evaluation.setdefault("reasons", [value.strip() for value in str(row.get("reason", "")).split(";") if value.strip()])
    evaluation.setdefault("has_contact", bool(row.get("email") or row.get("phone")))
    evaluation.setdefault("email", row.get("email", ""))
    evaluation.setdefault("phone", row.get("phone", ""))
    evaluation.setdefault("email_verification", row.get("email_verification", ""))
    evaluation.setdefault("email_failed", "email_gate_failed" in str(row.get("reason", "")))
    return evaluation


def _apply_current_policy(row: dict[str, Any], *, legacy_downgrade_only: bool) -> dict[str, Any]:
    old_eligible = row.get("publication_eligible") is True
    decision = publication_policy.evaluate(str(row.get("company", "")), _evaluation(row), str(row.get("status", "")), minimum_safety_score=int(config.PUBLICATION_POLICY_MIN_SAFETY_SCORE))
    if legacy_downgrade_only and not old_eligible:
        decision["eligible"] = False
        decision["action"] = "retain_legacy_abstention"
    row["publication_policy_version"] = decision["policy_version"]
    row["publication_policy_action"] = decision["action"]
    row["publication_eligible"] = bool(decision["eligible"])
    row["publication_safety_score"] = int(decision["safety_score"])
    row["publication_risk_index"] = int(decision["risk_index"])
    row["publication_risk_tier"] = decision["risk_tier"]
    blockers = [value.strip() for value in str(row.get("publication_blockers", "")).replace(",", ";").split(";") if value.strip()]
    blockers.extend(str(value) for value in decision.get("hard_blockers", []))
    row["publication_blockers"] = "; ".join(sorted(set(blockers)))
    return row


def _remove_marker(value: str) -> str:
    return "; ".join(item.strip() for item in str(value).replace(",", ";").split(";") if item.strip() and item.strip() != MARKER)


def _remove_temporary_markers(value: str) -> str:
    temporary = {MARKER, "paid_pending_approval"}
    return "; ".join(
        item.strip()
        for item in str(value).replace(",", ";").split(";")
        if item.strip() and item.strip() not in temporary
    )


def _configure_output(destination: Path) -> None:
    config.OUTPUT_DIR = destination
    for name in ("CONTACTS_FILE", "VERIFIED_CONTACTS_FILE", "REVIEW_QUEUE_FILE", "FAILED_FILE", "CANDIDATES_FILE", "REPORT_FILE", "LOG_FILE", "EVIDENCE_FILE", "ENTITY_RELATIONSHIPS_FILE", "TELEMETRY_FILE", "DISCOVERY_COVERAGE_FILE", "QUALITY_AUDIT_FILE", "REPLAY_SNAPSHOT_FILE", "ALL_RESULTS_FILE", "MANIFEST_FILE"):
        if hasattr(config, name):
            filename = Path(getattr(config, name)).name
            setattr(config, name, destination / filename)


@contextlib.contextmanager
def _network_block() -> Iterator[None]:
    original_connect = socket.socket.connect
    original_create = socket.create_connection
    def deny(*_args, **_kwargs):
        raise RuntimeError("network access is forbidden during reconciliation")
    socket.socket.connect = deny
    socket.create_connection = deny
    try:
        yield
    finally:
        socket.socket.connect = original_connect
        socket.create_connection = original_create


def _validate_remaining(remaining_plan: Path, remaining_run: Path, selected_ids: list[str]) -> tuple[dict, dict[int, dict]]:
    raw_plan = json.loads(remaining_plan.read_text(encoding="utf-8"))
    planned_workers = raw_plan.get("effective_config", {}).get("effective_settings", {}).get("max_workers")
    previous_workers = getattr(config, "MAX_WORKERS", None)
    if planned_workers is not None:
        config.MAX_WORKERS = int(planned_workers)
    try:
        plan = validate_remaining_plan(remaining_plan.parent / "remaining_159_fresh.xlsx", remaining_plan)
    finally:
        if previous_workers is not None:
            config.MAX_WORKERS = previous_workers
    ordered = plan["selection"]["ordered_index_id"]
    if [str(item["source_record_id"]) for item in ordered] != selected_ids:
        raise ValueError("remaining plan selection does not match recovery selection")
    manifest = json.loads((remaining_run / "manifest.json").read_text(encoding="utf-8"))
    if manifest.get("run_id") != remaining_run.name or manifest.get("run_id") != plan.get("expected_run_id") or manifest.get("input_sha256") != plan["workbook"]["sha256"]:
        raise ValueError("remaining run identity mismatch")
    if manifest.get("config_sha256") != hashlib.sha256(_canonical(plan["effective_config"]).encode()).hexdigest() and manifest.get("config_sha256") != plan.get("exact_command_effective_config_sha256"):
        raise ValueError("remaining run config hash mismatch")
    if manifest.get("runtime_source_tree_sha256") != plan["runtime_source_tree_sha256"] or manifest.get("ordered_source_record_ids") != selected_ids:
        raise ValueError("remaining run source-tree or ID order mismatch")
    return plan, manifest


def _load_remaining(
    remaining_run: Path,
    selected_ids: list[str],
    plan: dict,
) -> tuple[dict[int, dict], dict[int, dict], int]:
    db = remaining_run / "state" / "progress.sqlite3"
    with _ro(db) as connection:
        rows = connection.execute("SELECT item_index,source_record_id,free_state,paid_required,paid_state FROM run_items ORDER BY item_index").fetchall()
        payloads = {int(index): json.loads(payload) for index, payload in connection.execute("SELECT item_index,payload FROM results ORDER BY item_index")}
        calls = int(connection.execute("SELECT COUNT(*) FROM provider_calls").fetchone()[0])
        usage = connection.execute("SELECT provider,configured_limit,effective_limit,reserved,completed,failed FROM provider_usage").fetchall()
    scheduler = {
        int(row[0]): {
            "item_index": int(row[0]),
            "source_record_id": str(row[1]),
            "free_state": str(row[2]),
            "paid_required": bool(row[3]),
            "paid_state": str(row[4]),
        }
        for row in rows
    }
    if (
        calls != 0
        or len(rows) != EXPECTED_SELECTED_COUNT
        or [int(row[0]) for row in rows] != list(range(EXPECTED_SELECTED_COUNT))
        or [str(row[1]) for row in rows] != selected_ids
    ):
        raise ValueError("remaining run ID/call coverage mismatch")
    if any(any(int(value or 0) != 0 for value in row[1:]) for row in usage):
        raise ValueError("remaining run has nonzero paid usage")
    status_counts = Counter(
        (str(row[2]), bool(row[3]), str(row[4]))
        for row in rows
    )
    expected_status_counts = Counter({
        ("DONE", False, "NOT_REQUIRED"): 16,
        ("DONE", True, "PENDING"): 143,
    })
    if status_counts != expected_status_counts:
        raise ValueError(f"remaining scheduler status mismatch: {status_counts}")
    pending = sum(
        scheduler_row["paid_required"] is True
        and scheduler_row["paid_state"] in NONTERMINAL_PAID
        for scheduler_row in scheduler.values()
    )
    manifest_pending = int(json.loads((remaining_run / "manifest.json").read_text(encoding="utf-8")).get("paid_pending", pending))
    if pending != manifest_pending:
        raise ValueError("remaining paid_pending manifest/SQLite mismatch")
    if pending != 143:
        raise ValueError("remaining pending count mismatch")
    return payloads, scheduler, pending


def _artifact_manifest(destination: Path, result: output_artifacts.ArtifactResult, *, lineage: dict[str, Any], counts: dict[str, int], pending: int) -> dict[str, Any]:
    files: dict[str, dict[str, Any]] = {}
    for path in sorted(destination.iterdir(), key=lambda p: p.name):
        if path.is_file() and path.name not in {"delivery_manifest.json", "FINALIZATION_COMPLETE.txt", "REMEDIATION_PENDING.txt"}:
            files[path.name] = _xlsx_info(path) if path.suffix.casefold() == ".xlsx" else {"sha256": _sha256(path), "bytes": path.stat().st_size, "rows": None, "columns": None}
    payload = {
        "coverage_complete": True,
        "remediation_complete": pending == 0,
        "complete": pending == 0,
        "counts": counts | {"pending": pending},
        "lineage": lineage,
        "files": files,
        "artifact_set_sha256": result.artifacts.get("artifact_set_sha256", ""),
    }
    return payload


def _verify_existing_delivery(destination: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    for name, info in dict(manifest.get("files", {})).items():
        path = destination / name
        if not path.is_file() or _sha256(path) != str(info.get("sha256")) or path.stat().st_size != int(info.get("bytes", -1)):
            raise RuntimeError(f"existing delivery artifact mismatch: {name}")
        if path.suffix.casefold() == ".xlsx":
            actual = _xlsx_info(path)
            if actual.get("rows") != info.get("rows") or actual.get("columns") != info.get("columns"):
                raise RuntimeError(f"existing delivery row/column mismatch: {name}")
    return manifest | {"verify_only": True}


def reconcile(*, original_input: Path, legacy_db: Path, recovery_run: Path, remaining_plan: Path, remaining_run: Path, destination: Path) -> dict[str, Any]:
    original_input, legacy_db, recovery_run, remaining_plan, remaining_run, destination = [Path(path).resolve() for path in (original_input, legacy_db, recovery_run, remaining_plan, remaining_run, destination)]
    existing_manifest = destination / "delivery_manifest.json"
    if destination.exists() and existing_manifest.exists():
        return _verify_existing_delivery(destination, json.loads(existing_manifest.read_text(encoding="utf-8")))
    if destination.exists() and any(destination.iterdir()):
        raise RuntimeError("destination already contains unverified content")
    records = excel.read_company_records(original_input)
    if len(records) != 893 or _sha256(original_input) != EXPECTED_INPUT_SHA256:
        raise ValueError("original input coverage/hash mismatch")
    parent = _validate_recovery_parent(recovery_run, original_input)
    legacy_run_id, legacy_rows, legacy_info = _validate_legacy(legacy_db)
    mapping_rows, identities = _validate_mapping(parent["artifact_dir"], records, legacy_rows)
    forensic = _validate_forensic(parent["artifact_dir"], legacy_rows)
    recovery_items, recovery_payloads = _validate_recovery_db(parent["artifact_dir"], parent["manifest"])
    selected = [row for row in recovery_items if row["free_state"] == "PENDING" or (row["paid_required"] and row["paid_state"] in NONTERMINAL_PAID)]
    selected.sort(key=lambda row: row["item_index"])
    selected_ids = [row["source_record_id"] for row in selected]
    plan, remaining_manifest = _validate_remaining(remaining_plan, remaining_run, selected_ids)
    remaining_payloads, remaining_scheduler, actual_pending = _load_remaining(remaining_run, selected_ids, plan)
    remaining_payloads_by_source_index = {
        int(selection["item_index"]): remaining_payloads[ordinal]
        for ordinal, selection in enumerate(plan["selection"]["ordered_index_id"])
    }
    remaining_scheduler_by_source_index = {
        int(selection["item_index"]): remaining_scheduler[ordinal]
        for ordinal, selection in enumerate(plan["selection"]["ordered_index_id"])
    }

    rows: list[dict[str, Any]] = []
    selected_indexes = {row["item_index"] for row in selected}
    for index, record in enumerate(records):
        identity = identities[index]
        if index in selected_indexes:
            scheduler = remaining_scheduler_by_source_index.get(index)
            if scheduler is None or scheduler["source_record_id"] != identity["source_record_id"]:
                raise ValueError(f"remaining scheduler identity mismatch at index {index}")
            row = dict(remaining_payloads_by_source_index.get(index, {}))
            if str(row.get("source_record_id")) != identity["source_record_id"] or str(row.get("company", "")).strip().casefold() != identity["company"].strip().casefold():
                raise ValueError(f"remaining payload identity mismatch at index {index}")
            row.update({"company": identity["company"], "legacy_index": index, "source_record_id": identity["source_record_id"], "source_record_id_quality": identity["quality"], "free_state": str(scheduler["free_state"]), "paid_required": bool(scheduler["paid_required"]), "paid_state": str(scheduler["paid_state"])})
            is_pending = scheduler["paid_required"] is True and scheduler["paid_state"] in NONTERMINAL_PAID
            if is_pending:
                row["delivery_state"] = "REMEDIATION_PENDING"
                row["publication_eligible"] = False
                row["publication_blockers"] = "; ".join(sorted(set(filter(None, [_remove_marker(row.get("publication_blockers", "")), "paid_pending_approval"]))))
                output_artifacts.suppress_all_contacts(row, "paid_pending_approval")
            else:
                row["quarantine_state"] = ""
                row["quarantine_status"] = ""
                row["publication_blockers"] = _remove_temporary_markers(row.get("publication_blockers", ""))
                _apply_current_policy(row, legacy_downgrade_only=False)
                row["delivery_state"] = "DELIVERABLE" if row.get("publication_eligible") else "REVIEW"
        else:
            forensic_row = dict(forensic[index]["payload"])
            if str(forensic_row.get("company", "")).strip().casefold() != identity["company"].strip().casefold():
                raise ValueError(f"forensic company mismatch at index {index}")
            row = forensic_row
            row.update({"company": identity["company"], "legacy_index": index, "source_record_id": identity["source_record_id"], "source_record_id_quality": identity["quality"], "free_state": "DONE", "paid_required": False, "paid_state": "NOT_REQUIRED", "quarantine_state": "", "quarantine_status": ""})
            row["publication_blockers"] = _remove_marker(row.get("publication_blockers", ""))
            _apply_current_policy(row, legacy_downgrade_only=True)
            row["delivery_state"] = "DELIVERABLE" if row.get("publication_eligible") else "REVIEW"
            if identity["company"].strip().casefold() in {"akel", "akbarkod"}:
                row["collision_reason"] = "LEGACY_MERGED_COLLISION"
                row["publication_eligible"] = False
                row["delivery_state"] = "REVIEW"
                output_artifacts.suppress_all_contacts(row, "LEGACY_MERGED_COLLISION")
        rows.append(row)
    if len(rows) != 893 or len({str(row.get("source_record_id")) for row in rows}) != 893:
        raise ValueError("reconciliation row identity coverage mismatch")
    reconciliation_run_id = hashlib.sha256(
        _canonical({
            "original_input_sha256": _sha256(original_input),
            "recovery_run_id": recovery_run.name,
            "remaining_run_id": remaining_run.name,
        }).encode()
    ).hexdigest()
    for index, row in enumerate(rows):
        row["run_id"] = reconciliation_run_id
        row["original_index"] = index
    output_artifacts.apply_global_identity_collision_gate(rows)
    for row in rows:
        if row.get("delivery_state") != "REMEDIATION_PENDING" and not output_artifacts.is_publishable_row(row):
            row["delivery_state"] = "REVIEW"
    destination.mkdir(parents=True, exist_ok=False)
    _configure_output(destination)
    with _network_block():
        result = output_artifacts.write_outputs(rows, None, telemetry_snapshot={"generated_at": "unknown", "counters": {}})
    published = len(result.entity_memory_rows)
    counts = {"total": len(rows), "published": published, "review": len(rows) - published}
    lineage = {
        "original_input": {"path": str(original_input), "sha256": _sha256(original_input), "rows": 893},
        "legacy_db": {"path": str(legacy_db), "sha256": _sha256(legacy_db), "run_id": legacy_run_id, "run_signature_sha256": hashlib.sha256(legacy_info["run_signature"].encode()).hexdigest()},
        "recovery_run": {"path": str(recovery_run), "run_id": recovery_run.name, "manifest_sha256": parent["manifest_sha256"], "artifact_set_sha256": parent["manifest"]["artifact_set_sha256"]},
        "remaining_plan": {"path": str(remaining_plan), "sha256": _sha256(remaining_plan), "run_id": plan["expected_run_id"], "ordered_selected_source_record_ids_sha256": hashlib.sha256(_canonical(selected_ids).encode()).hexdigest()},
        "remaining_run": {"path": str(remaining_run), "run_id": remaining_run.name, "manifest_sha256": _sha256(remaining_run / "manifest.json"), "config_sha256": remaining_manifest.get("config_sha256"), "source_tree_sha256": remaining_manifest.get("runtime_source_tree_sha256")},
    }
    manifest = _artifact_manifest(destination, result, lineage=lineage, counts=counts, pending=actual_pending)
    manifest["reconciliation_run_id"] = reconciliation_run_id
    manifest_path = destination / "delivery_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    if actual_pending:
        (destination / "REMEDIATION_PENDING.txt").write_text(f"pending={actual_pending}\n", encoding="utf-8")
    else:
        (destination / "FINALIZATION_COMPLETE.txt").write_text("complete=true\n", encoding="utf-8")
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--original-input", type=Path, required=True)
    parser.add_argument("--legacy-db", type=Path, required=True)
    parser.add_argument("--recovery-run", type=Path, required=True)
    parser.add_argument("--remaining-plan", type=Path, required=True)
    parser.add_argument("--remaining-run", type=Path, required=True)
    parser.add_argument("--destination", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(reconcile(**vars(args)), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

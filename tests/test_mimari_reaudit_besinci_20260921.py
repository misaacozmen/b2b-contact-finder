from __future__ import annotations

import hashlib
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from openpyxl import load_workbook

import config
from modules import checkpoint, output_artifacts, publication_policy, runtime


def _budgets(provider: str, limit: int) -> dict[str, int]:
    return {name: (limit if name == provider else 0) for name in checkpoint.CANONICAL_PROVIDERS}


def _dispatch_run(tmp_path: Path, monkeypatch, *, count: int, budget: int, run_id: str) -> list[dict]:
    database = tmp_path / f"{run_id}.sqlite3"
    monkeypatch.setattr("config.PROGRESS_DB_FILE", database)
    checkpoint.initialize_schema(database)
    items = [
        {
            "item_index": index,
            "source_record_id": f"source:{index:03d}",
            "company": f"Company {index:03d}",
            "free_state": "DONE",
            "paid_required": True,
            "paid_state": "RUNNING",
        }
        for index in range(count)
    ]
    checkpoint.initialize_run(
        run_id=run_id, input_hash="input", run_signature="dispatch",
        context={"phase": "PAID"}, budgets=_budgets("hunter", budget), items=items,
    )
    return items


def _runtime_for(run_id: str, budget: int, item_index: int, source: str, round_ordinal: int) -> None:
    runtime.reset()
    runtime.configure_durable_run(run_id, _budgets("hunter", budget))
    runtime.set_phase("PAID")
    runtime.set_item_context(item_index, "paid")
    runtime.set_source_record_id(source)
    runtime.set_provider_dispatch_rounds({"hunter": round_ordinal})


def _ensure_job(run_id: str, item_index: int, source: str, operation: str, request: str, need: str) -> dict:
    return checkpoint.ensure_provider_work_item(
        run_id=run_id, item_index=item_index, source_record_id=source,
        provider="hunter", operation=operation, request_fingerprint=request,
        need_class=need,
    )


def _job_candidate(job: dict) -> dict:
    return {
        "item_index": job["item_index"], "source_record_id": job["source_record_id"],
        "need_class": job["need_class"], "job_fingerprint": job["job_fingerprint"],
        "operation": job["operation"], "request_fingerprint": job["request_fingerprint"],
        "query_fingerprint": job["query_fingerprint"], "plan_version": job["plan_version"],
    }


def test_g01_allocation_is_one_immutable_job_and_retry_uses_next_round(tmp_path: Path, monkeypatch) -> None:
    items = _dispatch_run(tmp_path, monkeypatch, count=2, budget=3, run_id="g01")
    first_job = _ensure_job("g01", 0, "source:000", "website", "job-000-1", "website")
    second_job = _ensure_job("g01", 0, "source:000", "website", "job-000-2", "website")
    other_job = _ensure_job("g01", 1, "source:001", "website", "job-001-1", "website")
    candidates = [_job_candidate(first_job), _job_candidate(other_job)]
    checkpoint.reserve_provider_dispatch_round(
        run_id="g01", provider="hunter", round_ordinal=0, candidates=candidates, cap=2,
    )
    _runtime_for("g01", 3, 0, "source:000", 0)
    first = runtime.reserve_api("hunter", operation="website", request_fingerprint="job-000-1")
    assert first.accepted
    runtime.start_api(first)
    runtime.complete_api(first, "FAILED")

    second_same_source = runtime.reserve_api(
        "hunter", operation="website", request_fingerprint="job-000-2",
    )
    assert not second_same_source and second_same_source.reason == "dispatch_not_allocated"

    _runtime_for("g01", 3, 1, "source:001", 0)
    other_company = runtime.reserve_api(
        "hunter", operation="website", request_fingerprint="job-001-1",
    )
    assert other_company.accepted
    runtime.start_api(other_company)
    runtime.complete_api(other_company, "DONE")

    round_zero = checkpoint.load_provider_dispatch_round("g01", "hunter", 0)
    assert {row["state"] for row in round_zero["allocations"]} == {"FAILED", "DONE"}
    assert checkpoint.next_provider_dispatch_round("g01", "hunter") == 1
    round_zero_identity = (
        round_zero["need_snapshot_sha256"], round_zero["selected_work_hash"],
        [(row["item_index"], row["job_fingerprint"], row["state"]) for row in round_zero["allocations"]],
    )
    retry_job = _ensure_job("g01", 0, "source:000", "website", "job-000-retry", "website")
    checkpoint.reserve_provider_dispatch_round(
        run_id="g01", provider="hunter", round_ordinal=1,
        candidates=[_job_candidate(retry_job)], cap=1,
    )
    _runtime_for("g01", 3, 0, "source:000", 1)
    retry = runtime.reserve_api("hunter", operation="website", request_fingerprint="job-000-retry")
    assert retry.accepted
    checkpoint.reserve_provider_dispatch_round(
        run_id="g01", provider="hunter", round_ordinal=0, candidates=candidates, cap=2,
    )
    round_zero_after_retry = checkpoint.load_provider_dispatch_round("g01", "hunter", 0)
    assert (
        round_zero_after_retry["need_snapshot_sha256"], round_zero_after_retry["selected_work_hash"],
        [(row["item_index"], row["job_fingerprint"], row["state"]) for row in round_zero_after_retry["allocations"]],
    ) == round_zero_identity


def test_g02_resume_loads_open_receipt_and_never_recomputes_selected_work(tmp_path: Path, monkeypatch) -> None:
    _dispatch_run(tmp_path, monkeypatch, count=2, budget=2, run_id="g02")
    first_job = _ensure_job("g02", 0, "source:000", "identity", "job-000", "identity")
    second_job = _ensure_job("g02", 1, "source:001", "contact", "job-001", "contact")
    candidates = [_job_candidate(first_job), _job_candidate(second_job)]
    first = checkpoint.reserve_provider_dispatch_round(
        run_id="g02", provider="hunter", round_ordinal=0, candidates=candidates, cap=2,
    )
    _runtime_for("g02", 2, 0, "source:000", 0)
    call = runtime.reserve_api("hunter", operation="identity", request_fingerprint="job-000")
    assert call.accepted

    resumed = checkpoint.reserve_provider_dispatch_round(
        run_id="g02", provider="hunter", round_ordinal=0,
        candidates=list(reversed(candidates)), cap=2,
    )
    receipt = checkpoint.load_provider_dispatch_round("g02", "hunter", 0)
    assert [
        (row["item_index"], row["source_record_id"], row["need_class"])
        for row in resumed
    ] == [
        (row["item_index"], row["source_record_id"], row["need_class"])
        for row in first
    ]
    assert checkpoint.next_provider_dispatch_round("g02", "hunter") == 0
    assert receipt["state"] == "IN_PROGRESS"
    consumed = next(row for row in receipt["allocations"] if row["item_index"] == 0)
    assert consumed["state"] == "CONSUMED" and consumed["consumed_call_id"] == call.call_id
    receipt_identity = (
        receipt["need_snapshot_sha256"], receipt["selected_work_hash"],
        [(row["item_index"], row["source_record_id"], row["need_class"], row["job_fingerprint"], row["consumed_call_id"]) for row in receipt["allocations"]],
    )
    with pytest.raises(checkpoint.ResumeInvariant):
        checkpoint.reserve_provider_dispatch_round(
            run_id="g02", provider="hunter", round_ordinal=1,
            candidates=[candidates[0]], cap=1,
        )
    drift_job = _ensure_job("g02", 0, "source:000", "identity-v2", "job-000-drift", "identity")
    drifted = [_job_candidate(drift_job), candidates[1]]
    with pytest.raises(checkpoint.ResumeInvariant, match="drift"):
        checkpoint.reserve_provider_dispatch_round(
            run_id="g02", provider="hunter", round_ordinal=0,
            candidates=drifted, cap=2,
        )

    runtime.start_api(call)
    runtime.complete_api(call, "FAILED")
    assert checkpoint.release_provider_dispatch_allocation(
        run_id="g02", provider="hunter", round_ordinal=0, item_index=1,
        source_record_id="source:001", reason="no_call_needed",
    )
    assert checkpoint.load_provider_dispatch_round("g02", "hunter", 0)["state"] == "COMPLETE"
    assert checkpoint.next_provider_dispatch_round("g02", "hunter") == 1
    checkpoint.reserve_provider_dispatch_round(
        run_id="g02", provider="hunter", round_ordinal=0,
        candidates=list(reversed(candidates)), cap=2,
    )
    receipt_after_terminal = checkpoint.load_provider_dispatch_round("g02", "hunter", 0)
    assert (
        receipt_after_terminal["need_snapshot_sha256"], receipt_after_terminal["selected_work_hash"],
        [(row["item_index"], row["source_record_id"], row["need_class"], row["job_fingerprint"], row["consumed_call_id"]) for row in receipt_after_terminal["allocations"]],
    ) == receipt_identity
    assert [(row["state"], row["consumed_call_id"]) for row in receipt_after_terminal["allocations"]] == [
        ("FAILED", call.call_id), ("RELEASED", ""),
    ]


def test_g02_hundred_sources_three_workers_have_stable_resume_allocation(tmp_path: Path, monkeypatch) -> None:
    items = _dispatch_run(tmp_path, monkeypatch, count=100, budget=100, run_id="g02-scale")
    jobs = [
        _ensure_job(
            "g02-scale", item["item_index"], item["source_record_id"],
            "website", f"job-{item['item_index']}", "website",
        )
        for item in items
    ]
    candidates = [_job_candidate(job) for job in jobs]
    checkpoint.reserve_provider_dispatch_round(
        run_id="g02-scale", provider="hunter", round_ordinal=0, candidates=candidates, cap=100,
    )

    def worker(batch: list[dict]) -> list[str]:
        accepted = []
        for item in batch:
            runtime.set_item_context(item["item_index"], "paid")
            runtime.set_source_record_id(item["source_record_id"])
            runtime.set_provider_dispatch_rounds({"hunter": 0})
            reservation = runtime.reserve_api(
                "hunter", operation="website", request_fingerprint=f"job-{item['item_index']}",
            )
            if reservation.accepted:
                accepted.append(item["source_record_id"])
        return accepted

    runtime.reset()
    runtime.configure_durable_run("g02-scale", _budgets("hunter", 100))
    runtime.set_phase("PAID")
    batches = [items[index::3] for index in range(3)]
    with ThreadPoolExecutor(max_workers=3) as executor:
        accepted = [value for batch in executor.map(worker, batches) for value in batch]
    assert sorted(accepted) == sorted(item["source_record_id"] for item in items)
    with sqlite3.connect(config.PROGRESS_DB_FILE) as connection:
        assert connection.execute("SELECT COUNT(*) FROM provider_calls WHERE run_id='g02-scale'").fetchone()[0] == 100
        assert connection.execute("SELECT COUNT(DISTINCT source_record_id) FROM provider_dispatch_allocations WHERE run_id='g02-scale'").fetchone()[0] == 100

    # A process reset only replays the immutable fingerprints; it cannot create new calls.
    runtime.reset()
    runtime.configure_durable_run("g02-scale", _budgets("hunter", 100))
    runtime.set_phase("PAID")
    for item in items:
        runtime.set_item_context(item["item_index"], "paid")
        runtime.set_source_record_id(item["source_record_id"])
        runtime.set_provider_dispatch_rounds({"hunter": 0})
        duplicate = runtime.reserve_api(
            "hunter", operation="website", request_fingerprint=f"job-{item['item_index']}",
        )
        assert not duplicate and duplicate.reason == "duplicate_request"
    assert checkpoint.reserve_provider_dispatch_round(
        run_id="g02-scale", provider="hunter", round_ordinal=0, candidates=candidates, cap=100,
    )
    assert checkpoint.next_provider_dispatch_round("g02-scale", "hunter") == 0


def _content(source_id: str, website: str, *, email: str = "", phone: str = "", email_allowed: bool = False, phone_allowed: bool = False, conflict: bool = False) -> tuple[dict, list[dict]]:
    digest = hashlib.sha256(website.encode()).hexdigest()
    evidence_id = hashlib.sha256(f"{source_id}:identity".encode()).hexdigest()
    records = [{
        "evidence_id": evidence_id, "target_source_record_id": source_id,
        "evidence_source_record_id": f"page:{source_id}", "observed_business": "Fixture Company",
        "url": website, "final_url": website, "content_sha256": digest,
        "retrieval_method": "http", "observation_type": "legal_name",
        "observation_value": "Fixture Company", "location": {"kind": "page_text"},
        "relation": "first_party_identity", "identity_route": "LEGAL_NAME", "first_party": True,
        "contact_fields": [field for field, value in (("email", email), ("phone", phone)) if value],
        "observed_contacts": {key: value for key, value in (("email", email), ("phone", phone)) if value},
    }]
    content = publication_policy.ContentDecision(
        verified=not conflict, website_allowed=not conflict,
        email_allowed=bool(email_allowed and not conflict), phone_allowed=bool(phone_allowed and not conflict),
        reason_codes=("identity_conflict",) if conflict else (),
        support_evidence_ids=(evidence_id,), conflict_evidence_ids=(evidence_id,) if conflict else (),
        identity_status="unverified" if conflict else "verified", identity_route="LEGAL_NAME",
        complete_contact=bool(email_allowed and phone_allowed and not conflict),
        missing_evidence=tuple(value for value, present in (("email_evidence_missing", not email), ("phone_evidence_missing", not phone or not phone_allowed)) if present),
        next_actions=tuple(value for value, present in (("acquire_email", not email or not email_allowed), ("acquire_phone", not phone or not phone_allowed)) if present),
        target_source_record_id=source_id,
        evidence_fingerprint=publication_policy.evidence_fingerprint(records, target_source_record_id=source_id, website=website, email=email, phone=phone),
        bound_website=publication_policy.scorer.normalize_domain(website),
        bound_email=email.casefold(), bound_phone=phone.casefold(),
    ).as_dict()
    return content, records


def test_g03_partial_policy_and_primary_review_export_keep_field_gap(tmp_path: Path, monkeypatch) -> None:
    partial, partial_records = _content(
        "source:partial", "https://partial.example", email="info@partial.example", email_allowed=True,
    )
    decision = publication_policy.finalize_publication(partial, {
        "free_state": "DONE", "paid_required": False, "paid_state": "NOT_REQUIRED",
        "content_status": "OK_HIGH_CONFIDENCE",
    })
    assert decision.publishable
    assert not decision.content_decision.complete_contact
    assert "phone_evidence_missing" in decision.content_decision.missing_evidence
    assert "missing_evidence:phone_evidence_missing" not in decision.blockers

    invalid_phone, invalid_records = _content(
        "source:invalid", "https://invalid.example", email="info@invalid.example", phone="+999", email_allowed=True,
    )
    invalid_decision = publication_policy.finalize_publication(invalid_phone, {
        "free_state": "DONE", "paid_required": False, "paid_state": "NOT_REQUIRED",
        "content_status": "OK_HIGH_CONFIDENCE",
    })
    assert invalid_decision.publishable and not invalid_decision.content_decision.phone_allowed

    conflict, conflict_records = _content(
        "source:conflict", "https://conflict.example", email="info@conflict.example", email_allowed=True,
        conflict=True,
    )
    conflict_decision = publication_policy.finalize_publication(conflict, {
        "free_state": "DONE", "paid_required": False, "paid_state": "NOT_REQUIRED",
        "content_status": "OK_HIGH_CONFIDENCE",
    })
    assert not conflict_decision.publishable

    database = tmp_path / "g03.sqlite3"
    monkeypatch.setattr("config.PROGRESS_DB_FILE", database)
    checkpoint.initialize_schema(database)
    checkpoint.initialize_run(
        run_id="g03", input_hash="input", run_signature="partial",
        context={"phase": "FREE"}, budgets=_budgets("hunter", 0), items=[
                {"item_index": 0, "source_record_id": "source:partial", "free_state": "RUNNING", "paid_required": False, "paid_state": "NOT_REQUIRED"},
                {"item_index": 1, "source_record_id": "source:invalid", "free_state": "RUNNING", "paid_required": False, "paid_state": "NOT_REQUIRED"},
                {"item_index": 2, "source_record_id": "source:conflict", "free_state": "RUNNING", "paid_required": False, "paid_state": "NOT_REQUIRED"},
        ],
    )
    rows = [
        {"company": "Fixture Company", "source_record_id": "source:partial", "free_state": "DONE", "paid_required": False, "paid_state": "NOT_REQUIRED", "status": "OK_HIGH_CONFIDENCE", "score": 95, "website": "https://partial.example", "email": "info@partial.example", "phone": "", "email_publication_status": "allowed", "phone_publication_status": "suppressed", "identity_assessment": {"publishable": True}, "content_decision": partial, "content_evidence_records": partial_records},
        {"company": "Invalid Company", "source_record_id": "source:invalid", "free_state": "DONE", "paid_required": False, "paid_state": "NOT_REQUIRED", "status": "OK_HIGH_CONFIDENCE", "score": 95, "website": "https://invalid.example", "email": "info@invalid.example", "phone": "+999", "email_publication_status": "allowed", "phone_publication_status": "suppressed", "identity_assessment": {"publishable": True}, "content_decision": invalid_phone, "content_evidence_records": invalid_records},
        {"company": "Conflict Company", "source_record_id": "source:conflict", "free_state": "DONE", "paid_required": False, "paid_state": "NOT_REQUIRED", "status": "OK_HIGH_CONFIDENCE", "score": 95, "website": "https://conflict.example", "email": "info@conflict.example", "phone": "", "email_publication_status": "allowed", "phone_publication_status": "suppressed", "identity_assessment": {"publishable": False, "conflicts": ["ownership"]}, "content_decision": conflict, "content_evidence_records": conflict_records},
    ]
    for index, row in enumerate(rows):
        checkpoint.save_item_transaction(
            run_id="g03", item_index=index, source_record_id=row["source_record_id"], payload=row,
            free_state="DONE", paid_state="NOT_REQUIRED", paid_required=False,
            free_attempts=1, paid_attempts=0,
        )
    loaded = list(checkpoint.load_results_by_id("g03").values())
    assert len(loaded) == 3
    output_root = tmp_path / "output"
    monkeypatch.setattr("config.OUTPUT_DIR", output_root)
    for name in (
        "CONTACTS_FILE", "VERIFIED_CONTACTS_FILE", "REVIEW_QUEUE_FILE", "FAILED_FILE",
        "CANDIDATES_FILE", "EVIDENCE_FILE", "ENTITY_RELATIONSHIPS_FILE", "QUALITY_AUDIT_FILE",
        "DISCOVERY_COVERAGE_FILE", "REPORT_FILE", "TELEMETRY_FILE",
    ):
        monkeypatch.setattr("config." + name, output_root / getattr(config, name).name)
    output_artifacts.write_outputs(loaded, 0)
    for path in (config.CONTACTS_FILE, config.REVIEW_QUEUE_FILE, config.EVIDENCE_FILE):
        assert path.is_file()
    workbook = load_workbook(config.CONTACTS_FILE, read_only=True, data_only=True)
    values = list(workbook.active.values)
    workbook.close()
    headers, *published_rows = values
    assert len(published_rows) == 2
    phone_index = list(headers).index("phone")
    assert all(row[phone_index] in (None, "") for row in published_rows)
    workbook = load_workbook(config.REVIEW_QUEUE_FILE, read_only=True, data_only=True)
    review_values = list(workbook.active.values)
    workbook.close()
    assert len(review_values) == 2
    evidence_text = config.EVIDENCE_FILE.read_text(encoding="utf-8")
    assert "phone_evidence_missing" in evidence_text and "identity_conflict" in evidence_text

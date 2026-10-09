"""Talimat 39: fair, hall and stand columns in the delivery sheet, and the opt-out list."""

from openpyxl import load_workbook

import config
from modules import checkpoint, opt_out, pipeline_runner, run_report


ROWS = [
    {
        "company": "Alfa Makina", "website": "https://alfa-ornek.com.tr/", "email": "info@alfa-ornek.com.tr",
        "listing_source": "Örnek Fuarı 2026 katılımcı listesi", "listing_hall": "3", "listing_stand": "B-12",
    },
    {
        "company": "Beta Gıda", "website": "https://www.beta-ornek.com.tr/", "email": "satis@beta-ornek.com.tr",
        "listing_source": "Örnek Fuarı 2026 katılımcı listesi", "listing_hall": "", "listing_stand": "A-1",
    },
    {"company": "Gama Ltd", "website": "", "email": "", "listing_source": "", "listing_hall": "", "listing_stand": ""},
]


def _sheet_rows(book, name):
    return [list(row) for row in book[name].iter_rows(values_only=True)]


def test_contact_sheet_ends_with_fair_hall_and_stand(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    assert run_report.CONTACT_HEADERS[-3:] == ["Fuar", "Salon", "Stant"]
    assert run_report.write_run_report(
        ROWS, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=True, data_only=True)
    try:
        rows = _sheet_rows(book, "İletişim")
        assert rows[0] == run_report.CONTACT_HEADERS
        assert [row[4:] for row in rows[1:]] == [
            ["Örnek Fuarı 2026", "3", "B-12"],
            ["Örnek Fuarı 2026", None, "A-1"],
            [None, None, None],
        ]
        assert book["Özet"]["A6"].value == "Ret listesiyle çıkarılan firma"
        assert book["Özet"]["B6"].value == 0
    finally:
        book.close()


def test_opt_out_entries_match_domain_email_mail_domain_and_name():
    def hit(*entries, company="Alfa Makina A.Ş.", website="https://www.alfa-ornek.com.tr/", email="info@alfa-ornek.com.tr"):
        return opt_out.matches(entries, company=company, website=website, email=email)

    assert hit("alfa-ornek.com.tr")
    assert hit("info@alfa-ornek.com.tr")
    assert hit("@alfa-ornek.com.tr", website="")
    assert hit("alfa-ornek.com.tr", website="")
    assert hit("alfa makina a.ş.")
    assert not hit("beta-ornek.com.tr")
    assert not hit("satis@alfa-ornek.com.tr")
    assert not hit("@beta-ornek.com.tr")
    assert not hit("alfa makina")
    assert not hit()


def test_opt_out_rows_leave_every_sheet_and_summary_counts_them(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    assert run_report.write_run_report(
        ROWS, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
        opt_out_entries=("beta-ornek.com.tr", "gama ltd"),
    ) == {}
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=True, data_only=True)
    try:
        for name in ("İletişim", "Detaylar"):
            assert [row[0] for row in _sheet_rows(book, name)[1:]] == ["Alfa Makina"]
        assert book["Özet"]["B3"].value == 1
        assert book["Özet"]["B6"].value == 2
    finally:
        book.close()


def test_opt_out_file_is_created_once_and_comments_are_skipped(tmp_path):
    path = tmp_path / "state" / "ret_listesi.txt"
    assert opt_out.load(path) == ()
    assert opt_out.ensure_file(path) == path
    assert path.read_text(encoding="utf-8") == opt_out.HEADER
    assert opt_out.load(path) == ()
    path.write_text(opt_out.HEADER + "Alfa-Ornek.com.tr\n\n  Beta   Gıda  \n# not\n", encoding="utf-8")
    opt_out.ensure_file(path)
    assert opt_out.load(path) == ("alfa-ornek.com.tr", "beta gıda")


def test_run_report_takes_listing_fields_from_input_and_opt_out_file(tmp_path, monkeypatch):
    rows = [
        {"company": "Alfa Makina", "website": "https://alfa-ornek.com.tr/", "source_record_id": "ornek:alfa"},
        {"company": "Beta Gıda", "website": "https://beta-ornek.com.tr/", "source_record_id": "ornek:beta"},
    ]
    snapshots = {
        0: {"source_record_id": "ornek:alfa", "source": "Örnek Fuarı katılımcı listesi", "hall": "2", "stand": "C-7"},
        1: {"source_record_id": "ornek:beta", "source": "Örnek Fuarı katılımcı listesi", "hall": "2", "stand": "C-8"},
    }
    opt_out_file = tmp_path / "ret_listesi.txt"
    opt_out_file.write_text("beta-ornek.com.tr\n", encoding="utf-8")
    monkeypatch.setattr(config, "OPT_OUT_FILE", opt_out_file)
    monkeypatch.setattr(pipeline_runner, "_durable_output_rows", lambda _run_id, _fallback=None: [dict(row) for row in rows])
    monkeypatch.setattr(checkpoint, "load_input_snapshots", lambda _run_id: snapshots)
    monkeypatch.setattr(checkpoint, "derive_telemetry", lambda _run_id: {})
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    output = tmp_path / "run" / "output"
    assert pipeline_runner._write_run_report(
        run_id="t39", output_root=output, fallback=None,
        run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    book = load_workbook(output / "sonuclar.xlsx", read_only=True, data_only=True)
    try:
        assert _sheet_rows(book, "İletişim")[1:] == [
            ["Alfa Makina", "https://alfa-ornek.com.tr/", None, None, "Örnek Fuarı", "2", "C-7"],
        ]
        assert book["Özet"]["B6"].value == 1
    finally:
        book.close()


def test_listing_fields_fall_back_to_rows_when_snapshots_fail(monkeypatch):
    def broken(_run_id):
        raise RuntimeError("no snapshots")

    monkeypatch.setattr(checkpoint, "load_input_snapshots", broken)
    rows = [{"company": "Alfa Makina", "source_record_id": "ornek:alfa"}]
    assert pipeline_runner._with_listing_fields("t39", rows) == rows

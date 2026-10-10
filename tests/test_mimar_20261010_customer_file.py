"""Talimat 43: the customer copy of a delivery and the Turkish person-name rule."""

from datetime import date
from pathlib import Path

from openpyxl import load_workbook
import pytest

from modules import customer_file, legal_form, run_report
from tools import musteri_dosyasi


@pytest.mark.parametrize("name", [
    "ORNEK GIDA SANAYİ AHMET YILMAZ", "MEHMET KAYA", "ORNEK AJANS - MİNE KAYA",
    "ORNEK DÖŞEME-BEKİR YILMAZ", "ORNEK END. SOĞUTMA SİS. AHMET VOLKAN KAYA", "Beta Ornek Adi Ortaklığı",
    "Gamma Ornek S.C.",
])
def test_names_that_show_a_person(name):
    assert legal_form.names_a_person(name)


@pytest.mark.parametrize("name", [
    "Deniz Expo Fuarcılık", "Gizem Tasarım Atölyesi", "ORNEKPAK", "Ornek Türkiye", "Can Tekstil",
    "ORNEK PLASTİK SAN.VE TİC.LDT.ŞTİ.", "Ornek Mutfak İç ve Dış Tic. Şti.", "Ahmet Yılmaz Ltd. Şti.",
    "Pınar Süt", "",
])
def test_names_that_do_not_show_a_person(name):
    assert not legal_form.names_a_person(name)


@pytest.mark.parametrize("name", ["Ornek Mutfak Tic. Şti.", "Ornek Tarım Kooperatifi", "Ornek Holding"])
def test_turkish_company_forms(name):
    assert legal_form.is_company(name)


ROWS = [
    {
        "company": "ALFA ORNEK GIDA SAN. VE TİC. A.Ş.", "website": "https://alfa-ornek.com.tr/",
        "email": "info@alfa-ornek.com.tr", "phone": "+90 212 555 01 23", "status": "OK_HIGH_CONFIDENCE",
        "listing_source": "Ornek Fuarı katılımcı listesi", "listing_hall": "5", "listing_stand": "B12",
    },
    {
        "company": "Beta Ornek", "website": "https://beta-ornek.com.tr/", "email": "satis@beta-ornek.com.tr",
        "phone": "02125550456", "status": "NEEDS_REVIEW",
    },
    {
        "company": "ORNEK MUTFAK SANAYİ AHMET YILMAZ", "website": "https://gama-ornek.com.tr/",
        "email": "info@gama-ornek.com.tr", "status": "OK_HIGH_CONFIDENCE",
    },
    {
        "company": "Delta Ornek Ltd. Şti.", "website": "https://delta-ornek.com.tr/",
        "email": "info@delta-ornek.com.tr", "phone": "+90 216 555 07 89", "status": "OK_HIGH_CONFIDENCE",
    },
]
TURKISH_RULE = musteri_dosyasi._keep_firm("TR")


def test_customer_rows_keep_only_confident_values_and_leave_out_people_and_opt_outs():
    rows, left_out, opted_out = customer_file.customer_rows(
        ROWS, keep_firm=TURKISH_RULE, opt_out_entries=("delta-ornek.com.tr",),
    )
    assert (left_out, opted_out) == (1, 1)
    assert rows == [
        {
            "Firma": "ALFA ORNEK GIDA SAN. VE TİC. A.Ş.", "Fuar": "Ornek Fuarı", "Salon": "5", "Stant": "B12",
            "Web sitesi": "https://alfa-ornek.com.tr/", "E-posta": "info@alfa-ornek.com.tr",
            "Telefon": "+90 212 555 01 23",
        },
        {
            "Firma": "Beta Ornek", "Fuar": "", "Salon": "", "Stant": "", "Web sitesi": "", "E-posta": "",
            "Telefon": "",
        },
    ]


def test_a_given_fair_name_wins():
    rows, _left_out, _opted_out = customer_file.customer_rows(ROWS[:1], fair="Başka Fuar")
    assert rows[0]["Fuar"] == "Başka Fuar"


def test_polish_runs_keep_companies_only_and_other_runs_leave_out_people():
    assert musteri_dosyasi._keep_firm("PL") is legal_form.is_company
    assert TURKISH_RULE("Beta Ornek") and not TURKISH_RULE("MEHMET KAYA")


def test_run_value_with_a_fair_name():
    assert musteri_dosyasi._split_run(r'runs\abc="Ornek Fuarı"') == (Path(r"runs\abc"), "Ornek Fuarı")
    assert musteri_dosyasi._split_run(r"runs\abc") == (Path(r"runs\abc"), "")


def _sheets(path):
    book = load_workbook(path, read_only=True)
    sheets = {name: [list(row) for row in book[name].iter_rows(values_only=True)] for name in book.sheetnames}
    book.close()
    return sheets


def test_full_customer_file_has_one_list_an_info_sheet_and_a_csv(tmp_path):
    rows, _left_out, _opted_out = customer_file.customer_rows(ROWS, keep_firm=TURKISH_RULE)
    written = customer_file.write_customer_file(tmp_path / "liste.xlsx", rows, lang="tr", today=date(2026, 10, 10))
    assert written == [tmp_path / "liste.xlsx", tmp_path / "liste.csv"]
    sheets = _sheets(tmp_path / "liste.xlsx")
    assert list(sheets) == ["Firmalar", "Bilgi"]
    assert sheets["Firmalar"][0] == ["Firma", "Fuar", "Salon", "Stant", "Web sitesi", "E-posta", "Telefon"]
    assert [row[0] for row in sheets["Firmalar"][1:]] == [ROWS[0]["company"], "Beta Ornek", ROWS[3]["company"]]
    assert sheets["Bilgi"][0][0] == "Tarih: 10.10.2026"
    assert sheets["Bilgi"][1][0] == "Firma sayısı: 3"
    csv_lines = (tmp_path / "liste.csv").read_text(encoding="utf-8-sig").splitlines()
    assert csv_lines[0] == "Firma;Fuar;Salon;Stant;Web sitesi;E-posta;Telefon"
    assert len(csv_lines) == 4


def test_sample_hides_contacts_and_writes_no_csv(tmp_path):
    rows, _left_out, _opted_out = customer_file.customer_rows(ROWS, keep_firm=TURKISH_RULE)
    sample = customer_file.sample_rows(rows, 20)
    assert [row["Firma"] for row in sample] == [ROWS[0]["company"], ROWS[3]["company"]]
    assert sample[0]["E-posta"] == "i***@alfa-ornek.com.tr"
    assert sample[0]["Telefon"] == "*** *** *23"
    written = customer_file.write_customer_file(tmp_path / "ornek.xlsx", sample, lang="pl", sample=True)
    assert written == [tmp_path / "ornek.xlsx"]
    sheets = _sheets(tmp_path / "ornek.xlsx")
    assert list(sheets) == ["Firmy", "Informacje"]
    assert sheets["Firmy"][0][1] == "Targi"
    assert sheets["Informacje"][-1][0].startswith("Próbka:")


def test_delivery_filter_counts_both_rules():
    kept, left_out, opted_out = run_report.deliver(
        run_report.prepare_rows(ROWS), opt_out_entries=("beta-ornek.com.tr",), keep_firm=TURKISH_RULE,
    )
    assert [row["Firma"] for _source, row in kept] == [ROWS[0]["company"], ROWS[3]["company"]]
    assert (left_out, opted_out) == (1, 1)
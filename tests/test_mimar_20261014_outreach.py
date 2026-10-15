"""Talimat 45: cold e-mails from a tracking workbook, and invisible characters in pasted keys."""

from datetime import date, datetime, timezone
import json
import smtplib

from openpyxl import Workbook, load_workbook
import pytest

import config
from modules import api_configuration, customer_file, outreach
from tools import eposta_gonder

# Wednesday 14 October 2026, 10:00 in Istanbul and 08:00 in London.
MORNING = datetime(2026, 10, 14, 7, 0, tzinfo=timezone.utc)
TEMPLATE = (
    "Gönderen: Ornek Gonderen <satis@ornek-gonderen.example>\n"
    "Konu: Fuar öncesi müşteri listesi\n"
    "\n"
    "Merhaba {hitap} ekibi,\n"
    "\n"
    "Kısa metin.\n"
)


@pytest.fixture(autouse=True)
def log_file(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "OUTREACH_LOG_FILE", tmp_path / "eposta_gonderim_kaydi.jsonl")
    monkeypatch.setattr(config, "OUTREACH_DELAY_SEC", (0, 0))
    monkeypatch.setattr(config, "OPT_OUT_FILE", tmp_path / "ret_listesi.txt")
    return config.OUTREACH_LOG_FILE


@pytest.mark.parametrize("name, country, expected", [
    ("ORNEK REKLAM ORGANİZASYON SANAYİ VE TİCARET LİMİTED ŞİRKETİ", "TR", "Ornek Reklam Organizasyon"),
    ("ALFA ORNEK GIDA SAN. VE TİC. A.Ş.", "TR", "Alfa Ornek Gida"),
    ("Delta Ornek Ltd. Şti.", "TR", "Delta Ornek"),
    ("İNCİ ORNEK MİMARLIK", "TR", "İnci Ornek Mimarlık"),
    ("AKYA ORNEK HIZMETLERI IC VE DIS TIC.LTD.STI.", "TR", "Akya Ornek Hizmetleri"),
    ("ORNEK INTERNATIONAL DIS TICARET LTD", "TR", "Ornek International"),
    ("ORNEK DISPLAYS LIMITED", "GB", "Ornek Displays"),
    ("Ornek2Design", "GB", "Ornek2Design"),
    ("ABC Ornek & Co", "GB", "ABC Ornek"),
    ("Ltd", "GB", "Ltd"),
])
def test_short_names_for_the_greeting(name, country, expected):
    assert outreach.salutation(name, country) == expected


def _tracker_file(path, rows):
    book = Workbook()
    sheet = book.active
    sheet.append(["Firma", "Web sitesi", "E-posta", "Telefon", "Gönderim tarihi", "Yanıt", "Yanıt tarihi", "Not"])
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


ROWS = [
    ["ALFA ORNEK GIDA SAN. VE TİC. A.Ş.", "https://alfa-ornek.com.tr/", "info@alfa-ornek.com.tr", "", None, None, None, None],
    ["Beta Ornek", "https://beta-ornek.com.tr/", "ahmet@beta-ornek.com.tr", "", None, None, None, None],
    ["Gama Ornek", "", "satis@gama-ornek.com.tr", "", "hata veriyor", None, None, None],
    ["Delta Ornek", "", "info@delta-ornek.com.tr", "", None, None, None, "Siteyi kontrol edin"],
    ["Epsilon Ornek", "", "info@epsilon-ornek.com.tr", "", None, "hayır", None, None],
    ["Zeta Ornek", "", "info@zeta-ornek.com.tr", "", None, None, None, None],
    ["Zeta Ornek Satış", "", "satis@zeta-ornek.com.tr", "", None, None, None, None],
    ["Eta Ornek", "", "eta-ornek@gmail.com", "", None, None, None, None],
    ["Theta Ornek", "", "info@theta-ornek.com.tr", "", None, None, None, None],
    ["Iota Ornek", "", "", "", None, None, None, None],
]


def test_older_tracker_gets_the_tool_columns_and_greetings(tmp_path):
    path = _tracker_file(tmp_path / "takip.xlsx", ROWS)
    tracker = outreach.Tracker(path)
    assert outreach.fill_salutations(tracker, "TR") == 10
    tracker.save()
    sheet = load_workbook(path).active
    header = [cell.value for cell in sheet[1]]
    assert header[-2:] == ["Hitap adı", "Gönderim sonucu"]
    assert sheet.cell(row=2, column=9).value == "Alfa Ornek Gida"


def test_which_firms_may_get_an_e_mail(tmp_path):
    tracker = outreach.Tracker(_tracker_file(tmp_path / "takip.xlsx", ROWS))
    log = [{"email": "info@theta-ornek.com.tr", "date": "2026-10-13", "result": "gönderildi"}]
    ready, reasons = outreach.next_firms(tracker, opt_out_entries=("@zeta-ornek.com.tr",), log=log)
    assert [row["Firma"] for row in ready] == ["ALFA ORNEK GIDA SAN. VE TİC. A.Ş.", "Eta Ornek"]
    assert reasons == {
        "kişi ya da kişisel veri adresi": 1, "daha önce gönderildi": 1, "Not sütunu dolu": 1, "yanıt var": 1,
        "ret listesinde": 2, "bu adrese daha önce yazıldı": 1, "e-posta yok": 1,
    }


def test_one_e_mail_per_company_domain(tmp_path):
    tracker = outreach.Tracker(_tracker_file(tmp_path / "takip.xlsx", ROWS))
    ready, reasons = outreach.next_firms(tracker, opt_out_entries=(), log=[])
    assert "Zeta Ornek Satış" not in [row["Firma"] for row in ready]
    assert reasons["bu firmaya daha önce yazıldı"] == 1


def test_template_rules(tmp_path):
    path = tmp_path / "sablon.txt"
    path.write_text(TEMPLATE, encoding="utf-8")
    template = outreach.load_template(path)
    assert (template["sender_name"], template["sender_address"]) == ("Ornek Gonderen", "satis@ornek-gonderen.example")
    assert outreach.render(template, {"Hitap adı": "Alfa", "Firma": "ALFA LTD"})[1].startswith("Merhaba Alfa ekibi,")
    for broken, message in (
        (TEMPLATE.replace("satis@ornek-gonderen.example", "[e-posta]"), "geçerli bir e-posta"),
        (TEMPLATE.replace("Kısa metin.", "Ben [Ad Soyad]."), "köşeli parantez"),
        (TEMPLATE.replace("{hitap}", "{firma}"), "{hitap} yok"),
        (TEMPLATE.replace("Konu: Fuar öncesi müşteri listesi\n", ""), "Konu satırı yok"),
    ):
        path.write_text(broken, encoding="utf-8")
        with pytest.raises(outreach.OutreachError, match=message):
            outreach.load_template(path)


def test_english_template_and_message_headers(tmp_path):
    path = tmp_path / "template.txt"
    path.write_text("From: Ornek Sender <hello@ornek-sender.example>\nSubject: Exhibitor lists\n\nHello {hitap} team,\n", encoding="utf-8")
    template = outreach.load_template(path)
    message = outreach.build_message(template, "info@ornek-displays.example", "Exhibitor lists", "Hello Ornek team,\n")
    assert message["From"] == "Ornek Sender <hello@ornek-sender.example>"
    assert message["List-Unsubscribe"] == "<mailto:hello@ornek-sender.example?subject=unsubscribe>"
    assert message["Message-ID"].endswith("@ornek-sender.example>")
    assert message.get_content_type() == "text/plain"


@pytest.mark.parametrize("moment, country, open_", [
    (MORNING, "TR", True),
    (MORNING, "GB", False),
    (datetime(2026, 10, 14, 9, 0, tzinfo=timezone.utc), "GB", True),
    (datetime(2026, 10, 14, 14, 30, tzinfo=timezone.utc), "TR", False),
    (datetime(2026, 10, 17, 9, 0, tzinfo=timezone.utc), "TR", False),
])
def test_office_hours_of_the_recipients(moment, country, open_):
    assert outreach.office_hours(moment, country) is open_


class FakeServer:
    sent: list = []
    refuse: set = set()
    fail: Exception | None = None
    logins: list = []

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def login(self, address, password):
        FakeServer.logins.append((address, password))

    def send_message(self, message):
        if FakeServer.fail:
            raise FakeServer.fail
        if message["To"] in FakeServer.refuse:
            raise smtplib.SMTPRecipientsRefused({message["To"]: (550, b"no such user")})
        FakeServer.sent.append(message)


@pytest.fixture
def server():
    FakeServer.sent, FakeServer.refuse, FakeServer.fail, FakeServer.logins = [], set(), None, []
    return FakeServer


def _template(tmp_path):
    path = tmp_path / "sablon.txt"
    path.write_text(TEMPLATE, encoding="utf-8")
    return outreach.load_template(path)


def test_sending_marks_the_workbook_and_the_log_and_respects_the_daily_limit(tmp_path, server):
    path = _tracker_file(tmp_path / "takip.xlsx", ROWS)
    tracker = outreach.Tracker(path)
    outreach.fill_salutations(tracker, "TR")
    server.refuse = {"eta-ornek@gmail.com"}
    result = outreach.send_batch(
        tracker, _template(tmp_path), country="TR", limit=20, password="synthetic-password", opt_out_entries=(),
        smtp_factory=FakeServer, now=lambda: MORNING, sleep=lambda _seconds: None, report=lambda _line: None,
    )
    assert result == {"sent": 3, "refused": 1, "stopped": ""}
    assert [message["To"] for message in server.sent] == [
        "info@alfa-ornek.com.tr", "info@zeta-ornek.com.tr", "info@theta-ornek.com.tr",
    ]
    assert server.sent[0].get_content().startswith("Merhaba Alfa Ornek Gida ekibi,")
    assert server.logins[0] == ("satis@ornek-gonderen.example", "synthetic-password")
    sheet = load_workbook(path).active
    columns = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(row=2, column=columns["Gönderim sonucu"]).value == "gönderildi"
    assert sheet.cell(row=2, column=columns["Gönderim tarihi"]).value.date() == MORNING.astimezone().date()
    assert sheet.cell(row=9, column=columns["Gönderim sonucu"]).value == "adres reddedildi (550)"
    entries = [json.loads(line) for line in config.OUTREACH_LOG_FILE.read_text(encoding="utf-8").splitlines()]
    assert [entry["result"] for entry in entries] == ["gönderildi", "gönderildi", "adres reddedildi (550)", "gönderildi"]
    assert "synthetic-password" not in config.OUTREACH_LOG_FILE.read_text(encoding="utf-8")
    again = outreach.send_batch(
        outreach.Tracker(path), _template(tmp_path), country="TR", limit=20, password="synthetic-password",
        opt_out_entries=(), smtp_factory=FakeServer, now=lambda: MORNING, sleep=lambda _seconds: None,
        report=lambda _line: None,
    )
    assert again == {"sent": 0, "refused": 0, "stopped": "gönderilecek firma yok"}


def test_daily_limit_counts_every_workbook_and_test_e_mails(tmp_path, server):
    today = MORNING.astimezone().date().isoformat()
    with config.OUTREACH_LOG_FILE.open("w", encoding="utf-8") as handle:
        for index in range(19):
            handle.write(json.dumps({"email": f"info@baska-{index}.example", "date": today, "result": "gönderildi"}) + "\n")
        handle.write(json.dumps({"email": "kendi@ornek-gonderen.example", "date": today, "result": "deneme"}) + "\n")
    path = _tracker_file(tmp_path / "takip.xlsx", ROWS)
    result = outreach.send_batch(
        outreach.Tracker(path), _template(tmp_path), country="TR", limit=21, password="p", opt_out_entries=(),
        smtp_factory=FakeServer, now=lambda: MORNING, sleep=lambda _seconds: None, report=lambda _line: None,
    )
    assert result == {"sent": 1, "refused": 0, "stopped": "günlük sınır doldu"}
    assert outreach.daily_room(outreach.load_log(), 500, MORNING.astimezone().date()) == config.OUTREACH_MAX_DAILY_LIMIT - 21


def test_outside_office_hours_nothing_is_sent(tmp_path, server):
    path = _tracker_file(tmp_path / "takip.xlsx", ROWS)
    result = outreach.send_batch(
        outreach.Tracker(path), _template(tmp_path), country="GB", limit=20, password="p", opt_out_entries=(),
        smtp_factory=FakeServer, now=lambda: MORNING, sleep=lambda _seconds: None, report=lambda _line: None,
    )
    assert result["stopped"] == "mesai saati dışında" and server.sent == []


@pytest.mark.parametrize("error, message", [
    (smtplib.SMTPAuthenticationError(535, b"bad"), "Uygulama şifresini"),
    (smtplib.SMTPDataError(554, b"blocked"), "kod 554"),
    (smtplib.SMTPServerDisconnected("gone"), "bağlanılamadı"),
])
def test_a_server_refusal_stops_the_run(tmp_path, server, error, message):
    server.fail = error
    path = _tracker_file(tmp_path / "takip.xlsx", ROWS)
    with pytest.raises(outreach.OutreachError, match=message):
        outreach.send_batch(
            outreach.Tracker(path), _template(tmp_path), country="TR", limit=20, password="p", opt_out_entries=(),
            smtp_factory=FakeServer, now=lambda: MORNING, sleep=lambda _seconds: None, report=lambda _line: None,
        )
    assert outreach.load_log() == []


def test_a_sent_e_mail_reaches_the_workbook_from_the_log(tmp_path):
    path = _tracker_file(tmp_path / "takip.xlsx", ROWS)
    outreach.append_log({"email": "info@alfa-ornek.com.tr", "date": "2026-10-14", "result": "gönderildi"})
    tracker = outreach.Tracker(path)
    assert outreach.sync_from_log(tracker, outreach.load_log()) == 1
    row = tracker.rows()[0]
    assert (row["Gönderim tarihi"], row["Gönderim sonucu"]) == (date(2026, 10, 14), "gönderildi")


def test_a_test_e_mail_goes_to_the_user_only(tmp_path, server):
    tracker = outreach.Tracker(_tracker_file(tmp_path / "takip.xlsx", ROWS))
    ready, _reasons = outreach.next_firms(tracker, opt_out_entries=(), log=[])
    outreach.send_test(_template(tmp_path), ready[0], to="kendi@ornek-gonderen.example", password="p",
                       smtp_factory=FakeServer, now=lambda: MORNING)
    assert [message["To"] for message in server.sent] == ["kendi@ornek-gonderen.example"]
    assert server.sent[0]["Subject"] == "[DENEME] Fuar öncesi müşteri listesi"
    assert outreach.written_addresses(outreach.load_log()) == (set(), set())


def test_new_tracker_from_a_customer_file(tmp_path):
    rows = [
        {"Firma": "ORNEK DISPLAYS LIMITED", "Fuar": "", "Salon": "", "Stant": "", "Web sitesi": "https://ornek-displays.example/",
         "E-posta": "info@ornek-displays.example", "Telefon": "020 7946 0123"},
        {"Firma": "Beta Ornek Ltd", "Fuar": "", "Salon": "", "Stant": "", "Web sitesi": "", "E-posta": "", "Telefon": ""},
    ]
    listing = tmp_path / "liste.xlsx"
    customer_file.write_customer_file(listing, rows, lang="en", today=date(2026, 10, 14))
    target = tmp_path / "takip.xlsx"
    assert outreach.new_tracker(target, outreach.read_customer_list(listing), "GB") == 1
    row = outreach.Tracker(target).rows()[0]
    assert (row["Firma"], row["Hitap adı"], row["E-posta"]) == ("ORNEK DISPLAYS LIMITED", "Ornek Displays", "info@ornek-displays.example")
    with pytest.raises(outreach.OutreachError, match="zaten var"):
        outreach.new_tracker(target, [], "GB")


def test_preview_command_writes_the_e_mails_and_sends_nothing(tmp_path, server, capsys):
    path = _tracker_file(tmp_path / "takip.xlsx", ROWS)
    template = tmp_path / "sablon.txt"
    template.write_text(TEMPLATE, encoding="utf-8")
    assert eposta_gonder.main(["onizle", "--takip", str(path), "--sablon", str(template), "--adet", "2"]) == 0
    text = (tmp_path / "takip_onizleme.txt").read_text(encoding="utf-8")
    assert "Kime: info@alfa-ornek.com.tr" in text and "Merhaba Alfa Ornek Gida ekibi," in text
    assert "Gönderilebilecek firma: 4" in capsys.readouterr().out
    assert server.sent == []


def test_pasted_keys_lose_invisible_characters():
    assert api_configuration.prompt_api_key("Ornek", lambda _prompt: "\x02abcd-1234\x16 ") == "abcd-1234"

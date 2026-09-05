from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

import config
from modules import publication_policy
from prepare_remediation_run import SOURCE_ONLY_FIELDS, prepare_remediation_run
from reconcile_runs import reconcile_runs


CONFIG_HASH = "a" * 64


def _input(path: Path, count: int = 893) -> list[str]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["company", "source_record_id", "sector", "listed_legal_name", "email", "website", "publication_eligible"])
    ids = []
    for index in range(count):
        source_id = f"src:{index:03d}"
        ids.append(source_id)
        sheet.append([f"Company {index:03d}", source_id, "textile", f"Company {index:03d} Ltd", "forbidden@example", "https://forbidden.example", True])
    workbook.save(path)
    workbook.close()
    return ids


def _envelope(source_id: str, run_id: str, publishable: bool) -> dict:
    row = {
        "source_record_id": source_id, "run_id": run_id, "config_sha256": CONFIG_HASH,
        "status": "REVIEW_NEEDED", "publication_eligible": False,
        "publication_advisory_eligible": False, "publication_blockers": "review",
        "reason": "review", "__evaluation": {},
    }
    decision = publication_policy.freeze_publication_decision(row, {}, config_sha256=CONFIG_HASH)
    if publishable:
        decision = dict(decision)
        decision.update({
            "publishable": True, "blockers": [], "advisory_eligible": True,
            "website_identity_verified": True, "allowed_contact_fields": ["email", "phone"],
        })
        decision["publication_decision_sha256"] = hashlib.sha256(
            json.dumps({key: value for key, value in decision.items() if key != "publication_decision_sha256"}, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
        ).hexdigest()
    return decision


def _bundle(tmp_path: Path, ids: list[str], *, name: str, publishable_ids: set[str] | None = None, status_by_id: dict[str, str] | None = None):
    publishable_ids = publishable_ids or set()
    status_by_id = status_by_id or {}
    run_id = f"{name}-run"
    result_rows = []
    evidence_rows = []
    for index, source_id in enumerate(ids):
        decision = _envelope(source_id, run_id, source_id in publishable_ids)
        status = status_by_id.get(source_id, "OK_HIGH_CONFIDENCE" if source_id in publishable_ids else "REVIEW_NEEDED")
        blockers = "; ".join(decision["blockers"])
        row = {
            "source_record_id": source_id, "run_id": run_id, "company": f"Company {index:03d}",
            "status": status, "publication_eligible": decision["publishable"],
            "publication_blockers": blockers, "publication_decision": decision,
            "config_sha256": CONFIG_HASH,
        }
        result_rows.append(row)
        evidence_rows.append({
            "source_record_id": source_id, "run_id": run_id, "config_sha256": CONFIG_HASH,
            "publication_decision": decision, "source_evidence": [],
        })
    results_path = tmp_path / f"{name}.json"
    evidence_path = tmp_path / f"{name}.evidence.jsonl"
    manifest_path = tmp_path / f"{name}.manifest.json"
    results_path.write_text(json.dumps(result_rows, ensure_ascii=False), encoding="utf-8")
    evidence_path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in evidence_rows), encoding="utf-8")
    manifest_path.write_text(json.dumps({
        "run_id": run_id, "source_record_ids": ids,
        "files": {
            results_path.name: {"sha256": hashlib.sha256(results_path.read_bytes()).hexdigest()},
            evidence_path.name: {"sha256": hashlib.sha256(evidence_path.read_bytes()).hexdigest()},
        },
    }), encoding="utf-8")
    return results_path, evidence_path, manifest_path


def test_prepare_remediation_is_source_only_and_covers_review_ids(tmp_path: Path):
    input_path = tmp_path / "input.xlsx"
    ids = _input(input_path)
    baseline = _bundle(tmp_path, ids, name="baseline", publishable_ids=set(ids[:416]))
    result = prepare_remediation_run(
        input_path=input_path, baseline_results=baseline[0], baseline_evidence=baseline[1], baseline_manifest=baseline[2], destination=tmp_path / "child",
    )
    assert result["task_count"] == 477
    plan = json.loads((tmp_path / "child" / "remediation_plan.json").read_text(encoding="utf-8"))
    assert [task["source_record_id"] for task in plan["tasks"]] == ids[416:]
    workbook = load_workbook(tmp_path / "child" / "remediation_input.xlsx", read_only=True)
    try:
        headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
        assert headers == list(SOURCE_ONLY_FIELDS)
        assert workbook.active.max_row == 478
        assert all(str(header).casefold() not in {"email", "website", "publication_eligible"} for header in headers)
    finally:
        workbook.close()


def test_reconcile_accepts_477_child_and_preserves_failed_baseline(tmp_path: Path):
    input_path = tmp_path / "input.xlsx"
    ids = _input(input_path)
    baseline = _bundle(tmp_path, ids, name="baseline", publishable_ids=set(ids[:416]))
    child_ids = ids[416:]
    plan_path = tmp_path / "remediation_plan.json"
    plan_path.write_text(json.dumps({"tasks": [{"source_record_id": value} for value in child_ids]}), encoding="utf-8")
    child = _bundle(tmp_path, child_ids, name="child", status_by_id={child_ids[0]: "FAILED", child_ids[1]: "UNKNOWN"})
    manifest = reconcile_runs(
        input_path=input_path, baseline_results=baseline[0], baseline_evidence=baseline[1], baseline_manifest=baseline[2],
        child_results=child[0], child_evidence=child[1], child_manifest=child[2], remediation_plan=plan_path,
        destination=tmp_path / "delivery",
    )
    assert manifest["counts"] == {"all_results": 893, "contacts": 416, "review": 477}
    coverage = json.loads((tmp_path / "delivery" / "discovery_coverage.json").read_text(encoding="utf-8"))
    assert coverage["source_count"] == 893
    assert set(coverage["statuses"]) == {"published", "review", "pending", "failed"}
    assert sum(coverage["statuses"].values()) == 893


def test_reconcile_rejects_child_order_or_size_mismatch(tmp_path: Path):
    input_path = tmp_path / "input.xlsx"
    ids = _input(input_path)
    baseline = _bundle(tmp_path, ids, name="baseline", publishable_ids=set(ids[:416]))
    plan_path = tmp_path / "remediation_plan.json"
    plan_path.write_text(json.dumps({"tasks": [{"source_record_id": value} for value in ids[416:]]}), encoding="utf-8")
    child = _bundle(tmp_path, ids[416:-1], name="child")
    import pytest
    with pytest.raises(ValueError, match="477|plan order"):
        reconcile_runs(
            input_path=input_path, baseline_results=baseline[0], baseline_evidence=baseline[1], baseline_manifest=baseline[2],
            child_results=child[0], child_evidence=child[1], child_manifest=child[2], remediation_plan=plan_path,
            destination=tmp_path / "delivery",
        )


def test_zero_publishable_baseline_derives_full_893_task_plan(tmp_path: Path):
    input_path = tmp_path / "input.xlsx"
    ids = _input(input_path)
    baseline = _bundle(tmp_path, ids, name="baseline_all_review", publishable_ids=set())
    result = prepare_remediation_run(
        input_path=input_path,
        baseline_results=baseline[0],
        baseline_evidence=baseline[1],
        baseline_manifest=baseline[2],
        destination=tmp_path / "child_all_review",
    )
    assert result["task_count"] == 893
    plan = json.loads((tmp_path / "child_all_review" / "remediation_plan.json").read_text(encoding="utf-8"))
    assert [task["source_record_id"] for task in plan["tasks"]] == ids

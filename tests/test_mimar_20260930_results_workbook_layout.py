from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pytest

from modules import checkpoint, run_report


def test_contact_and_detail_headers_partition_all_headers():
    assert set(run_report.CONTACT_HEADERS) | set(run_report.DETAIL_HEADERS) == set(run_report.HEADERS)
    assert set(run_report.CONTACT_HEADERS) & set(run_report.DETAIL_HEADERS) == {"Firma"}
    assert len(run_report.DETAIL_HEADERS) == 19


def test_sheet_order_and_contact_sheet_active(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    assert run_report.write_run_report(
        [], output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=False)
    try:
        assert book.sheetnames == ["İletişim", "Özet", "Detaylar"]
        assert book.active.title == "İletişim"
    finally:
        book.close()


def test_contact_and_detail_rows_keep_input_order(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    rows = [{"company": company} for company in ("Gama", "Alfa", "Beta")]
    assert run_report.write_run_report(
        rows, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=False)
    try:
        for name in ("İletişim", "Detaylar"):
            assert [book[name].cell(row, 1).value for row in range(2, 5)] == ["Gama", "Alfa", "Beta"]
    finally:
        book.close()


def test_low_confidence_value_is_yellow(tmp_path):
    rows = [
        {"Firma": "Alfa", "Web sitesi": "https://alfa.ornek.com.tr/", "Web güven": "LOW"},
        {"Firma": "Beta", "Web sitesi": "https://beta.ornek.com.tr/", "Web güven": "HIGH"},
        {"Firma": "Gama", "Web sitesi": "", "Web güven": "LOW"},
    ]
    path = tmp_path / "sonuclar.xlsx"
    run_report.write_results_workbook(
        path, rows, [], run_status="TAMAMLANDI", generated_at="2026-09-30T10:00:00+00:00",
    )
    book = load_workbook(path, read_only=False)
    try:
        sheet = book["İletişim"]
        assert sheet["B2"].fill.fgColor.rgb.endswith("FFF2CC")
        assert not sheet["B3"].fill.fgColor.rgb.endswith("FFF2CC")
        assert not sheet["B4"].fill.fgColor.rgb.endswith("FFF2CC")
    finally:
        book.close()


def test_website_hyperlink_only_for_http_urls(tmp_path):
    rows = [
        {"Firma": "Alfa", "Web sitesi": "https://alfa.ornek.com.tr/"},
        {"Firma": "Beta", "Web sitesi": "www.beta.ornek.com.tr"},
        {"Firma": "Gama", "Web sitesi": '=HYPERLINK("x")'},
    ]
    path = tmp_path / "sonuclar.xlsx"
    run_report.write_results_workbook(
        path, rows, [], run_status="TAMAMLANDI", generated_at="2026-09-30T10:00:00+00:00",
    )
    book = load_workbook(path, read_only=False)
    try:
        sheet = book["İletişim"]
        assert sheet["B2"].hyperlink.target == "https://alfa.ornek.com.tr/"
        assert sheet["B3"].hyperlink is None
        assert sheet["B4"].hyperlink is None
        assert sheet["B4"].value.startswith("'")
    finally:
        book.close()


def test_detail_labels_translate_codes_and_keep_unknown():
    assert run_report._detail_value("Web güven", "HIGH") == "Yüksek"
    assert run_report._detail_value("Web güven", "NONE") == ""
    assert run_report._detail_value("Web kaynağı", "REFERENCE_VERIFIED") == "Fuar listesi (doğrulandı)"
    assert run_report._detail_value("Durum", "YENI_KOD") == "YENI_KOD"
    assert run_report._detail_value("Eksik alanlar", "website;email") == "web sitesi, e-posta"
    assert run_report._detail_value("Referans sinyalleri", "S1_phone,S3_name") == "telefon, ad"
    assert run_report._detail_value("Yayına hazır", True) == "Evet"
    assert run_report._detail_value("Yayına hazır", False) == "Hayır"
    assert run_report._detail_value("Not", None) == ""


def test_technical_columns_hidden_in_outline(tmp_path):
    path = tmp_path / "sonuclar.xlsx"
    run_report.write_results_workbook(
        path, [{"Firma": "Alfa"}], [], run_status="TAMAMLANDI", generated_at="2026-09-30T10:00:00+00:00",
    )
    book = load_workbook(path, read_only=False)
    try:
        sheet = book["Detaylar"]
        for header in ("Durum", "Ücretli durum", "Not"):
            dimension = sheet.column_dimensions[get_column_letter(run_report.DETAIL_HEADERS.index(header) + 1)]
            assert dimension.hidden is True
            assert dimension.outlineLevel == 1
        company = sheet.column_dimensions[get_column_letter(run_report.DETAIL_HEADERS.index("Firma") + 1)]
        assert company.hidden is False
    finally:
        book.close()


def test_summary_counts_match_final_stage(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    rows = [
        {"company": "Alfa", "website": "https://alfa.ornek.com.tr/", "email": "info@alfa.ornek.com.tr", "status": "OK_HIGH_CONFIDENCE"},
        {"company": "Beta", "website": "https://beta.ornek.com.tr/", "phone": "0212 555 00 00", "status": "REVIEW_NEEDED"},
        {"company": "Gama", "website": "", "candidate_1_url": "https://aday.ornek.com.tr/", "status": "WEBSITE_NOT_FOUND"},
    ]
    assert run_report.write_run_report(
        rows, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=False)
    try:
        summary = book["Özet"]
        assert summary["B3"].value == 3
        assert summary["B8"].value == 1
        assert summary["C8"].value == pytest.approx(1 / 3)
        assert summary["C8"].number_format == "0.0%"
        assert summary["B9"].value == 1
        assert summary["B10"].value == 0
        assert summary["B11"].value == 1
        assert summary["B12"].value == 1
        assert summary["B8"].value == int(summary["B18"].value.split()[0])
        assert summary["B9"].value == int(summary["C18"].value.split()[0])
        assert summary["B10"].value == int(summary["D18"].value.split()[0])
        assert summary["A20"].value == "Açıklama"
        contact = book["İletişim"]
        assert contact["B3"].fill.fgColor.rgb.endswith("FFF2CC")
        assert contact["D3"].fill.fgColor.rgb.endswith("FFF2CC")
        assert not contact["B2"].fill.fgColor.rgb.endswith("FFF2CC")
    finally:
        book.close()

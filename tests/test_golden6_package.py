import json
import shutil
import tempfile
import unittest
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

from modules import scorer
import run_golden_6
from validate_benchmark_suite import validate_manifest
from validate_golden_xlsx import _sheet_rows, assertion_coverage, readiness_issues
from fixture_factory import benchmark_manifest, golden, path as fixture_path


ROOT = fixture_path()
GOLDEN6_DIR = golden(6).parent


def _names(path: Path, sheet: str | None = None, column: str = "Company") -> set[str]:
    return {
        scorer.normalize_text(str(row.get(column) or "")).strip()
        for row in _sheet_rows(path, sheet)
        if str(row.get(column) or "").strip()
    }


class Golden6PackageTests(unittest.TestCase):
    @staticmethod
    def _ids(path: Path, sheet: str | None = None) -> list[str]:
        return [
            str(row.get("source_record_id") or "").strip()
            for row in _sheet_rows(path, sheet)
        ]

    def test_golden6_has_twenty_unique_companies_with_no_prior_overlap_and_private_gate(self):
        manual = GOLDEN6_DIR / "golden_6_manual_validation_20_ready.xlsx"
        current = _names(manual, "Manual Report")
        prior = set()
        for path in (
            fixture_path("outputs/golden_manual_validation_20260713/golden_manual_validation_30.xlsx"),
            fixture_path("outputs/golden_2_20260714/golden_2_manual_validation_30.xlsx"),
            golden(3), golden(4), golden(5),
        ):
            prior.update(_names(path, "Manual Report"))
        self.assertEqual(len(current), 20)
        self.assertFalse(current & prior)

        splits = json.loads(benchmark_manifest().read_text(encoding="utf-8"))
        g6_set = next(s for s in splits["sets"] if s.get("name") == "golden_6")
        self.assertIs(g6_set.get("private_seen_check"), True)

    def test_source_assisted_input_keeps_official_fair_profiles_for_diagnostics(self):
        rows = _sheet_rows(GOLDEN6_DIR / "golden_6_pipeline_input_20.xlsx")
        self.assertEqual(len(rows), 20)
        profiles = set()
        for row in rows:
            self.assertEqual(row["source"], "automechanika_istanbul_2026")
            self.assertEqual(row["country"], "Türkiye")
            self.assertFalse(row.get("website"))
            self.assertTrue(row.get("sector"))
            self.assertIn("Automechanika Istanbul 2026", str(row.get("description") or ""))
            profile = str(row.get("profile_url") or "")
            self.assertEqual(urlparse(profile).hostname, "www.automechanikaistanbulplus.com")
            self.assertTrue(urlparse(profile).path.startswith("/company/"))
            profiles.add(profile)
        self.assertEqual(len(profiles), 20)

    def test_discovery_blind_input_removes_every_profile_shortcut(self):
        rows = _sheet_rows(GOLDEN6_DIR / "golden_6_discovery_blind_input_20.xlsx")
        self.assertEqual(len(rows), 20)
        for row in rows:
            self.assertEqual(row["source"], "automechanika_istanbul_2026")
            self.assertEqual(row["country"], "Türkiye")
            self.assertFalse(row.get("website"))
            self.assertFalse(row.get("profile_url"))
            self.assertTrue(row.get("sector"))

    def test_manual_report_and_discovery_blind_input_have_identical_companies(self):
        manual = _names(GOLDEN6_DIR / "golden_6_manual_validation_20_ready.xlsx", "Manual Report")
        pipeline = _names(GOLDEN6_DIR / "golden_6_discovery_blind_input_20.xlsx", column="company")
        self.assertEqual(manual, pipeline)

    def test_golden6_has_deterministic_ordered_ids_in_all_three_workbooks(self):
        expected_ids = self._ids(GOLDEN6_DIR / "golden_6_manual_validation_20_ready.xlsx", "Manual Report")
        self.assertEqual(expected_ids, [f"synthetic_golden6:{index:04d}" for index in range(20)])
        self.assertEqual(
            expected_ids,
            self._ids(GOLDEN6_DIR / "golden_6_pipeline_input_20.xlsx"),
        )
        self.assertEqual(
            expected_ids,
            self._ids(GOLDEN6_DIR / "golden_6_discovery_blind_input_20.xlsx"),
        )

    @staticmethod
    def _manifest_with_mutation(tmp_root: Path, mutate) -> list[str]:
        output = tmp_root / "outputs" / "golden_6_20260718"
        output.mkdir(parents=True)
        for filename in (
            "golden_6_manual_validation_20_ready.xlsx",
            "golden_6_pipeline_input_20.xlsx",
            "golden_6_discovery_blind_input_20.xlsx",
        ):
            shutil.copy2(GOLDEN6_DIR / filename, output / filename)
        mutate(output)
        manifest = tmp_root / "data" / "benchmark_splits.json"
        manifest.parent.mkdir()
        manifest.write_text(json.dumps({
            "version": 1,
            "policy": {"private_seen_expected_unique_companies": 71},
            "sets": [{
                "name": "golden_6",
                "role": "blind_golden6",
                "expected": "outputs/golden_6_20260718/golden_6_manual_validation_20_ready.xlsx",
                "pipeline_input": "outputs/golden_6_20260718/golden_6_discovery_blind_input_20.xlsx",
                "source_assisted_input": "outputs/golden_6_20260718/golden_6_pipeline_input_20.xlsx",
                "private_seen_check": True,
                "requires_source_record_id": True,
                "status": "manual_validation_pending",
            }],
        }), encoding="utf-8")
        _, issues = validate_manifest(manifest)
        return issues

    def test_required_ids_reject_missing_id(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            issues = self._manifest_with_mutation(root, lambda output: self._set_cell(
                output / "golden_6_manual_validation_20_ready.xlsx", "Manual Report", "H2", None
            ))
        self.assertTrue(any("incomplete source_record_id coverage" in issue for issue in issues))

    def test_required_ids_reject_duplicate_id(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            issues = self._manifest_with_mutation(root, lambda output: self._set_cell(
                output / "golden_6_discovery_blind_input_20.xlsx", "Pipeline Input", "H3",
                "synthetic_golden6:0000",
            ))
        self.assertTrue(any("duplicate source_record_id" in issue for issue in issues))

    def test_required_ids_reject_wrong_order(self):
        def swap(output):
            path = output / "golden_6_discovery_blind_input_20.xlsx"
            workbook = load_workbook(path)
            sheet = workbook["Pipeline Input"]
            sheet["H2"].value, sheet["H3"].value = sheet["H3"].value, sheet["H2"].value
            workbook.save(path)
            workbook.close()

        with tempfile.TemporaryDirectory() as directory:
            issues = self._manifest_with_mutation(Path(directory), swap)
        self.assertTrue(any("order does not match expected" in issue for issue in issues))

    def test_required_ids_reject_different_id(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            issues = self._manifest_with_mutation(root, lambda output: self._set_cell(
                output / "golden_6_manual_validation_20_ready.xlsx", "Manual Report", "H2",
                "synthetic_other:0000",
            ))
        self.assertTrue(any("order does not match expected" in issue for issue in issues))

    @staticmethod
    def _set_cell(path: Path, sheet: str, coordinate: str, value) -> None:
        workbook = load_workbook(path)
        workbook[sheet][coordinate].value = value
        workbook.save(path)
        workbook.close()

    def test_golden6_runner_uses_discovery_blind_input(self):
        self.assertEqual(
            run_golden_6.GOLDEN_INPUT.name,
            "golden_6_discovery_blind_input_20.xlsx",
        )

    def test_completed_manual_workbook_is_ready_for_blind_run(self):
        issues = readiness_issues(
            GOLDEN6_DIR / "golden_6_manual_validation_20_ready.xlsx"
        )
        self.assertEqual(issues, [])

    def test_manual_assertion_coverage_preserves_precision_gate(self):
        coverage = assertion_coverage(
            GOLDEN6_DIR / "golden_6_manual_validation_20_ready.xlsx"
        )
        self.assertEqual(coverage["website"], {"asserted": 20, "unknown": 0, "missing": 0})
        self.assertEqual(coverage["email"], {"asserted": 20, "unknown": 0, "missing": 0})
        self.assertEqual(coverage["phone"], {"asserted": 20, "unknown": 0, "missing": 0})

    def test_pending_blind_set_keeps_benchmark_manifest_valid(self):
        _, issues = validate_manifest(benchmark_manifest())
        self.assertEqual(issues, [])


if __name__ == "__main__":
    unittest.main()

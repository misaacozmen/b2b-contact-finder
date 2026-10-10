"""Talimat 41: the Polish profile delivers only firms whose name shows a company form."""

from openpyxl import load_workbook
import pytest

import config
from modules import checkpoint, country_profile, legal_form, pipeline_runner, run_report


@pytest.mark.parametrize("name", [
    "ALFA SPÓŁKA Z OGRANICZONĄ ODPOWIEDZIALNOŚCIĄ", "Beta Sp. z o.o.", "GAMMA S.A.", "DELTA SPÓŁKA JAWNA",
    "Epsilon sp. k.", "ZETA SPÓŁKA KOMANDYTOWO-AKCYJNA", "Eta GmbH", "Theta Ltd.", "Iota BV (MARKA)",
    "SAS KAPPA INDUSTRIE", "Lambda s.r.o.", "MU ODDZIAŁ W POLSCE", "OKRĘG WARSZAWSKI ZWIĄZKU ORNEK",
    "Nu Anonim Şirketi", "Xi Handelsgesellschaft mbH",
])
def test_company_forms(name):
    assert legal_form.is_company(name)


@pytest.mark.parametrize("name", [
    "JAN KOWALSKI", "PPHU ALFA Jan Kowalski", "ALFA SPÓŁKA CYWILNA JAN KOWALSKI, ADAM NOWAK",
    "Beta S.C.", "GAMMA SP. CYWILNA NOWAK, KOWALSKI", "Delta e.K.", "LAMINORNEK", "Instytut Urody Anna Nowak", "",
])
def test_sole_traders_and_unknown_forms(name):
    assert not legal_form.is_company(name)


ROWS = [
    {"company": "ALFA ORNEK SPÓŁKA Z OGRANICZONĄ ODPOWIEDZIALNOŚCIĄ", "website": "https://alfa-ornek.pl/"},
    {"company": "PPHU BETA Jan Kowalski", "website": "https://beta-ornek.pl/"},
    {"company": "GAMMA ORNEK S.C.", "website": "https://gamma-ornek.pl/"},
    {"company": "Delta Ornek GmbH", "website": "https://delta-ornek.de/"},
]


def _workbook(path):
    book = load_workbook(path / "sonuclar.xlsx", read_only=True, data_only=True)
    sheets = {name: [list(row) for row in book[name].iter_rows(values_only=True)] for name in book.sheetnames}
    book.close()
    return sheets


def test_companies_only_leaves_sole_traders_out_of_every_sheet(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    assert run_report.write_run_report(
        ROWS, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
        companies_only=True,
    ) == {}
    sheets = _workbook(tmp_path)
    for name in ("İletişim", "Detaylar"):
        assert [row[0] for row in sheets[name][1:]] == [ROWS[0]["company"], ROWS[3]["company"]]
    notes = [row[0] for row in sheets["Özet"] if row and row[0]]
    assert "Yalnız şirketler: şahıs firması ya da adi ortaklık olduğu için çıkarılan firma: 2." in notes
    assert sheets["Özet"][2][1] == 2


def test_stage_table_counts_only_the_delivered_firms(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    rows = [{**row, "status": "OK_HIGH_CONFIDENCE"} for row in ROWS]
    assert run_report.write_run_report(
        rows, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
        companies_only=True, opt_out_entries=("delta-ornek.de",),
    ) == {}
    summary = _workbook(tmp_path)["Özet"]
    assert summary[2][1] == 1
    final = next(row for row in summary if row[0] == "Final (A + Referans + Ücretli)")
    assert final[1] == "1 (100.0%)"


def test_without_companies_only_every_firm_stays(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    assert run_report.write_run_report(
        ROWS, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    sheets = _workbook(tmp_path)
    assert len(sheets["İletişim"]) == 5
    assert not any(str(row[0] or "").startswith("Yalnız şirketler") for row in sheets["Özet"])


def test_polish_profile_turns_the_rule_on_and_turkish_profile_off():
    try:
        country_profile.apply("PL")
        assert config.DELIVERY_COMPANIES_ONLY is True
    finally:
        country_profile.apply("TR")
    assert config.DELIVERY_COMPANIES_ONLY is False


def test_run_report_follows_the_profile_setting(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DELIVERY_COMPANIES_ONLY", True)
    monkeypatch.setattr(config, "OPT_OUT_FILE", tmp_path / "ret_listesi.txt")
    monkeypatch.setattr(pipeline_runner, "_durable_output_rows", lambda _run_id, _fallback=None: [dict(row) for row in ROWS])
    monkeypatch.setattr(checkpoint, "load_input_snapshots", lambda _run_id: {})
    monkeypatch.setattr(checkpoint, "derive_telemetry", lambda _run_id: {})
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    output = tmp_path / "run" / "output"
    assert pipeline_runner._write_run_report(
        run_id="t41", output_root=output, fallback=None,
        run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    assert [row[0] for row in _workbook(output)["İletişim"][1:]] == [ROWS[0]["company"], ROWS[3]["company"]]

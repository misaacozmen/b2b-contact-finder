from __future__ import annotations

import hashlib
import json
import shutil
import sqlite3
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook

import config
from modules import checkpoint, output_artifacts, run_context, runtime
from prepare_paid_continuation import prepare_paid_continuation


PROVIDERS = {name: 0 for name in checkpoint.CANONICAL_PROVIDERS}


def _payload(source_id: str, **extra: object) -> dict:
    return {"company": source_id, "source_record_id": source_id, "status": "OK", **extra}


def _payload_text(payload: dict) -> str:
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _bundle(tmp_path: Path, *, paid_pending: bool = False, quarantine: bool = False) -> tuple[Path, str, list[str], list[dict]]:
    root = tmp_path / "runs" / ("a" * 64)
    state = root / "state"
    output = root / "output"
    state.mkdir(parents=True)
    (output / "artifacts").mkdir(parents=True)
    source_ids = ["input:0", "input:1"]
    payloads = [_payload(source_id) for source_id in source_ids]
    items = []
    results = []
    for index, source_id in enumerate(source_ids):
        text = _payload_text(payloads[index])
        items.append({
            "item_index": index, "source_record_id": source_id,
            "free_state": "DONE", "paid_required": paid_pending,
            "paid_state": "PENDING" if paid_pending else "NOT_REQUIRED",
            "payload_sha256": hashlib.sha256(text.encode()).hexdigest(),
            "quarantine_state": "LEGACY_RECOVERY_PROVISIONAL_V2" if quarantine else "",
            "quarantine_status": "blocked" if quarantine else "",
            "publication_blockers": "legacy_recovery_provisional" if quarantine else "",
        })
        results.append({"item_index": index, "payload": text})
    db = state / "progress.sqlite3"
    with patch.object(config, "PROGRESS_DB_FILE", db):
        checkpoint.seed_recovered_run(
            path=db, run_id=root.name, input_hash="input-hash", run_signature="test",
            context={"phase": "PAID" if paid_pending else "FREE", "paid_query_limit_per_company": 3}, budgets=PROVIDERS,
            items=items, results=results,
        )
    return root, root.name, source_ids, results


def _write_manifest(root: Path, run_id: str, source_ids: list[str], *, paid_enabled: bool = False, phase: str = "FREE", provisional: bool = False, files: dict | None = None, artifact_hash: str = "") -> None:
    run_config = run_context.RunConfig.from_config(paid_enabled=paid_enabled)
    manifest = {
        "run_id": run_id, "input_sha256": "input-hash", "config_sha256": run_config.sha256,
        "complete": False, "phase": phase, "paid_enabled": paid_enabled,
        "paid_query_limit_per_company": 3,
        "run_config": run_config.as_dict(),
        "provisional": provisional, "item_count": len(source_ids),
        "ordered_source_record_ids": source_ids, "files": files or {},
        "artifact_set_sha256": artifact_hash,
    }
    (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")


def _freeze_artifact(root: Path, run_id: str, source_ids: list[str], *, paid_enabled: bool = False) -> None:
    db = root / "state" / "progress.sqlite3"
    digest = hashlib.sha256(db.read_bytes()).hexdigest()
    artifact_hash = hashlib.sha256(f"recovery_state.sqlite3:{digest}\n".encode()).hexdigest()
    artifact_dir = root / "output" / "artifacts" / artifact_hash
    artifact_dir.mkdir(parents=True)
    shutil.copy2(db, artifact_dir / "recovery_state.sqlite3")
    _write_manifest(root, run_id, source_ids, paid_enabled=paid_enabled, phase="PAID", files={"recovery_state.sqlite3": {"sha256": digest, "bytes": db.stat().st_size}}, artifact_hash=artifact_hash)


def test_active_resume_allows_unprocessed_missing_payload_but_detects_payload_identity_change(tmp_path: Path):
    root, run_id, source_ids, _ = _bundle(tmp_path)
    _write_manifest(root, run_id, source_ids)
    with patch.object(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3"):
        assert run_context.validate_run_bundle(root, profile="ACTIVE_RESUME")["item_count"] == 2
        with sqlite3.connect(config.PROGRESS_DB_FILE) as connection:
            connection.execute("DELETE FROM results WHERE run_id=? AND item_index=1", (run_id,))
            connection.execute("UPDATE run_items SET payload_sha256=? WHERE run_id=? AND item_index=0", ("wrong", run_id))
            connection.commit()
        with pytest.raises(checkpoint.ResumeInvariant, match="payload hash mismatch"):
            run_context.validate_run_bundle(root, profile="ACTIVE_RESUME")


def test_quarantine_is_item_durable_and_reapplied_on_save(tmp_path: Path):
    db = tmp_path / "progress.sqlite3"
    run_id = "b" * 64
    with patch.object(config, "PROGRESS_DB_FILE", db):
        checkpoint.seed_recovered_run(
            path=db, run_id=run_id, input_hash="h", run_signature="s", context={"phase": "FREE"}, budgets=PROVIDERS,
            items=[{"item_index": 0, "source_record_id": "input:0", "quarantine_state": "LEGACY_RECOVERY_PROVISIONAL_V2", "quarantine_status": "blocked", "publication_blockers": "legacy_recovery_provisional"}], results=[],
        )
        assert checkpoint.claim_item(run_id=run_id, item_index=0, phase="FREE")
        checkpoint.save_item_transaction(run_id=run_id, item_index=0, source_record_id="input:0", payload=_payload("input:0", email="should-not-publish", publication_eligible=True), free_state="DONE", paid_state="NOT_REQUIRED", paid_required=False, free_attempts=1, paid_attempts=0)
        row = checkpoint.load_results_by_id(run_id)[0]
        assert row["publication_eligible"] is False
        assert output_artifacts.is_quarantined_row(row)
        assert row["email"] == ""


def test_paid_attempt_is_durable_before_quality_selection_and_free_attempts_are_preserved(tmp_path: Path):
    db = tmp_path / "progress.sqlite3"
    run_id = "c" * 64
    with patch.object(config, "PROGRESS_DB_FILE", db):
        budgets = dict(PROVIDERS, brightdata=1)
        checkpoint.seed_recovered_run(path=db, run_id=run_id, input_hash="h", run_signature="s", context={"phase": "PAID"}, budgets=budgets, items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "DONE", "paid_required": True, "paid_state": "RUNNING", "free_attempts": 3}], results=[{"item_index": 0, "payload": _payload_text(_payload("input:0", email="old@example.com"))}])
        call_id = checkpoint.reserve_provider_call(run_id=run_id, provider="brightdata", item_index=0, phase="PAID", request_fingerprint="fp-1", configured_limit=1, effective_limit=1)
        checkpoint.record_paid_attempt(run_id=run_id, item_index=0, attempt_number=1, result="FAILED", reason="provider_timeout", call_id=call_id or "", request_fingerprint="fp-1", call_ids=[call_id or ""])
        attempts = checkpoint.load_paid_attempts(run_id)
        assert attempts[0]["result"] == "FAILED"
        assert attempts[0]["reason"] == "provider_timeout"
        assert attempts[0]["call_id"] == call_id
        assert checkpoint.load_paid_attempt_calls(run_id, 0, 1) == [call_id]


def test_finalization_intent_cas_and_memory_outbox_are_idempotent(tmp_path: Path):
    db = tmp_path / "progress.sqlite3"
    run_id = "d" * 64
    with patch.object(config, "PROGRESS_DB_FILE", db):
        payload = _payload_text(_payload("input:0"))
        checkpoint.seed_recovered_run(path=db, run_id=run_id, input_hash="h", run_signature="s", context={"phase": "FREE"}, budgets=PROVIDERS, items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "DONE", "payload_sha256": hashlib.sha256(payload.encode()).hexdigest()}], results=[{"item_index": 0, "payload": payload}])
        checkpoint.transition_phase(run_id, "FINALIZING", expected_count=1)
        checkpoint.begin_finalization_intent(run_id=run_id, generation="generation-1", input_snapshot_sha256="h")
        checkpoint.enqueue_memory_rows(run_id, [{"source_record_id": "input:0", "company": "A"}])
        with pytest.raises(RuntimeError, match="finalization CAS requires"):
            checkpoint.complete_finalization_phase(run_id, expected_count=1)
        assert checkpoint.load_run_state_by_id(run_id)["phase"] == "FINALIZING"
        assert len(checkpoint.load_memory_outbox(run_id)) == 1
        checkpoint.complete_memory_receipt(run_id, "input:0")
        checkpoint.enqueue_memory_rows(run_id, [{"source_record_id": "input:0", "company": "A"}])
        assert checkpoint.load_memory_outbox(run_id) == []


def test_effective_publication_setting_changes_run_config_hash(tmp_path: Path):
    with patch.object(config, "PUBLICATION_POLICY_MIN_SAFETY_SCORE", 75):
        first = run_context.RunConfig.from_config(paid_enabled=False)
    with patch.object(config, "PUBLICATION_POLICY_MIN_SAFETY_SCORE", 76):
        second = run_context.RunConfig.from_config(paid_enabled=False)
    assert first.sha256 != second.sha256
    assert first.as_dict()["effective_settings"]["publication_policy_min_safety_score"] == 75


def test_paid_continuation_uses_only_pending_paid_ids_and_is_one_time(tmp_path: Path):
    parent, run_id, source_ids, _ = _bundle(tmp_path, paid_pending=True)
    _freeze_artifact(parent, run_id, source_ids)
    parent_manifest = json.loads((parent / "manifest.json").read_text(encoding="utf-8"))
    auth = tmp_path / "approval.json"
    approval = {
        "approved": True,
        "limits": {name: 0 for name in checkpoint.CANONICAL_PROVIDERS},
        "parent_run_id": run_id,
        "parent_manifest_sha256": hashlib.sha256((parent / "manifest.json").read_bytes()).hexdigest(),
        "checkpoint_sha256": parent_manifest["files"]["recovery_state.sqlite3"]["sha256"],
        "paid_source_ids_sha256": hashlib.sha256(run_context.canonical_json(["input:0", "input:1"]).encode()).hexdigest(),
    }
    auth.write_text(json.dumps(approval), encoding="utf-8")
    destination = tmp_path / "child"
    with patch.object(config, "PROGRESS_DB_FILE", parent / "state" / "progress.sqlite3"):
        child = prepare_paid_continuation(parent, auth, destination)
        assert child["paid_enabled"] is True
        with pytest.raises(RuntimeError, match="another target"):
            prepare_paid_continuation(parent, auth, tmp_path / "other-child")


def test_real_runner_handoff_continuation_and_resume_without_provider_calls(tmp_path: Path):
    from modules import pipeline_runner

    input_file = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["company", "source_record_id", "website"])
    sheet.append(["Acme", "input:0", "https://acme.example"])
    workbook.save(input_file)
    workbook.close()

    def worker(_index, company, _logger, _website, record):
        return 0, _payload(record["source_record_id"], company=company, status="REVIEW_NEEDED", reason="needs_paid")

    def impossible_writer(_rows, _elapsed):
        raise AssertionError("handoff must not publish final artifacts")

    with patch.object(config, "SEARCH_PROVIDER", "brightdata"), patch.object(config, "RUNS_DIR", tmp_path / "parent-runs"):
        result = pipeline_runner.run_pipeline(
            input_file, allow_paid=False,
            process_company_fn=worker, write_outputs_fn=impossible_writer,
            set_output_dir_fn=lambda path: setattr(config, "OUTPUT_DIR", Path(path)),
            empty_result_fn=lambda company, status, reason: _payload(company, status=status, reason=reason),
        )
        assert result == "PAID_PENDING_APPROVAL"
    run_root = next((tmp_path / "parent-runs").iterdir())
    manifest = json.loads((run_root / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["handoff"] is True
    assert manifest["files"]["recovery_state.sqlite3"]

    auth = tmp_path / "approval.json"
    approval = {
        "approved": True,
        "limits": {name: 1 if name == "brightdata" else 0 for name in checkpoint.CANONICAL_PROVIDERS},
        "parent_run_id": manifest["run_id"],
        "parent_manifest_sha256": hashlib.sha256((run_root / "manifest.json").read_bytes()).hexdigest(),
        "checkpoint_sha256": manifest["checkpoint_sha256"],
        "paid_source_ids_sha256": hashlib.sha256(run_context.canonical_json(["input:0"]).encode()).hexdigest(),
    }
    auth.write_text(json.dumps(approval), encoding="utf-8")
    child_manifest = prepare_paid_continuation(run_root, auth, tmp_path / "child")
    child_root = tmp_path / "child" / "runs" / child_manifest["run_id"]

    def final_writer(rows, _elapsed, *, telemetry_snapshot=None):
        return output_artifacts.write_outputs(rows, _elapsed, telemetry_snapshot=telemetry_snapshot)

    with patch.object(config, "SEARCH_PROVIDER", "brightdata"), patch.object(config, "PROGRESS_DB_FILE", config.STATE_DIR / "handoff-test.sqlite3"):
        completed = pipeline_runner.run_pipeline(
            input_file, resume_run_dir=child_root, allow_paid=True,
            process_company_fn=lambda _index, company, _logger, _website, record: (0, dict(_payload(record["source_record_id"], company=company, status="OK_HIGH_CONFIDENCE", publication_eligible=True, website="https://acme.example"), known_website_evaluation={"status": "OK_HIGH_CONFIDENCE", "website": "https://acme.example"}, paid_attempt_result="NO_CALL_NEEDED", paid_attempt_reason="supplied_website_publishable_at_paid_entry", paid_evidence_ref="test:supplied-site")),
            write_outputs_fn=final_writer,
            set_output_dir_fn=lambda path: setattr(config, "OUTPUT_DIR", Path(path)),
            empty_result_fn=lambda company, status, reason: _payload(company, status=status, reason=reason),
        )
    assert "B2B Contact Finder" in completed
    completed_manifest = json.loads((child_root / "manifest.json").read_text(encoding="utf-8"))
    assert completed_manifest["complete"] is True
    artifact_dir = child_root / "output" / "artifacts" / completed_manifest["artifact_set_sha256"]
    workbook = __import__("openpyxl").load_workbook(artifact_dir / "all_results.xlsx", read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    assert rows[0][0] == "company"
    header = list(rows[0])
    assert any(row[header.index("company")] == "Acme" and row[header.index("source_record_id")] == "input:0" for row in rows[1:])
    artifact_hashes = {name: info["sha256"] for name, info in completed_manifest["files"].items()}
    resumed = pipeline_runner.run_pipeline(
        input_file, resume_run_dir=child_root, allow_paid=True,
        process_company_fn=lambda *_args: (_ for _ in ()).throw(AssertionError("resume must not call providers")),
        write_outputs_fn=lambda *_args: (_ for _ in ()).throw(AssertionError("complete resume must reuse artifacts")),
        set_output_dir_fn=lambda path: setattr(config, "OUTPUT_DIR", Path(path)),
        empty_result_fn=lambda company, status, reason: _payload(company, status=status, reason=reason),
    )
    assert resumed == "COMPLETE_RESUME_VERIFIED"
    resumed_manifest = json.loads((child_root / "manifest.json").read_text(encoding="utf-8"))
    assert {name: info["sha256"] for name, info in resumed_manifest["files"].items()} == artifact_hashes


def test_provider_fingerprint_duplicate_and_unknown_consume_durable_budget(tmp_path: Path):
    db = tmp_path / "progress.sqlite3"
    run_id = "e" * 64
    with patch.object(config, "PROGRESS_DB_FILE", db):
        budgets = dict(PROVIDERS, brightdata=8)
        payload = _payload_text(_payload("input:0"))
        checkpoint.seed_recovered_run(
            path=db, run_id=run_id, input_hash="h", run_signature="s", context={"phase": "PAID"},
            budgets=budgets, items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "DONE", "paid_required": True, "paid_state": "RUNNING", "free_attempts": 3, "payload_sha256": hashlib.sha256(payload.encode()).hexdigest()}],
            results=[{"item_index": 0, "payload": payload}],
        )
        runtime.reset()
        runtime.configure_durable_run(run_id, budgets)
        runtime.set_phase("PAID")
        runtime.set_item_context(0, "paid")
        first = runtime.reserve_api("brightdata", operation="search", request_fingerprint="same")
        duplicate = runtime.reserve_api("brightdata", operation="search", request_fingerprint="same")
        second = runtime.reserve_api("brightdata", operation="search", request_fingerprint="different")
        assert first and second
        assert not duplicate and duplicate.reason == "duplicate_request"
        assert checkpoint.provider_calls_for_item(run_id, 0) and len(checkpoint.provider_calls_for_item(run_id, 0)) == 2
        checkpoint.complete_provider_call(call_id=first.call_id, state="UNKNOWN")
        checkpoint.complete_provider_call(call_id=second.call_id, state="DONE")
        with sqlite3.connect(db) as connection:
            usage = connection.execute("SELECT reserved,completed,failed,unknown FROM provider_usage WHERE run_id=? AND provider='brightdata'", (run_id,)).fetchone()
        assert usage == (0, 1, 0, 1)


def test_provider_ledger_mismatch_stops_before_new_reservation(tmp_path: Path):
    db = tmp_path / "progress.sqlite3"
    run_id = "f" * 64
    with patch.object(config, "PROGRESS_DB_FILE", db):
        budgets = dict(PROVIDERS, brightdata=8)
        payload = _payload_text(_payload("input:0"))
        checkpoint.seed_recovered_run(path=db, run_id=run_id, input_hash="h", run_signature="s", context={"phase": "PAID"}, budgets=budgets, items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "DONE", "paid_required": True, "paid_state": "RUNNING", "payload_sha256": hashlib.sha256(payload.encode()).hexdigest()}], results=[{"item_index": 0, "payload": payload}])
        with sqlite3.connect(db) as connection:
            connection.execute("UPDATE provider_usage SET reserved=1 WHERE run_id=? AND provider='brightdata'", (run_id,))
            connection.commit()
        with pytest.raises(RuntimeError, match="inconsistent"):
            checkpoint.reserve_provider_call(run_id=run_id, provider="brightdata", item_index=0, phase="PAID", request_fingerprint="new", configured_limit=8, effective_limit=8)


def test_publication_policy_blocks_failed_paid_state_even_with_good_old_payload():
    row = _payload("input:0", status="OK_HIGH_CONFIDENCE", publication_eligible=True, website="https://acme.example", email="good@example", free_state="DONE", paid_state="FAILED")
    assert not output_artifacts.is_publishable_row(row)
    assert output_artifacts.partition_output_rows([row])[0] == []


def test_seed_rejects_item_payload_quarantine_mismatch(tmp_path: Path):
    db = tmp_path / "progress.sqlite3"
    with patch.object(config, "PROGRESS_DB_FILE", db):
        payload = _payload_text(_payload("input:0", quarantine_state="OTHER", quarantine_status="blocked", publication_eligible=False, publication_blockers="legacy_recovery_provisional"))
        with pytest.raises(checkpoint.ResumeInvariant, match="quarantine_state mismatch"):
            checkpoint.seed_recovered_run(
                path=db, run_id="1" * 64, input_hash="h", run_signature="s",
                context={"phase": "FREE", "provisional": True}, budgets=PROVIDERS,
                items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "DONE", "quarantine_state": "EXPECTED", "quarantine_status": "blocked", "publication_blockers": "legacy_recovery_provisional", "payload_sha256": hashlib.sha256(payload.encode()).hexdigest()}],
                results=[{"item_index": 0, "payload": payload}],
            )


def test_resume_rejects_payloadless_done_item(tmp_path: Path):
    root, run_id, source_ids, _ = _bundle(tmp_path)
    _write_manifest(root, run_id, source_ids)
    with patch.object(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3"):
        with sqlite3.connect(config.PROGRESS_DB_FILE) as connection:
            connection.execute("DELETE FROM results WHERE run_id=? AND item_index=0", (run_id,))
            connection.commit()
        with pytest.raises(checkpoint.ResumeInvariant, match="missing payload"):
            run_context.validate_run_bundle(root, profile="ACTIVE_RESUME")


def test_resume_rejects_payload_with_wrong_source_id(tmp_path: Path):
    root, run_id, source_ids, _ = _bundle(tmp_path)
    _write_manifest(root, run_id, source_ids)
    with patch.object(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3"):
        bad_payload = _payload_text(_payload("wrong:0"))
        with sqlite3.connect(config.PROGRESS_DB_FILE) as connection:
            connection.execute("UPDATE results SET payload=? WHERE run_id=? AND item_index=0", (bad_payload, run_id))
            connection.execute("UPDATE run_items SET payload_sha256=? WHERE run_id=? AND item_index=0", (hashlib.sha256(bad_payload.encode()).hexdigest(), run_id))
            connection.commit()
        with pytest.raises(checkpoint.ResumeInvariant, match="source ID mismatch"):
            run_context.validate_run_bundle(root, profile="ACTIVE_RESUME")


def test_resume_config_resolver_inherits_recorded_values_and_rejects_explicit_override():
    from modules import pipeline_runner

    with patch.object(config, "SEARCH_CACHE_MODE", "use"), patch.object(config, "CRAWL_CACHE_MODE", "refresh"):
        recorded = run_context.RunConfig.from_config(paid_enabled=False)
    manifest = {"run_config": recorded.as_dict()}
    resolved = pipeline_runner.resolve_run_config(manifest)
    assert resolved.as_dict() == recorded.as_dict()
    with pytest.raises(checkpoint.ResumeInvariant, match="search-cache override"):
        pipeline_runner.resolve_run_config(manifest, search_cache="replay")
    with pytest.raises(checkpoint.ResumeInvariant, match="paid-mode override"):
        pipeline_runner.resolve_run_config(manifest, allow_paid=True)


def test_google_places_adapter_distinguishes_timeout_budget_and_duplicate_offline():
    from modules import google_places
    from modules.runtime import Reservation

    runtime.reset()
    google_places.reset()
    with patch.object(config, "ENABLE_GOOGLE_PLACES", True), patch.object(config, "GOOGLE_PLACES_API_KEY", "test-key"), patch.object(config, "GOOGLE_PLACES_REQUEST_BUDGET", 1), patch.object(config, "SEARCH_CACHE_MODE", "off"):
        with patch.object(google_places.requests, "post", side_effect=TimeoutError("transport timeout")):
            timeout_result = google_places.search_company("Acme")
        assert timeout_result.result_state == "UNKNOWN"
        with patch.object(runtime, "reserve_api", return_value=Reservation(False, "google_places", "text_search", -1, "PAID", reason="budget_exhausted")):
            budget_result = google_places.search_company("Acme")
        assert budget_result.result_state == "BLOCKED_BUDGET"
        google_places.reset()
        with patch.object(runtime, "reserve_api", return_value=Reservation(False, "google_places", "text_search", -1, "PAID", reason="duplicate_request")):
            duplicate_result = google_places.search_company("Acme")
        assert duplicate_result.result_state == "DUPLICATE"


def test_finalization_artifact_ready_is_not_reset_or_reidentified(tmp_path: Path):
    db = tmp_path / "progress.sqlite3"
    run_id = "2" * 64
    with patch.object(config, "PROGRESS_DB_FILE", db):
        payload = _payload_text(_payload("input:0"))
        checkpoint.seed_recovered_run(path=db, run_id=run_id, input_hash="h", run_signature="s", context={"phase": "FREE"}, budgets=PROVIDERS, items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "DONE", "payload_sha256": hashlib.sha256(payload.encode()).hexdigest()}], results=[{"item_index": 0, "payload": payload}])
        checkpoint.transition_phase(run_id, "FINALIZING", expected_count=1)
        checkpoint.begin_finalization_intent(run_id=run_id, generation="generation-2", input_snapshot_sha256="result-snapshot", output_context={"elapsed_seconds": 1.25})
        checkpoint.mark_finalization_artifact(run_id=run_id, artifact_set_sha256="artifact-2", output_context={"files": {"all_results.xlsx": {"sha256": "x"}}})
        checkpoint.begin_finalization_intent(run_id=run_id, generation="generation-2", input_snapshot_sha256="result-snapshot", output_context={"elapsed_seconds": 99})
        intent = checkpoint.load_finalization_intent(run_id)
        assert intent["status"] == "ARTIFACT_READY"
        assert json.loads(intent["output_context_json"])["elapsed_seconds"] == 1.25
        with pytest.raises(RuntimeError, match="identity changed"):
            checkpoint.begin_finalization_intent(run_id=run_id, generation="different", input_snapshot_sha256="result-snapshot")


def test_continuation_preserves_mixed_parent_states_and_queues_only_pending(tmp_path: Path):
    root = tmp_path / "mixed" / ("3" * 64)
    (root / "state").mkdir(parents=True)
    (root / "output" / "artifacts").mkdir(parents=True)
    source_ids = [f"input:{index}" for index in range(4)]
    items, results = [], []
    states = [(False, "NOT_REQUIRED"), (True, "DONE"), (True, "PENDING"), (True, "UNKNOWN")]
    for index, source_id in enumerate(source_ids):
        payload = _payload_text(_payload(source_id))
        items.append({"item_index": index, "source_record_id": source_id, "free_state": "DONE", "paid_required": states[index][0], "paid_state": states[index][1], "payload_sha256": hashlib.sha256(payload.encode()).hexdigest()})
        results.append({"item_index": index, "payload": payload})
    db = root / "state" / "progress.sqlite3"
    with patch.object(config, "PROGRESS_DB_FILE", db):
        checkpoint.seed_recovered_run(path=db, run_id=root.name, input_hash="input-hash", run_signature="mixed", context={"phase": "PAID"}, budgets=PROVIDERS, items=items, results=results)
    _freeze_artifact(root, root.name, source_ids)
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    auth = tmp_path / "mixed-approval.json"
    approval = {"approved": True, "limits": {name: 0 for name in checkpoint.CANONICAL_PROVIDERS}, "parent_run_id": root.name, "parent_manifest_sha256": hashlib.sha256((root / "manifest.json").read_bytes()).hexdigest(), "checkpoint_sha256": manifest["files"]["recovery_state.sqlite3"]["sha256"], "paid_source_ids_sha256": hashlib.sha256(run_context.canonical_json(["input:2"]).encode()).hexdigest()}
    auth.write_text(json.dumps(approval), encoding="utf-8")
    child_manifest = prepare_paid_continuation(root, auth, tmp_path / "mixed-child")
    child_root = tmp_path / "mixed-child" / "runs" / child_manifest["run_id"]
    with sqlite3.connect(child_root / "state" / "progress.sqlite3") as connection:
        child_states = [row for row in connection.execute("SELECT paid_required,paid_state FROM run_items ORDER BY item_index")]
    assert child_states == [(0, "NOT_REQUIRED"), (1, "DONE"), (1, "PENDING"), (1, "UNKNOWN")]


def test_run_lease_contention_is_owner_scoped(tmp_path: Path):
    run_dir = tmp_path / "leased-run"
    owner = run_context.RunLease(run_dir)
    contender = run_context.RunLease(run_dir)
    owner.acquire()
    try:
        with pytest.raises(RuntimeError, match="already leased"):
            contender.acquire()
    finally:
        owner.release()
    contender.acquire()
    contender.release()


def test_finalization_atomic_commit_records_empty_plan_and_run_scoped_receipts(tmp_path: Path):
    db = tmp_path / "progress.sqlite3"
    run_id = "4" * 64
    with patch.object(config, "PROGRESS_DB_FILE", db):
        payload = _payload_text(_payload("input:0"))
        checkpoint.seed_recovered_run(
            path=db, run_id=run_id, input_hash="h", run_signature="s",
            context={"phase": "FREE"}, budgets=PROVIDERS,
            items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "DONE", "payload_sha256": hashlib.sha256(payload.encode()).hexdigest()}],
            results=[{"item_index": 0, "payload": payload}],
        )
        checkpoint.transition_phase(run_id, "FINALIZING", expected_count=1)
        telemetry = checkpoint.canonical_scheduler_receipt(run_id)
        checkpoint.begin_finalization_intent(
            run_id=run_id, generation="g", input_snapshot_sha256="r",
            result_snapshot_sha256="r", telemetry_snapshot=telemetry,
        )
        plan = checkpoint.mark_finalization_artifact_and_outbox(
            run_id=run_id, generation="g", result_snapshot_sha256="r",
            artifact_set_sha256="a", artifacts={"artifact_set_sha256": "a", "files": {}},
            memory_rows=[], counts={"input_count": 1},
            telemetry_snapshot=telemetry,
        )
        intent = checkpoint.load_finalization_intent(run_id)
        assert plan == {"memory_plan_sha256": plan["memory_plan_sha256"], "memory_plan_count": 0, "memory_plan_committed": True}
        assert intent["memory_plan_committed"] == 1
        assert intent["memory_plan_count"] == 0
        assert json.loads(intent["telemetry_snapshot_json"]) == telemetry
        assert checkpoint.load_memory_outbox_entries(run_id) == []


def test_entity_memory_receipt_is_idempotent_and_preserves_observation(tmp_path: Path):
    from modules import entity_memory

    path = tmp_path / "memory.jsonl"
    row = {
        "company": "Receipt Corp", "source_record_id": "input:0",
        "status": "OK_MEDIUM_CONFIDENCE", "publication_eligible": True,
        "free_state": "DONE", "paid_state": "NOT_REQUIRED", "paid_required": False,
        "website": "https://receipt.example", "email_source_url": "https://receipt.example/contact",
        "__memory_receipt_key": "run-1:input:0",
        "__evaluation": {
            "_identity_resolution": "candidate_resolved_by_target_fingerprint",
            "identity_assessment": {"conflicts": []},
            "crawl": {"pages": ["https://receipt.example/contact"]},
        },
    }
    with patch.object(config, "VERIFIED_ENTITY_MEMORY_FILE", path):
        assert entity_memory.remember([row]) == 1
        before = path.read_bytes()
        observed_at = json.loads(before.decode().splitlines()[0])["observed_at"]
        assert entity_memory.remember([row]) == 0
        assert path.read_bytes() == before
        assert json.loads(path.read_text(encoding="utf-8").splitlines()[0])["observed_at"] == observed_at


def test_provider_attempt_collector_uses_unknown_over_budget_and_failure():
    from modules import runtime
    from modules.pipeline_runner import paid_attempt_result

    runtime.reset()
    token = runtime.begin_provider_attempt()
    runtime.provider_result([], state="FAILED", reason="transport")
    runtime.provider_result([], state="BLOCKED_BUDGET", reason="budget_exhausted")
    runtime.provider_result([], state="UNKNOWN", reason="timeout")
    outcomes = runtime.end_provider_attempt(token)
    assert paid_attempt_result({"provider_results": outcomes, "status": "OK_HIGH_CONFIDENCE"}) == "UNKNOWN"


@pytest.mark.parametrize("fault_boundary", ["claim", "staging", "published"])
def test_paid_continuation_fault_boundaries_resume_same_child_and_artifact(tmp_path: Path, fault_boundary: str):
    import prepare_paid_continuation as continuation

    parent, parent_id, source_ids, _ = _bundle(tmp_path, paid_pending=True)
    _freeze_artifact(parent, parent_id, source_ids)
    parent_manifest = json.loads((parent / "manifest.json").read_text(encoding="utf-8"))
    authorization = tmp_path / "approval.json"
    authorization.write_text(json.dumps({
        "approved": True,
        "limits": {name: 0 for name in checkpoint.CANONICAL_PROVIDERS},
        "parent_run_id": parent_id,
        "parent_manifest_sha256": hashlib.sha256((parent / "manifest.json").read_bytes()).hexdigest(),
        "checkpoint_sha256": parent_manifest["files"]["recovery_state.sqlite3"]["sha256"],
        "paid_source_ids_sha256": hashlib.sha256(run_context.canonical_json(source_ids).encode()).hexdigest(),
    }), encoding="utf-8")
    destination = tmp_path / "child"
    original_claim = continuation._claim_authorization
    original_mark = continuation._mark_authorization_published
    original_replace = Path.replace
    fired = {"value": False}

    def claim_then_fail(*args, **kwargs):
        result = original_claim(*args, **kwargs)
        if not fired["value"]:
            fired["value"] = True
            raise RuntimeError("fault after authorization claim")
        return result

    def replace_then_fail(self, target):
        if (fault_boundary == "staging" and not fired["value"]
                and self.name.endswith(".staging") and len(Path(target).name) == 64):
            fired["value"] = True
            raise RuntimeError("fault before staging rename")
        return original_replace(self, target)

    def mark_then_fail(*args, **kwargs):
        if not fired["value"]:
            fired["value"] = True
            raise RuntimeError("fault after child rename")
        return original_mark(*args, **kwargs)

    with patch.object(continuation, "_claim_authorization", side_effect=claim_then_fail) if fault_boundary == "claim" else patch.object(Path, "replace", replace_then_fail) if fault_boundary == "staging" else patch.object(continuation, "_mark_authorization_published", side_effect=mark_then_fail):
        with pytest.raises(RuntimeError):
            continuation.prepare_paid_continuation(parent, authorization, destination)
    resumed = continuation.prepare_paid_continuation(parent, authorization, destination)
    child_root = destination / "runs" / resumed["run_id"]
    again = continuation.prepare_paid_continuation(parent, authorization, destination)
    assert again["run_id"] == resumed["run_id"]
    assert again["artifact_set_sha256"] == resumed["artifact_set_sha256"]
    assert (child_root / "manifest.json").is_file()


@pytest.mark.parametrize("fault_boundary", [
    "writer_return", "artifact_ready", "outbox_commit", "manifest_publish",
    "intent_complete", "db_complete", "memory_write", "receipt",
])
def test_finalization_fault_boundaries_resume_to_one_complete_receipt(tmp_path: Path, fault_boundary: str):
    from contextlib import ExitStack
    from modules import entity_memory, pipeline_runner, runtime_paths

    input_file = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["company", "source_record_id"])
    sheet.append(["Finalization Corp", "input:0"])
    workbook.save(input_file)
    workbook.close()
    output_root = tmp_path / "runs"

    def worker(_index, company, _logger, _website, record):
        return 0, {
            "company": company, "source_record_id": record["source_record_id"],
            "status": "OK_HIGH_CONFIDENCE", "publication_eligible": True,
            "website": "https://finalization.example", "email": "ops@finalization.example",
            "phone": "+90 212 000 00 00",
            "email_source_url": "https://finalization.example/contact",
            "phone_source_url": "https://finalization.example/contact",
            "__evaluation": {
                "_identity_resolution": "candidate_resolved_by_target_fingerprint",
                "identity_assessment": {"conflicts": []},
                "crawl": {"pages": ["https://finalization.example/contact"]},
            },
        }

    real_writer = output_artifacts.write_outputs
    fired = {"value": False}

    def fail_writer(rows, elapsed, **kwargs):
        result = real_writer(rows, elapsed, **kwargs)
        raise RuntimeError("fault after writer return")

    def fail_after(callable_value, message):
        def wrapper(*args, **kwargs):
            result = callable_value(*args, **kwargs)
            if not fired["value"]:
                fired["value"] = True
                raise RuntimeError(message)
            return result
        return wrapper

    with ExitStack() as stack:
        stack.enter_context(patch.object(config, "RUNS_DIR", output_root))
        stack.enter_context(patch.object(config, "VERIFIED_ENTITY_MEMORY_FILE", tmp_path / "verified_entity_memory.jsonl"))
        if fault_boundary == "writer_return":
            writer = fail_writer
        else:
            writer = real_writer
            if fault_boundary == "artifact_ready":
                original = checkpoint.mark_finalization_artifact
                stack.enter_context(patch.object(checkpoint, "mark_finalization_artifact", fail_after(original, fault_boundary)))
            elif fault_boundary == "outbox_commit":
                original = checkpoint.enqueue_memory_rows
                stack.enter_context(patch.object(checkpoint, "enqueue_memory_rows", fail_after(original, fault_boundary)))
            elif fault_boundary == "manifest_publish":
                original = output_artifacts.publish_manifest
                stack.enter_context(patch.object(output_artifacts, "publish_manifest", fail_after(original, fault_boundary)))
            elif fault_boundary == "intent_complete":
                original = checkpoint.complete_finalization_intent
                stack.enter_context(patch.object(checkpoint, "complete_finalization_intent", fail_after(original, fault_boundary)))
            elif fault_boundary == "db_complete":
                original = checkpoint.transition_phase
                def fail_db_complete(*args, **kwargs):
                    result = original(*args, **kwargs)
                    if kwargs.get("new_phase") == "COMPLETE" or len(args) > 1 and args[1] == "COMPLETE":
                        if not fired["value"]:
                            fired["value"] = True
                            raise RuntimeError(fault_boundary)
                    return result
                stack.enter_context(patch.object(checkpoint, "transition_phase", fail_db_complete))
            elif fault_boundary == "memory_write":
                original = entity_memory.remember
                stack.enter_context(patch.object(entity_memory, "remember", fail_after(original, fault_boundary)))
            elif fault_boundary == "receipt":
                original = checkpoint.complete_memory_receipt
                stack.enter_context(patch.object(checkpoint, "complete_memory_receipt", fail_after(original, fault_boundary)))
        with pytest.raises(RuntimeError):
            pipeline_runner.run_pipeline(
                input_file, allow_paid=False, process_company_fn=worker,
                write_outputs_fn=writer, set_output_dir_fn=runtime_paths.set_output_dir,
                empty_result_fn=lambda company, status, reason: {"company": company, "status": status, "reason": reason},
            )

    run_root = next(output_root.iterdir())
    resumed = pipeline_runner.run_pipeline(
        input_file, resume_run_dir=run_root, allow_paid=False,
        process_company_fn=lambda *_args: (_ for _ in ()).throw(AssertionError("finalization resume called provider")),
        write_outputs_fn=real_writer,
        set_output_dir_fn=runtime_paths.set_output_dir,
        empty_result_fn=lambda company, status, reason: {"company": company, "status": status, "reason": reason},
    )
    assert resumed in {"FINALIZATION_RESUME_RECONCILED", "COMPLETE_RESUME_VERIFIED"} or "B2B Contact Finder" in resumed
    manifest = json.loads((run_root / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["complete"] is True
    with patch.object(config, "PROGRESS_DB_FILE", run_root / "state" / "progress.sqlite3"):
        intent = checkpoint.load_finalization_intent(run_root.name)
        entries = checkpoint.load_memory_outbox_entries(run_root.name)
    assert intent["status"] == "COMPLETE"
    assert intent["memory_plan_committed"] == 1
    assert intent["memory_plan_count"] == 1
    assert len(entries) == 1
    assert all(entry["state"] == "DONE" for entry in entries)
    memory_lines = (config.VERIFIED_ENTITY_MEMORY_FILE).read_text(encoding="utf-8").splitlines()
    assert len(memory_lines) == 1

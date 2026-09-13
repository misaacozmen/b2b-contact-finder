from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import shutil
import sqlite3
import subprocess
import sys
import threading
from dataclasses import asdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
import requests
from openpyxl import Workbook

import config
import main
from modules import checkpoint, crawler, pipeline_runner, report, run_context, runtime, search
from prepare_paid_continuation import prepare_paid_continuation


def _assert_reparse_retarget(tmp_path, monkeypatch):
    import conftest

    root = tmp_path / "protected"
    first = root / "first"
    second = root / "second"
    link = root / "junction"
    first.mkdir(parents=True)
    second.mkdir()
    (first / "sentinel.txt").write_text("first", encoding="utf-8")
    (second / "sentinel.txt").write_text("second", encoding="utf-8")
    if sys.platform == "win32":
        env = {**os.environ, "B2B_JUNCTION_LINK": str(link), "B2B_JUNCTION_TARGET": str(first)}
        subprocess.run(["powershell", "-NoProfile", "-Command", "New-Item -ItemType Junction -Path $env:B2B_JUNCTION_LINK -Target $env:B2B_JUNCTION_TARGET | Out-Null"], check=True, env=env)
    else:
        os.symlink(first, link, target_is_directory=True)

    with monkeypatch.context() as scoped:
        scoped.setattr(conftest, "_REPO", tmp_path)
        scoped.setattr(conftest, "_PROTECTED_TEST_ROOTS", (root,))
        before = conftest._protected_manifest()
    os.rmdir(link) if sys.platform == "win32" else link.unlink()
    if sys.platform == "win32":
        env = {**os.environ, "B2B_JUNCTION_LINK": str(link), "B2B_JUNCTION_TARGET": str(second)}
        subprocess.run(["powershell", "-NoProfile", "-Command", "New-Item -ItemType Junction -Path $env:B2B_JUNCTION_LINK -Target $env:B2B_JUNCTION_TARGET | Out-Null"], check=True, env=env)
    else:
        os.symlink(second, link, target_is_directory=True)
    with monkeypatch.context() as scoped:
        scoped.setattr(conftest, "_REPO", tmp_path)
        scoped.setattr(conftest, "_PROTECTED_TEST_ROOTS", (root,))
        after = conftest._protected_manifest()

    assert before != after
    key = str(Path("protected") / "junction")
    assert before[key] != after[key]
    assert (first / "sentinel.txt").read_text(encoding="utf-8") == "first"
    assert (second / "sentinel.txt").read_text(encoding="utf-8") == "second"


def test_windows_reparse_retarget_changes_protected_manifest(tmp_path, monkeypatch):
    _assert_reparse_retarget(tmp_path, monkeypatch)


def test_windows_junction_retarget_changes_protected_manifest(tmp_path, monkeypatch):
    _assert_reparse_retarget(tmp_path, monkeypatch)


class Response:
    def __init__(self, status=200, payload=None, headers=None):
        self.status_code = status
        self.payload = payload
        self.text = json.dumps(payload) if payload is not None else ""
        self.headers = headers or {}

    def json(self):
        if self.payload is None:
            raise ValueError("not json")
        return self.payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"http {self.status_code}")


class RecordingPaidTransport:
    def __init__(self, outcomes):
        self._outcomes = iter(outcomes)
        self.journal = []

    def __call__(self, envelope):
        outcome = next(self._outcomes)
        entry = {**asdict(envelope),
            "invocation_ordinal": len(self.journal) + 1,
            "fake_response_class": type(outcome).__name__,
        }
        self.journal.append(entry)
        if isinstance(outcome, BaseException):
            raise outcome
        return outcome

    def post(self, _url, **_kwargs):
        raise AssertionError("paid adapter bypassed TransportEnvelope")


def _html_response(url: str):
    response = requests.Response()
    response.status_code = 200
    response.url = url
    response.headers = {"content-type": "text/html; charset=utf-8"}
    company = "FREE CO" if "freeco" in url else "ACME" if "acme" in url else "PAID CO"
    response._content = (f"""<html><head><title>{company}</title><script type='application/ld+json'>{{"@type":"Organization","name":"{company}","url":"{url}"}}</script></head><body><h1>{company}</h1><p>{company} resmi şirket sitesi Türkiye İstanbul</p><p>Telefon +90 212 555 12 12</p><a href='/contact'>İletişim</a></body></html>""").encode()
    response.encoding = "utf-8"
    return response


def _real_run_setup(tmp_path, monkeypatch, transport, *, workers=1):
    class DDGS:
        def __enter__(self): return self
        def __exit__(self, *_): return None
        def text(self, query, **_kwargs):
            if "free co" in str(query).casefold():
                return [{"href": "https://freeco.example", "title": "FREE CO", "body": "FREE CO resmi şirket sitesi Türkiye"}]
            return []
    monkeypatch.setattr(config, "RUNS_DIR", tmp_path / "runs")
    monkeypatch.setattr(config, "MAX_WORKERS", workers)
    monkeypatch.setattr(config, "SEARCH_PROVIDER", "brightdata")
    monkeypatch.setattr(config, "SEARCH_CACHE_MODE", "off")
    monkeypatch.setattr(config, "CRAWL_CACHE_MODE", "off")
    monkeypatch.setattr(config, "MIN_DELAY_SEC", 0)
    monkeypatch.setattr(config, "MAX_DELAY_SEC", 0)
    monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake")
    monkeypatch.setattr(config, "BRIGHTDATA_REQUEST_BUDGET", 8)
    monkeypatch.setattr(config, "BRIGHTDATA_REQUEST_HARD_CAP", 8)
    monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 0)
    monkeypatch.setattr(config, "GLOBAL_REQUESTS_PER_SECOND", 0)
    monkeypatch.setattr(config, "MAX_RETRIES", 0)
    monkeypatch.setattr(config, "MAX_SEARCH_QUERIES_PER_COMPANY", 1)
    monkeypatch.setattr(config, "DEFAULT_PAID_SEARCH_QUERY_LIMIT", 1)
    monkeypatch.setattr(config, "ENABLE_GOOGLE_PLACES", False)
    monkeypatch.setattr(config, "ENABLE_BRANDFETCH_DOMAIN_SEARCH", False)
    monkeypatch.setattr(config, "ENABLE_HUNTER_DOMAIN_FINDER", False)
    monkeypatch.setattr(config, "ENABLE_LINKEDIN_COMPANY_LOOKUP", False)
    monkeypatch.setattr(config, "ENABLE_LLM_ARBITER", False)
    monkeypatch.setattr(search, "DDGS", DDGS)
    monkeypatch.setattr(search.socket, "getaddrinfo", lambda *_a, **_k: [(2, 1, 6, "", ("203.0.113.10", 0))])
    monkeypatch.setattr(runtime, "_PAID_TRANSPORT", transport)
    monkeypatch.setattr(crawler, "_fetch", lambda url: _html_response(url))
    return tmp_path / "runs"


def _input_book(path: Path, companies):
    book = Workbook(); sheet = book.active; sheet.append(["company", "source_record_id"])
    for index, company in enumerate(companies):
        sheet.append([company, f"s{index}"])
    book.save(path); book.close()


def _capture_scenario(name: str, db_path: Path, *, run_root: Path | None = None, transport=None, outcome=None):
    evidence_root = os.environ.get("B2B_FINAL_SCENARIO_DIR", "").strip()
    if not evidence_root:
        return
    target = Path(evidence_root) / name
    target.mkdir(parents=True, exist_ok=True)
    db_path = Path(db_path)
    with sqlite3.connect(f"file:{db_path.resolve()}?mode=ro", uri=True) as source, sqlite3.connect(target / "sanitized_progress.sqlite3") as destination:
        source.backup(destination)
    with sqlite3.connect(db_path) as db:
        run_id, phase, context_json = db.execute("SELECT run_id,phase,context_json FROM runs ORDER BY updated_at DESC LIMIT 1").fetchone()
        transitions = [{"ordinal": row[0], "from": row[1], "to": row[2], "at": row[3]} for row in db.execute("SELECT ordinal,from_phase,to_phase,transitioned_at FROM run_phase_transitions WHERE run_id=? ORDER BY ordinal", (run_id,))]
        calls = [{"call_id": row[0], "provider": row[1], "item_index": row[2], "state": row[3], "attempt_ordinal": row[4], "flight_fingerprint": row[5], "http_started": bool(row[6])} for row in db.execute("SELECT call_id,provider,item_index,state,attempt_ordinal,flight_fingerprint,http_started_at FROM provider_calls WHERE run_id=? ORDER BY created_at,call_id", (run_id,))]
        consumers = [{"provider": row[0], "fingerprint": row[1], "item_index": row[2], "call_id": row[3], "relation": row[4]} for row in db.execute("SELECT provider,query_fingerprint,item_index,provider_call_id,relation FROM provider_query_flight_consumers WHERE run_id=? ORDER BY item_index", (run_id,))]
        authorized_provider_plan = [row[0] for row in db.execute("SELECT DISTINCT provider FROM paid_attempt_provider_plan WHERE run_id=? AND authorized=1 ORDER BY provider", (run_id,))]
    original_db = config.PROGRESS_DB_FILE
    config.PROGRESS_DB_FILE = db_path
    try:
        telemetry = checkpoint.derive_telemetry(str(run_id))
    finally:
        config.PROGRESS_DB_FILE = original_db
    journal = list(getattr(transport, "journal", []))
    if transport is None:
        raise AssertionError("scenario capture requires a recording transport")
    if run_root is None or not (Path(run_root) / "manifest.json").is_file():
        raise AssertionError("scenario capture requires the production manifest")
    manifest_source = Path(run_root) / "manifest.json"
    manifest = json.loads(manifest_source.read_text(encoding="utf-8"))
    if sum(1 for row in calls if row["http_started"]) != len(journal):
        raise AssertionError("transport journal and durable HTTP-start ledger differ")
    relation_counts = {
        relation: sum(1 for row in consumers if row["relation"] == relation)
        for relation in ("OWNER", "INHERITED")
    }
    actual = {
        "run_id": run_id,
        "phase": phase,
        "outcome": getattr(getattr(outcome, "status", None), "value", str(outcome or "")),
        "http_invocations": len(journal),
        "phase_path": [f"{row['from']}->{row['to']}" for row in transitions],
        "relation_counts": relation_counts,
        "authorized_provider_plan": authorized_provider_plan,
        "provider_budgets": telemetry.get("provider_budgets", {}),
        "free_queries": telemetry.get("free_queries", {}),
        "calls": calls,
        "consumers": consumers,
    }
    (target / "recording_transport_journal.json").write_text(json.dumps(journal, indent=2), encoding="utf-8")
    (target / "frozen_config.json").write_text(json.dumps(json.loads(context_json or "{}"), indent=2, sort_keys=True), encoding="utf-8")
    (target / "state_transition_trace.json").write_text(json.dumps(transitions, indent=2), encoding="utf-8")
    shutil.copy2(manifest_source, target / "manifest.json")
    with sqlite3.connect(db_path) as db:
        plan_rows = db.execute("SELECT item_index,plan_version,query_kind,round_ordinal,query_ordinal,normalized_query,query_sha256 FROM paid_query_plan_entries WHERE run_id=? ORDER BY item_index,plan_version,query_kind,round_ordinal,query_ordinal", (run_id,)).fetchall()
    (target / "paid_query_plan.json").write_text(json.dumps([list(row) for row in plan_rows], indent=2), encoding="utf-8")
    files = manifest.get("files", {})
    artifact_dir = Path(run_root) / "output" / "artifacts" / str(manifest.get("artifact_set_sha256", ""))
    if files:
        if not artifact_dir.is_dir():
            raise AssertionError("manifest artifact set is missing")
        for filename, info in files.items():
            source = artifact_dir / filename
            if not source.is_file() or hashlib.sha256(source.read_bytes()).hexdigest() != info.get("sha256"):
                raise AssertionError(f"manifest artifact missing or corrupt: {filename}")
            shutil.copy2(source, target / filename)
    elif manifest.get("complete"):
        raise AssertionError("completed scenario lacks immutable artifacts")
    observations = {"scenario": name, "observations": actual}
    if not manifest.get("complete"):
        observations["observed_absent"] = ["telemetry.json", "report.txt"]
    (target / "scenario_result.json").write_text(json.dumps(observations, indent=2), encoding="utf-8")


def init_run(path: Path, count=1, budget=20):
    config.PROGRESS_DB_FILE = path
    checkpoint.initialize_schema(path)
    checkpoint.initialize_run(
        run_id="run", input_hash="input", run_signature="sig",
        context={"phase": "PAID", "paid_query_limit_per_company": 3,
                 "budget_details": {"brightdata": {"population_count": count, "explicit_cap": budget}}},
        budgets={"brightdata": budget},
        items=[{"item_index": i, "source_record_id": f"s{i}", "free_state": "DONE",
                "paid_required": True, "paid_state": "RUNNING"} for i in range(count)],
    )
    runtime.reset()
    runtime.configure_durable_run("run", {"brightdata": budget})
    runtime.set_phase("PAID")


@pytest.fixture
def durable(tmp_path, monkeypatch):
    path = tmp_path / "progress.sqlite3"
    init_run(path)
    runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake")
    monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 0)
    monkeypatch.setattr(config, "GLOBAL_REQUESTS_PER_SECOND", 0)
    monkeypatch.setattr(config, "BRIGHTDATA_TIMEOUT_SEC", 2)
    monkeypatch.setattr(config, "MAX_RETRIES", 0)
    return path


def test_real_authorized_e2e_uses_main_process_company_and_recording_transport(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"
    _input_book(source, ["FREE CO", "PAID CO"])
    transport = RecordingPaidTransport([Response(200, {"organic": [{"link": "https://paidco.example", "title": "PAID CO", "description": "PAID CO resmi şirket sitesi Türkiye"}]})])
    runs = _real_run_setup(tmp_path, monkeypatch, transport)
    outcome = main.run(source, allow_paid=True)
    root = next(runs.iterdir())
    assert outcome.status is pipeline_runner.PipelineOutcomeStatus.COMPLETE
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["phase"] == "COMPLETE" and len(transport.journal) == 1
    monkeypatch.setattr(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3")
    telemetry = checkpoint.derive_telemetry(manifest["run_id"])["provider_budgets"]["brightdata"]
    assert {key: telemetry[key] for key in ("effective_limit", "reserved_total", "done", "failed", "unknown", "physical_http_attempts")} == {"effective_limit": 8, "reserved_total": 1, "done": 1, "failed": 0, "unknown": 0, "physical_http_attempts": 1}
    _capture_scenario("authorized_free_paid", root / "state" / "progress.sqlite3", run_root=root, transport=transport, outcome=outcome)


def test_real_fresh_run_without_authorization_seals_handoff(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"
    book = Workbook(); sheet = book.active; sheet.append(["company", "source_record_id"]); sheet.append(["PAID CO", "s0"]); book.save(source); book.close()
    transport = RecordingPaidTransport([])
    _real_run_setup(tmp_path, monkeypatch, transport)
    outcome = main.run(source, allow_paid=False)
    root = next((tmp_path / "runs").iterdir()); manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    assert outcome.status is pipeline_runner.PipelineOutcomeStatus.PAID_PENDING_APPROVAL and manifest["handoff"] and not transport.journal
    _capture_scenario("unauthorized_handoff", root / "state" / "progress.sqlite3", run_root=root, transport=transport, outcome=outcome)


def test_real_unknown_e2e_has_one_total_paid_call_without_test_side_break(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"
    book = Workbook(); sheet = book.active; sheet.append(["company", "source_record_id"]); sheet.append(["PAID CO", "s0"]); book.save(source); book.close()
    transport = RecordingPaidTransport([requests.ReadTimeout("post-send unknown")])
    _real_run_setup(tmp_path, monkeypatch, transport)
    monkeypatch.setattr(config, "ENABLE_GOOGLE_PLACES", True); monkeypatch.setattr(config, "GOOGLE_PLACES_API_KEY", "fake"); monkeypatch.setattr(config, "GOOGLE_PLACES_REQUEST_BUDGET", 2)
    monkeypatch.setattr(config, "ENABLE_HUNTER_DOMAIN_FINDER", True); monkeypatch.setattr(config, "HUNTER_API_KEY", "fake"); monkeypatch.setattr(config, "HUNTER_REQUEST_BUDGET", 2)
    monkeypatch.setattr(config, "ENABLE_BRANDFETCH_DOMAIN_SEARCH", True); monkeypatch.setattr(config, "BRANDFETCH_API_KEY", "fake", raising=False); monkeypatch.setattr(config, "BRANDFETCH_REQUEST_BUDGET", 2)
    monkeypatch.setattr(config, "ENABLE_LINKEDIN_COMPANY_LOOKUP", True); monkeypatch.setattr(config, "LINKEDIN_COMPANY_REQUEST_BUDGET", 2)
    monkeypatch.setattr(config, "ENABLE_LLM_ARBITER", True); monkeypatch.setattr(config, "LLM_ARBITER_BUDGET", 2)
    outcome = main.run(source, allow_paid=True)
    root = next((tmp_path / "runs").iterdir())
    assert outcome.status is pipeline_runner.PipelineOutcomeStatus.PAID_MANUAL_AUTHORIZATION_REVIEW_REQUIRED and len(transport.journal) == 1
    assert not (root / "output" / "artifacts").exists()
    _capture_scenario("unknown_item_stop", root / "state" / "progress.sqlite3", run_root=root, transport=transport, outcome=outcome)


def test_paid_provider_stop_stops_remaining_provider_queries(monkeypatch):
    monkeypatch.setattr(config, "SEARCH_PROVIDER", "brightdata")
    monkeypatch.setattr(search, "_brightdata_circuit_open", lambda: False)
    monkeypatch.setattr(runtime, "free_search_capacity", lambda *_: runtime.FreeCapacity(False, "discovery", 0, 6, 6, 6))
    calls = []
    monkeypatch.setattr(search, "_search_text", lambda q: (calls.append(q) or search.SearchResults([], "live", "brightdata", result_state="FAILED", stop_scope="PAID_PROVIDER")))
    monkeypatch.setattr(search, "_primary_queries", lambda *_: ["q1", "q2"])
    search.find_candidate_domains("ACME", {})
    assert calls == ["q1"]


def test_two_distinct_http_5xx_owner_queries_open_circuit_and_third_makes_no_post(tmp_path, monkeypatch):
    init_run(tmp_path / "circuit.sqlite3", count=3, budget=3)
    monkeypatch.setattr(config, "SEARCH_PROVIDER", "brightdata")
    monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake")
    monkeypatch.setattr(config, "BRIGHTDATA_CIRCUIT_FAILURE_THRESHOLD", 2)
    monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 0)
    monkeypatch.setattr(config, "GLOBAL_REQUESTS_PER_SECOND", 0)
    monkeypatch.setattr(config, "MAX_RETRIES", 0)
    monkeypatch.setattr(runtime, "free_search_capacity", lambda *_: runtime.FreeCapacity(False, "discovery", 0, 6, 6, 6))
    posts = []
    monkeypatch.setattr(search.requests, "post", lambda *_a, **_k: (posts.append(1) or Response(500, {"error": "down"})))
    for index, query in enumerate(("a", "b")):
        runtime.reset_item_stop_state(index)
        checkpoint.begin_paid_attempt(run_id="run", item_index=index, attempt_number=1)
        assert search._brightdata_text(query).result_state == "FAILED"
    runtime.reset_item_stop_state(2)
    result = search._safe_search_text("c")
    assert len(posts) == 2 and search._brightdata_circuit_open() and result.origin is search.SearchOrigin.CIRCUIT_BLOCK


def test_capacity_read_is_side_effect_free_in_volatile_and_durable_modes(tmp_path):
    runtime.reset()
    before = runtime.snapshot()["counters"]
    for _ in range(1000):
        assert runtime.free_search_capacity("discovery")
    assert runtime.snapshot()["counters"] == before
    assert runtime.reserve_free_physical_attempt("discovery", "fp", "ddgs")
    init_run(tmp_path / "db.sqlite3")
    before = checkpoint.derive_telemetry("run")["free_queries"]
    for _ in range(1000):
        assert checkpoint.free_search_capacity(run_id="run", item_index=0, bucket="discovery")["available"]
    assert checkpoint.derive_telemetry("run")["free_queries"] == before


def test_ddgs_every_physical_reservation_is_completed_exactly_once(tmp_path, monkeypatch):
    init_run(tmp_path / "db.sqlite3")
    runtime.set_phase("FREE"); runtime.set_item_context(0, "free")
    outcomes = iter(([{"href": "https://x"}], [], RuntimeError("boom")))
    class DDGS:
        def __enter__(self): return self
        def __exit__(self, *_): return None
        def text(self, *_a, **_k):
            value = next(outcomes)
            if isinstance(value, Exception): raise value
            return value
    monkeypatch.setattr(search, "DDGS", DDGS)
    monkeypatch.setattr(search, "PREFERRED_BACKENDS", ("one",))
    monkeypatch.setattr(search, "FALLBACK_BACKENDS", ())
    search._ddgs_text("a"); search._ddgs_text("b"); search._ddgs_text("c")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        assert db.execute("SELECT state,COUNT(*) FROM free_provider_attempts GROUP BY state ORDER BY state").fetchall() == [("DONE", 2), ("FAILED", 1)]


def test_brightdata_never_touches_free_attempt_ledger(durable, monkeypatch):
    monkeypatch.setattr(search.requests, "post", lambda *_a, **_k: Response(200, {"organic": []}))
    search._brightdata_text("q")
    with sqlite3.connect(durable) as db:
        assert db.execute("SELECT COUNT(*) FROM free_provider_attempts").fetchone()[0] == 0


def test_brightdata_200_error_header_overrides_valid_organic_body(durable, monkeypatch):
    monkeypatch.setattr(search.requests, "post", lambda *_a, **_k: Response(200, {"organic": [{"link": "https://x"}]}, {"X-Brd-Err-Code": "denied", "Proxy-Status": "proxy; error=denied"}))
    assert search._brightdata_text("q").result_state == "FAILED"


def test_request_slot_wait_occurs_before_every_brightdata_post(durable, monkeypatch):
    events = []
    monkeypatch.setattr(runtime, "wait_for_request_slot", lambda **_kwargs: events.append("wait"))
    monkeypatch.setattr(search.requests, "post", lambda *_a, **_k: (events.append("post") or Response(200, {"organic": []})))
    search._brightdata_text("q")
    assert events == ["wait", "post"]


def singleflight(tmp_path, monkeypatch, status):
    init_run(tmp_path / "flight.sqlite3", count=8, budget=8)
    monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake")
    monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 0)
    monkeypatch.setattr(config, "GLOBAL_REQUESTS_PER_SECOND", 0)
    monkeypatch.setattr(config, "BRIGHTDATA_TIMEOUT_SEC", 2)
    monkeypatch.setattr(config, "MAX_RETRIES", 0)
    posts = []
    monkeypatch.setattr(search.requests, "post", lambda *_a, **_k: (posts.append(1) or Response(status, {"organic": []} if status == 200 else {"error": "x"})))
    barrier = threading.Barrier(8)
    def worker(index):
        runtime.set_item_context(index, "paid")
        checkpoint.begin_paid_attempt(run_id="run", item_index=index, attempt_number=1)
        barrier.wait()
        result = search._brightdata_text("same")
        owned = bool(checkpoint.provider_calls_for_item("run", index))
        checkpoint.record_paid_attempt(run_id="run", item_index=index, attempt_number=1, result="COMPLETED" if status == 200 else "FAILED", call_ids=list(result.call_ids), call_relations={call_id: "OWNER" if owned else "INHERITED" for call_id in result.call_ids})
        with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
            db.execute("UPDATE run_items SET paid_state=? WHERE run_id='run' AND item_index=?", ("DONE" if status == 200 else "FAILED", index)); db.commit()
        return result
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(worker, range(8)))
    checkpoint.validate_paid_evidence("run")
    return posts, results


def test_singleflight_eight_workers_success_make_one_real_post(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"; _input_book(source, ["ACME"] * 8)
    transport = RecordingPaidTransport([Response(200, {"organic": [{"link": "https://acme.example", "title": "ACME", "description": "ACME resmi şirket sitesi Türkiye"}]})])
    runs = _real_run_setup(tmp_path, monkeypatch, transport, workers=8)
    outcome = main.run(source, allow_paid=True); root = next(runs.iterdir())
    monkeypatch.setattr(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        relations = db.execute("SELECT relation,COUNT(*) FROM provider_query_flight_consumers GROUP BY relation ORDER BY relation").fetchall()
    assert outcome.status is pipeline_runner.PipelineOutcomeStatus.COMPLETE and len(transport.journal) == 1 and relations == [("INHERITED", 7), ("OWNER", 1)]
    _capture_scenario("singleflight_success", config.PROGRESS_DB_FILE, run_root=root, transport=transport, outcome=outcome)


def test_single_failed_flight_with_eight_followers_counts_one_circuit_failure(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"; _input_book(source, ["ACME"] * 8)
    transport = RecordingPaidTransport([Response(400, {"error": "definite"})])
    runs = _real_run_setup(tmp_path, monkeypatch, transport, workers=8)
    outcome = main.run(source, allow_paid=True); root = next(runs.iterdir())
    monkeypatch.setattr(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        relations = db.execute("SELECT relation,COUNT(*) FROM provider_query_flight_consumers GROUP BY relation ORDER BY relation").fetchall()
    assert outcome.status is pipeline_runner.PipelineOutcomeStatus.COMPLETE and len(transport.journal) == 1 and relations == [("INHERITED", 7), ("OWNER", 1)] and search._BRIGHTDATA_CONSECUTIVE_FAILURES == 1
    _capture_scenario("singleflight_failure", config.PROGRESS_DB_FILE, run_root=root, transport=transport, outcome=outcome)


def test_heartbeat_covers_throttle_http_decode_and_retry_backoff(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"; _input_book(source, ["ACME", "ACME"])
    transport = RecordingPaidTransport([
        Response(429, {"error": "rate"}, {"Retry-After": "3"}),
        Response(200, {"organic": [{"link": "https://acme.example", "title": "ACME", "description": "ACME resmi şirket sitesi Türkiye"}]}),
    ])
    runs = _real_run_setup(tmp_path, monkeypatch, transport, workers=2)
    monkeypatch.setattr(config, "MAX_RETRIES", 1); monkeypatch.setattr(config, "MAX_RETRY_AFTER_SEC", 3)
    monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 600)
    monkeypatch.setattr(search, "_flight_lease_seconds", lambda: 1.0)
    monkeypatch.setattr(search, "_flight_heartbeat_interval", lambda: 0.2)
    outcome = main.run(source, allow_paid=True); root = next(runs.iterdir())
    monkeypatch.setattr(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        states = db.execute("SELECT paid_state,COUNT(*) FROM run_items GROUP BY paid_state").fetchall()
        relations = db.execute("SELECT relation,COUNT(*) FROM provider_query_flight_consumers GROUP BY relation ORDER BY relation").fetchall()
    assert outcome.status is pipeline_runner.PipelineOutcomeStatus.COMPLETE and len(transport.journal) == 2 and states == [("DONE", 2)] and relations == [("INHERITED", 2), ("OWNER", 2)]
    _capture_scenario("heartbeat_throttle_backoff", config.PROGRESS_DB_FILE, run_root=root, transport=transport, outcome=outcome)


def test_expired_potentially_charged_flight_becomes_unknown_without_reclaim(tmp_path):
    init_run(tmp_path / "db.sqlite3")
    runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="old", lease_seconds=1)
    reservation = runtime.reserve_api("brightdata", operation="search.attempt_1", request_fingerprint="x")
    checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="old", provider_call_id=reservation.call_id)
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        db.execute("UPDATE provider_query_flights SET lease_expires_at='2000-01-01T00:00:00+00:00'"); db.commit()
    claim = checkpoint.resolve_expired_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="new")
    assert not claim["leader"] and claim["state"] == "UNKNOWN" and claim["call_ids"] == [reservation.call_id]
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        assert db.execute("SELECT state FROM provider_calls WHERE call_id=?", (reservation.call_id,)).fetchone() == ("UNKNOWN",)
        assert db.execute("SELECT state,call_ids_json FROM provider_query_flight_terminals").fetchone() == ("UNKNOWN", json.dumps([reservation.call_id]))
        assert db.execute("SELECT reserved,unknown FROM provider_usage WHERE provider='brightdata'").fetchone() == (0, 1)


def test_nonexistent_paid_evidence_reference_is_rejected(durable):
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="NO_CALL_NEEDED", evidence_kind="supplied_website_publishable_at_paid_entry", input_snapshot_sha256="does-not-exist")
    with sqlite3.connect(durable) as db:
        db.execute("UPDATE run_items SET paid_state='DONE'"); db.commit()
    with pytest.raises(checkpoint.SchedulerInvariantError):
        checkpoint.validate_paid_evidence("run")


def test_current_paid_attempt_must_link_terminal_provider_call(durable):
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="FAILED")
    with sqlite3.connect(durable) as db:
        db.execute("UPDATE run_items SET paid_state='FAILED'"); db.commit()
    with pytest.raises(checkpoint.SchedulerInvariantError):
        checkpoint.validate_paid_evidence("run")


def test_real_fresh_handoff_resume_preserves_limit_and_plan(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"
    _input_book(source, ["PAID CO"])
    paid_response = {"organic": [{"link": "https://paidco.example", "title": "PAID CO", "description": "PAID CO resmi şirket sitesi Türkiye"}]}
    transport = RecordingPaidTransport([Response(200, paid_response) for _ in range(3)])
    runs = _real_run_setup(tmp_path, monkeypatch, transport)
    monkeypatch.setattr(config, "MAX_SEARCH_QUERIES_PER_COMPANY", 3)
    monkeypatch.setattr(config, "DEFAULT_PAID_SEARCH_QUERY_LIMIT", 3)

    fresh = main.run(source, allow_paid=False)
    parent = next(runs.iterdir())
    parent_manifest = json.loads((parent / "manifest.json").read_text(encoding="utf-8"))
    assert fresh.status is pipeline_runner.PipelineOutcomeStatus.PAID_PENDING_APPROVAL
    frozen_limit = parent_manifest["paid_query_limit_per_company"]
    source_id = parent_manifest["ordered_source_record_ids"][0]
    frozen_plan = search._primary_queries("PAID CO", {"source_record_id": source_id})[:frozen_limit]
    runtime.reset()
    search.reset_run_state()

    approval = {
        "approved": True,
        "limits": {provider: (3 if provider == "brightdata" else 0) for provider in checkpoint.CANONICAL_PROVIDERS},
        "parent_run_id": parent_manifest["run_id"],
        "parent_manifest_sha256": hashlib.sha256((parent / "manifest.json").read_bytes()).hexdigest(),
        "checkpoint_sha256": parent_manifest.get("_continuation_checkpoint_sha256") or parent_manifest.get("checkpoint_sha256") or parent_manifest["files"]["recovery_state.sqlite3"]["sha256"],
        "paid_source_ids_sha256": hashlib.sha256(run_context.canonical_json([source_id]).encode()).hexdigest(),
    }
    authorization = tmp_path / "approval.json"
    authorization.write_text(json.dumps(approval), encoding="utf-8")
    child_manifest = prepare_paid_continuation(parent, authorization, tmp_path / "continuation")
    child = tmp_path / "continuation" / "runs" / child_manifest["run_id"]
    assert child_manifest["paid_query_limit_per_company"] == frozen_limit

    monkeypatch.setattr(config, "MAX_SEARCH_QUERIES_PER_COMPANY", 99)
    completed = main.run(source, allow_paid=True, resume_run=child)
    final_manifest = json.loads((child / "manifest.json").read_text(encoding="utf-8"))
    assert completed.status is pipeline_runner.PipelineOutcomeStatus.COMPLETE
    assert final_manifest["paid_query_limit_per_company"] == frozen_limit
    assert search._primary_queries("PAID CO", {"source_record_id": source_id})[:frozen_limit] == frozen_plan
    with sqlite3.connect(child / "state" / "progress.sqlite3") as db:
        provider_plan = db.execute(
            "SELECT provider,authorized,effective_limit FROM paid_attempt_provider_plan ORDER BY plan_ordinal"
        ).fetchall()
    assert {provider: (authorized, limit) for provider, authorized, limit in provider_plan} == {
        provider: ((1, 3) if provider == "brightdata" else (0, 0))
        for provider in checkpoint.CANONICAL_PROVIDERS
    }
    _capture_scenario("fresh_handoff_resume", child / "state" / "progress.sqlite3", run_root=child, transport=transport, outcome=completed)


def test_paid_profile_candidate_does_not_bypass_paid_contract(monkeypatch):
    candidate = {"url": "https://profile.example", "domain": "profile.example", "score": 100}
    evaluation = {"candidate": candidate, "crawl_result": {"pages": [1]}, "identity_assessment": {"provisionally_publishable": True}, "reasons": []}
    monkeypatch.setattr(search, "find_profile_candidates", lambda *_: [candidate])
    monkeypatch.setattr(main, "_evaluate_candidate_with_stage", lambda *_a, **_k: evaluation)
    monkeypatch.setattr(main.entity_resolution, "resolve_profile_anchor", lambda *_: SimpleNamespace(status="resolved", selected=dict(evaluation), reason="profile"))
    monkeypatch.setattr(main, "_finalize_selected_evaluation", lambda *_: {"company": "ACME", "status": "OK_HIGH_CONFIDENCE", "publication_eligible": True, "reason": "profile"})
    called = []
    monkeypatch.setattr(search, "find_candidate_domains", lambda *_: (called.append(1) or search.CandidateList([])))
    main.process_company(0, "ACME", logging.getLogger("test"), metadata={}, execution_phase="PAID")
    assert called == [1]


def test_cli_maps_all_typed_outcomes_and_rejects_unknown_result(monkeypatch):
    args = argparse.Namespace(run_dir=None, resume_run=None, only_status="", from_run_manifest=None, non_interactive=True, input=Path("x"), allow_paid=False, search_cache=None, crawl_cache=None, brightdata_budget=None, google_places_budget=None, hunter_budget=None, brandfetch_budget=None, linkedin_company_budget=None, llm_budget=None, rerank_cache=False, companies="", replay_snapshot=None, replay_manifest=None)
    monkeypatch.setattr(main, "_ensure_safe_project_runtime", lambda *_: None)
    monkeypatch.setattr(main, "parse_args", lambda *_: args)
    monkeypatch.setattr(main, "_apply_cli_options", lambda *_: None)
    monkeypatch.setattr(main, "resolve_cli_run_config", lambda *_: None)
    monkeypatch.setattr(main, "_apply_saved_resolver_configuration", lambda: None)
    cases = [(pipeline_runner.PipelineOutcomeStatus.COMPLETE, 0), (pipeline_runner.PipelineOutcomeStatus.COMPLETE_RESUME_VERIFIED, 0), (pipeline_runner.PipelineOutcomeStatus.FINALIZATION_RESUME_RECONCILED, 0), (pipeline_runner.PipelineOutcomeStatus.PAID_PENDING_APPROVAL, 20), (pipeline_runner.PipelineOutcomeStatus.PAID_MANUAL_AUTHORIZATION_REVIEW_REQUIRED, 21), (pipeline_runner.PipelineOutcomeStatus.FINALIZATION_INVARIANT, 22)]
    for status, code in cases:
        monkeypatch.setattr(main, "run", lambda *_a, _status=status, **_k: pipeline_runner.PipelineOutcome(_status))
        assert main.cli([]) == code
    monkeypatch.setattr(main, "run", lambda *_a, **_k: None); assert main.cli([]) == 22
    monkeypatch.setattr(main, "run", lambda *_a, **_k: (_ for _ in ()).throw(OSError("boom"))); assert main.cli([]) == 1


def test_provider_and_free_ledger_equations_gate_finalization(durable, monkeypatch):
    monkeypatch.setattr(search.requests, "post", lambda *_a, **_k: Response(200, {"organic": []}))
    result = search._brightdata_text("q")
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="COMPLETED", call_ids=list(result.call_ids))
    assert checkpoint.validate_ledger_equations("run")["provider_budgets"]["brightdata"]["done"] == 1


def test_provider_budget_block_event_is_unique_per_scope(tmp_path):
    init_run(tmp_path / "db.sqlite3", budget=0)
    runtime.set_item_context(0, "paid")
    for _ in range(5):
        runtime.reserve_api("brightdata", operation="search", request_fingerprint="same")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        assert db.execute("SELECT COUNT(*) FROM provider_budget_blocks").fetchone()[0] == 1


def test_schema_migration_not_repeated_after_normal_write_wal_and_checkpoint(tmp_path, monkeypatch):
    path = tmp_path / "db.sqlite3"; checkpoint.initialize_schema(path)
    calls = []; original = checkpoint._migrate_schema
    monkeypatch.setattr(checkpoint, "_migrate_schema", lambda db: (calls.append(1), original(db))[-1])
    with sqlite3.connect(path) as db:
        db.execute("CREATE TABLE IF NOT EXISTS ordinary(x)"); db.commit(); db.execute("PRAGMA wal_checkpoint(TRUNCATE)")
    checkpoint.initialize_schema(path)
    assert calls == []


def test_schema_cache_invalidates_after_delete_and_recreate_same_path(tmp_path, monkeypatch):
    path = tmp_path / "db.sqlite3"; checkpoint.initialize_schema(path)
    path.unlink(); sqlite3.connect(path).close()
    calls = []; original = checkpoint._migrate_schema
    monkeypatch.setattr(checkpoint, "_migrate_schema", lambda db: (calls.append(1), original(db))[-1])
    checkpoint.initialize_schema(path)
    assert calls == [1]


def test_incomplete_v11_schema_migrates_transport_receipts_to_v12(tmp_path):
    path = tmp_path / "v11.sqlite3"; init_run(path)
    with sqlite3.connect(path) as db:
        for (trigger,) in db.execute("SELECT name FROM sqlite_master WHERE type='trigger'").fetchall():
            db.execute(f'DROP TRIGGER "{trigger}"')
        db.execute("ALTER TABLE provider_calls DROP COLUMN endpoint_sha256")
        db.execute("ALTER TABLE provider_calls DROP COLUMN request_shape_sha256")
        db.execute("PRAGMA user_version=11"); db.commit()
    checkpoint._SCHEMA_READY.clear(); checkpoint.initialize_schema(path)
    with sqlite3.connect(path) as db:
        columns = {row[1] for row in db.execute("PRAGMA table_info(provider_calls)")}
        triggers = {row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type='trigger'")}
        assert db.execute("PRAGMA user_version").fetchone() == (12,)
        assert {"endpoint_sha256", "request_shape_sha256"}.issubset(columns)
        assert {"provider_call_transport_receipt_update", "flight_result_receipt_delete", "flight_terminal_receipt_delete"}.issubset(triggers)


def test_default_zero_paid_budgets_remain_fail_closed_without_local_patch(monkeypatch):
    runtime.reset(); runtime.set_phase("PAID")
    monkeypatch.setattr(config, "BRIGHTDATA_REQUEST_BUDGET", 0)
    assert not runtime.reserve_api("brightdata", budget=0)


def test_protected_manifest_detects_same_size_same_mtime_content_change(tmp_path, monkeypatch):
    import conftest
    root = tmp_path / "protected"; root.mkdir()
    target = root / "x.bin"; target.write_bytes(b"AAAA")
    stamp = target.stat().st_mtime_ns
    monkeypatch.setattr(conftest, "_PROTECTED_TEST_ROOTS", (root,))
    monkeypatch.setattr(conftest, "_REPO", tmp_path)
    before = conftest._protected_manifest()
    monkeypatch.setattr(conftest, "_PROTECTED_TEST_ROOTS", ())
    target.write_bytes(b"BBBB"); os.utime(target, ns=(stamp, stamp))
    monkeypatch.setattr(conftest, "_PROTECTED_TEST_ROOTS", (root,))
    assert before != conftest._protected_manifest()


def test_unknown_from_primary_blocks_targeted_and_every_later_paid_resolver(durable, monkeypatch):
    runtime.reset_item_stop_state(0)
    calls = []
    monkeypatch.setattr(search.requests, "post", lambda *_a, **_k: (calls.append("brightdata") or (_ for _ in ()).throw(requests.ReadTimeout("post-send"))))
    result = search._brightdata_text("primary")
    assert result.stop_scope is runtime.StopScope.MANUAL_AUTHORIZATION
    for provider in ("brightdata", "google_places", "hunter", "brandfetch", "linkedin", "llm"):
        rejected = runtime.reserve_api(provider, budget=1, operation="later", request_fingerprint=provider)
        assert not rejected and rejected.reason == "item_stop_guard"
    assert calls == ["brightdata"]


def test_runtime_paid_stop_guard_prevents_calls_even_if_caller_forgets_to_stop(durable):
    runtime.mark_item_stop(runtime.StopScope.MANUAL_AUTHORIZATION, "brightdata", "uncertain", ("original",))
    before = checkpoint.derive_telemetry("run")["provider_budgets"]["brightdata"]["reserved_total"]
    rejected = runtime.reserve_api("brightdata", operation="forgotten", request_fingerprint="new")
    result = runtime.rejected_provider_result(rejected)
    after = checkpoint.derive_telemetry("run")["provider_budgets"]["brightdata"]["reserved_total"]
    assert not rejected and result.result_state == "UNKNOWN" and result.call_ids == ("original",) and before == after


def test_cache_hit_and_singleflight_follower_do_not_mutate_circuit(monkeypatch):
    search.reset_run_state()
    monkeypatch.setattr(config, "BRIGHTDATA_CIRCUIT_FAILURE_THRESHOLD", 2)
    search._observe_brightdata_owner(flight_fingerprint="cache", state="FAILED", http_attempted=False)
    search.SearchResults([], "cache_hit", "brightdata", result_state="FAILED", origin=search.SearchOrigin.CROSS_RUN_CACHE)
    search.SearchResults([], "singleflight", "brightdata", result_state="FAILED", origin=search.SearchOrigin.SINGLEFLIGHT_FOLLOWER)
    assert search._BRIGHTDATA_CONSECUTIVE_FAILURES == 0 and not search._brightdata_circuit_open()


def test_expired_no_call_or_all_failed_flight_is_safely_reclaimed(tmp_path):
    init_run(tmp_path / "reclaim.sqlite3", budget=2)
    first = checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="none", owner_token="old", lease_seconds=1)
    assert first["leader"]
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        db.execute("UPDATE provider_query_flights SET lease_expires_at='2000-01-01T00:00:00+00:00'"); db.commit()
    assert checkpoint.resolve_expired_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="none", owner_token="new")["leader"]
    runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="failed", owner_token="old", lease_seconds=1)
    reservation = runtime.reserve_api("brightdata", operation="x", request_fingerprint="failed-call")
    checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="failed", owner_token="old", provider_call_id=reservation.call_id)
    runtime.start_api(reservation); runtime.complete_api(reservation, "FAILED")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        db.execute("UPDATE provider_query_flights SET lease_expires_at='2000-01-01T00:00:00+00:00' WHERE query_fingerprint='failed'"); db.commit()
    assert checkpoint.resolve_expired_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="failed", owner_token="new")["leader"]


def test_retry_budget_rejection_preserves_prior_failed_call_ids(durable, monkeypatch):
    monkeypatch.setattr(config, "MAX_RETRIES", 1)
    runtime._DURABLE_BUDGETS["brightdata"] = 1
    with sqlite3.connect(durable) as db:
        db.execute("UPDATE provider_usage SET configured_limit=1,effective_limit=1 WHERE provider='brightdata'"); db.commit()
    monkeypatch.setattr(search.requests, "post", lambda *_a, **_k: Response(500, {"error": "retry"}))
    monkeypatch.setattr(search.time, "sleep", lambda *_: None)
    result = search._brightdata_text("retry")
    assert result.result_state in {"FAILED", "BLOCKED_BUDGET"} and len(result.call_ids) == 1
    with sqlite3.connect(durable) as db:
        flight_ids = json.loads(db.execute("SELECT call_ids_json FROM provider_query_flights").fetchone()[0])
    assert flight_ids == list(result.call_ids)


def test_same_run_cache_duplicate_writes_typed_inherited_consumer(tmp_path):
    init_run(tmp_path / "consumer.sqlite3", count=2, budget=2)
    runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    claim = checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner")
    reservation = runtime.reserve_api("brightdata", operation="x", request_fingerprint="owner")
    runtime.start_api(reservation); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=reservation.call_id)
    runtime.mark_api_http_started(reservation, 1, "fp"); checkpoint.bind_provider_call_transport_receipt(call_id=reservation.call_id, endpoint_sha256="a" * 64, request_shape_sha256="b" * 64)
    checkpoint.complete_provider_call_and_flight_success(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=reservation.call_id, result={"__search_result_state": "COMPLETED", "values": []}, call_ids=[reservation.call_id])
    runtime.set_item_context(1, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=1, attempt_number=1)
    checkpoint.record_paid_attempt(run_id="run", item_index=1, attempt_number=1, result="COMPLETED", call_ids=[reservation.call_id], call_relations={reservation.call_id: "INHERITED"})
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        assert db.execute("SELECT relation,item_index FROM provider_query_flight_consumers WHERE paid_attempt_id=(SELECT paid_attempt_id FROM paid_attempts WHERE item_index=1)").fetchone() == ("INHERITED", 1)


def test_cross_run_cache_cannot_satisfy_paid_attempt(durable):
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="COMPLETED", reason="cross_run_cache")
    with sqlite3.connect(durable) as db:
        db.execute("UPDATE run_items SET paid_state='DONE'"); db.commit()
    with pytest.raises(checkpoint.EvidenceInvariant): checkpoint.validate_paid_evidence("run")


def test_unrelated_same_run_cross_item_done_call_is_rejected(tmp_path):
    init_run(tmp_path / "cross.sqlite3", count=2, budget=2)
    runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    reservation = runtime.reserve_api("brightdata", operation="x", request_fingerprint="x"); runtime.start_api(reservation); runtime.complete_api(reservation, "DONE")
    runtime.set_item_context(1, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=1, attempt_number=1)
    with pytest.raises(checkpoint.EvidenceInvariant): checkpoint.record_paid_attempt(run_id="run", item_index=1, attempt_number=1, result="COMPLETED", call_ids=[reservation.call_id])


def test_owner_relation_requires_provider_call_item_match(tmp_path):
    test_unrelated_same_run_cross_item_done_call_is_rejected(tmp_path)


def test_inherited_relation_requires_matching_terminal_flight_fingerprint(tmp_path):
    init_run(tmp_path / "bad-fp.sqlite3", count=2, budget=2)
    runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    reservation = runtime.reserve_api("brightdata", operation="x", request_fingerprint="x"); runtime.start_api(reservation); runtime.complete_api(reservation, "DONE")
    runtime.set_item_context(1, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=1, attempt_number=1)
    with pytest.raises(checkpoint.EvidenceInvariant): checkpoint.record_paid_attempt(run_id="run", item_index=1, attempt_number=1, result="COMPLETED", call_ids=[reservation.call_id], call_relations={reservation.call_id: "INHERITED"})


def test_paid_state_and_current_attempt_result_must_match(durable):
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="FAILED")
    with sqlite3.connect(durable) as db:
        db.execute("UPDATE run_items SET paid_state='DONE'"); db.commit()
    with pytest.raises(checkpoint.EvidenceInvariant): checkpoint.validate_paid_evidence("run")


def _valid_no_call(path: Path):
    config.PROGRESS_DB_FILE = path; checkpoint.initialize_schema(path)
    checkpoint.initialize_run(run_id="run", input_hash="i", run_signature="s", context={"phase":"PAID"}, budgets={"brightdata":1}, items=[{"item_index":0,"source_record_id":"s0","company":"ACME","website":"https://acme.example","free_state":"DONE","paid_required":True,"paid_state":"RUNNING"}])
    runtime.reset(); runtime.configure_durable_run("run", {"brightdata":1}); runtime.set_phase("PAID"); runtime.set_item_context(0)
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    snap = checkpoint.immutable_input_snapshot_sha256("run", 0)
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="NO_CALL_NEEDED", evidence_kind="supplied_website_publishable_at_paid_entry", input_snapshot_sha256=snap)
    payload = {"company":"ACME","source_record_id":"s0","website":"https://acme.example","status":"OK_HIGH_CONFIDENCE","publication_eligible":True,"known_website_evaluation":{"status":"OK_HIGH_CONFIDENCE","website":"https://acme.example"}}
    checkpoint.save_item_transaction(run_id="run", item_index=0, source_record_id="s0", payload=payload, free_state="DONE", paid_state="DONE", paid_required=True, free_attempts=1, paid_attempts=1)
    checkpoint.record_paid_no_call_evidence(run_id="run", item_index=0, attempt_number=1)


def test_supplied_no_call_receipt_validates_input_website_evaluation_and_result(tmp_path):
    _valid_no_call(tmp_path / "no-call.sqlite3")
    checkpoint.validate_paid_evidence("run")


def test_forged_supplied_no_call_reason_and_snapshot_hash_are_rejected(durable):
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="NO_CALL_NEEDED", evidence_kind="supplied_website_publishable_at_paid_entry", input_snapshot_sha256=checkpoint.immutable_input_snapshot_sha256("run", 0))
    with sqlite3.connect(durable) as db:
        db.execute("UPDATE run_items SET paid_state='DONE'"); db.commit()
    with pytest.raises(checkpoint.EvidenceInvariant): checkpoint.validate_paid_evidence("run")


def test_blocked_budget_requires_current_attempt_frozen_provider_plan(durable):
    checkpoint.record_provider_budget_block(run_id="run", item_index=0, provider="hunter", block_kind="budget_exhausted")
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="BLOCKED_BUDGET")
    with sqlite3.connect(durable) as db:
        db.execute("UPDATE run_items SET paid_state='BLOCKED_BUDGET'"); db.commit()
    with pytest.raises(checkpoint.EvidenceInvariant): checkpoint.validate_paid_evidence("run")


def test_free_block_is_unique_across_backends_and_never_counted_as_paid(tmp_path):
    init_run(tmp_path / "free-block.sqlite3")
    for backend in ("bing", "duckduckgo", "yahoo"):
        checkpoint.record_provider_budget_block(run_id="run", item_index=0, provider=backend, bucket="discovery", block_kind="physical", backend=backend)
    telemetry = checkpoint.derive_telemetry("run")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        assert db.execute("SELECT COUNT(*),MIN(provider) FROM provider_budget_blocks").fetchone() == (1, "ddgs")
    assert telemetry["provider_budgets"]["brightdata"]["budget_blocked_items"] == 0


def test_incomplete_legacy_physical_usage_cannot_gain_new_slots(tmp_path):
    path = tmp_path / "legacy.sqlite3"; init_run(path)
    with sqlite3.connect(path) as db:
        db.execute("INSERT OR REPLACE INTO free_query_usage(run_id,item_index,quota,logical_used,physical_used,discovery_physical_used,physical_completed) VALUES('run',0,10,6,6,6,6)"); db.execute("DELETE FROM free_provider_attempts"); db.commit()
    checkpoint._SCHEMA_READY.clear()
    checkpoint.initialize_schema(path)
    runtime.set_phase("FREE"); runtime.set_item_context(0); runtime.set_search_bucket("discovery")
    assert not runtime.reserve_free_physical_attempt("discovery", "seventh", "bing")


def test_process_company_rethrows_scheduler_invariant(monkeypatch):
    monkeypatch.setattr(search, "find_profile_candidates", lambda *_: [])
    monkeypatch.setattr(search, "find_candidate_domains", lambda *_: (_ for _ in ()).throw(checkpoint.LedgerInvariant("bad")))
    with pytest.raises(checkpoint.LedgerInvariant): main.process_company(0, "ACME", logging.getLogger("test"), metadata={})


def test_worker_collector_rethrows_scheduler_invariant_without_free_retry(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"; _input_book(source, ["ACME"])
    monkeypatch.setattr(config, "RUNS_DIR", tmp_path / "runs")
    attempts = []
    def worker(*_a, **_k): attempts.append(1); raise checkpoint.LedgerInvariant("bad")
    with pytest.raises(checkpoint.LedgerInvariant): pipeline_runner.run_pipeline(source, process_company_fn=worker, write_outputs_fn=main._write_outputs, set_output_dir_fn=main._set_output_dir, empty_result_fn=main._empty_result)
    assert attempts == [1]


def test_resume_prevalidation_invariant_returns_cli_22(monkeypatch):
    args = argparse.Namespace(run_dir=None, resume_run=Path("run"), only_status="", from_run_manifest=None, non_interactive=True, input=Path("x"), allow_paid=False, search_cache=None, crawl_cache=None, brightdata_budget=None, google_places_budget=None, hunter_budget=None, brandfetch_budget=None, linkedin_company_budget=None, llm_budget=None, rerank_cache=False, companies="", replay_snapshot=None, replay_manifest=None)
    monkeypatch.setattr(main, "_ensure_safe_project_runtime", lambda *_: None); monkeypatch.setattr(main, "parse_args", lambda *_: args); monkeypatch.setattr(main, "_apply_cli_options", lambda *_: None); monkeypatch.setattr(main, "resolve_cli_run_config", lambda *_: None)
    monkeypatch.setattr(pipeline_runner, "validate_resume_before_credentials", lambda *_a, **_k: (_ for _ in ()).throw(checkpoint.ResumeInvariant("drift")))
    assert main.cli([]) == 22


def test_invalid_pipeline_outcome_status_returns_cli_22_without_attribute_error(monkeypatch):
    args = argparse.Namespace(run_dir=None, resume_run=None, only_status="", from_run_manifest=None, non_interactive=True, input=Path("x"), allow_paid=False, search_cache=None, crawl_cache=None, brightdata_budget=None, google_places_budget=None, hunter_budget=None, brandfetch_budget=None, linkedin_company_budget=None, llm_budget=None, rerank_cache=False, companies="", replay_snapshot=None, replay_manifest=None)
    monkeypatch.setattr(main, "_ensure_safe_project_runtime", lambda *_: None); monkeypatch.setattr(main, "parse_args", lambda *_: args); monkeypatch.setattr(main, "_apply_cli_options", lambda *_: None); monkeypatch.setattr(main, "resolve_cli_run_config", lambda *_: None); monkeypatch.setattr(main, "_apply_saved_resolver_configuration", lambda: None)
    monkeypatch.setattr(main, "run", lambda *_a, **_k: pipeline_runner.PipelineOutcome("bad"))
    assert main.cli([]) == 22


@pytest.mark.parametrize("column", ("reserved_total", "reserved", "completed", "failed", "unknown"))
def test_provider_usage_corruption_in_each_aggregate_column_blocks_finalization(tmp_path, column):
    path = tmp_path / f"{column}.sqlite3"; init_run(path)
    with sqlite3.connect(path) as db:
        db.execute(f"UPDATE provider_usage SET {column}={column}+1 WHERE provider='brightdata'"); db.commit()
    with pytest.raises(checkpoint.LedgerInvariant): checkpoint.validate_ledger_equations("run")


def test_reservation_without_http_start_is_not_physical_attempt(durable):
    reservation = runtime.reserve_api("brightdata", operation="reserve", request_fingerprint="reserve")
    assert reservation and checkpoint.derive_telemetry("run")["provider_budgets"]["brightdata"]["physical_http_attempts"] == 0


def test_http_start_marker_is_exactly_once(durable):
    reservation = runtime.reserve_api("brightdata", operation="start", request_fingerprint="start"); runtime.start_api(reservation); runtime.mark_api_http_started(reservation, 1, "fp")
    with pytest.raises(checkpoint.LedgerInvariant): runtime.mark_api_http_started(reservation, 1, "fp")
    with pytest.raises(checkpoint.LedgerInvariant, match="SHA-256"):
        checkpoint.bind_provider_call_transport_receipt(call_id=reservation.call_id, endpoint_sha256="z" * 64, request_shape_sha256="b" * 64)
    checkpoint.bind_provider_call_transport_receipt(call_id=reservation.call_id, endpoint_sha256="a" * 64, request_shape_sha256="b" * 64)
    with pytest.raises(checkpoint.LedgerInvariant, match="exactly once"):
        checkpoint.bind_provider_call_transport_receipt(call_id=reservation.call_id, endpoint_sha256="a" * 64, request_shape_sha256="b" * 64)
    with sqlite3.connect(durable) as db, pytest.raises(sqlite3.IntegrityError, match="transport receipt immutable"):
        db.execute("UPDATE provider_calls SET endpoint_sha256=? WHERE call_id=?", ("c" * 64, reservation.call_id))


def test_retry_count_uses_attempt_ordinal_not_operation_string(durable):
    reservation = runtime.reserve_api("brightdata", operation="no_retry_word", request_fingerprint="ordinal"); runtime.start_api(reservation); runtime.mark_api_http_started(reservation, 2, "fp")
    telemetry = checkpoint.derive_telemetry("run")["provider_budgets"]["brightdata"]
    assert telemetry["retry_attempts"] == 1


def test_report_manifest_and_checkpoint_use_identical_provider_values(durable):
    reservation = runtime.reserve_api("brightdata", operation="report", request_fingerprint="report"); runtime.start_api(reservation); runtime.mark_api_http_started(reservation, 1, "fp"); runtime.complete_api(reservation, "DONE")
    telemetry = checkpoint.derive_telemetry("run"); text = report.build_report([], 0, runtime_snapshot={"durable_scheduler": telemetry})
    values = telemetry["provider_budgets"]["brightdata"]
    assert all(f"{key}={values[key]}" in text for key in ("effective_limit", "reserved_total", "done", "failed", "unknown", "physical_http_attempts"))


def test_free_telemetry_has_bucket_state_and_unique_block_breakdowns(tmp_path):
    init_run(tmp_path / "free-telemetry.sqlite3"); runtime.set_phase("FREE"); runtime.set_item_context(0)
    for bucket in ("discovery", "targeted"):
        runtime.set_search_bucket(bucket); reservation = runtime.reserve_free_physical_attempt(bucket, bucket, "bing"); runtime.complete_free_physical_attempt(reservation.attempt_id, bucket == "discovery")
        checkpoint.record_provider_budget_block(run_id="run", item_index=0, provider="bing", bucket=bucket, block_kind="logical", backend="bing")
        checkpoint.record_provider_budget_block(run_id="run", item_index=0, provider="bing", bucket=bucket, block_kind="physical", backend="bing")
    buckets = checkpoint.derive_telemetry("run")["free_queries"]["buckets"]
    assert buckets["discovery"]["done"] == 1 and buckets["targeted"]["failed"] == 1 and all(set(("logical_accepted","physical_attempted","done","failed","reserved","unique_logical_blocks","unique_physical_blocks")) <= set(value) for value in buckets.values())


def test_real_free_caps_are_six_four_and_ten(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"
    book = Workbook(); sheet = book.active
    sheet.append(["company", "source_record_id", "sector", "listed_address"])
    sheet.append(["NOVA MAKINA", "s0", "endustriyel makine", "Istanbul Turkiye"])
    book.save(source); book.close()
    transport = RecordingPaidTransport([])
    runs = _real_run_setup(tmp_path, monkeypatch, transport)

    class BoundedDDGS:
        discovery_calls = 0
        calls_by_bucket = {"discovery": 0, "targeted": 0}
        def __enter__(self): return self
        def __exit__(self, *_): return None
        def text(self, *_a, **_k):
            bucket = runtime.search_bucket()
            type(self).calls_by_bucket[bucket] += 1
            if bucket == "discovery":
                type(self).discovery_calls += 1
                if type(self).discovery_calls == 6:
                    return [{"href": "https://candidate.example", "title": "NOVA MAKINA", "body": "NOVA MAKINA endustriyel makine Istanbul Turkiye"}]
            return []

    def generic_fetch(url):
        response = requests.Response(); response.status_code = 200; response.url = url
        response.headers = {"content-type": "text/html; charset=utf-8"}
        response._content = b"<html><head><title>Welcome</title></head><body>General information</body></html>"
        response.encoding = "utf-8"
        return response

    monkeypatch.setattr(search, "DDGS", BoundedDDGS)
    monkeypatch.setattr(search, "PREFERRED_BACKENDS", ("bing",))
    monkeypatch.setattr(search, "FALLBACK_BACKENDS", ())
    monkeypatch.setattr(
        search.socket,
        "getaddrinfo",
        lambda host, *_a, **_k: [(2, 1, 6, "", ("203.0.113.10", 0))]
        if str(host).casefold() == "candidate.example"
        else (_ for _ in ()).throw(search.socket.gaierror("offline missing")),
    )
    monkeypatch.setattr(crawler, "_fetch", generic_fetch)
    monkeypatch.setattr(config, "MAX_TARGETED_QUERIES_PER_ROUND", 5)

    outcome = main.run(source, allow_paid=False)
    root = next(runs.iterdir())
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    monkeypatch.setattr(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3")
    free = checkpoint.derive_telemetry(manifest["run_id"])["free_queries"]
    assert outcome.status is pipeline_runner.PipelineOutcomeStatus.PAID_PENDING_APPROVAL
    assert BoundedDDGS.calls_by_bucket == {"discovery": 6, "targeted": 4}
    assert free["buckets"]["discovery"]["logical_accepted"] == 6
    assert free["buckets"]["targeted"]["logical_accepted"] == 4
    assert free["buckets"]["discovery"]["unique_logical_blocks"] == 1
    assert free["buckets"]["targeted"]["unique_logical_blocks"] == 1
    assert free["buckets"]["discovery"]["physical_attempted"] == 6 and free["buckets"]["targeted"]["physical_attempted"] == 4 and free["physical_attempted"] == 10
    _capture_scenario("free_caps", config.PROGRESS_DB_FILE, run_root=root, transport=transport, outcome=outcome)


def test_runtime_unknown_terminalization_is_canonical_for_every_provider(tmp_path):
    for provider in sorted(checkpoint.CANONICAL_PROVIDERS):
        init_run(tmp_path / f"{provider}.sqlite3", budget=1)
        with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
            db.execute("UPDATE provider_usage SET configured_limit=1,effective_limit=1 WHERE run_id='run' AND provider=?", (provider,)); db.commit()
        runtime.configure_durable_run("run", {name: (1 if name == provider else 0) for name in checkpoint.CANONICAL_PROVIDERS})
        runtime.set_item_context(0, "paid")
        checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
        reservation = runtime.reserve_api(provider, operation="unknown", request_fingerprint=f"{provider}:unknown")
        runtime.start_api(reservation); runtime.mark_api_http_started(reservation, 1); runtime.complete_api(reservation, "UNKNOWN", "post_send")
        stop = runtime.item_stop_state()
        assert stop.scope is runtime.StopScope.MANUAL_AUTHORIZATION and stop.call_ids == (reservation.call_id,)
        assert all(not runtime.reserve_api(later, operation="later", request_fingerprint=f"{provider}:{later}") for later in checkpoint.CANONICAL_PROVIDERS)


@pytest.mark.parametrize("provider", sorted(checkpoint.CANONICAL_PROVIDERS))
def test_unknown_each_paid_adapter_stops_all_later_paid_calls(tmp_path, monkeypatch, provider):
    from modules import company_resolvers, google_places, hunter, linkedin_company, llm_arbiter
    _init_provider_run(tmp_path / f"unknown-{provider}.sqlite3", provider)
    checkpoint.freeze_paid_query_plan(run_id="run", item_index=0, queries=["blocked by unknown stop"])
    recorder = RecordingPaidTransport([requests.ReadTimeout("post-send")]); runtime.set_paid_transport(recorder)
    monkeypatch.setattr(config, "SEARCH_CACHE_MODE", "off"); monkeypatch.setattr(config, "SEARCH_PROVIDER", "brightdata")
    monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake"); monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 0); monkeypatch.setattr(config, "GLOBAL_REQUESTS_PER_SECOND", 0); monkeypatch.setattr(config, "MAX_RETRIES", 0)
    if provider == "brightdata": result = search._brightdata_text("ACME")
    elif provider == "google_places":
        monkeypatch.setattr(config, "ENABLE_GOOGLE_PLACES", True); monkeypatch.setattr(config, "GOOGLE_PLACES_API_KEY", "fake"); result = google_places.search_company("ACME")
    elif provider == "brandfetch":
        monkeypatch.setattr(config, "ENABLE_BRANDFETCH_DOMAIN_SEARCH", True); monkeypatch.setattr(config, "BRANDFETCH_CLIENT_ID", "fake"); result = company_resolvers.brandfetch_domains("ACME")
    elif provider == "hunter":
        monkeypatch.setattr(config, "ENABLE_HUNTER_FALLBACK", True); monkeypatch.setattr(config, "HUNTER_API_KEY", "fake"); result = hunter.find_domain_emails("acme.example")
    elif provider == "linkedin":
        monkeypatch.setattr(config, "LINKEDIN_COMPANY_DATASET_ID", "fake"); result = linkedin_company._scrape("https://linkedin.com/company/acme")
    else:
        try: llm_arbiter.OpenRouterClient("fake", "model", 2).generate("prompt", {})
        except requests.ReadTimeout: pass
        result = runtime.provider_result([], state="UNKNOWN", provider="llm", call_ids=runtime.item_stop_state().call_ids)
    assert result.result_state == "UNKNOWN" and len(recorder.journal) == 1
    main.process_company(0, "ACME", logging.getLogger("unknown-adapter"), "", {"source_record_id": "s0"}, execution_phase="PAID")
    stop = runtime.item_stop_state(); assert stop.scope is runtime.StopScope.MANUAL_AUTHORIZATION and len(recorder.journal) == 1
    assert all(not runtime.reserve_api(later, operation="later", request_fingerprint=f"{provider}:{later}") for later in checkpoint.CANONICAL_PROVIDERS)
    call_id = stop.call_ids[0]
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="UNKNOWN", call_ids=[call_id])
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        db.execute("UPDATE run_items SET paid_state='UNKNOWN'"); db.commit()
        assert db.execute("SELECT state FROM provider_calls").fetchall() == [("UNKNOWN",)]
        assert db.execute("SELECT result FROM paid_attempts").fetchone()[0] == "UNKNOWN"


@pytest.mark.parametrize("error_type", [
    checkpoint.StateTransitionInvariant, checkpoint.LedgerInvariant,
    checkpoint.EvidenceInvariant, checkpoint.ResumeInvariant,
    checkpoint.OutcomeInvariant,
])
def test_llm_scheduler_invariant_is_never_converted_to_failed_dict(durable, error_type):
    from modules import llm_arbiter
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        db.execute("UPDATE provider_usage SET configured_limit=1,effective_limit=1 WHERE run_id='run' AND provider='llm'")
    runtime.configure_durable_run("run", {"brightdata": 20, "llm": 1})
    class Client:
        def generate(self, *_args):
            raise error_type("typed-invariant")
    with pytest.raises(error_type, match="typed-invariant") as raised:
        llm_arbiter.arbitrate("ACME", "", "sector", "acme.example", "a sufficiently descriptive page summary", client=Client())
    assert type(raised.value) is error_type


@pytest.mark.parametrize("mutation", ("missing", "extra"))
def test_provider_set_missing_or_extra_row_blocks_telemetry_and_finalization(tmp_path, mutation):
    init_run(tmp_path / f"provider-{mutation}.sqlite3")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        if mutation == "missing":
            db.execute("DELETE FROM provider_usage WHERE run_id='run' AND provider='llm'")
        else:
            db.execute("INSERT INTO provider_usage(run_id,provider,configured_limit,effective_limit) VALUES('run','extra',0,0)")
        db.commit()
    with pytest.raises(checkpoint.LedgerInvariant): checkpoint.derive_telemetry("run")


@pytest.mark.parametrize("aggregate", ["provider", "free"])
def test_current_schema_process_restart_does_not_heal_provider_aggregate_corruption(tmp_path, aggregate):
    path = tmp_path / "restart.sqlite3"; init_run(path)
    with sqlite3.connect(path) as db:
        if aggregate == "provider": db.execute("UPDATE provider_usage SET reserved_total=7 WHERE run_id='run' AND provider='brightdata'")
        else: db.execute("INSERT INTO free_query_usage(run_id,item_index,used,quota,physical_used,discovery_physical_used) VALUES('run',0,0,10,2,2)")
        db.commit()
    before = path.read_bytes()
    code = "from pathlib import Path; import config; from modules import checkpoint; config.PROGRESS_DB_FILE=Path(__import__('sys').argv[1]); checkpoint._SCHEMA_READY.clear(); checkpoint.initialize_schema(config.PROGRESS_DB_FILE);\ntry: checkpoint.validate_ledger_equations('run')\nexcept checkpoint.LedgerInvariant: raise SystemExit(17)\nraise SystemExit(0)"
    completed = subprocess.run([sys.executable, "-c", code, str(path)], cwd=Path(__file__).parents[1])
    assert completed.returncode == 17 and path.read_bytes() == before
    with sqlite3.connect(path) as db:
        if aggregate == "provider": assert db.execute("SELECT reserved_total FROM provider_usage WHERE run_id='run' AND provider='brightdata'").fetchone()[0] == 7
        else: assert db.execute("SELECT physical_used FROM free_query_usage WHERE run_id='run'").fetchone()[0] == 2
    with pytest.raises(checkpoint.LedgerInvariant): checkpoint.validate_ledger_equations("run")


@pytest.mark.parametrize("version", (checkpoint.SCHEDULER_SCHEMA_VERSION, 10), ids=("current", "legacy"))
def test_complete_checkpoint_schema_open_is_byte_for_byte_read_only(tmp_path, version):
    def sidecars(database):
        values = {}
        for suffix in ("", "-wal", "-shm"):
            candidate = Path(str(database) + suffix)
            if candidate.exists():
                values[suffix] = (candidate.read_bytes(), candidate.stat().st_mtime_ns)
        return values
    path = tmp_path / f"complete-v{version}.sqlite3"; init_run(path)
    with sqlite3.connect(path) as db:
        db.execute("UPDATE runs SET phase='COMPLETE'")
        db.execute(f"PRAGMA user_version={version}"); db.commit()
    before = sidecars(path)
    checkpoint._SCHEMA_READY.clear()
    if version == checkpoint.SCHEDULER_SCHEMA_VERSION:
        checkpoint.initialize_schema(path)
    else:
        with pytest.raises(checkpoint.StateTransitionInvariant, match="read-only"):
            checkpoint.initialize_schema(path)
    after = sidecars(path)
    assert after == before


@pytest.mark.parametrize("mutation", (
    "orphan_attempt_call", "orphan_flight_consumer", "orphan_budget_block",
    "update_attempt_provider", "update_consumer_generation", "update_block_provider",
))
def test_relation_tables_reject_orphan_and_cross_scope_rows_with_foreign_keys(durable, mutation):
    orphan_statements = {
        "orphan_attempt_call": "INSERT INTO paid_attempt_calls(run_id,item_index,attempt_number,phase,call_id,paid_attempt_id,provider_call_id,provider,query_fingerprint,execution_generation,relation) VALUES('run',0,1,'PAID','x','missing','missing','brightdata','',1,'OWNER')",
        "orphan_flight_consumer": "INSERT INTO provider_query_flight_consumers(run_id,provider,query_fingerprint,execution_generation,paid_attempt_id,item_index,provider_call_id,relation,linked_at) VALUES('run','brightdata','missing',1,'missing',0,'missing','OWNER','now')",
        "orphan_budget_block": "INSERT INTO paid_attempt_block_links(run_id,item_index,provider,paid_attempt_id,block_id) VALUES('run',0,'brightdata','missing','missing')",
    }
    if mutation in orphan_statements:
        with sqlite3.connect(durable) as db:
            db.execute("PRAGMA foreign_keys=ON")
            with pytest.raises(sqlite3.IntegrityError): db.execute(orphan_statements[mutation])
            assert db.execute("PRAGMA foreign_key_check").fetchall() == []
        return
    runtime.set_item_context(0, "paid")
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="scope-fp", owner_token="scope")
    call = runtime.reserve_api("brightdata", operation="scope", request_fingerprint="scope-call")
    runtime.start_api(call); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="scope-fp", owner_token="scope", provider_call_id=call.call_id)
    runtime.mark_api_http_started(call, 1, "scope-fp"); runtime.complete_api(call, "FAILED")
    checkpoint.finish_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="scope-fp", owner_token="scope", state="FAILED", result={}, call_ids=[call.call_id])
    checkpoint.record_provider_budget_block(run_id="run", item_index=0, provider="brightdata", bucket="paid", block_kind="budget_exhausted")
    with sqlite3.connect(durable) as db:
        db.execute("PRAGMA foreign_keys=ON")
        attempt_id = db.execute("SELECT paid_attempt_id FROM paid_attempts").fetchone()[0]
        db.execute("INSERT INTO provider_query_flight_consumers(run_id,provider,query_fingerprint,execution_generation,paid_attempt_id,item_index,provider_call_id,relation,linked_at) VALUES('run','brightdata','scope-fp',1,?,0,?,'OWNER','now')", (attempt_id, call.call_id))
        block_id = db.execute("SELECT block_id FROM provider_budget_blocks WHERE run_id='run' AND item_index=0 AND provider='brightdata'").fetchone()[0]
        db.execute("INSERT OR IGNORE INTO paid_attempt_block_links(run_id,item_index,provider,paid_attempt_id,block_id) VALUES('run',0,'brightdata',?,?)", (attempt_id, block_id))
        updates = {
            "update_attempt_provider": ("UPDATE paid_attempt_calls SET provider='hunter' WHERE provider_call_id=?", (call.call_id,), "paid_attempt_calls_owner_scope"),
            "update_consumer_generation": ("UPDATE provider_query_flight_consumers SET execution_generation=2 WHERE provider_call_id=?", (call.call_id,), "flight_consumer_(?:scope|identity)"),
            "update_block_provider": ("UPDATE paid_attempt_block_links SET provider='hunter' WHERE paid_attempt_id=?", (attempt_id,), "paid_attempt_block_scope"),
        }
        statement, parameters, message = updates[mutation]
        with pytest.raises(sqlite3.IntegrityError, match=message):
            db.execute(statement, parameters)
        assert db.execute("PRAGMA foreign_key_check").fetchall() == []


def _mixed_paid_evidence(path: Path, final_result: str, paid_state: str):
    init_run(path, budget=2); runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    calls = []
    for state in ("DONE", "UNKNOWN"):
        reservation = runtime.reserve_api("brightdata", operation=state, request_fingerprint=state)
        runtime.start_api(reservation); runtime.mark_api_http_started(reservation, 1); runtime.complete_api(reservation, state)
        calls.append(reservation.call_id)
        if state == "UNKNOWN": runtime.reset_item_stop_state(0)
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result=final_result, call_ids=calls)
    with sqlite3.connect(path) as db:
        db.execute("UPDATE run_items SET paid_state=? WHERE run_id='run' AND item_index=0", (paid_state,)); db.commit()


def test_paid_evidence_unknown_has_precedence_over_done(tmp_path):
    _mixed_paid_evidence(tmp_path / "precedence.sqlite3", "COMPLETED", "DONE")
    with pytest.raises(checkpoint.EvidenceInvariant): checkpoint.validate_paid_evidence("run")


def test_paid_evidence_done_rejects_blocked_budget_outcome(tmp_path):
    init_run(tmp_path / "done-block.sqlite3"); runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    reservation = runtime.reserve_api("brightdata", operation="done", request_fingerprint="done"); runtime.start_api(reservation); runtime.complete_api(reservation, "DONE")
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="BLOCKED_BUDGET", call_ids=[reservation.call_id])
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db: db.execute("UPDATE run_items SET paid_state='BLOCKED_BUDGET'"); db.commit()
    with pytest.raises(checkpoint.EvidenceInvariant): checkpoint.validate_paid_evidence("run")


def test_paid_no_call_evidence_requires_exact_zero_call_or_inherited_receipt(tmp_path):
    path = tmp_path / "no-call.sqlite3"; init_run(path); runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    reservation = runtime.reserve_api("brightdata", operation="physical", request_fingerprint="physical")
    runtime.start_api(reservation); runtime.mark_api_http_started(reservation, 1); runtime.complete_api(reservation, "DONE")
    snapshot = checkpoint.immutable_input_snapshot_sha256("run", 0)
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="NO_CALL_NEEDED", evidence_kind="supplied_website_publishable_at_paid_entry", input_snapshot_sha256=snapshot)
    with sqlite3.connect(path) as db: db.execute("UPDATE run_items SET paid_state='DONE'"); db.commit()
    with pytest.raises(checkpoint.EvidenceInvariant, match="zero paid physical"):
        checkpoint.validate_paid_evidence("run")


def test_expired_flight_failed_then_done_receipt_reconciles_done(tmp_path):
    init_run(tmp_path / "mixed-flight.sqlite3", budget=2); runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner")
    failed = runtime.reserve_api("brightdata", operation="first", request_fingerprint="first"); runtime.start_api(failed); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=failed.call_id); runtime.complete_api(failed, "FAILED")
    done = runtime.reserve_api("brightdata", operation="second", request_fingerprint="second"); runtime.start_api(done); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=done.call_id); runtime.mark_api_http_started(done, 2, "fp"); checkpoint.bind_provider_call_transport_receipt(call_id=done.call_id, endpoint_sha256="a" * 64, request_shape_sha256="b" * 64)
    checkpoint.complete_provider_call_and_flight_success(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=done.call_id, result={"values": []}, call_ids=[failed.call_id, done.call_id])
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db: db.execute("UPDATE provider_query_flights SET state='RUNNING',lease_expires_at='2000-01-01T00:00:00+00:00'"); db.commit()
    reconciled = checkpoint.resolve_expired_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="new")
    assert reconciled["state"] == "DONE" and reconciled["call_ids"] == [failed.call_id, done.call_id]


def test_reclaimed_flight_preserves_all_prior_call_ids_append_only(tmp_path):
    init_run(tmp_path / "append.sqlite3", budget=3); runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    first = checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="same", owner_token="old", lease_seconds=.1)
    old = runtime.reserve_api("brightdata", operation="old", request_fingerprint="old"); runtime.start_api(old)
    checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="same", owner_token="old", provider_call_id=old.call_id); runtime.complete_api(old, "FAILED")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db: db.execute("UPDATE provider_query_flights SET lease_expires_at='2000-01-01T00:00:00+00:00'"); db.commit()
    reclaimed = checkpoint.resolve_expired_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="same", owner_token="new")
    assert reclaimed["leader"] and reclaimed["execution_generation"] == first["execution_generation"] and reclaimed["call_ids"] == [old.call_id]
    new = runtime.reserve_api("brightdata", operation="new", request_fingerprint="new"); runtime.start_api(new)
    checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="same", owner_token="new", provider_call_id=new.call_id); runtime.complete_api(new, "FAILED")
    with pytest.raises(checkpoint.EvidenceInvariant):
        checkpoint.finish_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="same", owner_token="new", state="FAILED", result={}, call_ids=[new.call_id])
    checkpoint.finish_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="same", owner_token="new", state="FAILED", result={}, call_ids=[old.call_id, new.call_id])
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        assert json.loads(db.execute("SELECT call_ids_json FROM provider_query_flights").fetchone()[0]) == [old.call_id, new.call_id]


def test_provider_query_execution_generation_is_monotonic_and_relational(tmp_path):
    path = tmp_path / "generation.sqlite3"; init_run(path, budget=2); runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    first = checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one")
    call = runtime.reserve_api("brightdata", operation="first-generation", request_fingerprint="same-request", flight_fingerprint="fp", execution_generation=1)
    with sqlite3.connect(path) as db:
        assert db.execute("SELECT flight_fingerprint FROM provider_calls WHERE call_id=?", (call.call_id,)).fetchone() == ("fp",)
        assert db.execute("SELECT query_fingerprint,execution_generation FROM paid_attempt_calls WHERE provider_call_id=?", (call.call_id,)).fetchone() == ("fp", 1)
    runtime.start_api(call); runtime.mark_api_http_started(call, 1, "fp"); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one", provider_call_id=call.call_id); runtime.complete_api(call, "FAILED")
    checkpoint.finish_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one", state="FAILED", result={}, call_ids=[call.call_id])
    second = checkpoint.start_new_provider_query_execution(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="two")
    assert first["execution_generation"] == 1 and second["execution_generation"] == 2 and second["call_ids"] == []
    call_two = runtime.reserve_api("brightdata", operation="second-generation", request_fingerprint="same-request", flight_fingerprint="fp", execution_generation=2)
    assert call_two.accepted
    runtime.start_api(call_two); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="two", provider_call_id=call_two.call_id); runtime.mark_api_http_started(call_two, 1, "fp"); runtime.complete_api(call_two, "FAILED")
    checkpoint.finish_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="two", state="FAILED", result={}, call_ids=[call_two.call_id])
    with sqlite3.connect(path) as db:
        rows = db.execute("SELECT execution_generation,state,call_ids_json FROM provider_query_flight_terminals ORDER BY execution_generation").fetchall()
        assert [(row[0], row[1], json.loads(row[2])) for row in rows] == [(1, "FAILED", [call.call_id]), (2, "FAILED", [call_two.call_id])]
        assert json.loads(db.execute("SELECT call_ids_json FROM provider_query_flights").fetchone()[0]) == [call_two.call_id]
        with pytest.raises(sqlite3.IntegrityError):
            db.execute("UPDATE paid_attempt_calls SET execution_generation=2 WHERE provider_call_id=?", (call.call_id,))


def test_new_provider_query_generation_requires_matching_terminal_receipt(tmp_path):
    path = tmp_path / "generation-terminal.sqlite3"; init_run(path, budget=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one")
    with sqlite3.connect(path) as db:
        db.execute("UPDATE provider_query_flights SET state='FAILED'"); db.commit()
    with pytest.raises((checkpoint.StateTransitionInvariant, checkpoint.EvidenceInvariant), match="reconciled terminal|terminal_flight_receipt"):
        checkpoint.start_new_provider_query_execution(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="two")
    with sqlite3.connect(path) as db:
        db.execute("INSERT INTO provider_query_flight_terminals(run_id,provider,query_fingerprint,execution_generation,state,provider_call_id,result_sha256,call_ids_json,created_at) VALUES('run','brightdata','fp',1,'FAILED','','forged','[]','now')"); db.commit()
    with pytest.raises((checkpoint.StateTransitionInvariant, checkpoint.EvidenceInvariant), match="reconciled terminal|flight_terminal_hash"):
        checkpoint.start_new_provider_query_execution(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="two")


def test_new_provider_query_generation_accepts_reconciled_done_receipt(tmp_path):
    path = tmp_path / "generation-done.sqlite3"; init_run(path, budget=2); runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one")
    call = runtime.reserve_api("brightdata", operation="done-generation", request_fingerprint="one")
    runtime.start_api(call); runtime.mark_api_http_started(call, 1, "fp"); checkpoint.bind_provider_call_transport_receipt(call_id=call.call_id, endpoint_sha256="a" * 64, request_shape_sha256="b" * 64); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one", provider_call_id=call.call_id)
    checkpoint.complete_provider_call_and_flight_success(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one", provider_call_id=call.call_id, result={"values": []}, call_ids=[call.call_id])
    second = checkpoint.start_new_provider_query_execution(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="two")
    assert second["execution_generation"] == 2 and second["call_ids"] == []


def test_non_atomic_flight_finish_rejects_done(tmp_path):
    path = tmp_path / "non-atomic-done.sqlite3"; init_run(path, budget=1); runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner")
    call = runtime.reserve_api("brightdata", operation="legacy-done", request_fingerprint="request")
    runtime.start_api(call); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=call.call_id)
    runtime.mark_api_http_started(call, 1, "fp"); checkpoint.bind_provider_call_transport_receipt(call_id=call.call_id, endpoint_sha256="a" * 64, request_shape_sha256="b" * 64)
    with pytest.raises(checkpoint.StateTransitionInvariant, match="conflicts with provider call states"):
        checkpoint.finish_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", state="FAILED", result={}, call_ids=[call.call_id])
    runtime.complete_api(call, "DONE")
    with pytest.raises(checkpoint.StateTransitionInvariant, match="atomic provider-call success"):
        checkpoint.finish_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", state="DONE", result={"values": []}, call_ids=[call.call_id])
    with sqlite3.connect(path) as db:
        assert db.execute("SELECT state FROM provider_query_flights").fetchone() == ("RUNNING",)
        assert db.execute("SELECT COUNT(*) FROM provider_query_flight_results").fetchone() == (0,)
        assert db.execute("SELECT COUNT(*) FROM provider_query_flight_terminals").fetchone() == (0,)


@pytest.mark.parametrize("receipt_table", ["provider_query_flight_results", "provider_query_flight_terminals"])
def test_generation_receipts_reject_cross_generation_mutation(tmp_path, receipt_table):
    path = tmp_path / f"generation-{receipt_table}.sqlite3"; init_run(path, budget=2); runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one")
    call = runtime.reserve_api("brightdata", operation="generation-bound", request_fingerprint="request")
    runtime.start_api(call); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one", provider_call_id=call.call_id); runtime.mark_api_http_started(call, 1, "fp"); checkpoint.bind_provider_call_transport_receipt(call_id=call.call_id, endpoint_sha256="a" * 64, request_shape_sha256="b" * 64)
    checkpoint.complete_provider_call_and_flight_success(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="one", provider_call_id=call.call_id, result={"values": []}, call_ids=[call.call_id])
    checkpoint.start_new_provider_query_execution(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="two")
    trigger = "flight_result_identity_update" if receipt_table.endswith("results") else "flight_terminal_identity_update"
    with sqlite3.connect(path) as db:
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            db.execute(f"UPDATE {receipt_table} SET execution_generation=2 WHERE execution_generation=1")
        db.rollback(); db.execute(f"DROP TRIGGER {trigger}"); db.execute(f"DROP TRIGGER {'flight_result_scope_update' if receipt_table.endswith('results') else 'flight_terminal_scope_update'}"); db.execute(f"DROP TRIGGER {'flight_result_receipt_update' if receipt_table.endswith('results') else 'flight_terminal_receipt_update'}"); db.execute(f"UPDATE {receipt_table} SET execution_generation=2 WHERE execution_generation=1"); db.commit()
    checkpoint._SCHEMA_READY.clear()
    with pytest.raises(checkpoint.EvidenceInvariant, match="generation|receipt_triggers"):
        checkpoint.derive_telemetry("run")


def test_atomic_flight_success_rejects_unbound_provider_call(tmp_path):
    path = tmp_path / "unbound-flight.sqlite3"; init_run(path, budget=1); runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner")
    call = runtime.reserve_api("brightdata", operation="unbound", request_fingerprint="request")
    runtime.start_api(call); runtime.mark_api_http_started(call, 1, "fp"); checkpoint.bind_provider_call_transport_receipt(call_id=call.call_id, endpoint_sha256="a" * 64, request_shape_sha256="b" * 64)
    with pytest.raises(checkpoint.StateTransitionInvariant, match="bound current-generation"):
        checkpoint.complete_provider_call_and_flight_success(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=call.call_id, result={"values": []}, call_ids=[call.call_id])


def test_volatile_same_query_independent_failures_open_circuit(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SEARCH_PROVIDER", "brightdata"); monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake"); monkeypatch.setattr(config, "BRIGHTDATA_REQUEST_BUDGET", 20)
    monkeypatch.setattr(config, "BRIGHTDATA_CIRCUIT_FAILURE_THRESHOLD", 2); monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 0)
    monkeypatch.setattr(config, "GLOBAL_REQUESTS_PER_SECOND", 0); monkeypatch.setattr(config, "MAX_RETRIES", 0); monkeypatch.setattr(config, "SEARCH_CACHE_MODE", "off")
    recorder = RecordingPaidTransport([Response(500, {"error": "down"}), Response(500, {"error": "down"})])
    runtime.set_paid_transport(recorder)
    monkeypatch.setattr(search, "_ddgs_text", lambda _query: search.SearchResults([], "offline", "ddgs", result_state="EMPTY"))
    assert search._brightdata_text("same query").result_state == "FAILED"
    assert search._brightdata_text("same query").result_state == "FAILED"
    blocked = search._safe_search_text("same query")
    assert len(recorder.journal) == 2 and blocked.origin is search.SearchOrigin.CIRCUIT_BLOCK
    search.reset_run_state(); recorder = RecordingPaidTransport([Response(200, {"organic": []}), Response(500, {"error": "down"})])
    runtime.set_paid_transport(recorder)
    assert search._brightdata_text("same query").result_state == "EMPTY"
    assert search._brightdata_text("same query").result_state == "FAILED"
    assert len(recorder.journal) == 2 and not search._brightdata_circuit_open()


def test_real_fresh_handoff_resume_preserves_durable_ordered_query_plan(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"; _input_book(source, ["PAID CO"])
    transport = RecordingPaidTransport([Response(200, {"organic": [{"link": "https://paidco.example", "title": "PAID CO", "description": "PAID CO resmi şirket sitesi Türkiye"}]})])
    runs = _real_run_setup(tmp_path, monkeypatch, transport)
    pending = main.run(source, allow_paid=False); parent = next(runs.iterdir()); parent_manifest = json.loads((parent / "manifest.json").read_text(encoding="utf-8"))
    assert pending.status is pipeline_runner.PipelineOutcomeStatus.PAID_PENDING_APPROVAL
    with sqlite3.connect(parent / "state" / "progress.sqlite3") as db:
        stored = db.execute("SELECT item_index,plan_version,query_kind,round_ordinal,query_ordinal,normalized_query,query_sha256 FROM paid_query_plan_entries ORDER BY item_index,plan_version,query_kind,round_ordinal,query_ordinal").fetchall()
    material = json.dumps([list(row) for row in stored], ensure_ascii=False, separators=(",", ":"))
    assert stored and parent_manifest["paid_query_plan_sha256"] == hashlib.sha256(material.encode()).hexdigest()
    source_id = parent_manifest["ordered_source_record_ids"][0]
    approval = {"approved": True, "limits": {provider: (1 if provider == "brightdata" else 0) for provider in checkpoint.CANONICAL_PROVIDERS}, "parent_run_id": parent_manifest["run_id"], "parent_manifest_sha256": hashlib.sha256((parent / "manifest.json").read_bytes()).hexdigest(), "checkpoint_sha256": parent_manifest.get("_continuation_checkpoint_sha256") or parent_manifest["files"]["recovery_state.sqlite3"]["sha256"], "paid_source_ids_sha256": hashlib.sha256(run_context.canonical_json([source_id]).encode()).hexdigest()}
    authorization = tmp_path / "approval.json"; authorization.write_text(json.dumps(approval), encoding="utf-8")
    child_manifest = prepare_paid_continuation(parent, authorization, tmp_path / "continuation"); child = tmp_path / "continuation" / "runs" / child_manifest["run_id"]
    monkeypatch.setattr(search, "_primary_queries", lambda *_a, **_k: (_ for _ in ()).throw(AssertionError("resume recomputed primary plan")))
    completed = main.run(source, allow_paid=True, resume_run=child)
    with sqlite3.connect(child / "state" / "progress.sqlite3") as db:
        resumed = db.execute("SELECT item_index,plan_version,query_kind,round_ordinal,query_ordinal,normalized_query,query_sha256 FROM paid_query_plan_entries ORDER BY item_index,plan_version,query_kind,round_ordinal,query_ordinal").fetchall()
    assert completed.status is pipeline_runner.PipelineOutcomeStatus.COMPLETE and resumed[:len(stored)] == stored and len(transport.journal) == 1


def test_paid_targeted_query_plan_is_append_only_across_real_rounds_and_resume(tmp_path, monkeypatch):
    path = tmp_path / "targeted-plan.sqlite3"; init_run(path, budget=6); runtime.set_item_context(0, "paid")
    monkeypatch.setattr(config, "MAX_SEARCH_QUERIES_PER_COMPANY", 6)
    monkeypatch.setattr(search, "_safe_search_text", lambda _query: search.SearchResults([], "live", "brightdata", result_state="EMPTY"))
    expected = []
    for ordinal, queries in ((1, ["acme contact"]), (2, ["acme legal title"]), (3, ["acme sector proof"])):
        search.find_targeted_candidates("ACME", {"source_record_id": "s0"}, queries, round_ordinal=ordinal, limit=1)
        expected.append((ordinal, queries[0]))
    search.find_targeted_candidates("ACME", {"source_record_id": "s0"}, ["acme legal title"], round_ordinal=2, limit=1)
    with pytest.raises(checkpoint.ResumeInvariant):
        search.find_targeted_candidates("ACME", {"source_record_id": "s0"}, ["different meaning"], round_ordinal=2, limit=1)
    with sqlite3.connect(path) as db:
        rows = db.execute("SELECT round_ordinal,normalized_query FROM paid_query_plan_entries WHERE query_kind='targeted' ORDER BY round_ordinal,query_ordinal").fetchall()
    assert rows == expected and checkpoint.paid_query_plan_receipt("run")["paid_query_plan_count"] == 3


@pytest.mark.parametrize("mutation", ["missing", "extra", "bool", "negative", "float"])
def test_run_config_requires_exact_canonical_provider_budget_set(tmp_path, mutation, monkeypatch):
    monkeypatch.setattr(main, "_ensure_safe_project_runtime", lambda *_a, **_k: None)
    payload = run_context.RunConfig.from_config(paid_enabled=True).as_dict()
    if mutation == "missing": payload["budgets"].pop("llm")
    elif mutation == "extra": payload["budgets"]["extra"] = 0
    elif mutation == "bool": payload["budgets"]["llm"] = True
    elif mutation == "negative": payload["budgets"]["llm"] = -1
    else: payload["budgets"]["llm"] = 1.5
    with pytest.raises(checkpoint.ResumeInvariant, match="exact canonical provider budgets"):
        run_context.RunConfig.from_dict(payload)
    source = tmp_path / "input.xlsx"; _input_book(source, ["ACME"])
    root = tmp_path / "invalid-resume"; root.mkdir(); (root / "manifest.json").write_text(json.dumps({"run_config": payload}), encoding="utf-8")
    assert main.cli(["--input", str(source), "--resume-run", str(root), "--allow-paid", "--non-interactive"]) == 22


@pytest.mark.parametrize("mutation", ["missing", "extra", "field"])
def test_report_rejects_missing_provider_set(mutation):
    providers = {provider: {"population_count": 0, "ratio": None, "explicit_cap": 0, "configured_limit": 0, "effective_limit": 0, "reserved_total": 0, "done": 0, "failed": 0, "unknown": 0, "reserved": 0, "running": 0, "physical_http_attempts": 0, "retry_attempts": 0, "inherited_uses": 0, "budget_blocked_items": 0} for provider in checkpoint.CANONICAL_PROVIDERS}
    if mutation == "missing": providers.pop("llm")
    elif mutation == "extra": providers["extra"] = dict(next(iter(providers.values())))
    else: providers["llm"].pop("reserved_total")
    with pytest.raises(checkpoint.EvidenceInvariant): report.build_report([], 0, runtime_snapshot={"durable_scheduler": {"provider_budgets": providers}})


@pytest.mark.parametrize("mutation", ["journal", "manifest", "artifact", "artifact_hash"])
def test_evidence_capture_rejects_missing_journal_manifest_or_artifact(tmp_path, monkeypatch, mutation, production_trace_call):
    monkeypatch.setenv("B2B_FINAL_SCENARIO_DIR", str(tmp_path / "evidence"))
    db = tmp_path / "capture.sqlite3"; init_run(db)
    run_root = tmp_path / "run"; run_root.mkdir(); transport = RecordingPaidTransport([])
    if mutation in {"artifact", "artifact_hash"}:
        payload = b"valid" if mutation == "artifact_hash" else b""
        digest = hashlib.sha256(b"different" if mutation == "artifact_hash" else payload).hexdigest()
        (run_root / "manifest.json").write_text(json.dumps({"complete": True, "artifact_set_sha256": "set", "files": {"telemetry.json": {"sha256": digest, "bytes": len(payload)}}}), encoding="utf-8")
        if mutation == "artifact_hash":
            artifact = run_root / "output" / "artifacts" / "set" / "telemetry.json"; artifact.parent.mkdir(parents=True); artifact.write_bytes(payload)
    expected = "recording transport" if mutation == "journal" else "production manifest" if mutation == "manifest" else "missing or corrupt" if mutation == "artifact_hash" else "missing"
    with pytest.raises(AssertionError, match=expected):
        production_trace_call(_capture_scenario, "unknown_item_stop", db, run_root=run_root, transport=None if mutation == "journal" else transport)


def test_recording_transport_uses_production_attempt_and_fingerprints(durable, monkeypatch):
    runtime.set_item_context(0, "paid"); monkeypatch.setattr(config, "MAX_RETRIES", 1); monkeypatch.setattr(search, "_retry_delay", lambda *_: 0)
    recorder = RecordingPaidTransport([Response(200, {"organic": []}), Response(500, {"error": "retry"}), Response(200, {"organic": []})]); runtime.set_paid_transport(recorder)
    assert search._brightdata_text("first independent query").result_state == "EMPTY"
    assert search._brightdata_text("second independent query").result_state == "EMPTY"
    with sqlite3.connect(durable) as db:
        calls = db.execute("SELECT run_id,call_id,provider,item_index,request_fingerprint,attempt_ordinal,flight_fingerprint,state,endpoint_sha256,request_shape_sha256 FROM provider_calls ORDER BY rowid").fetchall()
        attempt_id = db.execute("SELECT paid_attempt_id FROM paid_attempts").fetchone()[0]
    assert len(recorder.journal) == len(calls) == 3 and [row[5] for row in calls] == [1, 1, 2]
    by_id = {row[1]: row for row in calls}
    for envelope in recorder.journal:
        call = by_id[envelope["provider_call_id"]]
        assert (envelope["run_id"], envelope["provider"], envelope["item_index"], envelope["request_fingerprint"], envelope["attempt_ordinal"], envelope["flight_fingerprint"], envelope["paid_attempt_id"], envelope["execution_generation"]) == (call[0], call[2], call[3], call[4], call[5], call[6], attempt_id, 1)
        assert hashlib.sha256(envelope["endpoint"].encode()).hexdigest() == call[8]
        assert envelope["request_shape_sha256"] == call[9] and len(call[9]) == 64 and call[7] in {"DONE", "FAILED"}


def _init_provider_run(path: Path, provider: str):
    init_run(path, budget=1)
    with sqlite3.connect(path) as db:
        db.execute("UPDATE provider_usage SET configured_limit=0,effective_limit=0 WHERE run_id='run'")
        db.execute("UPDATE provider_usage SET configured_limit=1,effective_limit=1 WHERE run_id='run' AND provider=?", (provider,)); db.commit()
    runtime.reset(); runtime.configure_durable_run("run", {name: (1 if name == provider else 0) for name in checkpoint.CANONICAL_PROVIDERS}); runtime.set_phase("PAID"); runtime.set_item_context(0, "paid"); checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)


@pytest.mark.parametrize(("surface", "payload", "expected"), [
    ("brightdata", {"organic": []}, "EMPTY"), ("brightdata", {"organic": [{"link": 7}]}, "FAILED"),
    ("google_places", {"places": []}, "EMPTY"), ("google_places", {"places": [{"websiteUri": {}}]}, "FAILED"),
    ("hunter_email", {"data": {"emails": []}}, "EMPTY"), ("hunter_email", {"data": {"emails": [{"value": 7, "confidence": 90}]}}, "FAILED"),
    ("brandfetch", [], "EMPTY"), ("brandfetch", [{"domain": {}, "name": "ACME"}], "FAILED"),
    ("hunter_finder", {"data": []}, "EMPTY"), ("hunter_finder", {"data": [{"domain": "acme.example", "claimed": "yes"}]}, "FAILED"),
    ("linkedin_serp", {"organic": []}, "EMPTY"), ("linkedin_serp", {"malformed": True}, "FAILED"),
    ("linkedin_scrape", [], "EMPTY"), ("linkedin_scrape", [{"name": 7}] , "FAILED"),
    ("openrouter", {"choices": [{"message": {"content": '{"verdict":"match","reason":"a sufficiently detailed concrete reason","detected_sector":"metal","expected_sector":"metal"}'}}], "usage": {}}, "COMPLETED"),
    ("openrouter", {"choices": [{"message": {"content": '{"verdict":"match"}'}}]}, "FAILED"),
], ids=lambda value: str(value)[:40])
def test_paid_adapters_do_not_mark_done_before_semantic_validation(tmp_path, monkeypatch, surface, payload, expected):
    from modules import company_resolvers, google_places, hunter, linkedin_company
    provider = "hunter" if surface.startswith("hunter") else "linkedin" if surface.startswith("linkedin") else "llm" if surface == "openrouter" else surface
    _init_provider_run(tmp_path / f"semantic-{surface}.sqlite3", provider)
    recorder = RecordingPaidTransport([Response(200, payload)])
    monkeypatch.setattr(runtime, "_PAID_TRANSPORT", recorder)
    monkeypatch.setattr(config, "SEARCH_CACHE_MODE", "off")
    if surface == "brightdata":
        monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake"); monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 0); monkeypatch.setattr(config, "GLOBAL_REQUESTS_PER_SECOND", 0); monkeypatch.setattr(config, "MAX_RETRIES", 0)
        result = search._brightdata_text("ACME")
    elif surface == "google_places":
        monkeypatch.setattr(config, "ENABLE_GOOGLE_PLACES", True); monkeypatch.setattr(config, "GOOGLE_PLACES_API_KEY", "fake")
        result = google_places.search_company("ACME")
    elif surface == "hunter_email":
        monkeypatch.setattr(config, "ENABLE_HUNTER_FALLBACK", True); monkeypatch.setattr(config, "HUNTER_API_KEY", "fake")
        result = hunter.find_domain_emails("acme.example")
    elif surface == "brandfetch":
        monkeypatch.setattr(config, "ENABLE_BRANDFETCH_DOMAIN_SEARCH", True); monkeypatch.setattr(config, "BRANDFETCH_CLIENT_ID", "fake"); monkeypatch.setattr(company_resolvers, "_cached", lambda *_: None)
        result = company_resolvers.brandfetch_domains("ACME")
    elif surface == "hunter_finder":
        monkeypatch.setattr(config, "ENABLE_HUNTER_DOMAIN_FINDER", True); monkeypatch.setattr(config, "HUNTER_API_KEY", "fake"); monkeypatch.setattr(company_resolvers, "_cached", lambda *_: None)
        result = company_resolvers.hunter_domains("ACME")
    elif surface == "linkedin_serp":
        monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake")
        result = linkedin_company._find_company_url("ACME")
    elif surface == "linkedin_scrape":
        monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake"); monkeypatch.setattr(config, "LINKEDIN_COMPANY_DATASET_ID", "fake")
        result = linkedin_company._scrape("https://linkedin.com/company/acme")
    else:
        from modules import llm_arbiter
        client = llm_arbiter.OpenRouterClient("fake", "model", 2)
        try:
            value = client.generate("prompt", {})
            result = runtime.provider_result(value, state="COMPLETED", call_ids=value["provider_call_ids"])
        except ValueError:
            result = runtime.provider_result([], state="FAILED")
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        states = db.execute("SELECT state,COUNT(*) FROM provider_calls GROUP BY state").fetchall()
    terminal = "DONE" if expected in {"EMPTY", "COMPLETED"} else "FAILED"
    assert result.result_state == expected and states == [(terminal, 1)] and len(recorder.journal) == 1


@pytest.mark.parametrize("content,usage", [
    ("not-json", {}),
    ('{"verdict":"match"}', {}),
    ('{"verdict":"match","reason":"short","detected_sector":"metal","expected_sector":"metal"}', {}),
    ('{"verdict":"match","reason":"a sufficiently detailed concrete reason","detected_sector":"metal","expected_sector":"metal"}', {"prompt_tokens": True}),
    ('{"verdict":"match","reason":"a sufficiently detailed concrete reason","detected_sector":"metal","expected_sector":"metal"}', {"total_tokens": -1}),
])
def test_openrouter_invalid_structured_verdict_is_failed_not_done(tmp_path, monkeypatch, content, usage):
    from modules import llm_arbiter
    _init_provider_run(tmp_path / "invalid-llm.sqlite3", "llm")
    recorder = RecordingPaidTransport([Response(200, {"choices": [{"message": {"content": content}}], "usage": usage})])
    monkeypatch.setattr(runtime, "_PAID_TRANSPORT", recorder)
    client = llm_arbiter.OpenRouterClient("fake", "model", 2)
    with pytest.raises(ValueError): client.generate("prompt", {})
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db:
        assert db.execute("SELECT state FROM provider_calls").fetchall() == [("FAILED",)] and len(recorder.journal) == 1


@pytest.mark.parametrize("surface", ["google_places", "brandfetch", "hunter_finder", "openrouter"])
def test_paid_adapter_cache_write_failure_does_not_change_terminal_result(tmp_path, monkeypatch, surface):
    from modules import cache_store, company_resolvers, google_places, llm_arbiter
    provider = "google_places" if surface == "google_places" else "brandfetch" if surface == "brandfetch" else "hunter" if surface == "hunter_finder" else "llm"
    _init_provider_run(tmp_path / f"cache-write-{surface}.sqlite3", provider)
    monkeypatch.setattr(config, "SEARCH_CACHE_MODE", "refresh")
    if surface == "google_places":
        recorder = RecordingPaidTransport([Response(200, {"places": []})]); runtime.set_paid_transport(recorder)
        monkeypatch.setattr(config, "ENABLE_GOOGLE_PLACES", True); monkeypatch.setattr(config, "GOOGLE_PLACES_API_KEY", "fake")
        monkeypatch.setattr(cache_store, "save", lambda *_a, **_k: (_ for _ in ()).throw(OSError("cache")))
        result = google_places.search_company("ACME"); counter = "api.google_places.cache_write_error"
    elif surface in {"brandfetch", "hunter_finder"}:
        recorder = RecordingPaidTransport([Response(200, [] if surface == "brandfetch" else {"data": []})]); runtime.set_paid_transport(recorder)
        monkeypatch.setattr(company_resolvers, "_save", lambda *_a, **_k: (_ for _ in ()).throw(OSError("cache")))
        if surface == "brandfetch":
            monkeypatch.setattr(config, "ENABLE_BRANDFETCH_DOMAIN_SEARCH", True); monkeypatch.setattr(config, "BRANDFETCH_CLIENT_ID", "fake")
            result = company_resolvers.brandfetch_domains("ACME"); counter = "resolver.brandfetch.cache_write_error"
        else:
            monkeypatch.setattr(config, "ENABLE_HUNTER_DOMAIN_FINDER", True); monkeypatch.setattr(config, "HUNTER_API_KEY", "fake")
            result = company_resolvers.hunter_domains("ACME"); counter = "resolver.hunter.cache_write_error"
    else:
        valid = '{"verdict":"match","reason":"a sufficiently detailed concrete reason","detected_sector":"metal","expected_sector":"metal"}'
        recorder = RecordingPaidTransport([Response(200, {"choices": [{"message": {"content": valid}}], "usage": {}})]); runtime.set_paid_transport(recorder)
        monkeypatch.setattr(config, "ENABLE_LLM_ARBITER", True); monkeypatch.setattr(config, "OPENROUTER_API_KEY", "fake")
        monkeypatch.setattr(cache_store, "save", lambda *_a, **_k: (_ for _ in ()).throw(OSError("cache")))
        result = llm_arbiter.arbitrate("ACME", "", "metal", "acme.example", "ACME manufactures metal machinery and industrial equipment"); counter = "api.llm_arbiter.cache_write_error"
    with sqlite3.connect(config.PROGRESS_DB_FILE) as db: states = db.execute("SELECT state FROM provider_calls").fetchall()
    state = result.get("provider_result") if isinstance(result, dict) else result.result_state
    assert state in {"EMPTY", "COMPLETED"} and states == [("DONE",)] and len(recorder.journal) == 1 and runtime.snapshot()["counters"].get(counter) == 1


def test_unknown_checkpoint_active_resume_is_valid_and_no_network_occurs(tmp_path, monkeypatch):
    source = tmp_path / "input.xlsx"; _input_book(source, ["PAID CO"])
    recorder = RecordingPaidTransport([requests.ReadTimeout("post-send")])
    runs = _real_run_setup(tmp_path, monkeypatch, recorder)
    first = main.run(source, allow_paid=True); root = next(runs.iterdir())
    assert first.status is pipeline_runner.PipelineOutcomeStatus.PAID_MANUAL_AUTHORIZATION_REVIEW_REQUIRED
    runtime.reset(); search.reset_run_state()
    resumed = main.run(source, allow_paid=True, resume_run=root)
    assert resumed.status is pipeline_runner.PipelineOutcomeStatus.PAID_MANUAL_AUTHORIZATION_REVIEW_REQUIRED and len(recorder.journal) == 1


@pytest.mark.parametrize("mutation", ["positive", "cross", "excess", "missing_attribution", "cap"])
def test_legacy_free_migration_reconciles_exact_bucket_state_multiset(tmp_path, mutation):
    path = tmp_path / f"legacy-free-{mutation}.sqlite3"; init_run(path)
    with sqlite3.connect(path) as db:
        counters = {"positive": (2,1,0,2,0), "cross": (2,1,0,1,1), "excess": (1,1,0,1,0), "missing_attribution": (2,0,0,1,0), "cap": (11,0,0,6,5)}[mutation]
        db.execute("INSERT INTO free_query_usage(run_id,item_index,used,quota,physical_used,physical_completed,physical_failed,discovery_physical_used,targeted_physical_used) VALUES('run',0,0,10,?,?,?,?,?)", counters)
        rows = [("existing", "discovery", "DONE")]
        if mutation == "cross": rows = [("existing", "targeted", "DONE")]
        if mutation == "excess": rows.append(("excess", "discovery", "DONE"))
        db.executemany("INSERT INTO free_provider_attempts(attempt_id,run_id,item_index,bucket,provider,query_fingerprint,attempt_ordinal,state,reserved_at) VALUES(?,'run',0,?,'ddgs',?,1,?,'legacy')", [(identity, bucket, identity, state) for identity,bucket,state in rows])
        db.execute("PRAGMA user_version=7"); db.commit()
    checkpoint._SCHEMA_READY.clear()
    if mutation != "positive":
        with pytest.raises(checkpoint.LedgerInvariant): checkpoint.initialize_schema(path)
        return
    checkpoint.initialize_schema(path)
    if mutation == "positive":
        with sqlite3.connect(path) as db:
            assert db.execute("SELECT bucket,state,COUNT(*) FROM free_provider_attempts GROUP BY bucket,state ORDER BY bucket,state").fetchall() == [("discovery", "DONE", 1), ("discovery", "RESERVED", 1)]
            assert db.execute("PRAGMA user_version").fetchone()[0] == checkpoint.SCHEDULER_SCHEMA_VERSION


def test_singleflight_clocked_heartbeat_covers_every_wait_region(tmp_path):
    path = tmp_path / "clocked.sqlite3"; init_run(path)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", lease_seconds=1)
    moments = [datetime(2030, 1, 1, 0, 0, second, tzinfo=timezone.utc) for second in range(4)]
    receipts = [checkpoint.heartbeat_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", lease_seconds=2, now_fn=lambda value=value: value) for value in moments]
    assert receipts == sorted(receipts) and len(set(receipts)) == 4


@pytest.mark.parametrize("mutation", ["paid_mode", "manifest_plan_hash", "db_plan_query", "db_plan_hash", "budget_missing", "budget_extra", "budget_value", "semantic_config", "replay_snapshot", "replay_manifest"])
def test_resume_query_plan_or_config_drift_returns_cli_22_before_network(tmp_path, monkeypatch, mutation, production_trace_call):
    source = tmp_path / "input.xlsx"; _input_book(source, ["PAID CO"])
    recorder = RecordingPaidTransport([]); runs = _real_run_setup(tmp_path, monkeypatch, recorder)
    assert main.run(source, allow_paid=False).status is pipeline_runner.PipelineOutcomeStatus.PAID_PENDING_APPROVAL
    root = next(runs.iterdir()); manifest_path = root / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if mutation == "manifest_plan_hash":
        manifest["paid_query_plan_sha256"] = "0" * 64
    elif mutation in {"db_plan_query", "db_plan_hash"}:
        with sqlite3.connect(root / "state" / "progress.sqlite3") as db:
            column = "normalized_query" if mutation == "db_plan_query" else "query_sha256"
            db.execute(f"UPDATE paid_query_plan_entries SET {column}=? WHERE rowid=(SELECT MIN(rowid) FROM paid_query_plan_entries)", ("tampered",)); db.commit()
    elif mutation == "budget_missing":
        manifest["run_config"]["budgets"].pop("llm")
    elif mutation == "budget_extra":
        manifest["run_config"]["budgets"]["extra"] = 1
    elif mutation == "budget_value":
        manifest["run_config"]["budgets"]["llm"] += 1
    elif mutation == "semantic_config":
        key = next(iter(manifest["run_config"]["effective_settings"]))
        value = manifest["run_config"]["effective_settings"][key]
        manifest["run_config"]["effective_settings"][key] = not value if isinstance(value, bool) else value + 1 if isinstance(value, (int, float)) else f"{value}-tampered"
    if mutation not in {"paid_mode", "db_plan_query", "db_plan_hash"}:
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    credential_loads = []
    monkeypatch.setattr(main, "_ensure_safe_project_runtime", lambda *_a, **_k: None)
    monkeypatch.setattr(main, "_apply_saved_resolver_configuration", lambda *_a, **_k: credential_loads.append(True))
    paid_flag = "--allow-paid" if mutation == "paid_mode" else "--no-allow-paid"
    replay_args = [f"--{mutation.replace('_', '-')}", str(tmp_path / "foreign-replay.json")] if mutation in {"replay_snapshot", "replay_manifest"} else []
    code = production_trace_call(main.cli, ["--input", str(source), "--resume-run", str(root), paid_flag, "--non-interactive", *replay_args])
    assert code == 22 and credential_loads == [] and recorder.journal == []


def test_resume_llm_budget_drift_is_rejected_before_saved_credentials(tmp_path, monkeypatch, production_trace_call):
    source = tmp_path / "input.xlsx"; _input_book(source, ["PAID CO"])
    recorder = RecordingPaidTransport([]); runs = _real_run_setup(tmp_path, monkeypatch, recorder)
    assert main.run(source, allow_paid=False).status is pipeline_runner.PipelineOutcomeStatus.PAID_PENDING_APPROVAL
    root = next(runs.iterdir()); manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    requested = int(manifest["run_config"]["budgets"]["llm"]) + 1
    credential_loads = []
    monkeypatch.setattr(main, "_ensure_safe_project_runtime", lambda *_a, **_k: None)
    monkeypatch.setattr(main, "_apply_saved_resolver_configuration", lambda *_a, **_k: credential_loads.append(True))
    code = production_trace_call(main.cli, ["--input", str(source), "--resume-run", str(root), "--no-allow-paid", "--llm-budget", str(requested), "--non-interactive"])
    assert code == 22 and credential_loads == [] and recorder.journal == []


@pytest.mark.parametrize("replica", ["intent", "manifest", "telemetry", "report"])
def test_finalization_rejects_each_telemetry_replica_drift(tmp_path, monkeypatch, replica, production_trace_call):
    source = tmp_path / "input.xlsx"; _input_book(source, ["PAID CO"])
    recorder = RecordingPaidTransport([Response(200, {"organic": [{"link": "https://paidco.example", "title": "PAID CO", "description": "PAID CO resmi şirket sitesi Türkiye"}]})])
    runs = _real_run_setup(tmp_path, monkeypatch, recorder); outcome = main.run(source, allow_paid=True); assert outcome.status is pipeline_runner.PipelineOutcomeStatus.COMPLETE
    root = next(runs.iterdir()); manifest_path = root / "manifest.json"; manifest = json.loads(manifest_path.read_text(encoding="utf-8")); monkeypatch.setattr(config, "PROGRESS_DB_FILE", root / "state" / "progress.sqlite3")
    artifact_dir = root / "output" / "artifacts" / manifest["artifact_set_sha256"]
    if replica == "intent":
        with sqlite3.connect(config.PROGRESS_DB_FILE) as db: db.execute("UPDATE finalization_intent SET telemetry_sha256='bad'"); db.commit()
    elif replica == "manifest":
        manifest["telemetry_sha256"] = "bad"; manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    elif replica == "telemetry":
        (artifact_dir / "telemetry.json").write_text("{}", encoding="utf-8")
    else:
        (artifact_dir / "report.txt").write_text("bad", encoding="utf-8")
    with pytest.raises(checkpoint.SchedulerInvariantError): production_trace_call(checkpoint.validate_finalization_contract, manifest["run_id"], root)


def test_k6_uses_real_negative_resume_drift_not_positive_resume_node():
    runner = (Path(__file__).parents[1] / "tools" / "closure_audit_runner.py")
    if not runner.is_file():
        pytest.fail("closure audit runner is missing")
    source = runner.read_text(encoding="utf-8")
    assert "test_resume_query_plan_or_config_drift_returns_cli_22_before_network" in source
    k6 = source.split("k6_negative_architecture", 1)[1]
    assert "test_real_fresh_handoff_resume_preserves_durable_ordered_query_plan" not in k6.split("run_command", 1)[0]
    assert checkpoint.file_hash(runner)


def _valid_audit_scenario(root: Path):
    from tools import closure_audit_runner as audit
    root.mkdir(); db_path = root / "sanitized_progress.sqlite3"; init_run(db_path)
    with sqlite3.connect(db_path) as db:
        db.execute("UPDATE run_phase_transitions SET to_phase='FREE' WHERE run_id='run' AND ordinal=0")
        db.execute("INSERT INTO run_phase_transitions(run_id,ordinal,from_phase,to_phase,transitioned_at) VALUES('run',1,'FREE','PAID','2026-01-01T00:00:00+00:00')")
        db.commit(); db.execute("PRAGMA wal_checkpoint(TRUNCATE)"); db.execute("PRAGMA journal_mode=DELETE")
        run_id, phase, context_json = db.execute("SELECT run_id,phase,context_json FROM runs").fetchone()
        trace = [{"ordinal": row[0], "from": row[1], "to": row[2], "at": row[3]} for row in db.execute("SELECT ordinal,from_phase,to_phase,transitioned_at FROM run_phase_transitions ORDER BY ordinal")]
    plan_hash = hashlib.sha256(b"[]").hexdigest()
    recovery = root / "recovery_state.sqlite3"; recovery.write_bytes(b"immutable recovery snapshot")
    recovery_sha = hashlib.sha256(recovery.read_bytes()).hexdigest()
    recovery_files = {"recovery_state.sqlite3": {"sha256": recovery_sha, "bytes": recovery.stat().st_size}}
    artifact_set_sha256 = hashlib.sha256(f"recovery_state.sqlite3:{recovery_sha}\n".encode()).hexdigest()
    files = {
        "scenario_result.json": {"scenario": "unauthorized_handoff", "observations": {"outcome": "PAID_PENDING_APPROVAL"}, "observed_absent": ["telemetry.json", "report.txt"]},
        "recording_transport_journal.json": [], "manifest.json": {"run_id": run_id, "phase": phase, "complete": False, "handoff": True, "paid_enabled": False, "files": recovery_files, "artifact_set_sha256": artifact_set_sha256, "paid_query_plan_count": 0, "paid_query_plan_sha256": plan_hash},
        "frozen_config.json": json.loads(context_json), "paid_query_plan.json": [], "state_transition_trace.json": trace,
        "command.json": {"command_id": "meta-command", "gate": "K5", "name": "k5_unauthorized_handoff", "argv": [sys.executable, "-m", "pytest", "tests/test_search_phase_regressions.py::test_real_fresh_run_without_authorization_seals_handoff"], "started_at": "2026-01-01T00:00:00+00:00", "ended_at": "2026-01-01T00:00:01+00:00", "duration_seconds": 1.0, "exit_code": 0, "timed_out": False, "source_stable": True},
        "duration.json": {"started_at": "2026-01-01T00:00:00+00:00", "ended_at": "2026-01-01T00:00:01+00:00", "duration_seconds": 1.0, "exit_code": 0, "timed_out": False},
    }
    for name, value in files.items(): (root / name).write_text(json.dumps(value), encoding="utf-8")
    (root / "pytest_reports.jsonl").write_text(json.dumps({"command_id": "meta-command", "nodeid": "tests/test_search_phase_regressions.py::test_real_fresh_run_without_authorization_seals_handoff", "phase": "call", "outcome": "passed", "wasxfail": "", "report_ordinal": 1}) + "\n", encoding="utf-8")
    receipts = [audit.file_receipt(path, relative_to=root) for path in sorted(root.iterdir()) if path.is_file() and path.name != "artifact_hashes.json"]
    audit.dump(root / "artifact_hashes.json", receipts)
    validation = audit.validate_scenario(root)
    assert validation["pass"], validation


def test_audit_scenario_validator_rejects_each_tampered_required_artifact(tmp_path):
    from tools import closure_audit_runner as audit
    original = tmp_path / "original"; _valid_audit_scenario(original)
    for name in ("frozen_config.json", "paid_query_plan.json", "state_transition_trace.json", "command.json", "duration.json", "pytest_reports.jsonl", "sanitized_progress.sqlite3", "recording_transport_journal.json", "manifest.json", "scenario_result.json", "artifact_hashes.json"):
        target = tmp_path / hashlib.sha256(name.encode()).hexdigest()[:10]; shutil.copytree(original, target)
        path = target / name
        if path.suffix == ".sqlite3": path.write_bytes(path.read_bytes() + b"tamper")
        elif path.suffix == ".jsonl": path.write_text("{}\n", encoding="utf-8")
        else: path.write_text("{}", encoding="utf-8")
        assert not audit.validate_scenario(target)["pass"], name


def test_audit_scenario_validator_rejects_rehashed_semantic_command_forgery(tmp_path):
    from tools import closure_audit_runner as audit
    root = tmp_path / "semantic"; _valid_audit_scenario(root)
    command = json.loads((root / "command.json").read_text(encoding="utf-8"))
    command["argv"][-1] = "tests/test_search_phase_regressions.py::test_unrelated"
    (root / "command.json").write_text(json.dumps(command), encoding="utf-8")
    receipts = [audit.file_receipt(path, relative_to=root) for path in sorted(root.iterdir()) if path.is_file() and path.name != "artifact_hashes.json"]
    audit.dump(root / "artifact_hashes.json", receipts)
    assert not audit.validate_scenario(root)["pass"]


def _valid_paid_audit_scenario(root: Path):
    from tools import closure_audit_runner as audit
    root.mkdir(); db_path = root / "sanitized_progress.sqlite3"; init_run(db_path, budget=1); runtime.set_item_context(0, "paid")
    checkpoint.begin_paid_attempt(run_id="run", item_index=0, attempt_number=1)
    checkpoint.claim_provider_query_flight(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner")
    call = runtime.reserve_api("brightdata", operation="scenario", request_fingerprint="request")
    runtime.start_api(call); checkpoint.bind_provider_query_flight_call(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=call.call_id); runtime.mark_api_http_started(call, 1, "fp")
    envelope = runtime.transport_envelope(call, endpoint="https://api.brightdata.com/request", attempt_ordinal=1, flight_fingerprint="fp", request_shape={"method": "POST", "json": {"query": "ACME"}}, timeout=2)
    result_payload = {"__search_result_state": "COMPLETED", "values": []}
    checkpoint.complete_provider_call_and_flight_success(run_id="run", provider="brightdata", query_fingerprint="fp", owner_token="owner", provider_call_id=call.call_id, result=result_payload, call_ids=[call.call_id])
    checkpoint.record_paid_attempt(run_id="run", item_index=0, attempt_number=1, result="COMPLETED", call_ids=[call.call_id], call_relations={call.call_id: "OWNER"})
    with sqlite3.connect(db_path) as db:
        db.execute("UPDATE run_items SET paid_state='DONE' WHERE run_id='run'"); db.execute("UPDATE runs SET phase='COMPLETE' WHERE run_id='run'")
        db.execute("DELETE FROM run_phase_transitions WHERE run_id='run'")
        transitions = [(0, "", "FREE"), (1, "FREE", "PAID"), (2, "PAID", "FINALIZING"), (3, "FINALIZING", "COMPLETE")]
        db.executemany("INSERT INTO run_phase_transitions(run_id,ordinal,from_phase,to_phase,transitioned_at) VALUES('run',?,?,?,'2026-01-01T00:00:00+00:00')", transitions)
        db.commit(); db.execute("PRAGMA wal_checkpoint(TRUNCATE)"); db.execute("PRAGMA journal_mode=DELETE")
        context_json = db.execute("SELECT context_json FROM runs WHERE run_id='run'").fetchone()[0]
    plan_hash = hashlib.sha256(b"[]").hexdigest(); (root / "telemetry.json").write_text("{}", encoding="utf-8"); (root / "report.txt").write_text("ok", encoding="utf-8")
    artifact_info = {name: {"bytes": (root / name).stat().st_size, "sha256": hashlib.sha256((root / name).read_bytes()).hexdigest()} for name in ("telemetry.json", "report.txt")}
    artifact_set_sha256 = hashlib.sha256("".join(f"{name}:{artifact_info[name]['sha256']}\n" for name in sorted(artifact_info)).encode()).hexdigest()
    nodeid = "tests/test_search_phase_regressions.py::test_real_authorized_e2e_uses_main_process_company_and_recording_transport"
    command = {"command_id": "meta-paid", "gate": "K5", "name": "k5_authorized_free_paid", "argv": [sys.executable, "-m", "pytest", nodeid], "started_at": "2026-01-01T00:00:00+00:00", "ended_at": "2026-01-01T00:00:01+00:00", "duration_seconds": 1.0, "exit_code": 0, "timed_out": False, "source_stable": True}
    files = {
        "scenario_result.json": {"scenario": "authorized_free_paid", "observations": {"outcome": "COMPLETE"}},
        "recording_transport_journal.json": [{**asdict(envelope), "invocation_ordinal": 1, "fake_response_class": "Response"}],
        "manifest.json": {"run_id": "run", "phase": "COMPLETE", "complete": True, "handoff": False, "paid_enabled": True, "files": artifact_info, "artifact_set_sha256": artifact_set_sha256, "paid_query_plan_count": 0, "paid_query_plan_sha256": plan_hash},
        "frozen_config.json": json.loads(context_json), "paid_query_plan.json": [],
        "state_transition_trace.json": [{"ordinal": ordinal, "from": source, "to": target, "at": "2026-01-01T00:00:00+00:00"} for ordinal, source, target in transitions],
        "command.json": command,
        "duration.json": {key: command[key] for key in ("started_at", "ended_at", "duration_seconds", "exit_code", "timed_out")},
    }
    for name, value in files.items(): (root / name).write_text(json.dumps(value), encoding="utf-8")
    (root / "pytest_reports.jsonl").write_text(json.dumps({"command_id": "meta-paid", "nodeid": nodeid, "phase": "call", "outcome": "passed", "wasxfail": "", "report_ordinal": 1}) + "\n", encoding="utf-8")
    with sqlite3.connect(db_path) as db:
        db.execute("INSERT INTO finalization_intent(run_id,generation,input_snapshot_sha256,started_at,artifact_set_sha256,manifest_sha256,completed_at,status) VALUES('run','meta','input','2026-01-01T00:00:00+00:00',?,?,'2026-01-01T00:00:01+00:00','COMPLETE')", (artifact_set_sha256, hashlib.sha256((root / "manifest.json").read_bytes()).hexdigest()))
        db.commit(); db.execute("PRAGMA wal_checkpoint(TRUNCATE)"); db.execute("PRAGMA journal_mode=DELETE")
    audit.dump(root / "artifact_hashes.json", [audit.file_receipt(path, relative_to=root) for path in sorted(root.iterdir()) if path.is_file() and path.name != "artifact_hashes.json"])
    validation = audit.validate_scenario(root); assert validation["pass"], validation


@pytest.mark.parametrize("mutation", [
    "endpoint_provider", "request_shape", "consumer_generation", "terminal_calls",
    "invalid_terminal_json", "result_hash", "delete_receipt_chain", "manifest_empty", "manifest_escape",
])
def test_audit_scenario_validator_rejects_rehashed_transport_and_relational_forgery(tmp_path, mutation):
    from tools import closure_audit_runner as audit
    root = tmp_path / mutation; _valid_paid_audit_scenario(root)
    if mutation in {"endpoint_provider", "request_shape"}:
        journal_path = root / "recording_transport_journal.json"; journal = json.loads(journal_path.read_text(encoding="utf-8"))
        journal[0]["endpoint" if mutation == "endpoint_provider" else "request_shape_sha256"] = "https://api.hunter.io/v2/domain-search" if mutation == "endpoint_provider" else "f" * 64
        journal_path.write_text(json.dumps(journal), encoding="utf-8")
        if mutation == "endpoint_provider":
            with sqlite3.connect(root / "sanitized_progress.sqlite3") as db:
                trigger_sql = db.execute("SELECT sql FROM sqlite_master WHERE type='trigger' AND name='provider_call_transport_receipt_update'").fetchone()[0]
                db.execute("DROP TRIGGER provider_call_transport_receipt_update")
                db.execute("UPDATE provider_calls SET endpoint_sha256=?", (hashlib.sha256(journal[0]["endpoint"].encode()).hexdigest(),))
                db.execute(trigger_sql)
                db.commit(); db.execute("PRAGMA wal_checkpoint(TRUNCATE)"); db.execute("PRAGMA journal_mode=DELETE")
    elif mutation in {"manifest_empty", "manifest_escape"}:
        manifest_path = root / "manifest.json"; manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if mutation == "manifest_empty":
            manifest["files"] = {}
        else:
            outside = root.parent / "outside"; outside.write_text("outside", encoding="utf-8")
            digest = hashlib.sha256(outside.read_bytes()).hexdigest()
            manifest["files"] = {"../outside": {"bytes": outside.stat().st_size, "sha256": digest}}
        manifest["artifact_set_sha256"] = hashlib.sha256("".join(f"{name}:{manifest['files'][name]['sha256']}\n" for name in sorted(manifest["files"])).encode()).hexdigest()
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        with sqlite3.connect(root / "sanitized_progress.sqlite3") as db:
            db.execute("UPDATE finalization_intent SET artifact_set_sha256=?,manifest_sha256=?", (manifest["artifact_set_sha256"], hashlib.sha256(manifest_path.read_bytes()).hexdigest()))
            db.commit(); db.execute("PRAGMA wal_checkpoint(TRUNCATE)"); db.execute("PRAGMA journal_mode=DELETE")
    else:
        with sqlite3.connect(root / "sanitized_progress.sqlite3") as db:
            if mutation == "consumer_generation":
                trigger_sql = [db.execute("SELECT sql FROM sqlite_master WHERE type='trigger' AND name=?", (name,)).fetchone()[0] for name in ("flight_consumer_identity_update", "flight_consumer_scope_update")]
                db.execute("DROP TRIGGER flight_consumer_identity_update"); db.execute("DROP TRIGGER flight_consumer_scope_update"); db.execute("UPDATE provider_query_flight_consumers SET execution_generation=2")
                for statement in trigger_sql: db.execute(statement)
            elif mutation in {"terminal_calls", "invalid_terminal_json"}:
                trigger_sql = [db.execute("SELECT sql FROM sqlite_master WHERE type='trigger' AND name=?", (name,)).fetchone()[0] for name in ("flight_terminal_scope_update", "flight_terminal_receipt_update")]
                db.execute("DROP TRIGGER flight_terminal_scope_update"); db.execute("DROP TRIGGER flight_terminal_receipt_update"); db.execute("UPDATE provider_query_flight_terminals SET call_ids_json='[]'")
                if mutation == "invalid_terminal_json": db.execute("UPDATE provider_query_flight_terminals SET call_ids_json='{'")
                for statement in trigger_sql: db.execute(statement)
            elif mutation == "delete_receipt_chain":
                trigger_sql = [db.execute("SELECT sql FROM sqlite_master WHERE type='trigger' AND name=?", (name,)).fetchone()[0] for name in ("flight_result_receipt_delete", "flight_terminal_receipt_delete")]
                db.execute("DROP TRIGGER flight_result_receipt_delete"); db.execute("DROP TRIGGER flight_terminal_receipt_delete")
                db.execute("DELETE FROM provider_query_flight_consumers"); db.execute("DELETE FROM provider_query_flight_terminals"); db.execute("DELETE FROM provider_query_flight_results"); db.execute("DELETE FROM provider_query_flights")
                for statement in trigger_sql: db.execute(statement)
            else:
                trigger_sql = db.execute("SELECT sql FROM sqlite_master WHERE type='trigger' AND name='flight_result_receipt_update'").fetchone()[0]
                db.execute("DROP TRIGGER flight_result_receipt_update")
                db.execute("UPDATE provider_query_flight_results SET result_sha256=?", ("f" * 64,))
                db.execute(trigger_sql)
            db.commit(); db.execute("PRAGMA wal_checkpoint(TRUNCATE)"); db.execute("PRAGMA journal_mode=DELETE")
    audit.dump(root / "artifact_hashes.json", [audit.file_receipt(path, relative_to=root) for path in sorted(root.iterdir()) if path.is_file() and path.name != "artifact_hashes.json"])
    assert not audit.validate_scenario(root)["pass"]


def test_root_artifact_hash_is_a_marker_gate_and_includes_nested_manifests(tmp_path):
    from tools import closure_audit_runner as audit
    for index in range(8):
        nested = tmp_path / f"scenario-{index}"; nested.mkdir(); (nested / "artifact_hashes.json").write_text("[]", encoding="utf-8")
    report_path = tmp_path / "final_report.md"; report_path.write_text("SEARCH_FLOW_FIX_COMPLETE_CLOSURE_AUDIT", encoding="utf-8")
    audit.write_artifact_hashes(tmp_path); valid = audit.validate_artifact_hashes(tmp_path)
    assert valid["pass"] and len(valid["nested_manifests"]) == 8
    report_path.write_text("tampered", encoding="utf-8")
    assert not audit.validate_artifact_hashes(tmp_path)["pass"]


def test_contract_matrix_has_no_fallback_and_replays_every_receipt(tmp_path):
    from tools import closure_audit_runner as audit
    nodeid = "tests/test_search_phase_regressions.py::node"
    report = tmp_path / "pytest_unit.jsonl"; report.write_text(json.dumps({"command_id": "command", "nodeid": nodeid, "phase": "call", "outcome": "passed", "wasxfail": ""}) + "\n", encoding="utf-8")
    trace = tmp_path / "traces" / "command.jsonl"; trace.parent.mkdir(); trace.write_text(json.dumps({"command_id": "command", "nodeid": nodeid, "function_entries": ["modules/checkpoint.py:file_hash"]}) + "\n", encoding="utf-8")
    commands = [{"command_id": "command", "exit_code": 0, "timed_out": False}]
    base = {"R1": {"node": "node", "collected_nodeids": [nodeid], "command_ids": ["command"], "outcomes": ["passed"], "production_function_entries": [{"command_id": "command", "nodeid": nodeid, "trace_path": "traces/command.jsonl", "function_entries": ["modules/checkpoint.py:file_hash"]}], "receipt_paths": ["pytest_unit.jsonl", "traces/command.jsonl"], "pass": True}}
    matrix = tmp_path / "matrix.json"; matrix.write_text(json.dumps(base), encoding="utf-8")
    entrypoints = {"node": {"modules/checkpoint.py:file_hash"}}
    assert audit.validate_contract_matrix(matrix, tmp_path, expected_nodes=["node"], expected_entrypoints=entrypoints, command_records=commands)["pass"]
    for field in ("collected_nodeids", "command_ids", "outcomes", "production_function_entries", "receipt_paths"):
        mutated = json.loads(json.dumps(base)); mutated["R1"][field] = []
        matrix.write_text(json.dumps(mutated), encoding="utf-8")
        assert not audit.validate_contract_matrix(matrix, tmp_path, expected_nodes=["node"], expected_entrypoints=entrypoints, command_records=commands)["pass"]
    forged = json.loads(json.dumps(base)); forged["R1"].update({"node": "fake", "collected_nodeids": ["fake"], "command_ids": ["fake"], "outcomes": ["passed"], "production_function_entries": [{"command_id": "fake", "nodeid": "fake", "trace_path": "traces/command.jsonl", "function_entries": ["modules/checkpoint.py:file_hash"]}], "pass": True})
    matrix.write_text(json.dumps(forged), encoding="utf-8")
    assert not audit.validate_contract_matrix(matrix, tmp_path, expected_nodes=["node"], expected_entrypoints=entrypoints, command_records=commands)["pass"]
    generic = ["modules/runtime.py:reset"]
    trace.write_text(json.dumps({"command_id": "command", "nodeid": nodeid, "function_entries": generic}) + "\n", encoding="utf-8")
    forged = json.loads(json.dumps(base)); forged["R1"]["production_function_entries"][0]["function_entries"] = generic
    matrix.write_text(json.dumps(forged), encoding="utf-8")
    assert not audit.validate_contract_matrix(matrix, tmp_path, expected_nodes=["node"], expected_entrypoints=entrypoints, command_records=commands)["pass"]


def test_offline_paid_network_guard_fails_on_every_default_transport_fallthrough(tmp_path, monkeypatch):
    from modules import company_resolvers, google_places, hunter, linkedin_company, llm_arbiter
    receipt = Path(os.environ.get("B2B_SOCKET_DENY_JSONL", "")) if os.environ.get("B2B_SOCKET_DENY_JSONL", "").strip() else tmp_path / "socket-deny.jsonl"
    monkeypatch.setenv("B2B_SOCKET_DENY_JSONL", str(receipt))
    monkeypatch.setattr(config, "SEARCH_CACHE_MODE", "off"); monkeypatch.setattr(config, "GLOBAL_REQUESTS_PER_SECOND", 0); monkeypatch.setattr(config, "BRIGHTDATA_REQUESTS_PER_MINUTE", 0); monkeypatch.setattr(config, "MAX_RETRIES", 0)
    invocations = {
        "brightdata": lambda: search._brightdata_text("ACME"),
        "google_places": lambda: google_places.search_company("ACME"),
        "brandfetch": lambda: company_resolvers.brandfetch_domains("ACME"),
        "hunter": lambda: hunter.find_domain_emails("acme.example"),
        "linkedin": lambda: linkedin_company._scrape("https://linkedin.com/company/acme"),
        "llm": lambda: llm_arbiter.OpenRouterClient("fake", "model", 1).generate("prompt", {}),
    }
    observed = {}
    for provider, invoke in invocations.items():
        runtime.reset(); runtime.set_paid_transport(None)
        monkeypatch.setattr(config, "BRIGHTDATA_API_KEY", "fake"); monkeypatch.setattr(config, "BRIGHTDATA_REQUEST_BUDGET", 1)
        monkeypatch.setattr(config, "ENABLE_GOOGLE_PLACES", True); monkeypatch.setattr(config, "GOOGLE_PLACES_API_KEY", "fake"); monkeypatch.setattr(config, "GOOGLE_PLACES_REQUEST_BUDGET", 1)
        monkeypatch.setattr(config, "ENABLE_BRANDFETCH_DOMAIN_SEARCH", True); monkeypatch.setattr(config, "BRANDFETCH_CLIENT_ID", "fake"); monkeypatch.setattr(config, "BRANDFETCH_REQUEST_BUDGET", 1)
        monkeypatch.setattr(config, "ENABLE_HUNTER_FALLBACK", True); monkeypatch.setattr(config, "HUNTER_API_KEY", "fake"); monkeypatch.setattr(config, "HUNTER_REQUEST_BUDGET", 1)
        monkeypatch.setattr(config, "LINKEDIN_COMPANY_DATASET_ID", "fake"); monkeypatch.setattr(config, "LINKEDIN_COMPANY_REQUEST_BUDGET", 1)
        monkeypatch.setattr(config, "LLM_ARBITER_BUDGET", 1)
        before = len(receipt.read_text(encoding="utf-8").splitlines()) if receipt.exists() else 0
        try: invoke()
        except Exception: pass
        after = len(receipt.read_text(encoding="utf-8").splitlines()) if receipt.exists() else 0
        observed[provider] = after - before
        with receipt.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps({"kind": "fallthrough_probe", "command_id": os.environ.get("B2B_COMMAND_ID", ""), "nodeid": "tests/test_search_phase_regressions.py::test_offline_paid_network_guard_fails_on_every_default_transport_fallthrough", "provider": provider, "blocked_events": observed[provider]}, sort_keys=True) + "\n")
    assert set(observed) == checkpoint.CANONICAL_PROVIDERS and all(count >= 1 for count in observed.values())


def test_network_receipt_validator_derives_zero_external_calls_and_rejects_forgery(tmp_path):
    from tools import closure_audit_runner as audit
    command_id = "k6-network"; nodeid = "tests/test_search_phase_regressions.py::test_offline_paid_network_guard_fails_on_every_default_transport_fallthrough"
    commands = [{"command_id": command_id, "argv": [sys.executable, "-m", "pytest", nodeid]}]
    reports = [{"command_id": command_id, "nodeid": nodeid, "phase": "call", "outcome": "passed", "wasxfail": ""}]
    records = [{"kind": "guard_armed", "command_id": command_id, "pid": 1}]
    for provider in ("brightdata", "google_places", "brandfetch", "hunter", "linkedin", "llm"):
        records.extend((
            {"kind": "blocked_network", "command_id": command_id, "nodeid": nodeid, "event": "socket.connect", "args": []},
            {"kind": "fallthrough_probe", "command_id": command_id, "nodeid": nodeid, "provider": provider, "blocked_events": 1},
        ))
    receipt = tmp_path / "network" / f"{command_id}.jsonl"; receipt.parent.mkdir()
    receipt.write_text("\n".join(json.dumps(row) for row in records) + "\n", encoding="utf-8")
    valid = audit.validate_network_receipts(tmp_path, commands, reports, {"scenario": {"paid_endpoint_hits": 2}})
    assert valid["pass"] and valid["actual"]["real_external_paid_calls"] == 0 and valid["actual"]["recorder_interceptions"] == 2
    records.append({"kind": "fallthrough_probe", "command_id": "wrong", "nodeid": "wrong", "provider": "evil", "blocked_events": 1})
    receipt.write_text("\n".join(json.dumps(row) for row in records) + "\n", encoding="utf-8")
    assert not audit.validate_network_receipts(tmp_path, commands, reports, {})["pass"]
    records.pop()
    records[1]["nodeid"] = "tests/test_other.py::unexpected"
    receipt.write_text("\n".join(json.dumps(row) for row in records) + "\n", encoding="utf-8")
    assert not audit.validate_network_receipts(tmp_path, commands, reports, {})["pass"]


def test_k6_receipt_validator_rejects_missing_duplicate_and_forged_cases(tmp_path):
    from tools import closure_audit_runner as audit
    nodeid = "tests/test_search_phase_regressions.py::negative_case[value]"; command_id = "k6-command"
    disposable = tmp_path / "disposable"; disposable.mkdir()
    empty_manifest = []; empty_hash = hashlib.sha256(json.dumps(empty_manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    before = {
        "artifact_root": str(disposable.resolve()),
        "artifact_manifest": empty_manifest,
        "artifact_sha256": empty_hash,
        "recorder_journal": [],
        "recorder_count": 0,
        "recorder_sha256": hashlib.sha256(b"[]").hexdigest(),
        "network_receipt_path": "", "network_receipt_count": 0,
        "network_receipt_sha256": hashlib.sha256(b"[]").hexdigest(),
        "socket_blocked": 0,
    }
    proof = disposable / "mutation.json"; proof.write_text('{"rejected":true}', encoding="utf-8"); proof_data = proof.read_bytes(); proof_stat = proof.stat()
    manifest = [{"path": "mutation.json", "absolute_path": str(proof.resolve()), "bytes": len(proof_data), "sha256": hashlib.sha256(proof_data).hexdigest(), "mtime_ns": proof_stat.st_mtime_ns}]
    after = dict(before, artifact_manifest=manifest, artifact_sha256=hashlib.sha256(json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest())
    case_id = hashlib.sha256(nodeid.encode()).hexdigest()
    receipt = {
        "schema_version": 2, "command_id": command_id, "nodeid": nodeid,
        "mutation_id": "value", "case_id": case_id, "parameters": {"mutation": "value"},
        "disposable_root": str(disposable.resolve()), "before": before, "after": after,
        "expected": {"pytest_outcome": "passed", "mutation_observed": True, "exception_class": ["NONE"], "reason_pattern": [], "cli_exit": None},
        "actual": {"pytest_outcome": "passed", "mutation_observed": True, "exception_class": ["NONE"], "reason": [], "cli_exit": None},
        "rejections": [], "stable_reason": "artifact_state_changed",
        "mutation_evidence": {"artifact_changed": True, "recorder_changed": False, "socket_blocked_delta": 0, "matched_rejections": 0, "cli_exit_observed": False},
        "report_outcome": "passed", "wasxfail": "", "evidence_path": f"k6_receipts/{case_id}.json",
    }
    case_path = tmp_path / receipt["evidence_path"]; case_path.parent.mkdir(); case_path.write_text(json.dumps(receipt), encoding="utf-8")
    receipts = tmp_path / "k6.jsonl"; receipts.write_text(json.dumps(receipt) + "\n", encoding="utf-8")
    reports = [{"command_id": command_id, "nodeid": nodeid, "phase": "call", "outcome": "passed", "wasxfail": ""}]
    assert audit.validate_k6_receipts(receipts, tmp_path, reports, command_id)["pass"]
    invalid_regex = json.loads(json.dumps(receipt))
    invalid_regex["rejections"] = [{"expected_classes": ["builtins.ValueError"], "class": "builtins.ValueError", "reason": "rejected", "source": "test:1", "expected_reason_pattern": "["}]
    invalid_regex["expected"]["exception_class"] = ["builtins.ValueError"]; invalid_regex["expected"]["reason_pattern"] = ["["]
    invalid_regex["actual"]["exception_class"] = ["builtins.ValueError"]; invalid_regex["actual"]["reason"] = ["rejected"]
    invalid_regex["mutation_evidence"]["matched_rejections"] = 1; invalid_regex["stable_reason"] = "builtins.ValueError:rejected"
    case_path.write_text(json.dumps(invalid_regex), encoding="utf-8"); receipts.write_text(json.dumps(invalid_regex) + "\n", encoding="utf-8")
    assert not audit.validate_k6_receipts(receipts, tmp_path, reports, command_id)["pass"]
    forged_reason = json.loads(json.dumps(receipt)); forged_reason["stable_reason"] = "forged"
    case_path.write_text(json.dumps(forged_reason), encoding="utf-8"); receipts.write_text(json.dumps(forged_reason) + "\n", encoding="utf-8")
    assert not audit.validate_k6_receipts(receipts, tmp_path, reports, command_id)["pass"]
    receipts.write_text("", encoding="utf-8")
    assert not audit.validate_k6_receipts(receipts, tmp_path, reports, command_id)["pass"]
    receipts.write_text(json.dumps(receipt) + "\n" + json.dumps(receipt) + "\n", encoding="utf-8")
    assert not audit.validate_k6_receipts(receipts, tmp_path, reports, command_id)["pass"]
    forged = json.loads(json.dumps(receipt)); forged["actual"]["mutation_observed"] = False
    case_path.write_text(json.dumps(forged), encoding="utf-8"); receipts.write_text(json.dumps(forged) + "\n", encoding="utf-8")
    assert not audit.validate_k6_receipts(receipts, tmp_path, reports, command_id)["pass"]
    case_path.write_text(json.dumps(receipt), encoding="utf-8"); receipts.write_text(json.dumps(receipt) + "\n", encoding="utf-8"); proof.write_text('{"rejected":false}', encoding="utf-8")
    assert not audit.validate_k6_receipts(receipts, tmp_path, reports, command_id)["pass"]

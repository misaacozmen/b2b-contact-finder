"""Talimat 44: the British profile and the company check of British customer files."""

from datetime import date

from openpyxl import load_workbook
import pytest
import requests

import config
from modules import (
    company_register, country_profile, customer_file, discovery_rules, email_kind, legal_form, phone,
    query_planner, reference_resolution, scorer, tax_identity,
)
from tools import api_anahtarlari, musteri_dosyasi


@pytest.fixture
def britain():
    country_profile.apply("GB")
    try:
        yield
    finally:
        country_profile.apply("TR")


def test_british_profile_switches_and_turkey_restores(britain):
    assert config.TARGET_COUNTRY == "GB"
    assert config.PHONE_ALLOWED_COUNTRIES == ["GB", "GG", "JE", "IM"]
    assert "uk" not in config.FOREIGN_COUNTRY_TLDS and "tr" in config.FOREIGN_COUNTRY_TLDS
    assert config.DELIVERY_COMPANIES_ONLY and config.COMPANY_REGISTER == "companies_house"
    assert config.TAX_ID_FORMAT == ""
    assert country_profile.LABELS["GB"] == "Birleşik Krallık"
    country_profile.apply("TR")
    assert config.TARGET_COUNTRY == "TR" and config.COMPANY_REGISTER == "" and not config.DELIVERY_COMPANIES_ONLY


def test_british_queries_use_english_words(britain):
    queries = discovery_rules.fallback_queries("Ornek Stands Ltd", {})
    assert '"ornek stands ltd" UK official website' in queries
    assert all("Turkiye" not in query and "iletisim" not in query for query in queries)
    assert query_planner.query_intent('"ornek" contact UK') == "contact"
    assert scorer.search_display_core("ORNEK STANDS LTD T/A ORNEK EVENTS") == "ornek stands"


def test_british_phones_domains_and_directories(britain):
    assert phone.normalize_phone("+44 (0)20 7946 0123") == "02079460123"
    assert phone.normalize_phone("0113 496 0123") == "01134960123"
    assert phone.normalize_phone("0212 555 00 00") == ""
    assert not scorer.is_foreign_country_domain("ornek-stands.co.uk")
    assert scorer.is_foreign_country_domain("ornek.com.tr")
    assert scorer.is_excluded_domain("yell.com") and scorer.is_excluded_domain("find-and-update.company-information.service.gov.uk")


def test_british_contact_selection_prefers_a_generic_mailbox(britain):
    obs = {"emails": ["emma@ornek-stands.co.uk", "enquiries@ornek-stands.co.uk"], "phones": ["02079460123"], "listed_email": ""}
    contacts = reference_resolution.select_contacts(obs, "ornek-stands.co.uk", "")
    assert contacts["email"] == "enquiries@ornek-stands.co.uk"
    assert contacts["phone"] == "02079460123"


@pytest.mark.parametrize("name", [
    "Ornek Stands Ltd", "ORNEK DISPLAYS LIMITED", "Ornek Group PLC", "Ornek Group P.L.C.", "Ornek Partners LLP",
    "Ornek Events L.L.P.", "Ornek Community Interest Company", "Ornek Arts CIC", "Ornek Cyf",
    "Ornek Cyfyngedig", "Ornek Trust CIO", "University of Ornek", "Smith & Co Ltd", "Ornek Trading Co",
])
def test_british_company_forms(name):
    assert legal_form.is_company(name)


@pytest.mark.parametrize("name", [
    "Ornek Stands", "J Smith & Co", "Ornek and Company", "Ornek & Co. (Stand 4B)", "Ornek & Sons",
    "Ornek Partnership", "John Smith",
])
def test_british_sole_traders_and_partnerships(name):
    assert not legal_form.is_company(name)


@pytest.mark.parametrize("email, personal", [
    ("emma@ornek-stands.co.uk", True), ("chris.taylor@ornek-stands.co.uk", True), ("j.taylor@ornek-stands.co.uk", True),
    ("enquiries@ornek-stands.co.uk", False), ("hello@ornek-stands.co.uk", False), ("accounts@ornek-stands.co.uk", False),
    ("events.team@ornek-stands.co.uk", False),
])
def test_british_first_name_mailboxes(email, personal):
    assert email_kind.is_personal(email) is personal


@pytest.mark.parametrize("text, numbers", [
    ("© 2026 Ornek Exhibitions Ltd. Registered in England and Wales No. 99000001.", ["99000001"]),
    ("Registered in England & Wales, Company No 9900002", ["09900002"]),
    ("Registered in Scotland under company number SC 990001", ["SC990001"]),
    ("Company Registration Number: OC999901", ["OC999901"]),
    ("Registered office: 1 Ornek Street. Registered number 99000003", ["99000003"]),
    ("CRN 99000004 · Reg. No. 99000005", ["99000004", "99000005"]),
    ("VAT Reg No. 123456789", []),
    ("VAT Registration Number 99000006", []),
    ("Charity Registration No. 1990001", []),
    ("Registered Charity Number 1990001", []),
    ("Company No 0123456789", []),
    ("Call 020 7946 0123", []),
])
def test_company_numbers_on_a_page(text, numbers):
    assert company_register.page_numbers(text) == numbers


def test_legal_names_and_name_relations():
    text = "Home About Contact © 2026 Ornek Exhibitions Ltd. All rights reserved. Beta Web Design Limited"
    assert company_register.legal_names(text) == ["2026 Ornek Exhibitions Ltd", "Beta Web Design Limited"]
    assert company_register.related("ORNEK EXHIBITIONS LIMITED", "Ornek Stands")
    assert not company_register.related("BETA EXHIBITIONS LIMITED", "Ornek Exhibitions")
    assert not company_register.related("Home About Contact Ltd", "Home Displays")
    assert company_register.related("ORNEKSTAND LTD", "Gamma Displays", "https://www.ornekstand.co.uk/")
    assert company_register.same_name("ORNEK STANDS LTD", "Ornek Stands")
    assert company_register.same_name("THE ORNEK & CO LIMITED", "Ornek and Co")
    assert not company_register.same_name("ORNEK STANDS UK LTD", "Ornek Stands")


SITE = "https://www.ornek-stands.co.uk/"
FOOTER = "<footer>&copy; 2026 Ornek Exhibitions Ltd. Registered in England &amp; Wales No. 99000001</footer>"


class FakeRegister:
    def __init__(self, companies=None, results=None, error=False):
        self.companies = companies or {}
        self.results = results or []
        self.error = error
        self.calls = []

    def company(self, number):
        self.calls.append(("company", number))
        if self.error:
            raise company_register.RegisterError("http_500")
        return self.companies.get(number, {})

    def search(self, name):
        self.calls.append(("search", name))
        if self.error:
            raise company_register.RegisterError("http_500")
        return self.results


def _pages(mapping):
    fetched = []

    def fetch(url):
        fetched.append(url)
        return mapping.get(url, "")

    return fetch, fetched


def test_a_name_with_a_company_form_needs_no_check():
    fetch, fetched = _pages({})
    register = FakeRegister()
    result = company_register.check_firm("Ornek Stands Ltd", SITE, register=register, fetch_page=fetch)
    assert (result["company"], result["evidence"]) == (True, "name")
    assert fetched == [] and register.calls == []


def test_a_site_number_that_companies_house_confirms():
    fetch, _fetched = _pages({SITE: f"<html><body><p>Stands</p>{FOOTER}</body></html>"})
    register = FakeRegister(companies={"99000001": {
        "company_name": "ORNEK EXHIBITIONS LIMITED", "company_status": "active", "type": "ltd",
    }})
    result = company_register.check_firm("Ornek Stands", SITE, register=register, fetch_page=fetch)
    assert (result["company"], result["evidence"], result["number"]) == (True, "register_number", "99000001")
    assert register.calls == [("company", "99000001")]


def test_a_brand_of_the_company_that_the_page_names():
    footer = "<footer>Ornek Displays is a trading brand of Beta Gamma. Registered in England No 99000010</footer>"
    fetch, _fetched = _pages({SITE: f"<html><body>{footer}</body></html>"})
    register = FakeRegister(companies={
        "99000010": {"company_name": "BETA GAMMA LIMITED", "company_status": "active", "type": "ltd"},
    })
    result = company_register.check_firm("Ornek Displays", SITE, register=register, fetch_page=fetch)
    assert (result["company"], result["evidence"], result["register_name"]) == (True, "register_number", "BETA GAMMA LIMITED")
    fetch, _fetched = _pages({SITE: "<footer>Company No 99000010</footer>"})
    assert not company_register.check_firm("Ornek Displays", SITE, register=register, fetch_page=fetch)["company"]


def test_companies_house_wins_over_the_site():
    fetch, _fetched = _pages({SITE: f"<html><body>{FOOTER}</body></html>"})
    register = FakeRegister(companies={"99000001": {
        "company_name": "ORNEK EXHIBITIONS LIMITED", "company_status": "dissolved", "type": "ltd",
    }})
    result = company_register.check_firm("Ornek Stands", SITE, register=register, fetch_page=fetch)
    assert not result["company"]
    assert register.calls == [("company", "99000001"), ("search", "Ornek Stands")]


@pytest.mark.parametrize("company_type, company", [("ltd", True), ("llp", True), ("limited-partnership", False)])
def test_a_firm_with_exactly_its_name_at_companies_house(company_type, company):
    fetch, _fetched = _pages({})
    register = FakeRegister(results=[
        {"title": "ORNEK DISPLAYS UK LTD", "company_status": "active", "company_type": "ltd", "company_number": "99000007"},
        {"title": "ORNEK DISPLAYS LTD", "company_status": "active", "company_type": company_type, "company_number": "99000002"},
    ])
    result = company_register.check_firm("Ornek Displays", "", register=register, fetch_page=fetch)
    assert result["company"] is company
    if company:
        assert (result["evidence"], result["number"]) == ("register_name", "99000002")


def test_without_a_key_the_site_must_show_the_number_and_a_related_legal_name():
    fetch, _fetched = _pages({SITE: f"<html><body>{FOOTER}</body></html>"})
    result = company_register.check_firm("Ornek Stands", SITE, register=None, fetch_page=fetch)
    assert (result["company"], result["evidence"], result["number"]) == (True, "site", "99000001")
    fetch, _fetched = _pages({SITE: "<footer>Site by Beta Web Design Ltd. Company No 99000008</footer>"})
    assert not company_register.check_firm("Ornek Stands", SITE, register=None, fetch_page=fetch)["company"]
    fetch, _fetched = _pages({SITE: "<footer>© 2026 Ornek Exhibitions Ltd</footer>"})
    assert not company_register.check_firm("Ornek Stands", SITE, register=None, fetch_page=fetch)["company"]


def test_legal_pages_of_the_own_site_are_read_until_one_shows_a_number():
    home = (
        '<html><body><a href="/privacy-policy">Privacy</a><a href="/terms">Terms</a>'
        '<a href="https://www.yell.com/privacy">Directory</a><p>Ornek Stands</p></body></html>'
    )
    privacy = "https://www.ornek-stands.co.uk/privacy-policy"
    fetch, fetched = _pages({SITE: home, privacy: f"<html><body>{FOOTER}</body></html>"})
    result = company_register.check_firm("Ornek Stands", SITE, register=None, fetch_page=fetch)
    assert result["evidence"] == "site"
    assert fetched == [SITE, privacy]


def test_a_register_error_falls_back_to_the_site_rule():
    fetch, _fetched = _pages({SITE: f"<html><body>{FOOTER}</body></html>"})
    result = company_register.check_firm("Ornek Stands", SITE, register=FakeRegister(error=True), fetch_page=fetch)
    assert (result["company"], result["evidence"], result["register_error"]) == (True, "site", True)


class FakeResponse:
    def __init__(self, status, payload=None):
        self.status_code = status
        self._payload = payload

    def json(self):
        return self._payload


class FakeSession:
    def __init__(self, responses):
        self.responses = list(responses)
        self.requests = []

    def get(self, url, **kwargs):
        self.requests.append((url, kwargs))
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response


def test_companies_house_client_requests_and_errors():
    session = FakeSession([
        FakeResponse(200, {"company_name": "ORNEK STANDS LIMITED"}), FakeResponse(404),
        FakeResponse(200, {"items": [{"title": "ORNEK STANDS LIMITED"}, "bad"]}), FakeResponse(429),
        requests.ConnectionError("synthetic-key-0000 failed"),
    ])
    client = company_register.CompaniesHouse("synthetic-key-0000", session=session, min_interval=0)
    assert client.company("99000001") == {"company_name": "ORNEK STANDS LIMITED"}
    assert client.company("99000009") == {}
    assert client.search("Ornek Stands") == [{"title": "ORNEK STANDS LIMITED"}]
    url, kwargs = session.requests[0]
    assert url == f"{config.COMPANIES_HOUSE_API_URL}/company/99000001"
    assert kwargs["auth"] == ("synthetic-key-0000", "")
    assert session.requests[2][1]["params"] == {"q": "Ornek Stands", "items_per_page": 20}
    for expected in ("http_429", "ConnectionError"):
        with pytest.raises(company_register.RegisterError) as error:
            client.company("99000001")
        assert str(error.value) == expected and "synthetic-key" not in str(error.value)


def test_checks_are_kept_and_asked_again_with_a_key(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "COMPANY_CHECK_CACHE_FILE", tmp_path / "company_checks.json")
    fetch, fetched = _pages({SITE: f"<html><body>{FOOTER}</body></html>"})
    firms = [("Ornek Stands", SITE), ("Beta Ornek Ltd", ""), ("Ornek Stands", SITE)]
    first = company_register.check_firms(firms, fetch_page=fetch, key="", today=date(2026, 10, 14))
    assert first["Ornek Stands"]["evidence"] == "site" and first["Beta Ornek Ltd"]["evidence"] == "name"
    assert fetched == [SITE]
    company_register.check_firms(firms, fetch_page=fetch, key="", today=date(2026, 11, 1))
    assert fetched == [SITE]
    company_register.check_firms(firms, fetch_page=fetch, key="", today=date(2026, 11, 20))
    assert fetched == [SITE, SITE]
    register = FakeRegister(companies={"99000001": {
        "company_name": "ORNEK EXHIBITIONS LIMITED", "company_status": "active", "type": "ltd",
    }})
    monkeypatch.setattr(company_register, "CompaniesHouse", lambda key: register)
    third = company_register.check_firms(firms, fetch_page=fetch, key="synthetic-key-0000", today=date(2026, 11, 20))
    assert third["Ornek Stands"]["evidence"] == "register_number" and register.calls == [("company", "99000001")]
    summary = company_register.summary(third, register_on=True)
    assert summary == {
        "companies_house": "on", "name": 1, "register_number": 1, "register_name": 0, "site": 0,
        "left_out": 0, "register_errors": 0,
    }
    assert "synthetic-key" not in (tmp_path / "company_checks.json").read_text(encoding="utf-8")


ROWS = [
    {"company": "Alfa Ornek Ltd", "website": "https://alfa-ornek.co.uk/", "email": "info@alfa-ornek.co.uk",
     "phone": "020 7946 0123", "status": "OK_HIGH_CONFIDENCE"},
    {"company": "Ornek Stands", "website": SITE, "email": "hello@ornek-stands.co.uk", "status": "OK_HIGH_CONFIDENCE"},
    {"company": "Gamma Ornek", "website": "https://gamma-ornek.co.uk/", "email": "info@gamma-ornek.co.uk",
     "status": "NEEDS_REVIEW"},
]


def test_british_customer_rows_follow_the_company_check():
    assert customer_file.firm_sites(ROWS) == [
        ("Alfa Ornek Ltd", "https://alfa-ornek.co.uk/"), ("Ornek Stands", SITE), ("Gamma Ornek", ""),
    ]
    checks = {"Alfa Ornek Ltd": {"company": True}, "Ornek Stands": {"company": True}, "Gamma Ornek": {"company": False}}
    rows, left_out, opted_out = customer_file.customer_rows(ROWS, keep_firm=musteri_dosyasi._keep_firm("GB", checks))
    assert [row["Firma"] for row in rows] == ["Alfa Ornek Ltd", "Ornek Stands"]
    assert (left_out, opted_out) == (1, 0)
    assert not musteri_dosyasi._keep_firm("GB")("Alfa Ornek Ltd")
    assert musteri_dosyasi._keep_firm("PL") is legal_form.is_company


def test_a_companies_only_file_says_so(tmp_path):
    rows, _left_out, _opted_out = customer_file.customer_rows(ROWS[:1])
    customer_file.write_customer_file(tmp_path / "list.xlsx", rows, lang="en", companies_only=True, today=date(2026, 10, 14))
    book = load_workbook(tmp_path / "list.xlsx", read_only=True)
    info = [row[0] for row in book["Information"].iter_rows(values_only=True)]
    book.close()
    assert customer_file.COMPANIES_ONLY_LINE["en"] in info


def test_the_key_tool_asks_for_the_companies_house_key():
    assert ("companies_house", "Companies House (ücretsiz, İngiltere şirket sicili)") in api_anahtarlari.PROVIDERS


def test_legal_links_take_their_own_words():
    pages = [{"url": SITE, "html": '<a href="/privacy-policy">Privacy</a><a href="/shop">Shop</a>'}]
    assert tax_identity.legal_links(pages, 3) == []
    assert tax_identity.legal_links(pages, 3, words=("privacy",)) == ["https://www.ornek-stands.co.uk/privacy-policy"]

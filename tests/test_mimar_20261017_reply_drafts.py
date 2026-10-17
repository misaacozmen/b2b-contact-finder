"""Talimat 51: what a reply wants (Jev), a draft in Zoho from the templates, and the morning summary."""

from datetime import date, datetime
from email import message_from_bytes, policy
from email.message import EmailMessage
import io
import json
import urllib.error

from openpyxl import load_workbook
import pytest

import config
from modules import customer_file, intent, outreach, replies, reply_drafts, sales_db
from tools import eposta_gonder


ACCOUNT = "satis@ornek-gonderen.example"
OWNER = "sahip@ornek-kisisel.example"
FIRMS = [
    ("Alfa Ornek", "info@alfa-ornek.com.tr", "<id-alfa@ornek-gonderen.example>"),
    ("Beta Ornek", "info@beta-ornek.com.tr", "<id-beta@ornek-gonderen.example>"),
    ("Gama Ornek", "info@gama-ornek.com.tr", "<id-gama@ornek-gonderen.example>"),
]
TEMPLATES = """# deneme şablonu
[imza]
İyi çalışmalar,
Ornek Gonderen

[ilgileniyor]
Merhaba,

İlginiz için teşekkürler.

{sss}

Hangi fuar?

[sss_bicim]
Excel ve CSV.

[sss_sutunlar]
Firma, stand, site, e-posta, telefon.

[sss_kisi_bilgisi]
Kişi bilgisi yok.

[sss_teslim]
Aynı gün.

[sss_fark]
Site, e-posta ve telefon ekleniyor.

[fiyat_fuarsiz]
Merhaba,

Fiyat fuara göre. Hangi fuar?

[teklif]
Merhaba,

{fuar} listesinde {turk_firma} Türk firması var. En çok {en_cok_tutar} TL.

[teklif_ornek]
Örnek ekte.

[liste_yok]
Merhaba,

{fuar} listesi henüz yayımlanmadı.

[ornek_fuarsiz]
Merhaba,

Hangi fuar için örnek?

[soru_adres]
Merhaba,

Adresiniz sitenizden alındı.

[soru_kvkk]
Merhaba,

Listeler fuar listesine dayanıyor.

[simdi_degil]
Merhaba,

Teşekkürler.

[istemiyor]
Merhaba,

Adresiniz çıkarıldı.

[yanlis_kisi]
Firmanızda bu konuyla sizin ilgilendiğinizi öğrendim.
"""
COLD = "Gönderen: Ornek Gonderen <satis@ornek-gonderen.example>\nKonu: Fuar listeleri\n\nMerhaba {hitap} ekibi,\n\nKısa metin.\n"
QUOTE = "\n\nOrnek Gonderen <satis@ornek-gonderen.example>, 15 Eki 2026 tarihinde şunu yazdı:\n> eski metin\n"


@pytest.fixture(autouse=True)
def files(tmp_path, monkeypatch):
    state = tmp_path / "state"
    for name, file_name in (
        ("OUTREACH_LOG_FILE", "eposta_gonderim_kaydi.jsonl"), ("REPLY_LOG_FILE", "yanit_kaydi.jsonl"),
        ("REPLY_STATE_FILE", "durum.json"), ("REPLY_SETTINGS_FILE", "ayarlar.json"),
        ("REPLY_RUN_LOG_FILE", "calisma.log"), ("OPT_OUT_FILE", "ret_listesi.txt"), ("SALES_DB_FILE", "satis.sqlite3"),
        ("SAVED_API_KEYS_FILE", "api_keys.json"),
    ):
        monkeypatch.setattr(config, name, state / file_name)
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    state.mkdir()
    (tmp_path / "takip").mkdir()
    with config.OUTREACH_LOG_FILE.open("w", encoding="utf-8") as handle:
        for company, address, message_id in FIRMS:
            handle.write(json.dumps({
                "time": "2026-10-15T09:00:00+00:00", "date": "2026-10-15", "tracker": "Ornek_takip.xlsx",
                "company": company, "email": address, "message_id": message_id, "result": outreach.SENT,
            }) + "\n")
    outreach.new_tracker(tmp_path / "takip" / "Ornek_takip.xlsx", [
        {"Firma": company, "E-posta": address} for company, address, _id in FIRMS
    ], "TR")
    (tmp_path / "sablon.txt").write_text(TEMPLATES, encoding="utf-8")
    (tmp_path / "soguk.txt").write_text(COLD, encoding="utf-8")
    contacts = [
        {"Firma": f"Bebek Ornek {index}", "Fuar": "Ornek Bebek Fuarı 2026", "Salon": "1", "Stant": f"A{index}",
         "Web sitesi": f"https://bebek{index}.ornek.example", "E-posta": f"info@bebek{index}.ornek.example" if index % 3 else "",
         "Telefon": "" if index % 3 else "02120000000"}
        for index in range(1, 31)
    ] + [{"Firma": "Bos Ornek", "Fuar": "", "Salon": "", "Stant": "", "Web sitesi": "", "E-posta": "", "Telefon": ""}]
    customer_file.write_customer_file(tmp_path / "bebek.xlsx", contacts, lang="tr")
    sales_db.create(config.SALES_DB_FILE)
    connection = sales_db.connect(config.SALES_DB_FILE)
    sales_db.save_fair(connection, ad="Ornek Bebek Fuarı 2026", diger_adlar="bebek fuarı", turk_firma=85,
                       liste_dosyasi="", iletisim_dosyasi=str(tmp_path / "bebek.xlsx"), durum="hazir")
    sales_db.save_fair(connection, ad="Ornek Sera Fuarı 2026", diger_adlar="sera fuarı", turk_firma=0,
                       liste_dosyasi="", iletisim_dosyasi="", durum="yok")
    connection.close()
    return tmp_path


def _mail(sender: str, body: str, **headers) -> bytes:
    message = EmailMessage()
    message["From"] = sender
    message["To"] = ACCOUNT
    message["Subject"] = headers.pop("subject", "Re: Fuar listeleri")
    message["Date"] = headers.pop("date", "Fri, 16 Oct 2026 10:00:00 +0300")
    message["Message-ID"] = headers.pop("message_id", "<cevap@alfa-ornek.com.tr>")
    for name, value in headers.items():
        message[name.replace("_", "-")] = value
    message.set_content(body)
    return message.as_bytes()


ANSWERS = {
    "ne kadar": {"niyet": ("fiyat", 0.95), "fiyat": 0.9, "ornek": 0.1, "fuar": ("f1", 0.9), "sss": "yok"},
    "örnek": {"niyet": ("ornek", 0.92), "fiyat": 0.2, "ornek": 0.9, "fuar": ("f1", 0.9), "sss": "yok"},
    "Ödemeyi": {"niyet": ("siparis", 0.93), "fiyat": 0.2, "ornek": 0.1, "fuar": ("yok", 0.9), "sss": "yok"},
    "İlgilenmiyoruz": {"niyet": ("olumsuz", 0.96), "fiyat": 0.0, "ornek": 0.0, "fuar": ("yok", 0.99), "sss": "yok"},
    "sera": {"niyet": ("ilgileniyor", 0.9), "fiyat": 0.3, "ornek": 0.1, "fuar": ("f2", 0.88), "sss": "yok"},
    "Excel": {"niyet": ("ilgileniyor", 0.9), "fiyat": 0.1, "ornek": 0.1, "fuar": ("yok", 0.9), "sss": "bicim"},
    "satinalma": {"niyet": ("yanlis_kisi", 0.9), "fiyat": 0.0, "ornek": 0.0, "fuar": ("yok", 0.9), "sss": "yok"},
    "mobilya": {"niyet": ("fiyat", 0.9), "fiyat": 0.9, "ornek": 0.0, "fuar": ("katalog_disi", 0.8), "sss": "yok"},
}


class FakeJev:
    def __init__(self):
        self.states = []

    def __call__(self, request, timeout):
        body = json.loads(request.data.decode("utf-8"))
        self.states.append(body["state"])
        assert request.headers["Authorization"] == "Bearer synthetic-jev-key"
        answer = next(value for word, value in ANSWERS.items() if word in body["state"])
        payload = {"model": "jev-test", "answers": {
            "niyet": {"type": "choice", "choice": answer["niyet"][0], "confidence": answer["niyet"][1]},
            "fiyat_istiyor": {"type": "noul", "noul": answer["fiyat"]},
            "ornek_istiyor": {"type": "noul", "noul": answer["ornek"]},
            "fuar": {"type": "choice", "choice": answer["fuar"][0], "confidence": answer["fuar"][1]},
            "sss": {"type": "choice", "choice": answer["sss"], "confidence": 0.9},
        }, "usage": {"input_tokens": 10}}
        return io.BytesIO(json.dumps(payload).encode("utf-8"))


class FakeImap:
    def __init__(self, folders):
        self.folders = folders
        self.current = None
        self.appended = []

    def login(self, user, password):
        pass

    def select(self, folder, readonly=False):
        assert readonly is True
        name = folder.strip('"')
        if name not in self.folders:
            return "NO", [b"no such folder"]
        self.current = name
        return "OK", [str(len(self.folders[name])).encode()]

    def response(self, code):
        return code, [b"7"]

    def uid(self, command, *args):
        box = self.folders[self.current]
        if command == "SEARCH":
            if args[1] == "UID":
                start = int(args[2].split(":")[0])
                uids = [uid for uid in box if uid >= start] or ([max(box)] if box else [])
            else:
                uids = list(box)
            return "OK", [" ".join(str(uid) for uid in uids).encode()]
        uid = int(args[0])
        return "OK", [(f"{uid} (UID {uid} BODY[] {{1}}".encode(), box[uid]), b")"]

    def append(self, mailbox, flags, date_time, message):
        assert mailbox == '"Drafts"' and "\\Draft" in flags
        self.appended.append(message_from_bytes(message, policy=policy.default))
        return "OK", [b"[APPENDUID 1 1] Success"]

    def logout(self):
        pass


class FakeSmtp:
    sent = []

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def login(self, user, password):
        pass

    def send_message(self, message):
        FakeSmtp.sent.append(message)


@pytest.fixture
def smtp():
    FakeSmtp.sent = []
    return FakeSmtp


def _drafting(files, jev=None, key="synthetic-jev-key") -> reply_drafts.Drafting:
    return reply_drafts.Drafting(
        templates=reply_drafts.load_templates(files / "sablon.txt"), sender_name="Ornek Gonderen",
        since=date(2026, 10, 16), db_path=config.SALES_DB_FILE, jev_key=key,
        cold_template=outreach.load_template(files / "soguk.txt"), opener=jev or FakeJev(),
    )


def _check(files, imap, drafting, now=datetime(2026, 10, 16, 7, 30)):
    return replies.check(account=ACCOUNT, password="synthetic-password", notify_to=OWNER, tracker_dir=files / "takip",
                         imap_factory=lambda: imap, smtp_factory=FakeSmtp, drafting=drafting, now=now)


def _decision(niyet="fiyat", guven=0.9, fiyat=False, ornek=False, fuar="yok", fuar_guven=0.9, sss="yok"):
    return intent.Decision(niyet, guven, fiyat, ornek, fuar, fuar_guven, sss)


def _fairs():
    connection = sales_db.connect(config.SALES_DB_FILE)
    try:
        return sales_db.fairs(connection)
    finally:
        connection.close()


def test_personal_data_is_removed_before_the_model_sees_the_reply():
    text = ("Merhaba Ornek Bey,\nListe ne kadar? info@alfa-ornek.com.tr ya da 0212 555 00 00, www.alfa-ornek.com.tr\n"
            "\nSaygılarımla\nAli Ornekoglu\nSatın Alma")
    assert intent.anonymize(text) == "Liste ne kadar? [e-posta] ya da [telefon], [web adresi]"


def test_the_model_request_and_answer(files):
    jev = FakeJev()
    decision = intent.decide("Bebek fuarı listesi ne kadar?", _fairs(), key="synthetic-jev-key", opener=jev)
    assert decision == _decision("fiyat", 0.95, True, False, "f1", 0.9, "yok")
    body = intent.request_body("x", _fairs())
    assert set(body["questions"]) == {"niyet", "fiyat_istiyor", "ornek_istiyor", "fuar", "sss"}
    assert set(body["questions"]["fuar"]["criteria"]) == {"f1", "f2", "katalog_disi", "yok"}
    with pytest.raises(intent.IntentError, match="anahtarı"):
        intent.decide("x", [], key="")

    def refused(request, timeout):
        raise urllib.error.HTTPError(request.full_url, 402, "Payment Required", {}, None)

    with pytest.raises(intent.IntentError, match="402"):
        intent.decide("x", [], key="synthetic-jev-key", opener=refused)


def test_templates_must_be_complete(tmp_path):
    path = tmp_path / "eksik.txt"
    path.write_text("[imza]\nOrnek\n", encoding="utf-8")
    with pytest.raises(reply_drafts.ReplyDraftError, match="eksik bölüm"):
        reply_drafts.load_templates(path)
    path.write_text(TEMPLATES.replace("{fuar} listesi henüz", "{fuarr} listesi henüz"), encoding="utf-8")
    with pytest.raises(reply_drafts.ReplyDraftError, match="fuarr"):
        reply_drafts.load_templates(path)


@pytest.mark.parametrize("count, rate", [(85, 6.0), (150, 6.0), (151, 5.5), (400, 5.5), (1000, 5.0), (1001, 4.0)])
def test_price_per_firm_by_the_number_of_turkish_exhibitors(count, rate):
    assert reply_drafts.rate_for(count) == rate
    assert reply_drafts.quote({"turk_firma": count})["en_cok"] == count * rate


def test_money_is_written_the_turkish_way():
    assert (reply_drafts.money(510), reply_drafts.money(1017.5), reply_drafts.money(5000)) == ("510", "1.017,50", "5.000")


@pytest.mark.parametrize("decision, template, reason", [
    (None, "", "niyet belirlenemedi"),
    (_decision(guven=0.4), "", "niyet belirsiz"),
    (_decision("siparis"), "", "sipariş, ödeme ya da pazarlık"),
    (_decision("olumsuz"), "", "olumsuz yanıt; firma kapatıldı"),
    (_decision("soru_adres"), "soru_adres", ""),
    (_decision("soru_kvkk"), "soru_kvkk", ""),
    (_decision("simdi_degil"), "simdi_degil", ""),
    (_decision("fiyat", fuar="katalog_disi"), "", "katalogda olmayan fuar"),
    (_decision("fiyat", fuar="f1"), "teklif", ""),
    (_decision("fiyat", fuar="f1", fuar_guven=0.3), "", "fuar belirsiz"),
    (_decision("ilgileniyor", fuar="f2"), "liste_yok", ""),
    (_decision("ornek"), "ornek_fuarsiz", ""),
    (_decision("fiyat"), "fiyat_fuarsiz", ""),
    (_decision("ilgileniyor", fiyat=True), "fiyat_fuarsiz", ""),
    (_decision("ilgileniyor"), "ilgileniyor", ""),
])
def test_which_template_answers_a_reply(files, decision, template, reason):
    assert reply_drafts.plan({"kind": replies.REPLY}, decision, _fairs()) == (template, reason)
    assert reply_drafts.plan({"kind": replies.OPT_OUT}, None, _fairs()) == ("istemiyor", "")


def _entry(text, **extra):
    return {"key": "INBOX:7:1", "kind": replies.REPLY, "from": "info@alfa-ornek.com.tr", "email": "info@alfa-ornek.com.tr",
            "company": "Alfa Ornek", "tracker": "Ornek_takip.xlsx", "date": "2026-10-16",
            "received": "2026-10-16T10:00+03:00", "subject": "Re: Fuar listeleri", "text": text,
            "message_id": "<cevap@alfa-ornek.com.tr>", "references": "<id-alfa@ornek-gonderen.example>", **extra}


def test_a_quote_draft_answers_in_the_same_thread(files, tmp_path):
    drafting = _drafting(files)
    drafting.attachments_dir = tmp_path
    message, template, _ = reply_drafts.build(_entry("Liste ne kadar?"), _decision("fiyat", fuar="f1"), _fairs(),
                                              drafting=drafting, account=ACCOUNT)
    assert template == "teklif"
    assert message["To"] == "info@alfa-ornek.com.tr" and message["Subject"] == "Re: Fuar listeleri"
    assert message["In-Reply-To"] == "<cevap@alfa-ornek.com.tr>"
    assert message["References"] == "<id-alfa@ornek-gonderen.example> <cevap@alfa-ornek.com.tr>"
    body = message.get_body().get_content()
    assert "Ornek Bebek Fuarı 2026 listesinde 85 Türk firması var. En çok 510 TL." in body
    assert "İyi çalışmalar,\nOrnek Gonderen" in body and "> Liste ne kadar?" in body
    assert "Örnek ekte." not in body and not list(message.iter_attachments())


def test_a_sample_goes_only_when_asked(files, tmp_path):
    drafting = _drafting(files)
    drafting.attachments_dir = tmp_path
    message, _, _ = reply_drafts.build(_entry("Örnek görebilir miyiz?"), _decision("ornek", ornek=True, fuar="f1"),
                                       _fairs(), drafting=drafting, account=ACCOUNT)
    assert "Örnek ekte." in message.get_body().get_content()
    attachment = next(message.iter_attachments())
    sheet = load_workbook(io.BytesIO(attachment.get_content())).worksheets[0]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 20 and all(row[5] or row[6] for row in rows)


def test_a_faq_answer_and_a_fair_without_a_list(files):
    drafting = _drafting(files)
    message, _, _ = reply_drafts.build(_entry("Excel mi?"), _decision("ilgileniyor", sss="bicim"), _fairs(),
                                       drafting=drafting, account=ACCOUNT)
    assert "İlginiz için teşekkürler.\n\nExcel ve CSV.\n\nHangi fuar?" in message.get_body().get_content()
    message, _, _ = reply_drafts.build(_entry("Sera fuarı?"), _decision("ilgileniyor", fuar="f2"), _fairs(),
                                       drafting=drafting, account=ACCOUNT)
    assert "Ornek Sera Fuarı 2026 listesi henüz yayımlanmadı." in message.get_body().get_content()


def test_a_new_generic_address_gets_our_first_e_mail(files):
    drafting = _drafting(files)
    message, template, _ = reply_drafts.build(
        _entry("Lütfen satinalma@alfa-ornek.com.tr adresine yazın."), _decision("yanlis_kisi"), _fairs(),
        drafting=drafting, account=ACCOUNT,
    )
    assert template == "yanlis_kisi" and message["To"] == "satinalma@alfa-ornek.com.tr"
    assert message["Subject"] == "Fuar listeleri" and message["In-Reply-To"] is None
    assert message.get_body().get_content().startswith(
        "Merhaba Alfa Ornek ekibi,\n\nFirmanızda bu konuyla sizin ilgilendiğinizi öğrendim.\n\nKısa metin.")
    message, _, reason = reply_drafts.build(_entry("Lütfen mehmet.ornekoglu@alfa-ornek.com.tr adresine yazın."),
                                            _decision("yanlis_kisi"), _fairs(), drafting=drafting, account=ACCOUNT)
    assert message is None and reason == "yeni adres yok ya da kişiye ait"


def test_a_check_writes_drafts_notices_and_one_summary(files, smtp):
    jev = FakeJev()
    drafting = _drafting(files, jev)
    imap = FakeImap({
        "INBOX": {
            1: _mail("Alfa Ornek <info@alfa-ornek.com.tr>",
                     "Merhaba,\nBebek fuarı listesi ne kadar?\n\nSaygılarımla\nAli Ornekoglu\n0212 555 00 00" + QUOTE,
                     In_Reply_To="<id-alfa@ornek-gonderen.example>"),
            2: _mail("Beta <info@beta-ornek.com.tr>", "Ödemeyi nasıl yapalım?" + QUOTE, message_id="<b@beta.example>"),
            3: _mail("Gama <info@gama-ornek.com.tr>", "İlgilenmiyoruz." + QUOTE, message_id="<g@gama.example>"),
            4: _mail("Alfa Ornek <info@alfa-ornek.com.tr>", "Eski yanıt, ne kadar?", message_id="<eski@alfa.example>",
                     date="Thu, 15 Oct 2026 10:00:00 +0300"),
        },
        "Sent": {},
    })
    result = _check(files, imap, drafting)
    assert (result["taslak"], result["size kalan"], result["özet"]) == (1, 2, 0)
    assert len(imap.appended) == 1 and imap.appended[0]["To"] == "info@alfa-ornek.com.tr"
    assert not any("0212" in state or "Ornekoglu" in state for state in jev.states)
    connection = sales_db.connect(config.SALES_DB_FILE)
    outcomes = {row["eposta"]: (row["niyet"], row["sonuc"]) for row in connection.execute("SELECT * FROM yanitlar")}
    stages = {row["eposta"]: row["asama"] for row in connection.execute("SELECT * FROM firmalar")}
    connection.close()
    assert outcomes == {
        "info@alfa-ornek.com.tr": ("fiyat", sales_db.DRAFT), "info@beta-ornek.com.tr": ("siparis", sales_db.USER),
        "info@gama-ornek.com.tr": ("olumsuz", sales_db.CLOSED),
    }
    assert stages["info@gama-ornek.com.tr"] == "kapandı"
    subjects = sorted(message["Subject"] for message in smtp.sent)
    assert "Taslak hazır (fiyat): Alfa Ornek" in subjects
    assert "Size kalan yanıt (sipariş, ödeme ya da pazarlık): Beta Ornek" in subjects

    smtp.sent.clear()
    answer = EmailMessage()
    answer["From"], answer["To"], answer["Subject"] = ACCOUNT, "info@alfa-ornek.com.tr", "Re: Fuar listeleri"
    answer["Date"], answer["In-Reply-To"] = "Fri, 16 Oct 2026 12:00:00 +0300", "<cevap@alfa-ornek.com.tr>"
    answer.set_content("Yanıtımız")
    imap.folders["Sent"][1] = answer.as_bytes()
    result = _check(files, imap, drafting, now=datetime(2026, 10, 17, 8, 5))
    assert (result["yanıtlanan"], result["özet"], result["taslak"]) == (1, 1, 0)
    summary = smtp.sent[-1]
    assert summary["Subject"] == "Günlük özet — 17.10.2026" and summary["To"] == OWNER
    text = summary.get_content()
    assert "Zoho'da gönderilmeyi bekleyen taslak: 0" in text
    assert "Taslak yazılmayan, sizin yanıtlayacağınız: 1" in text and "Beta Ornek — sipariş" in text
    assert "Gönderdiğiniz yanıt: 1" in text and "Günlük gönderim hakkı: 20" in text
    assert _check(files, imap, drafting, now=datetime(2026, 10, 17, 8, 35))["özet"] == 0


def test_without_the_model_the_reply_still_reaches_the_user(files, smtp):
    drafting = _drafting(files, key="")
    imap = FakeImap({"INBOX": {1: _mail("Alfa <info@alfa-ornek.com.tr>", "Bebek fuarı listesi ne kadar?")}, "Sent": {}})
    result = _check(files, imap, drafting)
    assert (result["niyet hatası"], result["size kalan"], result["bildirim"]) == (1, 1, 1)
    assert imap.appended == [] and smtp.sent[0]["Subject"] == "Size kalan yanıt (niyet belirlenemedi): Alfa Ornek"
    assert "HATA: Jev anahtarı kayıtlı değil" in config.REPLY_RUN_LOG_FILE.read_text(encoding="utf-8")


def test_the_tool_sets_up_drafts_and_fairs(files, capsys):
    assert eposta_gonder.main(["yanit-ayar", "--hesap", ACCOUNT, "--bildirim", OWNER, "--klasor", str(files / "takip")]) == 0
    assert eposta_gonder.main(["taslak-ayar", "--yanit-sablon", str(files / "sablon.txt"),
                               "--soguk-sablon", str(files / "soguk.txt"), "--ornek", str(files / "bebek.xlsx")]) == 0
    settings = replies.load_settings()
    assert settings["taslak_baslangic"] and settings["ornek"].endswith("bebek.xlsx")
    assert eposta_gonder.main(["fuar-ekle", "--ad", "Ornek Yapi Fuari 2026", "--turk", "420",
                               "--iletisim", str(files / "bebek.xlsx")]) == 0
    capsys.readouterr()
    assert eposta_gonder.main(["fuarlar"]) == 0
    assert "Ornek Yapi Fuari 2026 | durum hazir | Türk firma 420 | en çok 2.100 TL" in capsys.readouterr().out
    assert eposta_gonder.main(["ozet"]) == 0
    assert "SİZİ BEKLEYENLER" in capsys.readouterr().out

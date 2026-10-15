"""Talimat 46: prospect lists of firms that serve fair exhibitors, found by web search."""

from openpyxl import Workbook, load_workbook
import pytest

import config
from modules import country_profile, prospect_search
from tools import hedef_bul


@pytest.fixture(autouse=True)
def turkey():
    country_profile.apply("TR")
    yield
    country_profile.apply("TR")


@pytest.mark.parametrize("title, website, expected", [
    ("Fuar Taşımacılığı — Fuar Lojistiği — Odak Ornek", "https://www.odakornek.com/", "Odak Ornek"),
    ("Ornek Catering | İstanbul Fuar Catering", "https://ornekcatering.com.tr/", "Ornek Catering"),
    ("Kiralık Fuar Mobilyası Kiralama; Fuar ve Etkinlikleriniz İçin", "https://bizdenornek.com/", "Bizdenornek"),
    ("", "https://ornekfoto.com/", "Ornekfoto"),
])
def test_firm_name_from_the_page_title(title, website, expected):
    assert prospect_search.firm_name(title, website) == expected


@pytest.mark.parametrize("domain, title, categories, reason", [
    ("ornekcatering.com.tr", "Ornek Catering | Fuar Catering", ["Catering"], ""),
    ("armut.com", "En İyi 40 İstanbul Fuar Catering Firması | Armut", ["Catering"], "rehber, pazar yeri ya da büyük site"),
    ("ornekliste.com", "En İyi 10 Fuar Hostes Ajansı", ["Hostes"], "liste ya da rehber sayfası"),
    ("ornekfuari.com", "Uluslararası Ornek Gıda Fuarı", ["Lojistik"], "fuar ya da etkinlik sayfası"),
    ("ornekexpo.com", "Seramik Fuarı | ORNEK EXPO İstanbul 2026", ["Lojistik"], "fuar ya da etkinlik sayfası"),
    ("ornekhaber.com", "Lojistik sektöründe yeni dönem", ["Lojistik"], "haber sitesi"),
    ("ornek.de", "Ornek Messebau", ["Lojistik"], "yabancı uzantı"),
    ("ornek.gov.tr", "Ornek Müdürlüğü", ["Lojistik"], "kamu sitesi"),
    ("ornekhub.com", "Ornek Hub", ["Hostes", "Catering", "Lojistik"], "çok kategoride çıktı"),
    ("eskiornek.com", "Eski Ornek", ["Catering"], "zaten listede"),
    ("ornektusid.com", "ORNEK İstanbul Fuarını Keşfedin", ["Catering"], "fuar ya da etkinlik sayfası"),
    ("ornekfoto.com", "İzmir Düğün Fotoğrafçısı | Ornek Foto", ["Fotoğraf ve video"], "düğün ya da özel gün sayfası"),
    ("ornekder.org.tr", "Uluslararası Ornek Nakliyeciler Derneği", ["Lojistik"], "dernek ya da oda"),
    ("ornekbaski.com", "Folyo Baskı Blog, Rehberler ve İpuçları", ["Promosyon ve baskı"], "liste ya da rehber sayfası"),
    ("ornek-sayfa.web.app", "Fuar Mobilya Kiralama", ["Mobilya ve ekipman kiralama"], "rehber, pazar yeri ya da büyük site"),
    ("ornekfuar.com", "Ornek Fuar İzmir Catering", ["Catering"], ""),
])
def test_rows_that_are_no_service_firm(domain, title, categories, reason):
    row = {"domain": domain, "title": title, "categories": categories}
    assert prospect_search.skip_reason(row, {"eskiornek.com"}) == reason


def test_queries_cover_every_category_and_city():
    pairs = prospect_search.queries("TR")
    assert len(pairs) == 6 * 2 * 4
    assert ("Lojistik", "fuar lojistiği firması İzmir") in pairs
    assert prospect_search.queries("GB", ("Logistics",))[0] == ("Logistics", "exhibition logistics company London")


HOME = "<html><body><h1>{title}</h1><p>{text}</p></body></html>"
PAGES = {
    "https://ornekcatering.com.tr/": HOME.format(title="Ornek Catering", text="Fuar ve etkinlik catering, kokteyl ikram."),
    "https://www.ornekcatering.com/": HOME.format(title="Ornek Catering", text="Fuar catering."),
    "https://ornekyemek.com/": HOME.format(title="Ornek Yemek", text="Ev yemekleri ve paket servis."),
    "https://ornekfoto.com/": HOME.format(title="Ornek Foto", text="Düğün albümleri."),
}


def fake_search(query):
    if "catering" in query:
        return [
            {"title": "Ornek Catering | Fuar Catering", "href": "https://ornekcatering.com.tr/hakkimizda"},
            {"title": "Ornek Catering", "href": "https://www.ornekcatering.com/"},
            {"title": "Ornek Yemek", "href": "https://ornekyemek.com/"},
            {"title": "En İyi 40 Fuar Catering Firması", "href": "https://armut.com/catering"},
        ]
    if "fotoğraf" in query:
        return [{"title": "Ornek Foto", "href": "https://ornekfoto.com/"}]
    return []


def test_search_keeps_one_row_per_firm_and_checks_its_home_page():
    lines = []
    rows, stats = prospect_search.find(
        "TR", search=fake_search, fetch_page=PAGES.get, sleep=lambda _seconds: None,
        categories=("Catering", "Fotoğraf ve video"), report=lines.append,
    )
    by_domain = {row["domain"]: row for row in rows}
    assert stats["queries"] == 16 and stats["answered"] == 12
    assert by_domain["ornekcatering.com.tr"]["ok"] and by_domain["ornekcatering.com.tr"]["website"] == "https://ornekcatering.com.tr/"
    assert by_domain["ornekcatering.com"]["reason"] == "aynı firma, başka uzantı"
    assert by_domain["ornekyemek.com"]["reason"] == "sitede fuar ya da etkinlik geçmiyor"
    assert by_domain["ornekfoto.com"]["reason"] == "sitede bu hizmet geçmiyor"
    assert by_domain["armut.com"]["reason"] == "rehber, pazar yeri ya da büyük site"
    assert [row["domain"] for row in rows if row["ok"]] == ["ornekcatering.com.tr"]
    assert stats["sites_checked"] == 3
    assert lines[-1] == "site kontrolü 3/3"


def test_site_that_does_not_open():
    row = {"website": "https://ornekyok.com/", "categories": ["Catering"]}
    assert prospect_search.site_reason(row, lambda _url: "") == "site açılmadı"


def test_domains_of_earlier_lists(tmp_path):
    path = tmp_path / "takip.xlsx"
    book = Workbook()
    book.active.append(["Firma", "Web sitesi", "E-posta"])
    book.active.append(["Eski Ornek", "https://www.eskiornek.com.tr/", "info@eskiornek.com"])
    book.active.append(["Gmail Ornek", "", "ornek@gmail.com"])
    book.save(path)
    assert prospect_search.domains_in([path]) == {"eskiornek.com.tr", "eskiornek.com"}


def test_candidate_and_input_workbooks(tmp_path):
    rows, _stats = prospect_search.find(
        "TR", search=fake_search, fetch_page=PAGES.get, sleep=lambda _seconds: None,
        categories=("Catering",), report=lambda _line: None,
    )
    review, run_input = tmp_path / "adaylar.xlsx", tmp_path / "girdi.xlsx"
    prospect_search.write_review(review, rows)
    assert prospect_search.write_input(run_input, rows, "TR") == 1
    sheet = load_workbook(review).active
    assert [cell.value for cell in sheet[1]] == list(prospect_search.REVIEW_COLUMNS)
    assert sheet.cell(row=2, column=1).value == "evet" and sheet.cell(row=2, column=4).value == "Ornek Catering"
    assert [list(row) for row in load_workbook(run_input).active.iter_rows(values_only=True)] == [
        ["company", "website", "country"], ["Ornek Catering", "https://ornekcatering.com.tr/", "Türkiye"],
    ]


def test_command_writes_both_files_and_refuses_to_overwrite(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(prospect_search, "free_search", lambda query, region="": fake_search(query))
    from modules import crawler

    monkeypatch.setattr(crawler, "fetch_page", PAGES.get)
    monkeypatch.setattr(prospect_search, "QUERY_PAUSE_SEC", 0)
    review, run_input = tmp_path / "adaylar.xlsx", tmp_path / "girdi.xlsx"
    assert hedef_bul.main(["--ulke", "TR", "--aday", str(review), "--girdi", str(run_input), "--kategori", "Catering"]) == 0
    out = capsys.readouterr().out
    assert "Uygun firma: 1" in out and review.exists() and run_input.exists()
    assert hedef_bul.main(["--ulke", "TR", "--aday", str(review), "--girdi", str(run_input)]) == 1
    assert "zaten var" in capsys.readouterr().out
    assert hedef_bul.main(["--ulke", "TR", "--aday", str(tmp_path / "a.xlsx"), "--girdi", str(tmp_path / "b.xlsx"),
                           "--kategori", "Uzay"]) == 1


def test_command_stops_when_no_engine_answers(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(prospect_search, "free_search", lambda query, region="": [])
    monkeypatch.setattr(prospect_search, "QUERY_PAUSE_SEC", 0)
    assert hedef_bul.main(["--ulke", "GB", "--aday", str(tmp_path / "a.xlsx"), "--girdi", str(tmp_path / "b.xlsx"),
                           "--kategori", "Logistics"]) == 1
    assert "sonuç vermedi" in capsys.readouterr().out
    assert not (tmp_path / "a.xlsx").exists()
    assert config.TARGET_COUNTRY == "GB"

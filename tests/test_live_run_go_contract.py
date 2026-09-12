from __future__ import annotations

import hashlib
import json
import inspect
import sqlite3
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

import pytest
import requests
from openpyxl import Workbook

import config
from modules import checkpoint, company_resolvers, google_places, hunter, linkedin_company, pipeline_runner, runtime, run_context, search


def _legacy_complete_fixture(tmp_path: Path) -> tuple[Path, str, Path]:
    run_id = "e" * 64
    root = tmp_path / "runs" / run_id
    (root / "state").mkdir(parents=True)
    artifact_root = root / "output" / "artifacts"
    artifact_root.mkdir(parents=True)
    db = root / "state" / "progress.sqlite3"
    with patch.object(config, "PROGRESS_DB_FILE", db):
        checkpoint.seed_recovered_run(
            path=db, run_id=run_id, input_hash="h", run_signature="s",
            context={"phase": "FREE", "lineage": {"type": "legacy_recovery"}},
            budgets={provider: 0 for provider in checkpoint.CANONICAL_PROVIDERS},
            items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "DONE", "quarantine_state": "LEGACY", "quarantine_status": "PROVISIONAL", "publication_blockers": "legacy_recovery_provisional"}],
            results=[{"item_index": 0, "payload": '{"company":"Legacy","source_record_id":"input:0","publication_eligible":false,"quarantine_state":"LEGACY","quarantine_status":"PROVISIONAL","publication_blockers":"legacy_recovery_provisional"}'}],
        )
        checkpoint.transition_phase(run_id, "FINALIZING", expected_count=1)
        checkpoint.begin_finalization_intent(run_id=run_id, generation="g", input_snapshot_sha256="r", result_snapshot_sha256="r")
        body = b"legacy-artifact"
        file_hash = hashlib.sha256(body).hexdigest()
        artifact_hash = hashlib.sha256(f"file.txt:{file_hash}\n".encode()).hexdigest()
        artifact_dir = artifact_root / artifact_hash
        artifact_dir.mkdir()
        (artifact_dir / "file.txt").write_bytes(body)
        manifest = root / "manifest.json"
        manifest.write_text(json.dumps({
            "complete": True, "phase": "COMPLETE", "artifact_set_sha256": artifact_hash,
            "files": {"file.txt": {"sha256": file_hash, "bytes": len(body)}},
        }), encoding="utf-8")
        with sqlite3.connect(db) as connection:
            connection.execute(
                "UPDATE finalization_intent SET status='SKIPPED_LEGACY_NO_PLAN', finalization_schema_version=0, artifact_set_sha256=?, manifest_sha256=? WHERE run_id=?",
                (artifact_hash, checkpoint.file_hash(manifest), run_id),
            )
            connection.execute("UPDATE runs SET phase='COMPLETE' WHERE run_id=?", (run_id,))
            connection.commit()
    return root, run_id, db


def test_non_complete_complete_manifest_is_rejected_immediately():
    with pytest.raises(RuntimeError, match="exactly COMPLETE"):
        pipeline_runner.require_complete_manifest_phase({"complete": True, "phase": "FINALIZING"})


def test_schema_zero_legacy_terminal_pass_pending_and_provider_fail(tmp_path: Path):
    root, run_id, db = _legacy_complete_fixture(tmp_path)
    with patch.object(config, "PROGRESS_DB_FILE", db):
        assert checkpoint.validate_finalization_contract(run_id, root)["memory_plan_count"] == 0
        with sqlite3.connect(db) as connection:
            connection.execute("UPDATE run_items SET free_state='PENDING' WHERE run_id=?", (run_id,))
            connection.commit()
        with pytest.raises(RuntimeError, match="scheduler/provider state is not terminal"):
            checkpoint.validate_finalization_contract(run_id, root)
        with sqlite3.connect(db) as connection:
            connection.execute("UPDATE run_items SET free_state='DONE' WHERE run_id=?", (run_id,))
            connection.execute("INSERT INTO provider_calls(run_id,call_id,provider,item_index,phase,request_fingerprint,state,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)", (run_id, "reserved", "llm", 0, "PAID", "fp", "RESERVED", "now", "now"))
            connection.commit()
        with pytest.raises(RuntimeError, match="scheduler/provider state is not terminal"):
            checkpoint.validate_finalization_contract(run_id, root)


def test_provider_fold_reserved_completed_empty_retry_success_and_duplicate_inheritance():
    assert pipeline_runner.paid_attempt_result({"provider_results": [{"result_state": "RESERVED", "call_ids": ["a"]}, {"result_state": "COMPLETED", "call_ids": ["a"]}]}) == "COMPLETED"
    assert pipeline_runner.paid_attempt_result({"provider_results": [{"result_state": "EMPTY", "call_ids": ["e"]}]}) == "COMPLETED"
    assert pipeline_runner.paid_attempt_result({"provider_results": [{"result_state": "FAILED", "call_ids": ["f"]}, {"result_state": "COMPLETED", "call_ids": ["s"]}]}) == "COMPLETED"
    assert pipeline_runner.paid_attempt_result({"provider_results": [{"result_state": "BLOCKED_BUDGET"}, {"result_state": "COMPLETED", "call_ids": ["s"]}]}) == "COMPLETED"
    inherited = runtime.rejected_provider_result(runtime.Reservation(False, "llm", "arbiter", 0, "PAID", call_id="done", reason="duplicate_request", inherited_state="DONE"))
    assert inherited.result_state == "COMPLETED" and inherited.call_ids == ("done",)


def test_free_scheduler_classifier_retries_once_and_never_promotes_failures():
    assert pipeline_runner.classify_scheduler_states({"status": "SEARCH_FAILED"}, attempt_number=1) == {"free_state": "PENDING", "paid_state": "NOT_REQUIRED", "paid_required": False}
    assert pipeline_runner.classify_scheduler_states({"status": "PROCESSING_FAILED"}, attempt_number=2) == {"free_state": "FAILED", "paid_state": "PENDING", "paid_required": True}
    assert pipeline_runner.classify_scheduler_states({"status": "OK_HIGH_CONFIDENCE", "publication_eligible": False}, attempt_number=1) == {"free_state": "DONE", "paid_state": "PENDING", "paid_required": True}
    assert pipeline_runner.classify_scheduler_states({"status": "OK_HIGH_CONFIDENCE", "publication_eligible": True}, attempt_number=2) == {"free_state": "DONE", "paid_state": "NOT_REQUIRED", "paid_required": False}


def test_free_query_quota_is_scoped_to_run_and_item(tmp_path: Path):
    db = tmp_path / "state.sqlite3"
    with patch.object(config, "PROGRESS_DB_FILE", db):
        checkpoint.initialize_schema(db)
        checkpoint.initialize_run(
            run_id="run-a", input_hash="h", run_signature="s", context={"phase": "FREE"},
            budgets={provider: 0 for provider in checkpoint.CANONICAL_PROVIDERS},
            items=[{"item_index": 0, "source_record_id": "a"}, {"item_index": 1, "source_record_id": "b"}],
        )
        runtime.configure_durable_run("run-a", {})
        runtime.set_item_context(0, "free")
        assert sum(runtime.reserve_search_query(0) for _ in range(10)) == 10
        assert not runtime.reserve_search_query(0)
        runtime.set_item_context(1, "free")
        assert runtime.reserve_search_query(0)


def test_handoff_quarantines_all_rows_until_atomic_complete_release(tmp_path: Path):
    db = tmp_path / "state.sqlite3"
    items = [{"item_index": index, "source_record_id": f"input:{index}", "free_state": "FAILED" if index == 0 else "DONE", "paid_required": index == 0, "paid_state": "PENDING" if index == 0 else "NOT_REQUIRED"} for index in range(159)]
    results = [{"item_index": index, "payload": json.dumps({"company": f"C{index}", "source_record_id": f"input:{index}", "publication_eligible": True, "website": "https://c.example" if index == 0 else "", "email": "ops@c.example" if index == 0 else ""}, separators=(",", ":"))} for index in range(159)]
    with patch.object(config, "PROGRESS_DB_FILE", db):
        checkpoint.seed_recovered_run(path=db, run_id="handoff", input_hash="h", run_signature="s", context={"phase": "FREE"}, budgets={provider: 0 for provider in checkpoint.CANONICAL_PROVIDERS}, items=items, results=results)
        assert checkpoint.mark_handoff_pending(run_id="handoff", expected_count=159) == {"items": 159, "quarantined": 159}
        rows = checkpoint.load_results_by_id("handoff")
    assert len(rows) == 159 and all(row["publication_eligible"] is False and row["quarantine_state"] == "HANDOFF_PENDING" for row in rows.values())
    assert rows[0]["website"] == "https://c.example" and rows[0]["email"] == "ops@c.example"


def test_real_pipeline_handoff_accepts_failed_free_item_as_terminal(tmp_path: Path):
    input_file = tmp_path / "remaining_159.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["company", "source_record_id"])
    for index in range(159):
        sheet.append([f"COMPANY {index:03d}", f"input:{index}"])
    workbook.save(input_file)
    workbook.close()

    attempts: dict[int, int] = {}

    def worker(index, company, _logger, _website, record):
        attempts[index] = attempts.get(index, 0) + 1
        if index == 0:
            return index, {
                "company": company, "source_record_id": record["source_record_id"],
                "status": "SEARCH_FAILED", "reason": "transient exhausted",
                "publication_eligible": False,
            }
        return index, {
            "company": company, "source_record_id": record["source_record_id"],
            "status": "OK_HIGH_CONFIDENCE", "publication_eligible": True,
        }

    runs_dir = tmp_path / "runs"
    with patch.object(config, "RUNS_DIR", runs_dir), patch.object(
        search, "preflight_source_profiles", return_value=[]
    ):
        result = pipeline_runner.run_pipeline(
            input_file, allow_paid=False,
            process_company_fn=worker,
            write_outputs_fn=lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("handoff must not publish")),
            set_output_dir_fn=lambda path: setattr(config, "OUTPUT_DIR", Path(path)),
            empty_result_fn=lambda company, status, reason: {"company": company, "status": status, "reason": reason, "publication_eligible": False},
        )

    assert result == "PAID_PENDING_APPROVAL"
    assert attempts[0] == 2
    run_root = next(runs_dir.iterdir())
    manifest = json.loads((run_root / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["handoff"] is True and manifest["paid_pending"] == 1
    assert {key: manifest["telemetry"][key] for key in (
        "total_items", "free_completed", "free_failed", "item_terminal",
        "paid_required", "paid_completed", "result_count", "manifest_count",
    )} == {
        "total_items": 159, "free_completed": 158, "free_failed": 1,
        "item_terminal": 159, "paid_required": 1, "paid_completed": 0,
        "result_count": 159, "manifest_count": 159,
    }
    with sqlite3.connect(run_root / "state" / "progress.sqlite3") as connection:
        assert connection.execute("SELECT COUNT(*) FROM run_items WHERE free_state IN ('DONE','FAILED','NOT_REQUIRED')").fetchone()[0] == 159
        assert connection.execute("SELECT COUNT(*) FROM run_items WHERE quarantine_state='HANDOFF_PENDING'").fetchone()[0] == 159
        assert connection.execute("SELECT COUNT(*) FROM results WHERE json_extract(payload, '$.publication_eligible')=0").fetchone()[0] == 159


def test_duplicate_failed_or_unknown_never_becomes_no_call_needed():
    for inherited, expected in (("FAILED", "FAILED"), ("UNKNOWN", "UNKNOWN")):
        result = runtime.rejected_provider_result(runtime.Reservation(False, "hunter", "domain", 0, "PAID", call_id="old", reason="duplicate_request", inherited_state=inherited))
        assert pipeline_runner.paid_attempt_result({"provider_results": [result]}) == expected


def test_handoff_release_is_atomic_and_recomputes_policy_without_contact_loss(tmp_path: Path):
    db = tmp_path / "state.sqlite3"
    payload = {
        "company": "ACME MAKINA", "source_record_id": "input:0", "status": "OK_HIGH_CONFIDENCE",
        "publication_eligible": True, "website": "https://acme-makina.example", "email": "info@acme-makina.example",
        "publication_blockers": "identity_not_publishable",
        "email_publication_status": "allowed", "identity_assessment": {"publishable": True, "conflicts": [], "support_count": 2},
        "__evaluation": {"candidate": {"url": "https://acme-makina.example"}, "has_contact": True, "email": "info@acme-makina.example", "reasons": ["page_identity_strong:1/1", "legal_name_phrase_match:2", "context_match:1/1", "country_identity_tr_text"], "identity_assessment": {"publishable": True, "conflicts": [], "support_count": 2}},
    }
    text_payload = json.dumps(payload, separators=(",", ":"))
    with patch.object(config, "PROGRESS_DB_FILE", db):
        checkpoint.seed_recovered_run(path=db, run_id="handoff-release", input_hash="h", run_signature="s", context={"phase": "FREE"}, budgets={provider: 0 for provider in checkpoint.CANONICAL_PROVIDERS}, items=[{"item_index": 0, "source_record_id": "input:0", "free_state": "FAILED", "paid_required": True, "paid_state": "PENDING"}], results=[{"item_index": 0, "payload": text_payload}])
        checkpoint.mark_handoff_pending(run_id="handoff-release", expected_count=1)
        checkpoint.transition_phase("handoff-release", "PAID", expected_count=1)
        with sqlite3.connect(db) as connection:
            connection.execute("UPDATE run_items SET paid_state='DONE' WHERE run_id=?", ("handoff-release",))
            connection.commit()
        checkpoint.release_handoff_pending("handoff-release", expected_count=1)
        released = checkpoint.load_results_by_id("handoff-release")[0]
    assert released["quarantine_state"] == "" and "HANDOFF_PENDING" not in released.get("publication_blockers", "")
    assert "identity_not_publishable" not in released.get("publication_blockers", "")
    assert released["website"] == payload["website"] and released["email"] == payload["email"]
    assert released["publication_eligible"] is True


def test_normalize_domain_malformed_url_and_item_skip_are_fail_closed():
    assert __import__("modules.scorer", fromlist=["normalize_domain"]).normalize_domain("https://[broken") == ""
    candidates = {}
    search._add_search_results(candidates, "Acme", "Acme official", [{"href": "https://[broken", "title": "Acme", "body": ""}])
    assert not candidates and runtime.snapshot()["counters"].get("search.malformed_link_items", 0) == 1


@pytest.mark.parametrize("provider", ["brightdata", "google_places", "brandfetch", "hunter", "linkedin", "llm"])
def test_six_adapter_failure_classification_matrix(provider: str):
    assert runtime.is_unknown_transport_error(TimeoutError(provider))
    assert runtime.is_unknown_transport_error(ConnectionResetError(provider))
    assert runtime.is_unknown_transport_error(RuntimeError("remote end closed connection"))
    assert not runtime.is_unknown_transport_error(ValueError("HTTP 422 validation error"))


def test_paid_disabled_hunter_fallback_is_blocked_without_network():
    with patch.object(config, "PAID_ENABLED", False), patch.object(config, "ENABLE_HUNTER_FALLBACK", True), patch.object(config, "HUNTER_API_KEY", "secret"), patch.object(hunter.requests, "get", side_effect=AssertionError("network")):
        result = hunter.find_domain_emails("example.com")
    assert result.result_state == "NOT_ENABLED"


def test_free_handoff_preflight_precedes_approval_gate():
    source = inspect.getsource(pipeline_runner._run_pipeline_impl_body)
    assert source.index("preflight_source_profiles") < source.index("if escalation and not allow_paid")


def test_two_concurrent_workers_have_one_physical_probe(tmp_path: Path):
    db = tmp_path / "state.sqlite3"
    with patch.object(config, "PROGRESS_DB_FILE", db):
        checkpoint.initialize_schema(db)
        with ThreadPoolExecutor(max_workers=2) as pool:
            claims = list(pool.map(lambda _: checkpoint.claim_source_probe(run_id="r", host="fair.example"), range(2)))
    assert sum(bool(claim["owner"]) for claim in claims) == 1


def test_owner_crash_takeover_run_isolation_and_snapshot_hydration(tmp_path: Path):
    db = tmp_path / "state.sqlite3"
    with patch.object(config, "PROGRESS_DB_FILE", db), patch.object(config, "SOURCE_PROBE_LEASE_SEC", 1):
        first = checkpoint.claim_source_probe(run_id="run-a", host="fair.example")
        with sqlite3.connect(db) as connection:
            connection.execute("UPDATE source_probes SET lease_expires_at='2000-01-01T00:00:00+00:00' WHERE run_id=? AND host=?", ("run-a", "fair.example"))
            connection.commit()
        takeover = checkpoint.claim_source_probe(run_id="run-a", host="fair.example")
        other = checkpoint.claim_source_probe(run_id="run-b", host="fair.example")
        assert first["owner"] and takeover["owner"] and other["owner"]
        checkpoint.finish_source_probe(run_id="run-a", host="fair.example", owner_token=takeover["owner_token"], snapshot={"host": "fair.example", "status": "available"})
        runtime.configure_durable_run("run-a", {provider: 0 for provider in checkpoint.CANONICAL_PROVIDERS})
        hydrated = search._hydrate_source_health("https://fair.example/profile", {"host": "fair.example", "status": "available"})
        assert hydrated["status"] == "available"


def test_exact_command_config_and_run_id_reconcile_and_negative_tamper_cases(tmp_path: Path):
    from fixture_factory import remaining_run_sources
    result = __import__("prepare_remaining_run").prepare_remaining_run(tmp_path, **remaining_run_sources())
    plan_path = Path(result["plan"])
    workbook_path = Path(result["workbook"])
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    command = plan["exact_command"]
    powershell_command = plan["powershell_command"]
    assert "main.py" in command and "<" not in command and ">" not in command
    assert powershell_command.startswith("& ")
    exact_argv, exact_resolved = __import__("prepare_remaining_run")._parse_exact_command(command)
    ps_argv, ps_resolved = __import__("prepare_remaining_run")._parse_exact_command(powershell_command)
    assert exact_argv == ps_argv
    assert exact_resolved["effective_config"] == ps_resolved["effective_config"] == plan["effective_config"]
    for required in ("--no-allow-paid", "--search-cache use", "--crawl-cache use", "--brightdata-budget 0", "--google-places-budget 0", "--linkedin-company-budget 0", "--non-interactive"):
        assert required in command
    assert plan["effective_config"]["paid_enabled"] is False
    assert all(value == 0 for value in plan["effective_config"]["budgets"].values())
    assert plan["expected_run_id"] == plan["exact_command_run_id"]
    assert plan["expected_run_id"] != "93dfec9a0eac6a5e1dae1f74c089c6d4cd743f0b9cbb0c4b54cb464b6282717e"
    assert __import__("prepare_remaining_run").validate_remaining_plan(workbook_path, plan_path)["selection"]["count"] == 159
    tampered = dict(plan)
    tampered["selection"] = dict(plan["selection"])
    tampered["selection"]["ordered_index_id"] = list(reversed(plan["selection"]["ordered_index_id"]))
    plan_path.write_text(json.dumps(tampered), encoding="utf-8")
    with pytest.raises(ValueError, match="payload hash"):
        __import__("prepare_remaining_run").validate_remaining_plan(workbook_path, plan_path)


@pytest.mark.parametrize("row_count", [158, 160])
def test_remaining_plan_rejects_158_or_160_rows(tmp_path: Path, row_count: int):
    from fixture_factory import remaining_run_sources
    result = __import__("prepare_remaining_run").prepare_remaining_run(tmp_path, **remaining_run_sources())
    workbook_path = Path(result["workbook"])
    plan_path = Path(result["plan"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["company"])
    for index in range(row_count):
        sheet.append([f"row-{index}"])
    tampered = tmp_path / f"tampered_{row_count}.xlsx"
    workbook.save(tampered)
    workbook.close()
    with pytest.raises(ValueError):
        __import__("prepare_remaining_run").validate_remaining_plan(tampered, plan_path)


class _MalformedProviderResponse:
    status_code = 200
    text = "not-json"
    headers = {}

    def raise_for_status(self):
        return None

    def json(self):
        raise ValueError("malformed json")


def test_brightdata_connection_error_is_unknown_before_any_success():
    runtime.reset()
    search.reset_source_health()
    runtime.set_phase("PAID")
    token = runtime.begin_provider_attempt()
    with patch.object(config, "BRIGHTDATA_API_KEY", "key"), patch.object(
        config, "BRIGHTDATA_REQUEST_BUDGET", 1
    ), patch.object(
        search.requests, "post", side_effect=requests.ConnectionError("connection reset")
    ):
        result = search._brightdata_text("Acme")
    assert result.result_state == "UNKNOWN"
    outcomes = runtime.end_provider_attempt(token)
    assert any(item["result_state"] == "UNKNOWN" for item in outcomes)


def test_google_and_hunter_malformed_json_are_failed_before_done():
    with patch.object(config, "ENABLE_GOOGLE_PLACES", True), patch.object(
        config, "GOOGLE_PLACES_API_KEY", "key"
    ), patch.object(config, "SEARCH_CACHE_MODE", "off"), patch.object(
        runtime, "reserve_api", return_value=runtime.Reservation(True, "google_places", "text", 0, "PAID", call_id="g")
    ), patch.object(google_places.requests, "post", return_value=_MalformedProviderResponse()):
        assert google_places.search_company("Acme").result_state == "FAILED"

    with patch.object(config, "PAID_ENABLED", True), patch.object(
        config, "ENABLE_HUNTER_FALLBACK", True
    ), patch.object(config, "HUNTER_REQUEST_BUDGET", 1
    ), patch.object(config, "HUNTER_API_KEY", "key"), patch.object(
        runtime, "reserve_api", return_value=runtime.Reservation(True, "hunter", "domain", 0, "PAID", call_id="h")
    ), patch.object(hunter.requests, "get", return_value=_MalformedProviderResponse()):
        assert hunter.find_domain_emails("example.com").result_state == "FAILED"


def test_six_real_adapters_route_duplicate_through_runtime_helper_without_network():
    duplicate = runtime.Reservation(False, "provider", "operation", 0, "PAID", call_id="old-call", reason="duplicate_request", inherited_state="DONE")
    with patch.object(config, "BRIGHTDATA_API_KEY", "key"), patch.object(
        runtime, "reserve_api", return_value=duplicate
    ), patch.object(search.requests, "post", side_effect=AssertionError("network")):
        result = search._brightdata_text("Acme")
        assert result.result_state == "COMPLETED" and result.call_ids == ("old-call",)

    with patch.object(config, "ENABLE_GOOGLE_PLACES", True), patch.object(
        config, "GOOGLE_PLACES_API_KEY", "key"
    ), patch.object(config, "SEARCH_CACHE_MODE", "off"), patch.object(
        runtime, "reserve_api", return_value=duplicate
    ), patch.object(google_places.requests, "post", side_effect=AssertionError("network")):
        result = google_places.search_company("Acme")
        assert result.result_state == "COMPLETED" and result.call_ids == ("old-call",)

    with patch.object(config, "ENABLE_BRANDFETCH_DOMAIN_SEARCH", True), patch.object(
        config, "BRANDFETCH_CLIENT_ID", "key"
    ), patch.object(config, "SEARCH_CACHE_MODE", "off"), patch.object(
        company_resolvers, "_cached", return_value=None
    ), patch.object(runtime, "reserve_api", return_value=duplicate), patch.object(
        company_resolvers.requests, "get", side_effect=AssertionError("network")
    ):
        result = company_resolvers.brandfetch_domains("Acme")
        assert result.result_state == "COMPLETED" and result.call_ids == ("old-call",)

    with patch.object(config, "ENABLE_HUNTER_DOMAIN_FINDER", True), patch.object(
        config, "HUNTER_API_KEY", "key"
    ), patch.object(config, "SEARCH_CACHE_MODE", "off"), patch.object(
        company_resolvers, "_cached", return_value=None
    ), patch.object(runtime, "reserve_api", return_value=duplicate), patch.object(
        company_resolvers.requests, "get", side_effect=AssertionError("network")
    ):
        result = company_resolvers.hunter_domains("Acme")
        assert result.result_state == "COMPLETED" and result.call_ids == ("old-call",)

    with patch.object(linkedin_company, "_reserve", return_value=duplicate), patch.object(
        linkedin_company.requests, "post", side_effect=AssertionError("network")
    ):
        result = linkedin_company._find_company_url("Acme")
        assert result.result_state == "COMPLETED" and result.call_ids == ("old-call",)

    with patch.object(config, "ENABLE_LLM_ARBITER", True), patch.object(
        config, "OPENROUTER_API_KEY", "key"
    ), patch.object(config, "SEARCH_CACHE_MODE", "off"), patch.object(
        runtime, "reserve_api", return_value=duplicate
    ):
        result = __import__("modules.llm_arbiter", fromlist=["arbitrate"]).arbitrate(
            "Acme", "Acme Ltd", "Packaging", "acme.example", "short summary"
        )
        assert result["provider_result"] == "COMPLETED"
        assert result["provider_call_ids"] == ["old-call"]


def test_cache_hit_is_success_and_makes_no_physical_call():
    from modules import cache_store

    with patch.object(config, "ENABLE_GOOGLE_PLACES", True), patch.object(
        config, "GOOGLE_PLACES_API_KEY", "key"
    ), patch.object(config, "SEARCH_CACHE_MODE", "use"), patch.object(
        cache_store, "load", return_value=[{"website": "https://acme.example"}]
    ), patch.object(google_places.requests, "post", side_effect=AssertionError("network")):
        result = google_places.search_company("Acme")
    assert result.result_state == "CACHE_HIT"


def test_hunter_central_authorization_gate_blocks_before_network():
    with patch.object(config, "PAID_ENABLED", True), patch.object(
        config, "ENABLE_HUNTER_FALLBACK", True
    ), patch.object(config, "HUNTER_API_KEY", "key"), patch.object(
        runtime, "paid_access_allowed", return_value=False
    ), patch.object(hunter.requests, "get", side_effect=AssertionError("network")):
        result = hunter.find_domain_emails("example.com")
    assert result.result_state == "NOT_ENABLED"


def test_golden_publication_regression_fixture():
    fixture = json.loads((Path(__file__).parent / "fixtures" / "golden_publication_regression_r5.json").read_text(encoding="utf-8"))
    for item in fixture["records"]:
        row = {
            "company": item["company"], "source_record_id": f"r3:{item['item_index']}",
            "free_state": "DONE", "paid_state": "NOT_REQUIRED", "paid_required": False,
            "status": item["status"], "score": item["score"],
            "publication_eligible": True, "identity_resolution": item["identity_resolution"],
            "website": item["website"], "identity_assessment": item["identity_assessment"],
            "reason": "; ".join(item["reasons"]), "email": item["email"], "phone": item["phone"],
            "email_publication_status": "allowed" if item["email"] else "suppressed",
            "phone_publication_status": "allowed" if item["phone"] else "suppressed",
            "__evaluation": {"candidate": item["candidate"], "reasons": item["reasons"], "identity_assessment": item["identity_assessment"]},
        }
        assert pipeline_runner.output_artifacts.is_publishable_row(row) is False

    safe = {
        "company": "ACME MAKINA", "source_record_id": "fixture:safe",
        "free_state": "DONE", "paid_state": "NOT_REQUIRED", "paid_required": False,
        "status": "OK_HIGH_CONFIDENCE", "score": 90, "publication_eligible": True,
        "website": "https://acme-makina.example", "email": "info@acme-makina.example",
        "email_publication_status": "allowed", "identity_assessment": {"publishable": True, "conflicts": [], "support_count": 2},
        "__evaluation": {"candidate": {"url": "https://acme-makina.example"}, "reasons": ["page_identity_strong:1/1", "legal_name_phrase_match:2", "context_match:1/1", "country_identity_tr_text"], "identity_assessment": {"publishable": True, "conflicts": [], "support_count": 2}},
    }
    assert pipeline_runner.output_artifacts.is_publishable_row(safe) is True


def test_slow_probe_heartbeats_and_only_one_owner_per_host(tmp_path: Path):
    db = tmp_path / "state.sqlite3"
    calls = {"probe": 0, "heartbeat": 0}
    real_heartbeat = checkpoint.heartbeat_source_probe

    def slow_probe(_url):
        calls["probe"] += 1
        time.sleep(0.45)

    def heartbeat(**kwargs):
        calls["heartbeat"] += 1
        return real_heartbeat(**kwargs)

    with patch.object(config, "PROGRESS_DB_FILE", db), patch.object(
        config, "SOURCE_PROBE_LEASE_SEC", 1
    ), patch.object(config, "SEARCH_CACHE_MODE", "off"), patch.object(
        search, "_profile_external_websites", side_effect=slow_probe
    ), patch.object(checkpoint, "heartbeat_source_probe", side_effect=heartbeat):
        checkpoint.initialize_schema(db)
        results = search.preflight_source_profiles(
            [{"profile_url": "https://fair.example/profile"}, {"profile_url": "https://fair.example/other"}],
            run_id="run-a",
        )
    assert calls["probe"] == 1
    assert calls["heartbeat"] >= 1
    assert results[0]["host"] == "fair.example"


def test_transient_exhaustion_preserves_original_error_and_finishes_each_claim_once(tmp_path: Path):
    db = tmp_path / "state.sqlite3"
    finish_calls = []
    real_finish = checkpoint.finish_source_probe

    def finish(**kwargs):
        finish_calls.append(kwargs["owner_token"])
        return real_finish(**kwargs)

    with patch.object(config, "PROGRESS_DB_FILE", db), patch.object(
        config, "SOURCE_PROFILE_MAX_TRANSIENT_RETRIES", 2
    ), patch.object(search, "_profile_external_websites", side_effect=ConnectionError("original transport")), patch.object(
        checkpoint, "finish_source_probe", side_effect=finish
    ):
        checkpoint.initialize_schema(db)
        with pytest.raises(RuntimeError, match="original transport"):
            search.preflight_source_profiles([{"profile_url": "https://fair.example/profile"}], run_id="run-a")
    assert len(finish_calls) == 2
    assert len(set(finish_calls)) == 2


def test_resume_identity_and_non_complete_manifest_reject_before_network_or_worker(tmp_path: Path):
    input_file = tmp_path / "input.xlsx"
    workbook = Workbook()
    workbook.active.append(["company"])
    workbook.active.append(["Acme"])
    workbook.save(input_file)
    workbook.close()
    run_config = run_context.RunConfig.from_config(paid_enabled=False)
    input_hash = checkpoint.file_hash(input_file)
    source_id = run_context.source_record_identity({"company": "Acme"})[0]
    run_id = run_context.canonical_run_id(
        input_sha256=input_hash, ordered_source_record_ids=[source_id],
        effective_config=run_config.as_dict(), runtime_source_tree_hash=run_context.source_tree_sha256(),
    )
    run_root = tmp_path / run_id
    run_root.mkdir()
    (run_root / "manifest.json").write_text(json.dumps({
        "complete": True, "phase": "FINALIZING", "run_id": run_id,
        "input_sha256": input_hash, "run_config": run_config.as_dict(),
        "ordered_source_record_ids": [source_id], "lineage": {"type": "fresh"},
    }), encoding="utf-8")
    with patch.object(search, "preflight_source_profiles", side_effect=AssertionError("network")):
        with pytest.raises(RuntimeError, match="exactly COMPLETE"):
            pipeline_runner.run_pipeline(
                input_file, resume_run_dir=run_root, allow_paid=False,
                process_company_fn=lambda *_: (_ for _ in ()).throw(AssertionError("worker")),
                write_outputs_fn=lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("writer")),
                set_output_dir_fn=lambda *_: None,
                empty_result_fn=lambda *_: {},
            )

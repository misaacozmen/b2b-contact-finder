from __future__ import annotations

from openpyxl import load_workbook
import pytest

from modules import field_merge, run_report


@pytest.mark.parametrize(
    ("row", "field", "expected"),
    [
        ({"website": "https://example.com", "website_source": "OWN_SEARCH", "status": "OK_HIGH_CONFIDENCE"}, "website", "HIGH"),
        ({"website": "https://example.com", "website_source": "OWN_SEARCH", "status": "OK_MEDIUM_CONFIDENCE"}, "website", "MEDIUM"),
        ({"website": "https://example.com", "website_source": "OWN_SEARCH+REFERENCE"}, "website", "HIGH"),
        ({"website": "https://example.com", "website_source": "REFERENCE_VERIFIED"}, "website", "HIGH"),
        ({"website": "https://example.com", "website_source": "REFERENCE_ACCEPTED"}, "website", "MEDIUM"),
        ({"website": "https://example.com", "website_source": "REFERENCE_UNREACHABLE"}, "website", "LOW"),
        ({"website": "https://example.com", "website_source": "REFERENCE_THIN"}, "website", "LOW"),
        ({"website": "https://example.com", "website_source": "PAID_BRIGHTDATA", "status": "OK_HIGH_CONFIDENCE"}, "website", "HIGH"),
        ({"website": "https://example.com", "website_source": "PAID_BRIGHTDATA", "status": "OK_MEDIUM_CONFIDENCE"}, "website", "MEDIUM"),
        ({"website": "https://example.com", "status": "OK_HIGH_CONFIDENCE"}, "website", "HIGH"),
        ({"website": "https://example.com", "website_source": "OWN_SEARCH", "status": "OK_HIGH_CONFIDENCE", "email": "info@example.com", "email_source_tier": "SITE"}, "email", "HIGH"),
        ({"website": "https://example.com", "website_source": "OWN_SEARCH", "status": "OK_MEDIUM_CONFIDENCE", "email": "info@example.com", "email_source_tier": "SITE"}, "email", "MEDIUM"),
        ({"website": "https://example.com", "website_source": "OWN_SEARCH", "status": "OK_HIGH_CONFIDENCE", "email": "person@gmail.com", "email_source_tier": "SITE_FREEMAIL"}, "email", "LOW"),
        ({"email": "person@gmail.com", "email_source_tier": "REFERENCE_LISTING"}, "email", "MEDIUM"),
        ({"email": "person@example.com", "email_source_tier": "PAID_HUNTER"}, "email", "MEDIUM"),
        ({"website": "https://example.com", "website_source": "OWN_SEARCH", "status": "OK_HIGH_CONFIDENCE", "phone": "02125550000", "phone_source_tier": "SITE"}, "phone", "HIGH"),
        ({"website": "https://example.com", "website_source": "REFERENCE_THIN", "phone": "02125550000", "phone_source_tier": "SITE"}, "phone", "MEDIUM"),
        ({"phone": "02125550000", "phone_source_tier": "REFERENCE_LISTING"}, "phone", "MEDIUM"),
        ({"phone": "02125550000", "phone_source_tier": "PAID_GOOGLE_PLACES"}, "phone", "MEDIUM"),
        ({"website": "", "website_source": "REFERENCE_VERIFIED"}, "website", "NONE"),
    ],
)
def test_confidence_table(row, field, expected):
    assert field_merge.field_confidence(row, field) == expected


def test_unknown_website_source_counts_as_own_search():
    row = {
        "website": "https://example.com/",
        "website_source": "Example official website contact",
        "status": "OK_MEDIUM_CONFIDENCE",
    }

    assert field_merge.field_confidence(row, "website") == "MEDIUM"


def test_field_gaps_counts_low_as_gap():
    row = {
        "website": "https://thin.example/", "website_source": "REFERENCE_THIN",
        "email": "person@gmail.com", "email_source_tier": "SITE_FREEMAIL",
        "phone": "02125550000", "phone_source_tier": "REFERENCE_LISTING",
    }
    field_merge.annotate(row, None)
    assert field_merge.field_gaps(row) == {"website", "email"}
    assert row["field_gaps"] == "website;email"


def test_merge_never_downgrades_or_empties():
    previous = {
        "website": "https://good.example/", "website_source": "REFERENCE_VERIFIED",
        "status": "OK_HIGH_CONFIDENCE", "email": "info@good.example",
        "email_source_tier": "SITE", "phone": "02125550000",
        "phone_source_tier": "REFERENCE_LISTING",
    }
    current = {
        "website": "", "website_source": "REFERENCE_THIN",
        "status": "WEBSITE_NOT_FOUND", "email": "",
        "email_source_tier": "PAID_HUNTER", "phone": "",
        "phone_source_tier": "SITE",
    }
    merged = field_merge.merge(previous, current)
    assert merged["website"] == previous["website"]
    assert merged["email"] == previous["email"]
    assert merged["phone"] == previous["phone"]
    assert merged["website_source"] == "REFERENCE_VERIFIED"
    assert merged["email_source_tier"] == "SITE"
    assert merged["phone_source_tier"] == "REFERENCE_LISTING"
    assert merged["status"] == "OK_HIGH_CONFIDENCE"


def test_merge_prefers_higher_confidence_and_moves_source():
    previous = {
        "website": "https://site.example/old", "website_source": "REFERENCE_ACCEPTED",
        "status": "OK_MEDIUM_CONFIDENCE", "email": "person@gmail.com",
        "email_source_tier": "SITE_FREEMAIL", "phone": "02125550000",
        "phone_source_tier": "REFERENCE_LISTING", "alternative_phones": "old-alt",
    }
    current = {
        "website": "https://site.example/new", "website_source": "OWN_SEARCH",
        "status": "OK_HIGH_CONFIDENCE", "email": "info@site.example",
        "email_source_tier": "SITE", "phone": "02125551111",
        "phone_source_tier": "PAID_GOOGLE_PLACES", "alternative_phones": "new-alt",
    }
    merged = field_merge.merge(previous, current)
    assert merged["website"] == current["website"]
    assert merged["website_source"] == "OWN_SEARCH"
    assert merged["email"] == current["email"]
    assert merged["email_source_tier"] == "SITE"
    # Both phones are MEDIUM; ties retain the earlier value and companions.
    assert merged["phone"] == previous["phone"]
    assert merged["phone_source_tier"] == "REFERENCE_LISTING"
    assert merged["alternative_phones"] == "old-alt"


def test_ready_for_publication_rule():
    assert field_merge.ready_for_publication({
        "website": "https://example.com", "website_source": "OWN_SEARCH",
        "status": "OK_HIGH_CONFIDENCE", "email": "info@example.com",
        "email_source_tier": "SITE",
    })
    assert not field_merge.ready_for_publication({
        "website": "https://example.com", "website_source": "REFERENCE_THIN",
        "phone": "02125550000", "phone_source_tier": "REFERENCE_LISTING",
    })
    assert not field_merge.ready_for_publication({
        "website": "https://example.com", "website_source": "REFERENCE_VERIFIED",
    })


def test_report_rates_table_three_stages(monkeypatch):
    from modules import checkpoint

    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    rows = [
        {
            "company": "A",
            "stage_a": {"website": "a", "website_confidence": "HIGH", "email": "a", "email_confidence": "MEDIUM", "phone_confidence": "NONE"},
            "stage_ab": {"website": "a", "website_confidence": "HIGH", "email_confidence": "LOW", "phone": "a", "phone_confidence": "MEDIUM"},
            "website": "a", "website_confidence": "HIGH", "email": "a", "email_confidence": "MEDIUM", "phone_confidence": "NONE",
        },
        {
            "company": "B",
            "stage_a": {"website": "b", "website_confidence": "MEDIUM", "email_confidence": "NONE", "phone": "b", "phone_confidence": "HIGH"},
            "stage_ab": {"email": "b", "email_confidence": "MEDIUM", "website_confidence": "NONE", "phone_confidence": "NONE"},
            "website_confidence": "NONE", "email": "b", "email_confidence": "HIGH", "phone": "b", "phone_confidence": "MEDIUM",
        },
        {
            "company": "C",
            "stage_a": {"website": "c", "website_confidence": "LOW", "email": "c", "email_confidence": "HIGH", "phone": "c", "phone_confidence": "MEDIUM"},
            "stage_ab": {"website": "c", "website_confidence": "MEDIUM", "email": "c", "email_confidence": "HIGH", "phone": "c", "phone_confidence": "HIGH"},
            "website": "c", "website_confidence": "MEDIUM", "email_confidence": "NONE", "phone": "c", "phone_confidence": "HIGH",
        },
    ]
    markdown = run_report._build_report(
        rows, run_id="fake-run", run_status="TAMAMLANDI", status_detail="ok",
        elapsed_seconds=3, telemetry=None, generated_at="2026-09-26T00:00:00+00:00",
        file_names=["sonuclar.xlsx"],
    )
    expected = "\n".join([
        "| Aşama | Web sitesi | E-posta | Telefon | Yayına hazır |",
        "|---|---:|---:|---:|---:|",
        "| A — Bağımsız arama (referanssız) | 2 (66.7%) | 2 (66.7%) | 2 (66.7%) | 2 (66.7%) |",
        "| A + Referans | 2 (66.7%) | 2 (66.7%) | 2 (66.7%) | 2 (66.7%) |",
        "| Final (A + Referans + Ücretli) | 2 (66.7%) | 2 (66.7%) | 2 (66.7%) | 2 (66.7%) |",
    ])
    assert expected in markdown


def test_report_splits_strict_and_brand_hit_rate(monkeypatch):
    from modules import checkpoint

    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    rows = [
        {
            "company": "Same domain",
            "stage_a": {"website": "https://acme.com/", "status": "OK_HIGH_CONFIDENCE"},
            "reference_website": "https://www.acme.com/",
        },
        {
            "company": "TLD variant",
            "stage_a": {"website": "https://acme.com/", "status": "OK_MEDIUM_CONFIDENCE"},
            "reference_website": "https://acme.com.tr/",
        },
        {
            "company": "Different brand",
            "stage_a": {"website": "https://other.net/", "status": "REVIEW_NEEDED"},
            "reference_website": "https://acme.com/",
        },
    ]

    markdown = run_report._build_report(
        rows, run_id="fake-run", run_status="TAMAMLANDI", status_detail="ok",
        elapsed_seconds=3, telemetry=None, generated_at="2026-09-26T00:00:00+00:00",
        file_names=["sonuclar.xlsx"],
    )

    assert "A isabeti — aynı alan adı (katı): 1 / 3 (33.3%)" in markdown
    assert "A isabeti — aynı marka çekirdeği (TLD farkı dahil, ör. firma.com ↔ firma.com.tr): 2 / 3 (66.7%)" in markdown
    assert "Yalnız doğrulanmış (OK) A satırları — katı: 1 / 2 (50.0%)" in markdown
    assert "Yalnız doğrulanmış (OK) A satırları — marka: 2 / 2 (100.0%)" in markdown


def test_report_input_status_distribution_from_row(tmp_path, monkeypatch):
    from modules import checkpoint

    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    rows = [
        {"company": "OK", "listed_website_status": "OK"},
        {"company": "Fixed", "listed_website_status": "FIXED"},
        {"company": "Invalid", "listed_website_status": "INVALID"},
        {"company": "Empty", "listed_website_status": "EMPTY"},
    ]

    errors = run_report.write_run_report(
        rows, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test",
        elapsed_seconds=1,
    )

    assert errors == {}
    markdown = (tmp_path / "rapor.md").read_text(encoding="utf-8")
    assert "Girdi web: EMPTY=1, FIXED=1, INVALID=1, OK=1" in markdown
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=True, data_only=True)
    try:
        headers = [cell.value for cell in next(book["Tüm firmalar"].iter_rows())]
        status_column = headers.index("Girdi web durumu")
        assert [book["Tüm firmalar"].cell(row, status_column + 1).value for row in range(2, 6)] == [
            "OK", "FIXED", "INVALID", "EMPTY",
        ]
    finally:
        book.close()


def test_excel_has_four_sheets_and_exact_headers(tmp_path):
    rows = [
        {
            "company": "Ready Co", "website": "https://ready.example/",
            "website_source": "OWN_SEARCH", "status": "OK_HIGH_CONFIDENCE",
            "email": "info@ready.example", "email_source_tier": "SITE",
            "phone": "", "reference_website": "https://ready.example/",
            "reference_tier": "REFERENCE_MATCHES_OWN_SEARCH",
        },
        {
            "company": "Thin Co", "website": "https://thin.example/",
            "website_source": "REFERENCE_THIN", "status": "WEBSITE_NOT_FOUND",
            "email": "", "phone": "", "reference_website": "https://thin.example/",
            "reference_tier": "REFERENCE_THIN",
        },
    ]
    result = run_report.write_run_report(
        rows, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test",
        elapsed_seconds=1, telemetry={"free_search_backend_health": {"bing": {"ok": 2, "empty": 1, "error": 0}}},
    )
    assert result == {}
    assert (tmp_path / "sonuclar.xlsx").is_file()
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["Tüm firmalar", "Yayına hazır", "Eksikler", "Özet"]
        for name in book.sheetnames[:3]:
            assert [cell.value for cell in next(book[name].iter_rows())] == run_report.HEADERS
        summary_headers = [cell.value for cell in next(book["Özet"].iter_rows())]
        assert summary_headers == ["Aşama", "Web sitesi", "E-posta", "Telefon", "Yayına hazır"]
        assert book["Yayına hazır"].max_row == 2
        assert book["Eksikler"].max_row == 3
    finally:
        book.close()

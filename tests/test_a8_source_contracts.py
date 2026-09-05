from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tools.acquire_independent_sources import (
    _OfficialRequestLimiter,
    _clean_string,
    _get,
    _labelled_website,
    _select_pool,
    _valid_company_website,
)
from tools.build_actual_inputs import build
from tools.run_source_review_pass import _labelled_company_website
from modules.source_adapters import SOURCE_ADAPTERS, classify_link


class _Response:
    def __init__(self, status: int, body: bytes = b"ok", headers: dict | None = None):
        self.status_code = status
        self.content = body
        self.headers = headers or {}
        self.url = "https://hometex.com.tr/test"

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class _Session:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = 0

    def get(self, url, **kwargs):
        self.calls += 1
        return self.responses.pop(0)


def test_actual_builder_is_selection_only_and_keeps_column_roles(tmp_path: Path):
    selection = tmp_path / "selection.json"
    selection.write_text(json.dumps({"records": [
        {"source_record_id": "zuchex_2026:a", "source": "zuchex_2026", "display_name": "A", "official_profile_url": "https://zuchex.com/detail/a", "listing_url": "https://zuchex.com/"},
        {"source_record_id": "hometex_2026:b", "source": "hometex_2026", "display_name": "B", "official_profile_url": "https://hometex.com.tr/en/b", "listed_website": "https://b.example"},
    ]}), encoding="utf-8")
    output = tmp_path / "actual.xlsx"
    build(selection, output)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        headers = [str(value) for value in next(workbook.active.iter_rows(values_only=True))]
        rows = [dict(zip(headers, row)) for row in workbook.active.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()
    assert not any(header.casefold() == "expected" for header in headers)
    assert all(not row["website"] for row in rows)
    assert not rows[0]["profile_url"]
    assert rows[0]["listing_url"] == "https://zuchex.com/"
    assert rows[1]["profile_url"].endswith("/b")


def test_website_role_classifier_rejects_non_company_values():
    for value in (
        "https://apps.apple.com/app/texhibition",
        "https://xing.com/spi/shares/new",
        "https://www.informa.com/organizer",
        "https://old.texhibitionist.com/storage/certificate.pdf",
        "http://-",
        "https://tobb.org.tr/Sayfalar/Eng/AnaSayfa.php",
    ):
        assert not _valid_company_website(value)


def test_four_source_adapters_and_generic_unknown_link_contract():
    assert set(SOURCE_ADAPTERS) == {"texhibition_2026", "zuchex_2026", "hometex_2026", "ambiente_2026"}
    assert classify_link("https://unknown.example", company_name="Different Company")["role"] == "unknown"
    assert classify_link("https://firma.example", label="Website", company_name="Firma")["role"] == "company_candidate"


def test_labelled_website_positive_and_footer_external_negative():
    from bs4 import BeautifulSoup

    html = """
    <div class='company'><span>Website</span><a href='https://firma.example'>Website</a></div>
    <footer><a href='https://organizer.example'>Organizer</a></footer>
    """
    soup = BeautifulSoup(html, "html.parser")
    assert _labelled_website(soup) == "https://firma.example"
    assert _labelled_company_website(soup) == "https://firma.example"
    footer_only = BeautifulSoup("<footer><a href='https://firma.example'>Company</a></footer>", "html.parser")
    assert _labelled_website(footer_only) == ""


def test_ambiente_html_entity_is_decoded_before_normalization():
    assert _clean_string("K&uuml;tahya") == "Kütahya"


def test_retry_429_then_success_is_bounded_and_telemetry_is_recorded(monkeypatch):
    session = _Session([_Response(429, headers={"Retry-After": "0"}), _Response(200)])
    telemetry = []
    session._a8_telemetry = telemetry
    session._a8_limiter = _OfficialRequestLimiter(0)
    monkeypatch.setattr("tools.acquire_independent_sources.time.sleep", lambda _: None)
    response = _get(session, "https://hometex.com.tr/test")
    assert response.status_code == 200
    assert session.calls == 2
    assert [entry["status"] for entry in telemetry] == [429, 200]


def test_retryable_only_429_and_5xx(monkeypatch):
    session = _Session([_Response(404)])
    session._a8_telemetry = []
    session._a8_limiter = _OfficialRequestLimiter(0)
    with pytest.raises(RuntimeError):
        _get(session, "https://hometex.com.tr/test")
    assert session.calls == 1


def test_duplicate_overlap_key_preserves_all_previous_ids():
    records = [
        {"source_record_id": "new:1", "legal_name": "Same Name", "brand": "", "website": ""},
        {"source_record_id": "new:2", "legal_name": "Different Name", "brand": "", "website": ""},
    ]
    known = [
        {"source_record_id": "old:1", "legal_name": "Same Name", "brand": "", "website": ""},
        {"source_record_id": "old:2", "legal_name": "Same Name", "brand": "", "website": ""},
    ]
    selected, excluded = _select_pool(records, known, "hometex_2026", limit=1)
    assert [item["source_record_id"] for item in selected] == ["new:2"]
    assert {item["matched_previous_id"] for item in excluded} == {"old:1", "old:2"}

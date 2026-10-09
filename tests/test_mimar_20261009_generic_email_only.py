"""Talimat 40: only generic company mailboxes reach a delivery file."""

from openpyxl import load_workbook
import pytest

from modules import checkpoint, email_kind, run_report


@pytest.mark.parametrize("email", [
    "jan.kowalski@firma-ornek.pl", "a.nowak@firma-ornek.pl", "KOWALSKI.JAN@firma-ornek.pl",
    "michal.k.kruk@firma-ornek.pl", "grzegorz@firma-ornek.pl", "Łukasz@firma-ornek.pl",
    "mehmet@firma-ornek.com.tr", "Bülent@firma-ornek.com.tr", "ayse.yilmaz@firma-ornek.com.tr",
])
def test_personal_mailboxes(email):
    assert email_kind.is_personal(email)
    assert email_kind.is_excluded(email)


@pytest.mark.parametrize("email", [
    "rodo@firma-ornek.pl", "iod@firma-ornek.pl", "ochrona.danych@firma-ornek.pl",
    "pl-m-inspektor-ochrony-danych@firma-ornek.pl", "kvkk@firma-ornek.com.tr",
    "abc.kvkk@firma-ornek.com.tr", "privacy@firma-ornek.com.tr",
])
def test_data_protection_mailboxes(email):
    assert email_kind.is_data_protection(email)
    assert email_kind.is_excluded(email)


@pytest.mark.parametrize("email", [
    "info@firma-ornek.pl", "biuro@firma-ornek.pl", "biuro.handlowe@firma-ornek.pl", "e-sklep@firma-ornek.pl",
    "stal.kolno@firma-ornek.pl", "sprzedaz2@firma-ornek.pl", "firmaornek@firma-ornek.pl",
    "jkowalski@firma-ornek.pl", "satis@firma-ornek.com.tr", "",
])
def test_generic_mailboxes_stay(email):
    assert not email_kind.is_excluded(email)


ROWS = [
    {
        "company": "Alfa", "website": "https://alfa-ornek.pl/", "email": "jan.kowalski@alfa-ornek.pl",
        "alternative_emails": "anna@alfa-ornek.pl; biuro@alfa-ornek.pl",
        "status": "OK_HIGH_CONFIDENCE",
        "stage_a": {
            "website": "https://alfa-ornek.pl/", "website_confidence": "HIGH",
            "email": "jan.kowalski@alfa-ornek.pl", "email_confidence": "HIGH",
        },
    },
    {"company": "Beta", "website": "https://beta-ornek.pl/", "email": "rodo@beta-ornek.pl", "status": "OK_HIGH_CONFIDENCE"},
    {
        "company": "Gama", "website": "https://gama-ornek.pl/", "email": "teresa@gama-ornek.pl",
        "alternative_emails": "biuro@baska-ornek.pl", "status": "OK_HIGH_CONFIDENCE",
    },
    {"company": "Delta", "website": "https://delta-ornek.pl/", "email": "biuro@delta-ornek.pl", "status": "OK_HIGH_CONFIDENCE"},
]


def test_report_swaps_in_a_generic_address_or_leaves_the_email_empty(tmp_path, monkeypatch):
    monkeypatch.setattr(checkpoint, "load_run_state_by_id", lambda _run_id: {})
    original = [dict(row) for row in ROWS]
    assert run_report.write_run_report(
        ROWS, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    assert ROWS == original
    assert ROWS[0]["stage_a"]["email"] == "jan.kowalski@alfa-ornek.pl"
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=True, data_only=True)
    try:
        contact = [list(row) for row in book["İletişim"].iter_rows(values_only=True)]
        assert [row[2] for row in contact[1:]] == ["biuro@alfa-ornek.pl", None, None, "biuro@delta-ornek.pl"]
        details = [list(row) for row in book["Detaylar"].iter_rows(values_only=True)]
        source = details[0].index("E-posta kaynağı")
        assert [row[source] for row in details[1:]] == ["Firma sitesi", None, None, "Firma sitesi"]
        assert book["Özet"]["B9"].value == 2
        assert book["Özet"]["C16"].value.startswith("0 ")
        assert book["Özet"]["E16"].value.startswith("0 ")
    finally:
        book.close()
    markdown = (tmp_path / "rapor.md").read_text(encoding="utf-8")
    assert "kowalski" not in markdown and "rodo@" not in markdown and "teresa" not in markdown

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

import prepare_remaining_run as builder


def test_prepare_remaining_run_exact_159_and_verify_only(tmp_path: Path):
    first = builder.prepare_remaining_run(tmp_path)
    workbook = load_workbook(first["workbook"], read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    plan = json.loads(Path(first["plan"]).read_text(encoding="utf-8"))
    assert len(rows) == 160
    assert len(rows[0]) == 17
    assert plan["selection"]["count"] == 159
    assert plan["selection"]["free_count"] == 2
    assert plan["selection"]["paid_count"] == 157
    assert plan["result_field_leakage"] == []
    second = builder.prepare_remaining_run(tmp_path)
    assert second["verify_only"] is True
    assert second["workbook_sha256"] == first["workbook_sha256"]
    assert second["plan_sha256"] == first["plan_sha256"]


def test_prepare_remaining_run_plan_has_no_recovery_result_payload():
    plan = json.loads((Path(__file__).parents[1] / "input" / "remaining_159_plan.json").read_text(encoding="utf-8")) if (Path(__file__).parents[1] / "input" / "remaining_159_plan.json").exists() else None
    if plan is not None:
        serialized = json.dumps(plan, ensure_ascii=False).casefold()
        assert '"website_source"' not in serialized
        assert '"email_source"' not in serialized
        assert '"phone_source"' not in serialized

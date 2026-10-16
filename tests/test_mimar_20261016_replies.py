"""Talimat 50: replies read from the mailbox, written to the workbooks and sent to the user as a notice."""

from datetime import date, datetime
from email.message import EmailMessage
import json

from openpyxl import load_workbook
import pytest

import config
from modules import outreach, replies
from tools import eposta_gonder


ACCOUNT = "satis@ornek-gonderen.example"
OWNER = "sahip@ornek-kisisel.example"
FIRMS = [
    ("Alfa Ornek", "info@alfa-ornek.com.tr", "<id-alfa@ornek-gonderen.example>"),
    ("Beta Ornek", "info@beta-ornek.com.tr", "<id-beta@ornek-gonderen.example>"),
    ("Gama Ornek", "info@gama-ornek.com.tr", "<id-gama@ornek-gonderen.example>"),
    ("Delta Ornek", "info@delta-ornek.com.tr", "<id-delta@ornek-gonderen.example>"),
]
QUOTE = (
    "\n\nOrnek Gonderen <satis@ornek-gonderen.example>, 15 Eki 2026 Per, 12:21 tarihinde şunu yazdı:\n"
    "> Benden tekrar e-posta almak istemiyorsanız \"istemiyorum\" yazmanız yeterli.\n"
)


@pytest.fixture(autouse=True)
def files(tmp_path, monkeypatch):
    for name, file_name in (
        ("OUTREACH_LOG_FILE", "eposta_gonderim_kaydi.jsonl"), ("REPLY_LOG_FILE", "eposta_yanit_kaydi.jsonl"),
        ("REPLY_STATE_FILE", "eposta_yanit_durumu.json"), ("REPLY_SETTINGS_FILE", "eposta_yanit_ayarlari.json"),
        ("REPLY_RUN_LOG_FILE", "eposta_yanit_calisma.log"), ("OPT_OUT_FILE", "ret_listesi.txt"),
    ):
        monkeypatch.setattr(config, name, tmp_path / "state" / file_name)
    (tmp_path / "state").mkdir()
    (tmp_path / "takip").mkdir()
    with config.OUTREACH_LOG_FILE.open("w", encoding="utf-8") as handle:
        for company, address, message_id in FIRMS:
            handle.write(json.dumps({
                "time": "2026-10-15T09:00:00+00:00", "date": "2026-10-15", "tracker": "Ornek_takip.xlsx",
                "company": company, "email": address, "message_id": message_id, "result": outreach.SENT,
            }) + "\n")
    outreach.new_tracker(tmp_path / "takip" / "Ornek_takip.xlsx", [
        {"Firma": company, "E-posta": address} for company, address, _id in FIRMS
    ], "TR")
    return tmp_path


def _mail(sender: str, body: str, **headers) -> bytes:
    message = EmailMessage()
    message["From"] = sender
    message["To"] = ACCOUNT
    message["Subject"] = headers.pop("subject", "Re: Fuar öncesi müşteri listesi")
    message["Date"] = headers.pop("date", "Fri, 16 Oct 2026 10:00:00 +0300")
    for name, value in headers.items():
        message[name.replace("_", "-")] = value
    message.set_content(body)
    return message.as_bytes()


def _bounce(address: str) -> bytes:
    return (
        "From: Mail Delivery System <MAILER-DAEMON@mx.ornek-posta.example>\r\n"
        f"To: {ACCOUNT}\r\nSubject: Undelivered Mail Returned to Sender\r\n"
        "Date: Fri, 16 Oct 2026 10:05:00 +0300\r\nMIME-Version: 1.0\r\n"
        'Content-Type: multipart/report; report-type=delivery-status; boundary="B"\r\n\r\n'
        "--B\r\nContent-Type: text/plain\r\n\r\nThe message could not be delivered.\r\n"
        "--B\r\nContent-Type: message/delivery-status\r\n\r\n"
        f"Reporting-MTA: dns; mx.ornek-posta.example\r\n\r\nFinal-Recipient: rfc822; {address}\r\nAction: failed\r\n"
        "--B--\r\n"
    ).encode()


MESSAGES = {
    "INBOX": {
        1: _mail("Alfa Ornek <info@alfa-ornek.com.tr>", "Merhaba,\nFiyatınız nedir?" + QUOTE,
                 In_Reply_To="<id-alfa@ornek-gonderen.example>"),
        2: _mail("Beta Satis <satis@beta-ornek.com.tr>", "Lütfen adresimizi listenizden çıkarın." + QUOTE),
        3: _mail("info@gama-ornek.com.tr", "Ofis dışındayım.", subject="Otomatik yanıt: Fuar öncesi",
                 Auto_Submitted="auto-replied"),
        4: _bounce("info@delta-ornek.com.tr"),
        5: _mail("bulten@baska-ornek.example", "Haftalık bülten."),
        6: _mail(f"Ornek Gonderen <{ACCOUNT}>", "Kendi e-postamız."),
    },
    "Spam": {},
}


class FakeImap:
    def __init__(self, folders):
        self.folders = folders
        self.current = None
        self.calls = []

    def login(self, user, password):
        self.calls.append(("login", user))

    def select(self, folder, readonly=False):
        assert readonly is True
        name = folder.strip('"')
        if name not in self.folders:
            return "NO", [b"no such folder"]
        self.current = name
        return "OK", [str(len(self.folders[name])).encode()]

    def response(self, code):
        return code, [b"7"]

    def uid(self, command, *args):
        self.calls.append((command, args))
        box = self.folders[self.current]
        if command == "SEARCH":
            if args[1] == "UID":
                start = int(args[2].split(":")[0])
                uids = [uid for uid in box if uid >= start] or ([max(box)] if box else [])
            else:
                uids = list(box)
            return "OK", [" ".join(str(uid) for uid in uids).encode()]
        if command == "FETCH":
            uid = int(args[0])
            assert "PEEK" in args[1]
            return "OK", [(f"{uid} (UID {uid} BODY[] {{1}}".encode(), box[uid]), b")"]
        raise AssertionError(command)

    def logout(self):
        self.calls.append(("logout",))


class FakeSmtp:
    sent = []
    fail = False

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def login(self, user, password):
        pass

    def send_message(self, message):
        if FakeSmtp.fail:
            raise OSError("down")
        FakeSmtp.sent.append(message)


@pytest.fixture
def smtp():
    FakeSmtp.sent = []
    FakeSmtp.fail = False
    return FakeSmtp


def _check(files, imap, **extra):
    return replies.check(
        account=ACCOUNT, password="synthetic-password", notify_to=OWNER, tracker_dir=files / "takip",
        imap_factory=lambda: imap, smtp_factory=FakeSmtp, **extra,
    )


@pytest.mark.parametrize("now, quiet", [
    (datetime(2026, 10, 16, 15, 30), True), (datetime(2026, 10, 16, 15, 19), False),
    (datetime(2026, 10, 16, 17, 15), False), (datetime(2026, 10, 17, 15, 30), False),
])
def test_quiet_window_on_weekdays_only(now, quiet):
    assert replies.quiet_now("15:20-17:15", now) is quiet
    assert replies.quiet_now("", now) is False


@pytest.mark.parametrize("body", [
    "Fiyat nedir?" + QUOTE,
    "Fiyat nedir?\n\n-----Original Message-----\nFrom: Ornek Gonderen\nistemiyorum",
    "Fiyat nedir?\n\nKimden: Ornek Gonderen\nKonu: Fuar\nistemiyorum",
    "Fiyat nedir?\n> istemiyorum",
])
def test_the_quoted_original_is_cut(body):
    message = replies.email.message_from_bytes(_mail("info@alfa-ornek.com.tr", body), policy=replies.policy.default)
    assert replies.reply_text(message, ACCOUNT) == "Fiyat nedir?"


def test_each_message_is_sorted_and_matched(files):
    index = replies.sent_index(outreach.load_log(), files / "takip")
    kinds = {}
    for uid, raw in MESSAGES["INBOX"].items():
        found = replies.classify(replies.email.message_from_bytes(raw, policy=replies.policy.default), index, ACCOUNT)
        kinds[uid] = (found["kind"], found["email"]) if found else None
    assert kinds == {
        1: (replies.REPLY, "info@alfa-ornek.com.tr"),
        2: (replies.OPT_OUT, "info@beta-ornek.com.tr"),
        3: (replies.AUTO, "info@gama-ornek.com.tr"),
        4: (replies.BOUNCE, "info@delta-ornek.com.tr"),
        5: None,
        6: None,
    }


def test_a_check_fills_the_workbook_the_opt_out_list_and_sends_notices_once(files, smtp):
    imap = FakeImap({name: dict(box) for name, box in MESSAGES.items()})
    result = _check(files, imap)
    assert result["yeni"] == 6 and result[replies.REPLY] == 1 and result[replies.OPT_OUT] == 1
    assert result[replies.AUTO] == 1 and result[replies.BOUNCE] == 1 and result["bildirim"] == 2
    assert not any(call[0] == "STORE" for call in imap.calls)
    sheet = load_workbook(files / "takip" / "Ornek_takip.xlsx").active
    columns = {cell.value: cell.column for cell in sheet[1]}
    rows = {sheet.cell(row=number, column=columns["E-posta"]).value: number for number in range(2, sheet.max_row + 1)}

    def value(address, column):
        return sheet.cell(row=rows[address], column=columns[column]).value

    assert value("info@alfa-ornek.com.tr", "Yanıt") == "Yanıt: Merhaba, Fiyatınız nedir?"
    assert value("info@alfa-ornek.com.tr", "Yanıt tarihi").date() == date(2026, 10, 16)
    assert value("info@beta-ornek.com.tr", "Yanıt").startswith("İstemiyor: Lütfen adresimizi")
    assert value("info@gama-ornek.com.tr", "Not") == replies.AUTO
    assert value("info@delta-ornek.com.tr", "Gönderim sonucu") == replies.BOUNCE
    assert "info@beta-ornek.com.tr" in config.OPT_OUT_FILE.read_text(encoding="utf-8")
    assert sorted(message["Subject"] for message in smtp.sent) == [
        "Yanıt geldi: Alfa Ornek", "İstemiyor (ret listesine eklendi): Beta Ornek",
    ]
    assert all(message["To"] == OWNER for message in smtp.sent)
    assert "Fiyatınız nedir?" in smtp.sent[0].get_content() or "Fiyatınız nedir?" in smtp.sent[1].get_content()

    again = _check(files, FakeImap({name: dict(box) for name, box in MESSAGES.items()}))
    assert again["yeni"] == 0 and again["bildirim"] == 0 and len(smtp.sent) == 2
    assert config.OPT_OUT_FILE.read_text(encoding="utf-8").count("info@beta-ornek.com.tr") == 1


def test_a_failed_notice_is_sent_on_the_next_run_and_old_replies_get_none(files, smtp):
    smtp.fail = True
    first = _check(files, FakeImap({name: dict(box) for name, box in MESSAGES.items()}))
    assert first["bildirim"] == 0 and first["bildirim hatası"] == 2
    smtp.fail = False
    second = _check(files, FakeImap({name: dict(box) for name, box in MESSAGES.items()}),
                    notify_since=date(2026, 10, 17))
    assert second["bildirim"] == 0 and smtp.sent == []
    third = _check(files, FakeImap({name: dict(box) for name, box in MESSAGES.items()}))
    assert third["bildirim"] == 2


def test_an_open_workbook_waits_for_the_next_run(files, smtp, monkeypatch):
    monkeypatch.setattr(outreach, "writable", lambda _path: False)
    result = _check(files, FakeImap({name: dict(box) for name, box in MESSAGES.items()}))
    assert result["açık dosya"] == 1
    monkeypatch.setattr(outreach, "writable", lambda _path: True)
    assert _check(files, FakeImap({name: dict(box) for name, box in MESSAGES.items()}))["yazılan"] == 2


def test_a_refused_login_explains_the_imap_setting(files):
    class Refusing(FakeImap):
        def login(self, user, password):
            raise replies.imaplib.IMAP4.error("AUTHENTICATIONFAILED")

    with pytest.raises(replies.ReplyError, match="IMAP erişimini"):
        _check(files, Refusing({}))


def test_the_tool_saves_settings_and_keeps_quiet_hours(files, monkeypatch, capsys):
    assert eposta_gonder.main(["yanitlar"]) == 1
    assert eposta_gonder.main([
        "yanit-ayar", "--hesap", ACCOUNT, "--bildirim", OWNER, "--klasor", str(files / "takip"), "--sessiz", "15:20-17:15",
    ]) == 0
    settings = replies.load_settings()
    assert settings["hesap"] == ACCOUNT and settings["sessiz"] == "15:20-17:15" and settings["baslangic"]
    assert config.REPLY_LOG_FILE.is_file() and config.REPLY_RUN_LOG_FILE.is_file() and config.REPLY_STATE_FILE.is_file()
    monkeypatch.setattr(replies, "quiet_now", lambda _window, _now: True)
    assert eposta_gonder.main(["yanitlar"]) == 0
    assert "Sessiz saat" in capsys.readouterr().out


def test_the_windows_task_runs_the_check_every_half_hour_without_a_window():
    arguments = replies.task_arguments()
    assert arguments[:2] == ["schtasks", "/Create"]
    assert arguments[arguments.index("/RI") + 1] == "30" and arguments[arguments.index("/DU") + 1] == "14:00"
    command = arguments[arguments.index("/TR") + 1]
    assert command.endswith('eposta_gonder.py" yanitlar') and len(command) < 262
    assert replies.task_arguments(remove=True)[:2] == ["schtasks", "/Delete"]
    seen = []

    class Done:
        returncode, stdout, stderr = 0, "SUCCESS", ""

    assert replies.schedule_task(runner=lambda args, **_kw: seen.append(args) or Done()) == (0, "SUCCESS")
    assert seen[0][1] == "/Create"

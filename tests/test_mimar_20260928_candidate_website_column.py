"""Talimat 15: sonuclar.xlsx shows the best unverified website candidate for manual review."""

from __future__ import annotations

from openpyxl import load_workbook

from modules import run_report


HEADER = "Aday web sitesi (kontrol edin)"


def test_candidate_column_follows_website_block():
    assert run_report.CANDIDATE_WEBSITE_HEADER == HEADER
    index = run_report.HEADERS.index(HEADER)
    assert run_report.HEADERS[index - 3:index] == ["Web sitesi", "Web kaynağı", "Web güven"]
    assert run_report.HEADERS.count(HEADER) == 1


def test_candidate_hidden_when_website_accepted():
    row = {"company": "Kabul", "website": "https://kabul.com.tr/", "candidate_1_url": "https://baska.com.tr/"}
    assert run_report._candidate_website(row) == ""


def test_candidate_skips_excluded_and_foreign_domains():
    row = {
        "company": "Deltakes",
        "website": "",
        "candidate_1_url": "https://www.facebook.com/deltakes",
        "candidate_2_url": "https://brobo.com.au/",
        "candidate_3_url": "http://deltakes.com.tr",
    }
    assert run_report._candidate_website(row) == "http://deltakes.com.tr"


def test_candidate_empty_when_only_excluded_or_missing():
    assert run_report._candidate_website({"company": "Yok", "website": ""}) == ""
    row = {"company": "Rehber", "website": "", "candidate_1_url": "https://www.firmatlas.com/x", "candidate_2_url": "https://kcnmarine.co.uk/"}
    assert run_report._candidate_website(row) == ""


def test_candidate_column_written_to_detail_sheet(tmp_path):
    rows = [
        {
            "company": "Aday Co", "website": "", "status": "REVIEW_NEEDED",
            "candidate_1_url": "https://aday.com.tr/", "email": "", "phone": "",
        },
        {
            "company": "Kabul Co", "website": "https://kabul.com.tr/", "website_source": "OWN_SEARCH",
            "status": "OK_HIGH_CONFIDENCE", "candidate_1_url": "https://kabul.com.tr/",
            "email": "info@kabul.com.tr", "email_source_tier": "SITE", "phone": "",
        },
    ]
    assert run_report.write_run_report(
        rows, output_root=tmp_path, run_status="TAMAMLANDI", status_detail="test", elapsed_seconds=1,
    ) == {}
    book = load_workbook(tmp_path / "sonuclar.xlsx", read_only=True, data_only=True)
    try:
        values = list(book["Detaylar"].iter_rows(values_only=True))
        column = values[0].index(HEADER)
        by_company = {row[0]: row[column] for row in values[1:]}
        assert by_company["Aday Co"] == "https://aday.com.tr/"
        if "Kabul Co" in by_company:
            assert by_company["Kabul Co"] in (None, "")
        contact_headers = next(book["İletişim"].iter_rows(values_only=True))
        assert HEADER not in contact_headers
    finally:
        book.close()

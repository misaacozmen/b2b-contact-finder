import tempfile
import unittest
import os
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import config
import main
from modules import excel, runtime, search
from modules.exhibitor_scraper import _metalexpo_list_rows, _texhibition_list_rows


class RunStateIsolationTests(unittest.TestCase):
    def test_cli_supports_non_interactive_mode(self):
        with patch("sys.argv", ["main.py", "--non-interactive"]):
            self.assertTrue(main.parse_args().non_interactive)

    def test_workbook_reader_rejects_oversized_expanded_archive(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "oversized.xlsx"
            workbook = Workbook()
            workbook.active.append(["company"])
            workbook.active.append(["A" * 1000])
            workbook.save(path)
            with patch.object(config, "MAX_WORKBOOK_UNCOMPRESSED_BYTES", 100):
                with self.assertRaisesRegex(ValueError, "expanded size limit"):
                    excel.read_company_records(path)

    def test_excel_writer_escapes_formula_prefixes_without_changing_other_types(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "safe.xlsx"
            excel.write_company_records(path, [{
                "company": "=SUM(1,1)",
                "website": "  +cmd|' /C calc'!A0",
                "listed_website": "\t-2+3",
                "source": " \n@SUM(A1:A2)",
                "country": "normal value",
                "hall": 42,
                "stand": True,
                "brands": "'=already-safe",
                "description": None,
            }])

            workbook = load_workbook(path, data_only=False)
            try:
                values = [cell.value for cell in workbook.active[2]]
                data_types = [cell.data_type for cell in workbook.active[2]]
            finally:
                workbook.close()

        self.assertEqual(values[0], "'=SUM(1,1)")
        self.assertEqual(values[1], "'  +cmd|' /C calc'!A0")
        self.assertEqual(values[2], "'\t-2+3")
        self.assertEqual(values[3], "' \n@SUM(A1:A2)")
        self.assertEqual(values[4], "normal value")
        self.assertEqual(values[10], 42)
        self.assertIs(values[11], True)
        self.assertEqual(values[12], "'=already-safe")
        self.assertIsNone(values[15])
        self.assertTrue(all(data_types[index] == "s" for index in range(5)))

    def test_run_restores_run_scoped_config_after_success(self):
        original = {
            name: getattr(config, name)
            for name in main._RUN_SCOPED_CONFIG_NAMES
        }

        def mutate_config(*_args, **_kwargs):
            for name in main._RUN_SCOPED_CONFIG_NAMES:
                setattr(config, name, object())
            return "done"

        with patch.object(main, "_run", side_effect=mutate_config):
            self.assertEqual(main.run(Path("input.xlsx")), "done")

        self.assertEqual(
            {name: getattr(config, name) for name in main._RUN_SCOPED_CONFIG_NAMES},
            original,
        )

    def test_run_restores_run_scoped_config_after_failure(self):
        original = {
            name: getattr(config, name)
            for name in main._RUN_SCOPED_CONFIG_NAMES
        }

        def mutate_then_fail(*_args, **_kwargs):
            for name in main._RUN_SCOPED_CONFIG_NAMES:
                setattr(config, name, object())
            raise RuntimeError("pipeline failed")

        with patch.object(main, "_run", side_effect=mutate_then_fail):
            with self.assertRaisesRegex(RuntimeError, "pipeline failed"):
                main.run(Path("input.xlsx"))

        self.assertEqual(
            {name: getattr(config, name) for name in main._RUN_SCOPED_CONFIG_NAMES},
            original,
        )

    def test_run_state_dir_rebinds_only_run_specific_state(self):
        with tempfile.TemporaryDirectory() as directory:
            state_dir = Path(directory) / "state"
            saved_keys = config.SAVED_API_KEYS_FILE
            resolver_settings = config.RESOLVER_SETTINGS_FILE
            with patch.multiple(
                config,
                STATE_DIR=config.STATE_DIR,
                PROGRESS_FILE=config.PROGRESS_FILE,
                PROGRESS_DB_FILE=config.PROGRESS_DB_FILE,
                SEARCH_CACHE_DIR=config.SEARCH_CACHE_DIR,
                CRAWL_CACHE_DIR=config.CRAWL_CACHE_DIR,
                EMAIL_CACHE_DIR=config.EMAIL_CACHE_DIR,
            ):
                main._set_run_state_dir(state_dir)
                self.assertEqual(config.PROGRESS_DB_FILE, state_dir / "progress.sqlite3")
                self.assertEqual(config.SEARCH_CACHE_DIR, state_dir / "search_cache")
                self.assertEqual(config.CRAWL_CACHE_DIR, state_dir / "crawl_cache")
                self.assertEqual(config.EMAIL_CACHE_DIR, state_dir / "email_cache")
                self.assertEqual(config.SAVED_API_KEYS_FILE, saved_keys)
                self.assertEqual(config.RESOLVER_SETTINGS_FILE, resolver_settings)

    def test_metalexpo_listing_website_is_discovery_only(self):
        html = """
        <a href="https://ornek-metal.com/">
          <div class="katilimci-text">
            <span class="text">ÖRNEK METAL</span>
            <span class="text">HALL 5 / 5A-12</span>
          </div>
        </a>
        """
        row = _metalexpo_list_rows(html)[0]
        self.assertEqual(row["website"], "")
        self.assertEqual(row["listed_website"], "https://ornek-metal.com/")
        self.assertEqual(row["hall"], "5")
        self.assertEqual(row["stand"], "5A-12")

        candidates = {}
        search._add_profile_candidates(candidates, row["company"], row)
        candidate = candidates["ornek-metal.com"]
        self.assertEqual(candidate["query"], "fair_listed_website")
        self.assertEqual(candidate["_source_profile_evidence"], 1)
        self.assertEqual(candidate["_official_query_evidence"], 0)

        brand_candidates = {}
        search._add_profile_candidates(
            brand_candidates,
            "N\u0130KEL PASLANMAZ",
            {"listed_website": "https://www.nikelpaslanmaz.com/"},
        )
        brand_candidate = brand_candidates["nikelpaslanmaz.com"]
        self.assertTrue(brand_candidate["_exact_brand_domain"])
        self.assertEqual(brand_candidate["_official_query_evidence"], 0)

    def test_company_record_round_trip_keeps_listing_provenance(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "firms.xlsx"
            rows = _metalexpo_list_rows("""
            <a href="https://ornek-metal.com/">
              <div class="katilimci-text">
                <span class="text">ÖRNEK METAL</span>
                <span class="text">HALL 5 / 5A-12</span>
              </div>
            </a>
            """)
            excel.write_company_records(path, rows)
            record = excel.read_company_records(path)[0]
            self.assertEqual(record["listed_website"], "https://ornek-metal.com/")
            self.assertIn("metalexpo.com.tr", record["listing_url"])

    def test_company_record_round_trip_keeps_fair_profile_contacts(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "firms.xlsx"
            excel.write_company_records(path, [{
                "company": "ÖRNEK GIDA",
                "listed_phone": "+90 212 555 00 00",
                "listed_email": "info@ornek.example",
                "listed_address": "İstanbul",
                "brands": "Örnek",
                "representations": "Example GmbH",
            }])
            record = excel.read_company_records(path)[0]
            self.assertEqual(record["listed_phone"], "+90 212 555 00 00")
            self.assertEqual(record["listed_email"], "info@ornek.example")
            self.assertEqual(record["listed_address"], "İstanbul")
            self.assertEqual(record["brands"], "Örnek")
            self.assertEqual(record["representations"], "Example GmbH")

    def test_country_column_is_not_used_as_website_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "firms.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["company", "country"])
            sheet.append(["ACME GIDA", "Türkiye"])
            workbook.save(path)

            record = excel.read_company_records(path)[0]

            self.assertEqual(record["website"], "")
            self.assertEqual(record["country"], "Türkiye")

    def test_texhibition_list_parser_keeps_profile_as_discovery(self):
        rows = _texhibition_list_rows("""
        <a href="https://www.texhibitionist.com/katilimcilar/ornek-tekstil">
          <div class="item">
            <div class="title">ÖRNEK TEKSTİL</div>
            <div class="category">Pamuk, Dokuma</div>
          </div>
        </a>
        """)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["website"], "")
        self.assertEqual(rows[0]["listed_website"], "")
        self.assertEqual(
            rows[0]["profile_url"],
            "https://www.texhibitionist.com/katilimcilar/ornek-tekstil",
        )
        self.assertEqual(rows[0]["sector"], "Pamuk, Dokuma")

    def test_runtime_budget_offset_survives_reset(self):
        with patch.dict(os.environ, {"BRIGHTDATA_REQUEST_OFFSET": "12"}):
            runtime.reset()
            self.assertEqual(
                runtime.snapshot()["counters"]["api.brightdata.requests"], 12
            )
        runtime.reset()

    def test_listed_domain_conflict_requires_strong_legal_identity(self):
        metadata = {"listed_website": "https://listed-brand.example"}
        weak = {
            "candidate": {"url": "https://different-brand.example"},
            "reasons": ["legal_name_phrase_missing:0/2"],
        }
        self.assertTrue(
            main._listed_domain_conflict_requires_review(
                "ÖRNEK METAL", weak, metadata
            )
        )
        strong = {
            **weak,
            "reasons": ["legal_name_ownership_match:2"],
        }
        self.assertFalse(
            main._listed_domain_conflict_requires_review(
                "ÖRNEK METAL", strong, metadata
            )
        )

    def test_weak_partial_brand_search_identity_requires_review(self):
        weak = {
            "candidate": {
                "url": "https://ornekgroup.example",
                "query": "ornek Turkiye official website",
            },
            "reasons": [
                "no_context_tokens",
                "legal_name_phrase_missing:0/3",
            ],
        }
        self.assertTrue(
            main._weak_search_identity_requires_review(
                "ÖRNEK DEMİR PROFİL", weak
            )
        )
        strong = {
            **weak,
            "reasons": ["no_context_tokens", "legal_name_phrase_match:3"],
        }
        self.assertFalse(
            main._weak_search_identity_requires_review(
                "ÖRNEK DEMİR PROFİL", strong
            )
        )


        resolved = {
            **weak,
            "_identity_resolution": "candidate_resolved_by_target_fingerprint",
        }
        self.assertTrue(
            main._weak_search_identity_requires_review(
                "ORNEK DEMIR PROFIL", resolved
            )
        )

    def test_discovery_route_wins_only_on_stronger_first_party_legal_identity(self):
        listed = {
            "candidate": {
                "url": "https://ornek-metal.example",
                "_source_profile_evidence": 1,
            },
            "reasons": ["legal_name_phrase_match:2"],
            "identity_assessment": {
                "provisionally_publishable": True,
                "conflicts": [],
            },
        }
        weak_search = {
            "candidate": {"url": "https://ornek-makine.example"},
            "reasons": ["legal_name_phrase_missing:0/2"],
            "identity_assessment": {
                "provisionally_publishable": True,
                "conflicts": [],
            },
        }
        self.assertIs(
            main._preferred_verified_discovery_route([weak_search, listed]),
            listed,
        )

        equally_verified = {
            **weak_search,
            "reasons": ["legal_name_phrase_match:2"],
        }
        self.assertIsNone(
            main._preferred_verified_discovery_route([listed, equally_verified])
        )

        conflicted = {
            **listed,
            "identity_assessment": {
                "provisionally_publishable": True,
                "conflicts": ["country_conflict"],
            },
        }
        self.assertIsNone(
            main._preferred_verified_discovery_route([conflicted, weak_search])
        )

    def test_public_body_domains_are_not_company_candidates(self):
        for url in (
            "https://ornek.gov.tr",
            "https://ornek.bel.tr",
            "https://ornek.meb.k12.tr",
            "https://ornek.edu.tr",
        ):
            self.assertEqual(
                search._candidate_role("Ã–RNEK METAL", url, "Ã–rnek Metal", ""),
                "public_body",
            )

    def test_fair_and_news_hosts_override_accidental_name_similarity(self):
        self.assertEqual(
            search._candidate_role(
                "CCN FOOD",
                "https://foodistexpo.com/brand/ccn-food",
                "Foodist Gida Fuari",
                "Salon 11 Stant 1132",
            ),
            "fair_profile",
        )
        self.assertEqual(
            search._candidate_role(
                "ELBISTAN DOGAN",
                "https://elbistaninsesi.com/haber/elbistan-dogan",
                "Elbistanin Sesi Haber",
                "Yerel haber",
            ),
            "news",
        )
    def test_language_subdomain_is_not_a_homonym(self):
        first = {"candidate": {"url": "https://www.example.com"}}
        second = {"candidate": {"url": "https://en.example.com"}}
        result = main._homonym_conflict("Example", first, second)
        self.assertFalse(result["ambiguous"])
        self.assertEqual(result["reason"], "same_registrable_domain")

    def test_pipeline_runner_restores_paid_query_limit_on_success(self):
        from modules import pipeline_runner
        search._RUN_PAID_QUERY_LIMIT = 42
        try:
            with patch.object(pipeline_runner, "_run_pipeline_impl", return_value="success"):
                res = pipeline_runner.run_pipeline(
                    Path("firms.xlsx"),
                    process_company_fn=lambda *a: None,
                    write_outputs_fn=lambda *a: "ok",
                    set_output_dir_fn=lambda *a: None,
                    empty_result_fn=lambda *a: {},
                )
                self.assertEqual(res, "success")
                self.assertEqual(search._RUN_PAID_QUERY_LIMIT, 42)
        finally:
            search._RUN_PAID_QUERY_LIMIT = None

    def test_pipeline_runner_restores_paid_query_limit_on_exception(self):
        from modules import pipeline_runner
        search._RUN_PAID_QUERY_LIMIT = 42
        try:
            with patch.object(pipeline_runner, "_run_pipeline_impl", side_effect=RuntimeError("boom")):
                with self.assertRaises(RuntimeError):
                    pipeline_runner.run_pipeline(
                        Path("firms.xlsx"),
                        process_company_fn=lambda *a: None,
                        write_outputs_fn=lambda *a: "ok",
                        set_output_dir_fn=lambda *a: None,
                        empty_result_fn=lambda *a: {},
                    )
                self.assertEqual(search._RUN_PAID_QUERY_LIMIT, 42)
        finally:
            search._RUN_PAID_QUERY_LIMIT = None

    def test_pipeline_runner_restores_paid_query_limit_on_keyboard_interrupt(self):
        from modules import pipeline_runner
        search._RUN_PAID_QUERY_LIMIT = 42
        try:
            with patch.object(pipeline_runner, "_run_pipeline_impl", side_effect=KeyboardInterrupt):
                with self.assertRaises(KeyboardInterrupt):
                    pipeline_runner.run_pipeline(
                        Path("firms.xlsx"),
                        process_company_fn=lambda *a: None,
                        write_outputs_fn=lambda *a: "ok",
                        set_output_dir_fn=lambda *a: None,
                        empty_result_fn=lambda *a: {},
                    )
                self.assertEqual(search._RUN_PAID_QUERY_LIMIT, 42)
        finally:
            search._RUN_PAID_QUERY_LIMIT = None

    def test_pipeline_runner_passes_callbacks_by_identity(self):
        from modules import pipeline_runner
        proc = lambda *a: None
        write = lambda *a: "ok"
        set_dir = lambda *a: None
        empty = lambda *a: {}
        with patch.object(pipeline_runner, "_run_pipeline_impl", return_value="ok") as mock_impl:
            pipeline_runner.run_pipeline(
                Path("firms.xlsx"),
                process_company_fn=proc,
                write_outputs_fn=write,
                set_output_dir_fn=set_dir,
                empty_result_fn=empty,
            )
            mock_impl.assert_called_once()
            kwargs = mock_impl.call_args.kwargs
            self.assertIs(kwargs["process_company_fn"], proc)
            self.assertIs(kwargs["write_outputs_fn"], write)
            self.assertIs(kwargs["set_output_dir_fn"], set_dir)
            self.assertIs(kwargs["empty_result_fn"], empty)


if __name__ == "__main__":
    unittest.main()

"""Offline, source-ID-only reconciliation of a baseline and remediation run."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from modules import excel, output_artifacts, publication_policy, redaction, run_context


FAILURE_STATUSES = frozenset({
    "PROCESSING_FAILED", "SEARCH_FAILED", "WEBSITE_FETCH_FAILED",
    "WEBSITE_NOT_FOUND", "UNKNOWN", "FAILED", "BLOCKED_BUDGET",
})
CONFLICT_MARKERS = ("collision", "conflict", "wrong_owner", "different_owner", "homonym")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _rows(path: Path) -> list[dict[str, Any]]:
    path = Path(path)
    if path.suffix.casefold() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()
        if not values:
            return []
        headers = [str(value or "").strip() for value in values[0]]
        return [
            {headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))}
            for row in values[1:]
            if any(value not in (None, "") for value in row)
        ]
    if path.suffix.casefold() == ".jsonl":
        return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload if isinstance(payload, list) else payload.get("records", payload.get("rows", []))


def _by_source(path: Path | None) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    if path is None:
        return result
    for row in _rows(Path(path)):
        source_id = str(row.get("source_record_id", "") or "").strip()
        if not source_id:
            raise ValueError(f"result row has no source_record_id: {path}")
        if source_id in result:
            raise ValueError(f"duplicate source_record_id in {path}: {source_id}")
        result[source_id] = dict(row)
    return result


def _evidence_by_source(path: Path | None) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    if path is None:
        return result
    for raw in Path(path).read_text(encoding="utf-8").splitlines():
        if not raw.strip():
            continue
        record = json.loads(raw)
        source_id = str(record.get("source_record_id", "") or "").strip()
        if not source_id:
            raise ValueError(f"evidence row has no source_record_id: {path}")
        if source_id in result:
            raise ValueError(f"duplicate evidence source_record_id: {source_id}")
        decision = record.get("publication_decision")
        if not isinstance(decision, dict):
            raise ValueError(f"evidence has no frozen publication_decision: {source_id}")
        publication_policy.verify_publication_decision(
            decision,
            source_record_id=source_id,
            run_id=str(record.get("run_id", "")) or None,
            config_sha256=str(record.get("config_sha256", "")) or None,
        )
        if str(decision.get("source_record_id", "")) != source_id:
            raise ValueError(f"evidence decision source ID mismatch: {source_id}")
        result[source_id] = record
    return result


def _manifest_file_entry(payload: dict[str, Any], path: Path) -> dict[str, Any]:
    files = payload.get("files")
    if not isinstance(files, dict):
        raise ValueError("bundle manifest must contain file hashes")
    for key in (path.name, str(path), str(path.resolve())):
        entry = files.get(key)
        if isinstance(entry, dict) and entry.get("sha256"):
            return entry
    raise ValueError(f"bundle manifest has no hash for {path.name}")


def _validate_manifest(
    path: Path | None,
    source_ids: list[str],
    result_path: Path | None,
    evidence_path: Path | None,
    *,
    label: str,
) -> dict[str, Any]:
    if path is None or result_path is None or evidence_path is None:
        raise ValueError(f"{label} results, evidence and manifest are required together")
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"{label} manifest must be an object")
    listed = payload.get("source_record_ids") or payload.get("ordered_source_record_ids")
    if list(listed or []) != list(source_ids):
        raise ValueError(f"{label} manifest source_record_id order mismatch: {path}")
    for artifact_path in (Path(result_path), Path(evidence_path)):
        entry = _manifest_file_entry(payload, artifact_path)
        if entry.get("sha256") != _sha256(artifact_path):
            raise ValueError(f"{label} manifest hash mismatch: {artifact_path.name}")
    return payload


def _hydrate_bundle(
    rows: dict[str, dict[str, Any]],
    evidence: dict[str, dict[str, Any]],
    *,
    label: str,
) -> dict[str, dict[str, Any]]:
    if not rows:
        return rows
    if not evidence:
        raise ValueError(f"{label} evidence bundle is required; no publication decision may be reconstructed")
    missing = set(rows) - set(evidence)
    extra = set(evidence) - set(rows)
    if missing or extra:
        raise ValueError(f"{label} evidence/source ID coverage mismatch: missing={len(missing)} extra={len(extra)}")
    hydrated = {}
    for source_id, row in rows.items():
        decision = evidence[source_id]["publication_decision"]
        if str(decision.get("run_id", "")) and str(row.get("run_id", "")) and str(decision["run_id"]) != str(row["run_id"]):
            raise ValueError(f"{label} decision run ID mismatch: {source_id}")
        hydrated_row = dict(row)
        publication_policy.apply_frozen_decision_fields(hydrated_row, decision)
        hydrated_row["source_evidence"] = evidence[source_id].get("source_evidence", evidence[source_id].get("field_evidence", []))
        hydrated[source_id] = hydrated_row
    return hydrated


def _is_failure(row: dict[str, Any]) -> bool:
    status = str(row.get("status", "")).strip().upper()
    reason = str(row.get("reason", "")).casefold()
    return status in FAILURE_STATUSES or any(marker in reason for marker in ("processing_failed", "provider_failed"))


def _is_conflict(row: dict[str, Any]) -> bool:
    text = " ".join(str(row.get(key, "") or "") for key in ("publication_blockers", "collision_reason", "reason")).casefold()
    return any(marker in text for marker in CONFLICT_MARKERS)


def _contact_count(row: dict[str, Any]) -> int:
    return sum(bool(str(row.get(field, "") or "").strip()) for field in ("email", "phone"))


def _choose(source_id: str, baseline: dict[str, Any] | None, child: dict[str, Any] | None) -> tuple[dict[str, Any], str]:
    if baseline is None and child is None:
        return {"source_record_id": source_id, "status": "REVIEW_NEEDED", "reason": "missing_result"}, "missing_result"
    if baseline is None:
        return dict(child), "child_only"
    if child is None or _is_failure(child):
        return dict(baseline), "baseline_retained_after_failed_child"
    if _is_conflict(child):
        return dict(child), "child_conflict_blocks_baseline"
    child_publishable = bool((child.get("publication_decision") or {}).get("publishable"))
    baseline_publishable = bool((baseline.get("publication_decision") or {}).get("publishable"))
    if child_publishable and not baseline_publishable:
        return dict(child), "child_verified_result"
    if baseline_publishable and not child_publishable:
        return dict(baseline), "baseline_retained_against_nonverified_child"
    child_score = (_contact_count(child), int(child.get("score", 0) or 0), bool(child.get("website")))
    base_score = (_contact_count(baseline), int(baseline.get("score", 0) or 0), bool(baseline.get("website")))
    return (dict(child), "child_improves_evidence") if child_score > base_score else (dict(baseline), "baseline_retained")


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(redaction.sanitize(payload), ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")


def _coverage_rows(final_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    coverage = []
    for row in sorted(final_rows, key=lambda item: int(item.get("original_index", 0))):
        decision = publication_policy.verify_publication_decision(row["publication_decision"])
        status = str(row.get("status", "")).upper()
        if decision["publishable"]:
            state, next_action = "published", "none"
        elif _is_failure(row):
            state, next_action = "failed", "retry_or_manual_review"
        elif status in {"PENDING", "PENDING_PAID", "PENDING_FREE_RETRY", "UNKNOWN", "BLOCKED_BUDGET"} or str(row.get("paid_state", "")).upper() in {"PENDING", "UNKNOWN", "BLOCKED_BUDGET"}:
            state, next_action = "pending", "bounded_remediation"
        else:
            state, next_action = "review", "manual_review"
        coverage.append({
            "source_record_id": str(row["source_record_id"]),
            "original_index": int(row.get("original_index", 0)),
            "company": row.get("company", ""),
            "status": state,
            "blockers": list(decision["blockers"]),
            "next_action": next_action,
            "publication_decision_sha256": decision["publication_decision_sha256"],
        })
    return coverage


def reconcile_runs(
    *,
    input_path: Path,
    destination: Path,
    baseline_results: Path | None = None,
    child_results: Path | None = None,
    baseline_evidence: Path | None = None,
    child_evidence: Path | None = None,
    baseline_manifest: Path | None = None,
    child_manifest: Path | None = None,
    remediation_plan: Path | None = None,
) -> dict[str, Any]:
    input_path, destination = Path(input_path).resolve(), Path(destination).resolve()
    source_records = excel.read_company_records(input_path)
    source_ids: list[str] = []
    for index, record in enumerate(source_records):
        source_id = str(record.get("source_record_id", "") or "").strip()
        if not source_id:
            raise ValueError(f"original input source_record_id is empty at row {index + 2}")
        source_ids.append(source_id)
        record["source_record_id"] = source_id
        record["original_index"] = index
    if len(source_ids) != len(set(source_ids)):
        raise ValueError("input source_record_id values are not unique")
    if len(source_ids) != 893:
        raise ValueError("original input must contain exactly 893 source_record_id values")
    if baseline_results is None or baseline_evidence is None or baseline_manifest is None:
        raise ValueError("baseline results, evidence and manifest are required")
    child_values = (child_results, child_evidence, child_manifest)
    if any(value is not None for value in child_values) and not all(value is not None for value in child_values):
        raise ValueError("child results, evidence and manifest are required together")
    baseline = _by_source(baseline_results)
    baseline_order = [str(row.get("source_record_id", "") or "").strip() for row in _rows(baseline_results)]
    if baseline_order != source_ids:
        raise ValueError("baseline must contain all original source IDs in original order")
    baseline_manifest_payload = _validate_manifest(
        baseline_manifest, source_ids, baseline_results, baseline_evidence, label="baseline",
    )
    baseline_evidence_rows = _evidence_by_source(baseline_evidence)
    baseline_review_ids = [
        source_id for source_id in source_ids
        if not bool(baseline_evidence_rows[source_id]["publication_decision"].get("publishable"))
    ]
    child = _by_source(child_results)
    child_order = [str(row.get("source_record_id", "") or "").strip() for row in _rows(child_results)] if child_results is not None else []
    child_manifest_payload = {}
    plan_ids: list[str] = []
    if remediation_plan is not None:
        plan_payload = json.loads(Path(remediation_plan).read_text(encoding="utf-8"))
        if not isinstance(plan_payload, dict):
            raise ValueError("remediation plan must be an object")
        plan_ids = [str(task.get("source_record_id", "") or "").strip() for task in plan_payload.get("tasks", [])]
        if not plan_ids or len(set(plan_ids)) != len(plan_ids):
            raise ValueError("remediation plan must contain unique nonempty task IDs")
        if plan_ids != baseline_review_ids:
            raise ValueError("remediation plan must exactly match verified baseline review IDs in original order")
    if child_results is not None:
        expected_child_ids = plan_ids or child_order
        if len(child_order) != len(set(child_order)) or child_order != expected_child_ids:
            raise ValueError("child must match remediation plan order")
        if not set(child_order).issubset(set(source_ids)):
            raise ValueError("child IDs must be a subset of original IDs")
        child_manifest_payload = _validate_manifest(
            child_manifest, child_order, child_results, child_evidence, label="child",
        )
    elif plan_ids:
        raise ValueError("remediation plan requires a child bundle")
    child_evidence_rows = _evidence_by_source(child_evidence)
    baseline = _hydrate_bundle(baseline, baseline_evidence_rows, label="baseline")
    child = _hydrate_bundle(child, child_evidence_rows, label="child")
    reconciliation_run_id = hashlib.sha256(_canonical({
        "input_sha256": _sha256(input_path), "baseline_sha256": _sha256(baseline_results) if baseline_results else None,
        "child_sha256": _sha256(child_results) if child_results else None, "source_ids": source_ids,
    }).encode("utf-8")).hexdigest()
    final_rows: list[dict[str, Any]] = []
    lineage: list[dict[str, Any]] = []
    for index, record in enumerate(source_records):
        source_id = source_ids[index]
        selected, selection_reason = _choose(source_id, baseline.get(source_id), child.get(source_id))
        selected["company"] = record.get("company", selected.get("company", ""))
        selected["source_record_id"] = source_id
        selected["original_index"] = index
        selected.setdefault("free_state", "DONE")
        selected.setdefault("paid_required", False)
        selected.setdefault("paid_state", "NOT_REQUIRED")
        selected.setdefault("status", "REVIEW_NEEDED")
        decision = selected.get("publication_decision")
        if not isinstance(decision, dict):
            # A source-only result without its evidence bundle is never
            # eligible for reconciliation.  _hydrate_bundle normally makes
            # this unreachable; keep the guard explicit for callers.
            raise ValueError(f"missing frozen publication_decision: {source_id}")
        selected["run_id"] = str(decision.get("run_id", "") or reconciliation_run_id)
        selected["reconciliation_run_id"] = reconciliation_run_id
        if _is_conflict(child.get(source_id, {})):
            decision = publication_policy.with_blocker(decision, "reconciliation_new_conflict")
            selected["publication_blockers"] = "; ".join(decision["blockers"])
            publication_policy.apply_frozen_decision_fields(selected, decision)
            output_artifacts.suppress_all_contacts(selected, "reconciliation_new_conflict")
            selected["delivery_state"] = "REVIEW"
        else:
            selected["delivery_state"] = "DELIVERABLE" if decision.get("publishable") is True else "REVIEW"
        final_rows.append(selected)
        before = baseline.get(source_id, {})
        after_decision = {
            "publishable": bool(decision.get("publishable", False)),
            "blockers": decision.get("blockers", []),
        }
        if _canonical({"before": bool((before.get("publication_decision") or {}).get("publishable")), "after": after_decision, "reason": selection_reason}) != _canonical({"before": after_decision, "after": after_decision, "reason": "baseline_retained"}):
            lineage.append({
                "source_record_id": source_id,
                "original_index": index,
                "selection": selection_reason,
                "before": {"publication_eligible": bool((before.get("publication_decision") or {}).get("publishable")), "website": before.get("website", ""), "email": bool(before.get("email")), "phone": bool(before.get("phone"))},
                "after": after_decision,
                "child_evidence": child.get(source_id, {}).get("source_evidence", child.get(source_id, {}).get("__evaluation", {})),
            })
    destination.mkdir(parents=True, exist_ok=False)
    coverage_rows = _coverage_rows(final_rows)
    if len(coverage_rows) != len(source_ids) or {row["source_record_id"] for row in coverage_rows} != set(source_ids):
        raise RuntimeError("reconciliation coverage is not complete")
    coverage = {
        "schema_version": 2,
        "run_id": reconciliation_run_id,
        "source_count": len(coverage_rows),
        "statuses": {state: sum(row["status"] == state for row in coverage_rows) for state in ("published", "review", "pending", "failed")},
        "source_record_ids": [row["source_record_id"] for row in coverage_rows],
        "records": coverage_rows,
    }
    coverage["source_record_ids_sha256"] = hashlib.sha256(_canonical(coverage["source_record_ids"]).encode("utf-8")).hexdigest()
    coverage["coverage_sha256"] = hashlib.sha256(_canonical(coverage_rows).encode("utf-8")).hexdigest()
    excel.write_contacts(destination / "all_results.xlsx", final_rows)
    contacts = [row for row in final_rows if publication_policy.frozen_publishable(row, require=True)]
    review = [row for row in final_rows if not publication_policy.frozen_publishable(row, require=True)]
    excel.write_contacts(destination / "contacts.xlsx", contacts)
    excel.write_contacts(destination / "review_queue.xlsx", review)
    from modules import evidence
    evidence.write_jsonl(destination / "evidence.jsonl", final_rows)
    _write_json(destination / "discovery_coverage.json", coverage)
    report = {
        "schema_version": 1,
        "status": "offline_reconciliation",
        "run_id": reconciliation_run_id,
        "input_count": len(final_rows), "published_count": len(contacts), "review_count": len(review),
        "provider_calls_new": 0, "physical_http_requests_new": 0,
        "elapsed_seconds": None, "cost": None,
        "unknown_historical_cost": True,
        "lineage_change_count": len(lineage),
        "coverage_sha256": coverage["coverage_sha256"],
        "source_record_ids_sha256": coverage["source_record_ids_sha256"],
    }
    _write_json(destination / "reconciliation_report.json", report)
    _write_json(destination / "reconciliation_lineage.json", lineage)
    files = {
        path.name: {"sha256": _sha256(path), "bytes": path.stat().st_size}
        for path in sorted(destination.iterdir()) if path.is_file()
    }
    manifest = {
        "schema_version": 1, "run_id": reconciliation_run_id,
        "coverage_complete": len(final_rows) == len(source_ids),
        "source_record_ids_sha256": coverage["source_record_ids_sha256"],
        "coverage_sha256": coverage["coverage_sha256"],
        "source_record_ids": source_ids,
        "counts": {"all_results": len(final_rows), "contacts": len(contacts), "review": len(review)},
        "parent_results": {"path": str(baseline_results) if baseline_results else None, "sha256": _sha256(baseline_results) if baseline_results else None, "manifest": str(baseline_manifest) if baseline_manifest else None, "manifest_sha256": _sha256(baseline_manifest) if baseline_manifest else None},
        "child_results": {"path": str(child_results) if child_results else None, "sha256": _sha256(child_results) if child_results else None, "manifest": str(child_manifest) if child_manifest else None, "manifest_sha256": _sha256(child_manifest) if child_manifest else None},
        "new_provider_calls": 0, "new_physical_http_requests": 0,
        "files": files,
    }
    _write_json(destination / "delivery_manifest.json", manifest)
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", "--original-input", dest="input_path", type=Path, required=True)
    parser.add_argument("--baseline", "--baseline-results", dest="baseline_results", type=Path)
    parser.add_argument("--child", "--child-results", dest="child_results", type=Path)
    parser.add_argument("--baseline-evidence", type=Path)
    parser.add_argument("--child-evidence", type=Path)
    parser.add_argument("--baseline-manifest", type=Path)
    parser.add_argument("--child-manifest", type=Path)
    parser.add_argument("--remediation-plan", type=Path)
    parser.add_argument("--destination", type=Path, required=True)
    args = parser.parse_args(argv)
    print(json.dumps(reconcile_runs(**vars(args)), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

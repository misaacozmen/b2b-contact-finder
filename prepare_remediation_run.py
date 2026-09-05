"""Prepare a source-only child run for every baseline review record.

The command is intentionally network-free.  It creates a new input package
from the original workbook and writes no previous result/contact value into
that package.  Baseline results are used only to explain why a task exists.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

import config
from modules import excel, run_context, scorer


SOURCE_ONLY_FIELDS = (
    "company", "source_record_id", "original_index", "source", "country",
    "profile_url", "listing_url", "listed_website", "listed_phone", "listed_email",
    "listed_address", "listed_legal_name", "source_detail_status", "source_detail_url",
    "source_detail_content_sha256", "hall", "stand", "brands", "representations",
    "sector", "description",
)
CONFLICT_MARKERS = (
    "collision", "conflict", "cross_entity", "homonym", "wrong_owner",
    "different_owner", "identity_conflict",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _read_table(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not rows:
        raise ValueError(f"empty workbook: {path}")
    headers = [str(value or "").strip() for value in rows[0]]
    folded_headers = [header.casefold() for header in headers]
    if len(folded_headers) != len(set(folded_headers)):
        raise ValueError(f"case-insensitive duplicate headers: {path}")
    return headers, [
        {headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))}
        for row in rows[1:]
        if any(value not in (None, "") for value in row)
    ]


def _read_records(path: Path) -> tuple[list[str], list[dict[str, Any]], list[dict[str, Any]]]:
    headers, raw_rows = _read_table(path)
    if len({header.casefold() for header in headers}) != len(headers):
        raise ValueError("case-insensitive duplicate input headers")
    records = excel.read_company_records(path)
    if len(records) != len(raw_rows):
        raise ValueError("input workbook row parsing mismatch")
    for index, (record, raw) in enumerate(zip(records, raw_rows)):
        record["original_index"] = index
        source_id = str(record.get("source_record_id", "") or "").strip()
        if not source_id:
            raise ValueError(f"input source_record_id is empty at row {index + 2}")
        raw["source_record_id"] = source_id
    ids = [str(record["source_record_id"]) for record in records]
    if len(ids) != len(set(ids)):
        raise ValueError("input source_record_id values are not unique")
    return headers, records, raw_rows


def _load_result_rows(path: Path) -> tuple[list[str], dict[str, dict[str, Any]]]:
    path = Path(path)
    if path.suffix.casefold() == ".xlsx":
        _headers, rows = _read_table(path)
    elif path.suffix.casefold() == ".jsonl":
        rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    else:
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload if isinstance(payload, list) else payload.get("records", payload.get("rows", []))
    result: dict[str, dict[str, Any]] = {}
    ordered: list[str] = []
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError(f"baseline result row is not an object: {path}")
        source_id = str(row.get("source_record_id", "") or "").strip()
        if not source_id:
            raise ValueError(f"baseline result source_record_id is empty: {path}")
        if source_id in result:
            raise ValueError(f"duplicate baseline source_record_id: {source_id}")
        result[source_id] = row
        ordered.append(source_id)
    return ordered, result


def _load_frozen_evidence(path: Path) -> tuple[list[str], dict[str, dict[str, Any]]]:
    from modules import publication_policy

    result: dict[str, dict[str, Any]] = {}
    ordered: list[str] = []
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        if not isinstance(row, dict):
            raise ValueError(f"baseline evidence row is not an object: {path}")
        source_id = str(row.get("source_record_id", "") or "").strip()
        if not source_id or source_id in result:
            raise ValueError(f"baseline evidence has empty or duplicate source_record_id: {path}")
        decision = row.get("publication_decision")
        if not isinstance(decision, dict):
            raise ValueError(f"baseline evidence has no frozen publication_decision: {source_id}")
        publication_policy.verify_publication_decision(decision, source_record_id=source_id)
        result[source_id] = row
        ordered.append(source_id)
    return ordered, result


def _manifest_file_hash(manifest: dict[str, Any], path: Path) -> str:
    files = manifest.get("files")
    if not isinstance(files, dict):
        raise ValueError("baseline manifest has no file hash map")
    for key in (path.name, str(path), str(path.resolve())):
        entry = files.get(key)
        if isinstance(entry, dict) and entry.get("sha256"):
            return str(entry["sha256"])
    raise ValueError(f"baseline manifest has no hash for {path.name}")


def _validate_baseline_bundle(
    *, input_ids: list[str], results_path: Path, evidence_path: Path, manifest_path: Path,
) -> tuple[dict[str, Any], dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    result_order, results = _load_result_rows(results_path)
    evidence_order, evidence = _load_frozen_evidence(evidence_path)
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    if not isinstance(manifest, dict):
        raise ValueError("baseline manifest must be an object")
    manifest_ids = list(manifest.get("source_record_ids") or manifest.get("ordered_source_record_ids") or [])
    if manifest_ids != input_ids:
        raise ValueError("baseline manifest must contain the original ordered source IDs")
    if result_order != input_ids or evidence_order != input_ids:
        raise ValueError("baseline results/evidence must cover all original IDs in the same order")
    for path in (results_path, evidence_path):
        if _manifest_file_hash(manifest, path) != _sha256(path):
            raise ValueError(f"baseline bundle hash mismatch: {path.name}")
    from modules import publication_policy
    for source_id in input_ids:
        decision = evidence[source_id]["publication_decision"]
        result_decision = results[source_id].get("publication_decision")
        if isinstance(result_decision, dict):
            publication_policy.verify_publication_decision(result_decision, source_record_id=source_id)
            if result_decision["publication_decision_sha256"] != decision["publication_decision_sha256"]:
                raise ValueError(f"baseline result/evidence decision mismatch: {source_id}")
    return manifest, results, evidence


def _split_reasons(row: dict[str, Any]) -> list[str]:
    values: list[str] = []
    for key in ("publication_blockers", "reason", "last_error", "collision_reason"):
        raw = row.get(key, "")
        if isinstance(raw, list):
            values.extend(str(value).strip() for value in raw)
        else:
            values.extend(value.strip() for value in str(raw or "").replace(",", ";").split(";"))
    return list(dict.fromkeys(value for value in values if value))


def _gaps(row: dict[str, Any], blockers: list[str], record: dict[str, Any]) -> list[str]:
    explicit = row.get("evidence_gaps", [])
    if isinstance(explicit, str):
        explicit = explicit.replace(",", ";").split(";")
    if explicit:
        return list(dict.fromkeys(str(value).strip() for value in explicit if str(value).strip()))
    gaps: list[str] = []
    joined = " ".join(blockers).casefold()
    if any(marker in joined for marker in ("legal", "ownership", "identity", "company_name")):
        gaps.append("missing_legal_identity")
    if any(marker in joined for marker in ("context", "sector", "relationship")):
        gaps.append("missing_context")
    if any(marker in joined for marker in ("contact", "email", "phone")):
        gaps.append("missing_contact")
    if any(marker in joined for marker in ("unreachable", "fetch", "timeout", "http_")):
        gaps.append("unreachable_candidates")
    if not any(str(record.get(key, "") or "").strip() for key in ("listed_legal_name", "brands", "sector", "description")):
        gaps.append("missing_metadata")
    return list(dict.fromkeys(gaps or ["missing_identity_coherence"]))


def _has_conflict(blockers: list[str]) -> bool:
    joined = " ".join(blockers).casefold()
    return any(marker in joined for marker in CONFLICT_MARKERS)


def _task(record: dict[str, Any], baseline: dict[str, Any] | None) -> dict[str, Any]:
    baseline = baseline or {}
    blockers = _split_reasons(baseline)
    gaps = _gaps(baseline, blockers, record)
    selected_domain = scorer.normalize_domain(
        str(baseline.get("website") or baseline.get("selected_website") or record.get("website") or "")
    )
    if _has_conflict(blockers):
        next_action = "alternative_candidate_research"
        capability, provider, max_requests = "identity_resolution", "ddgs", 4
        terminal_reason = "confirmed_identity_conflict_requires_alternative_candidate"
    elif "missing_metadata" in gaps or not any(str(record.get(key, "") or "").strip() for key in ("listed_legal_name", "brands", "sector", "description")):
        next_action = "metadata_acquisition"
        capability, provider, max_requests = "source_metadata", "organizer", 0
        terminal_reason = ""
    elif selected_domain:
        next_action = "first_party_targeted_evidence"
        capability, provider, max_requests = "static_http", "first_party", 0
        terminal_reason = ""
    elif any(marker in " ".join(blockers).casefold() for marker in ("javascript", "js_shell", "pdf", "ocr", "render")):
        next_action = "local_js_or_pdf_recovery"
        capability, provider, max_requests = "browser_or_ocr", "local", 0
        terminal_reason = ""
    elif any(gap in gaps for gap in ("no_candidates", "unreachable_candidates", "missing_identity_coherence")):
        next_action = "reserved_targeted_search"
        capability, provider, max_requests = "free_search", "ddgs", 4
        terminal_reason = ""
    else:
        next_action = "approved_external_provider"
        capability, provider, max_requests = "approved_external", "none", 0
        terminal_reason = "awaiting_explicit_provider_approval"
    return {
        "source_record_id": str(record["source_record_id"]),
        "original_index": int(record["original_index"]),
        "selected_domain": selected_domain,
        "blocking_reasons": blockers or ["baseline_review_required"],
        "evidence_gaps": gaps,
        "next_action": next_action,
        "required_capability": capability,
        "provider": provider,
        "max_requests": int(max_requests),
        "terminal_reason": terminal_reason,
    }


def _write_child_input(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    del headers
    output_headers = list(SOURCE_ONLY_FIELDS)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(output_headers)
    for row in rows:
        source_id = str(row.get("source_record_id", "") or "").strip()
        if not source_id:
            raise ValueError("child input source_record_id is empty")
        sheet.append([row.get(header, "") for header in output_headers])
    workbook.save(path)
    workbook.close()


def prepare_remediation_run(
    *,
    input_path: Path,
    destination: Path,
    baseline_results: Path,
    baseline_evidence: Path,
    baseline_manifest: Path,
    parent_manifest: Path | None = None,
) -> dict[str, Any]:
    input_path, destination = Path(input_path).resolve(), Path(destination).resolve()
    headers, records, raw_rows = _read_records(input_path)
    input_ids = [str(record["source_record_id"]) for record in records]
    if len(input_ids) != 893:
        raise ValueError("original input must contain exactly 893 source_record_id values")
    _baseline_manifest_payload, baseline, baseline_evidence_rows = _validate_baseline_bundle(
        input_ids=input_ids,
        results_path=Path(baseline_results).resolve(),
        evidence_path=Path(baseline_evidence).resolve(),
        manifest_path=Path(baseline_manifest).resolve(),
    )
    tasks = [
        _task(record, {
            **baseline.get(str(record["source_record_id"]), {}),
            "publication_eligible": baseline_evidence_rows[str(record["source_record_id"])]["publication_decision"]["publishable"],
            "publication_blockers": "; ".join(baseline_evidence_rows[str(record["source_record_id"])]["publication_decision"]["blockers"]),
        })
        for record in records
        if not baseline_evidence_rows[str(record["source_record_id"])]["publication_decision"]["publishable"]
    ]
    if not tasks:
        raise ValueError("baseline contains no remediation rows")
    task_ids = {task["source_record_id"] for task in tasks}
    child_rows = [record for record in records if str(record["source_record_id"]) in task_ids]
    if len(child_rows) != len(tasks):
        raise ValueError("remediation task/source input coverage mismatch")
    destination.mkdir(parents=True, exist_ok=False)
    child_input = destination / "remediation_input.xlsx"
    plan_path = destination / "remediation_plan.json"
    _write_child_input(child_input, headers, child_rows)
    child_input_hash = _sha256(child_input)
    parent = {}
    if parent_manifest is not None:
        parent_manifest = Path(parent_manifest).resolve()
        parent_payload = json.loads(parent_manifest.read_text(encoding="utf-8"))
        parent = {
            "manifest": str(parent_manifest),
            "manifest_sha256": _sha256(parent_manifest),
            "run_id": parent_payload.get("run_id", ""),
            "artifact_set_sha256": parent_payload.get("artifact_set_sha256", ""),
        }
    run_config = run_context.RunConfig.from_config(paid_enabled=False)
    lineage = {"type": "remediation_child", "parent_run_id": parent.get("run_id", "")}
    run_id = run_context.canonical_run_id(
        input_sha256=child_input_hash,
        ordered_source_record_ids=[task["source_record_id"] for task in tasks],
        effective_config=run_config.as_dict(),
        runtime_source_tree_hash=run_context.source_tree_sha256(),
        lineage=lineage,
    )
    plan = {
        "schema_version": 1,
        "plan_kind": "remediation",
        "coverage": {"input_count": len(records), "baseline_count": len(baseline), "task_count": len(tasks), "review_source_ids": len(tasks)},
        "tasks": tasks,
        "child": {
            "input": str(child_input), "input_sha256": child_input_hash,
            "rows": len(child_rows), "run_id": run_id,
            "config_sha256": run_config.sha256,
            "runtime_source_tree_sha256": run_context.source_tree_sha256(),
        },
        "parent": parent,
        "source": {
            "original_input": str(input_path), "original_input_sha256": _sha256(input_path),
            "baseline_results": str(baseline_results),
            "baseline_results_sha256": _sha256(baseline_results),
            "baseline_evidence": str(baseline_evidence),
            "baseline_evidence_sha256": _sha256(baseline_evidence),
            "baseline_manifest": str(baseline_manifest),
            "baseline_manifest_sha256": _sha256(baseline_manifest),
        },
        "baseline_bundle": {
            "source_record_ids": input_ids,
            "results_sha256": _sha256(baseline_results),
            "evidence_sha256": _sha256(baseline_evidence),
            "manifest_sha256": _sha256(baseline_manifest),
        },
    }
    plan["plan_payload_sha256"] = hashlib.sha256(_canonical(plan).encode("utf-8")).hexdigest()
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    return {"plan": str(plan_path), "input": str(child_input), "run_id": run_id, "task_count": len(tasks)}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", "--original-input", dest="input_path", type=Path, required=True)
    parser.add_argument("--baseline", "--baseline-results", dest="baseline_results", type=Path, required=True)
    parser.add_argument("--baseline-evidence", type=Path, required=True)
    parser.add_argument("--baseline-manifest", type=Path, required=True)
    parser.add_argument("--parent-manifest", type=Path)
    parser.add_argument("--destination", type=Path, required=True)
    args = parser.parse_args(argv)
    print(json.dumps(prepare_remediation_run(**vars(args)), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
